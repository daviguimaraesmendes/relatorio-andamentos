"""Teste de fumaça do painel local (Flask): GET em cada rota de leitura e POSTs
inofensivos (sempre com token), num relatório temporário, sem rede, sem subir
processo e sem tocar em projetos/, config.json ou nos dados reais.

Cada resposta é registrada (status, redirecionamento, tipo e um resumo do HTML
normalizado: token, pastas temporárias e datas do dia viram marcadores) e
comparada com tests/fixtures/painel_instantaneo.json. O instantâneo foi gerado
ANTES do fatiamento do revisao.py em src/painel/: se uma resposta mudar, o
fatiamento mudou o comportamento. Para regravar de propósito (tela alterada
de verdade): PAINEL_GRAVAR=1 python3 -m unittest tests/test_painel.py
Para ver o HTML normalizado de cada passo: PAINEL_DUMP=<pasta>.

    python3 -m unittest tests/test_painel.py -v
"""
import contextlib
import hashlib
import io
import json
import os
import re
import sys
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent))
from isolamento import TMP  # noqa: E402  (antes de tudo)
import comum  # noqa: E402

GOLDEN = Path(__file__).resolve().parent / "fixtures" / "painel_instantaneo.json"
PROC = "1234567-06.2026.8.06.0001"      # dígito verificador correto (número fictício)
PROC_2 = "1234568-72.2024.4.05.8100"
CLIENTE = "EMPRESA TESTE COMERCIO LTDA"
GLOBAIS = ("PROJETOS_DIR", "ATUAL_FILE", "CONFIG_FILE", "PROJETO", "PROJETO_DIR", "PROJETO_FILE", "DATA",
           "CARTEIRA_FILE", "CLIENTES_FILE", "EVENTOS_FILE", "ESTADO_FILE", "DOCS_DIR", "TEXTOS_DIR",
           "RELATORIOS_DIR", "DIAG_DIR", "PRINTS_DIR")
REGISTRO = {}


def _tarefa():
    """O dicionário da tarefa em andamento (muda de lugar quando o painel é fatiado)."""
    import revisao
    if hasattr(revisao, "TAREFA"):
        return revisao.TAREFA
    from painel import base
    return base.TAREFA


def _normalizar(texto, token):
    texto = texto.replace(token, "«TOKEN»").replace(str(TMP), "«TMP»").replace(str(comum.RAIZ), "«RAIZ»")
    texto = texto.replace(__import__("datetime").date.today().isoformat(), "«HOJE»")
    texto = re.sub(r"início \d{2}:\d{2}", "início «HH:MM»", texto)
    texto = re.sub(r"<b>\d{6}</b>", "<b>«CODIGO»</b>", texto)
    return re.sub(r"\d{8}-\d{6}-", "«CARIMBO»-", texto)


class ProcFalso:
    """Substitui subprocess.Popen: nada é executado, mas o pedido fica registrado."""
    chamadas = []
    codigo = None   # None = em andamento; número = terminou com esse código

    def __init__(self, args, **kw):
        ProcFalso.chamadas.append((list(args), kw))
        if kw.get("stdout"):
            kw["stdout"].close()   # o painel abre o log e entrega ao processo
        self.pid = 4242
        self.args = args

    def poll(self):
        return ProcFalso.codigo

    @property
    def returncode(self):
        return ProcFalso.codigo


class Painel(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.salvo = {n: getattr(comum, n) for n in GLOBAIS}
        comum.PROJETOS_DIR = TMP / "projetos"
        comum.ATUAL_FILE = comum.PROJETOS_DIR / ".projeto_atual"
        comum.CONFIG_FILE = TMP / "config-painel.json"
        comum.PROJETOS_DIR.mkdir(parents=True, exist_ok=True)
        cls.slug = comum.criar_projeto("Grupo Exemplo")
        comum.usar_projeto(cls.slug)
        for caminho in (comum.DATA, comum.CARTEIRA_FILE, comum.CLIENTES_FILE, comum.PROJETO_FILE, comum.CONFIG_FILE):
            if not str(Path(caminho).resolve()).startswith(str(TMP.resolve())):
                raise SystemExit(f"TESTE ABORTADO: {caminho} aponta para fora da pasta temporária.")
        import acesso
        cls.cofre = {}
        cls.patches = [
            mock.patch.object(acesso, "obter", lambda chave: cls.cofre.get(chave)),
            mock.patch.object(acesso, "guardar", lambda chave, valor: cls.cofre.__setitem__(chave, valor)),
            mock.patch("subprocess.Popen", ProcFalso),
        ]
        for p in cls.patches:
            p.start()
        import revisao
        cls.revisao = revisao
        cls.TOKEN = revisao.TOKEN
        cls.app = revisao.app
        cls.app.config["TESTING"] = True
        cls.c = cls.app.test_client()
        cls.dump = Path(os.environ["PAINEL_DUMP"]) if os.environ.get("PAINEL_DUMP") else None
        if cls.dump:
            cls.dump.mkdir(parents=True, exist_ok=True)
        ProcFalso.chamadas.clear()
        ProcFalso.codigo = None
        _tarefa().clear()
        REGISTRO.clear()

    @classmethod
    def tearDownClass(cls):
        for p in cls.patches:
            p.stop()
        _tarefa().clear()
        for n, v in cls.salvo.items():
            setattr(comum, n, v)

    # --- apoio ---

    def passo(self, nome, resp, contem=(), ausente=()):
        """Registra e confere uma resposta. `contem`: trechos-chave que têm de estar no corpo."""
        tipo = resp.headers.get("Content-Type", "").split(";")[0]
        info = {"status": resp.status_code, "tipo": tipo}
        if resp.headers.get("Location"):
            info["location"] = _normalizar(resp.headers["Location"], self.TOKEN)
        if resp.headers.get("Content-Disposition"):
            info["disposicao"] = _normalizar(resp.headers["Content-Disposition"], self.TOKEN)
        if tipo.startswith("text/") or tipo == "application/json":
            corpo = _normalizar(resp.get_data(as_text=True), self.TOKEN)
            info["tamanho"] = len(corpo)
            info["sha256"] = hashlib.sha256(corpo.encode()).hexdigest()[:16]
            for trecho in contem:
                self.assertTrue(trecho in corpo, f"{nome}: faltou o trecho {trecho!r}")
            for trecho in ausente:
                self.assertFalse(trecho in corpo, f"{nome}: não devia ter {trecho!r}")
            if self.dump:
                (self.dump / (re.sub(r"\W+", "_", nome) + ".txt")).write_text(corpo, encoding="utf-8")
        elif tipo != "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":  # xlsx: data dentro do zip
            info["bytes"] = len(resp.get_data())
        self.assertNotIn(nome, REGISTRO, f"passo repetido: {nome}")
        REGISTRO[nome] = info
        resp.close()
        return resp

    def get(self, caminho, nome=None, **kw):
        return self.passo(nome or f"GET {caminho}", self.c.get(caminho), **kw)

    def post(self, caminho, dados=None, nome=None, token=True, **kw):
        corpo = dict(dados or {})
        if token:
            corpo["token"] = self.TOKEN
        contem, ausente = kw.pop("contem", ()), kw.pop("ausente", ())
        resp = self.c.post(caminho, data=corpo, **kw)
        return self.passo(nome or f"POST {caminho}", resp, contem=contem, ausente=ausente)

    def seguir(self, resp, nome, **kw):
        """GET no endereço do redirecionamento (a mensagem aparece na tela seguinte)."""
        self.assertEqual(resp.status_code, 302, nome)
        return self.get(resp.headers["Location"], nome=nome, **kw)

    # --- as rotas e a proteção por token ---

    def test_01_rotas_registradas(self):
        regras = sorted(f"{r.rule} {','.join(sorted(r.methods - {'HEAD', 'OPTIONS'}))} -> {r.endpoint}"
                        for r in self.app.url_map.iter_rules() if r.endpoint != "static")
        REGISTRO["rotas"] = regras
        self.assertEqual(len([r for r in regras if r.startswith("/cadastro")]), 7)

    def test_02_sem_token_e_rotas_inexistentes(self):
        for caminho in ("/evento", "/tarefa", "/interromper", "/mostrar", "/novo", "/planilha", "/config", "/acesso",
                        "/cadastro/cliente", "/cadastro/importar", "/cadastro/processo", "/cadastro/completar",
                        "/cadastro/descobrir", "/cadastro/pendente"):
            self.post(caminho, {"id": "x", "acao": "salvar", "numero": "x", "indice": "novo"}, token=False,
                      nome=f"POST sem token {caminho}")
            self.assertEqual(REGISTRO[f"POST sem token {caminho}"]["status"], 403)
        self.get("/nao-existe", nome="GET rota inexistente")
        self.get("/evento", nome="GET em rota só de POST")

    # --- relatórios (abas) ---

    def test_03_projetos_e_cookie(self):
        self.get(f"/p/{self.slug}", nome="GET /p/<slug> troca de relatório")
        info = self.c.get_cookie("projeto")
        REGISTRO["cookie do relatório"] = f"{info.value};{info.max_age};{info.same_site}"
        self.assertEqual(info.value, self.slug)
        self.get("/p/inexistente", nome="GET /p/<inexistente>")
        self.get("/novo", contem=("<h1>Novo relatório</h1>", "name='nome'", "«TOKEN»"))
        self.seguir(self.post("/novo", {"nome": "  "}, nome="POST /novo sem nome"), "mensagem /novo sem nome",
                    contem=("Dê um nome ao relatório.",))

    def test_04_sem_relatorio_nenhum(self):
        """Sem relatório criado, tudo cai na tela de criar (exceto /novo, /acesso e as tarefas)."""
        dir_real, atual_real = comum.PROJETOS_DIR, comum.ATUAL_FILE
        vazio = TMP / "projetos-vazio"
        vazio.mkdir(exist_ok=True)
        comum.PROJETOS_DIR, comum.ATUAL_FILE = vazio, vazio / ".projeto_atual"
        try:
            for caminho in ("/", "/planilha", "/config", "/cadastro", "/documento", "/relatorio"):
                self.get(caminho, nome=f"GET {caminho} sem relatório (vai para /novo)")
                self.assertEqual(REGISTRO[f"GET {caminho} sem relatório (vai para /novo)"]["location"], "/novo")
            self.get("/novo", nome="GET /novo sem relatório", contem=("<h1>Novo relatório</h1>",))
            self.get("/acesso", nome="GET /acesso sem relatório", contem=("<h1>Acesso e escritório</h1>",))
            r = self.post("/novo", {"nome": "Relatório Dois"}, nome="POST /novo cria relatório")
            self.assertEqual(self.c.get_cookie("projeto").value, "relatorio-dois")
            self.assertTrue((vazio / "relatorio-dois" / "projeto.json").exists())
            self.seguir(r, "mensagem /cadastro após criar", contem=("Relatório criado: Relatório Dois.", "Clientes e processos"))
        finally:
            comum.PROJETOS_DIR, comum.ATUAL_FILE = dir_real, atual_real
            comum.usar_projeto(self.slug)
            self.c.set_cookie("projeto", self.slug)

    # --- Revisar (/ e /evento) ---

    def test_05_revisar(self):
        docs, prints = comum.DOCS_DIR / "teste", comum.PRINTS_DIR
        docs.mkdir(parents=True, exist_ok=True)
        prints.mkdir(parents=True, exist_ok=True)
        (docs / "decisao.html").write_text("<p>DECISÃO fictícia.</p>", encoding="utf-8")
        (prints / "decisao.png").write_bytes(b"\x89PNG-fake")
        (comum.DATA.parent / "fora.html").write_text("<p>fora da pasta de documentos</p>", encoding="utf-8")
        comum.save_json(comum.CARTEIRA_FILE, [{"numero": PROC, "cliente": CLIENTE, "ativo": True}])
        self.get("/", nome="GET / sem rascunhos", contem=("<h1>Revisar andamentos</h1>", "Nada para revisar.",
                                                         "<b>0</b>para revisar", "<b>1</b>processos acompanhados"))
        base = {"numero": PROC, "cliente": CLIENTE, "tipo_evento": "documento", "detectado_em": "2026-10-03T19:00:00",
                "status": "rascunho", "polo_cliente": "passivo", "efeito": "favoravel"}
        comum.salvar_eventos([
            {**base, "id": "t:1", "titulo": "1 - Decisão", "data": "01/10/2026", "grau": "1º grau", "modelo": "teste:3b",
             "frase": "Foi proferida decisão.", "conteudo": "negando a urgência", "prazo": "15 dias",
             "audiencia": "12/11/2026, às 9h", "trecho_origem": "INDEFIRO o pedido", "alertas": ["trecho de teste"],
             "arquivo": str(docs / "decisao.html"), "print": str(prints / "decisao.png")},
            {**base, "id": "t:2", "titulo": "Conclusos", "data": "02/10/2026", "frase": "Foi ao juiz.", "polo_cliente": "ativo",
             "efeito": "incerto"},
            {**base, "id": "t:3", "titulo": "Juntada", "data": "03/10/2026", "frase": "Houve juntada.", "cliente": "",
             "polo_cliente": None, "efeito": None},
            {**base, "id": "t:4", "titulo": "Descartável", "data": "04/10/2026", "frase": "Foi descartado."},
            {**base, "id": "t:5", "titulo": "Sem frase", "data": "05/10/2026", "frase": ""},
            {**base, "id": "t:6", "titulo": "Já aprovado", "data": "06/10/2026", "frase": "Já aprovado.", "status": "aprovado"},
            {**base, "id": "t:7", "titulo": "Fora", "data": "07/10/2026", "frase": "x", "status": "aprovado",
             "arquivo": str(comum.DATA.parent / "fora.html")},
        ])
        self.get("/", nome="GET / com rascunhos", contem=(
            "<b>5</b>para revisar", "<b>2</b>aprovados", "action='/evento'", "name='frase'", "abrir documento",
            "/documento?id=t:1", "/print?id=t:1", "modelo teste:3b", "efeito para o cliente: <b>favorável</b>",
            "(réu)", "(autor)", "polo não informado", "Trecho do documento", "<div class='alerta'>trecho de teste</div>",
            "<h2>sem cliente</h2>", f"<h2>{CLIENTE}</h2>", "Aprovar", "Descartar", "Salvar sem aprovar"))
        self.get("/documento?id=t:1", nome="GET /documento", contem=("DECISÃO fictícia.",))
        self.get("/print?id=t:1", nome="GET /print")
        self.get("/documento?id=t:2", nome="GET /documento sem arquivo")
        self.get("/documento?id=t:7", nome="GET /documento fora da pasta")
        self.get("/documento", nome="GET /documento sem id")
        self.get("/print?id=t:2", nome="GET /print sem print")
        self.get("/print?id=nao", nome="GET /print id inexistente")

        self.post("/evento", {"id": "nao", "acao": "salvar"}, nome="POST /evento inexistente")
        self.post("/evento", {"id": "t:6", "acao": "salvar"}, nome="POST /evento já aprovado")
        self.post("/evento", {"id": "t:5", "acao": "aprovar", "frase": " "}, nome="POST /evento aprovar sem frase")
        self.post("/evento", {"id": "t:1", "acao": "aprovar", "frase": " Foi proferida decisão. ", "conteudo": "negando",
                              "prazo": "", "audiencia": "12/11/2026"}, nome="POST /evento aprovar")
        self.post("/evento", {"id": "t:2", "acao": "salvar", "frase": "Foi ao juiz (editado).", "conteudo": "",
                              "prazo": "", "audiencia": ""}, nome="POST /evento salvar")
        self.post("/evento", {"id": "t:4", "acao": "descartar", "frase": "Foi descartado."}, nome="POST /evento descartar")
        estado = {e["id"]: (e["status"], e.get("frase"), e.get("conteudo"), e.get("prazo"), e.get("audiencia"),
                            bool(e.get("aprovado_por")), bool(e.get("aprovado_em")), e.get("motivo"))
                  for e in comum.eventos()}
        REGISTRO["eventos depois das ações"] = estado
        self.assertEqual(estado["t:1"][0], "aprovado")
        self.assertEqual(estado["t:2"][:2], ("rascunho", "Foi ao juiz (editado)."))
        self.assertEqual(estado["t:4"][0], "descartado")
        self.assertEqual(REGISTRO["POST /evento aprovar sem frase"]["status"], 400)
        self.get("/", nome="GET / depois das ações", contem=("<b>3</b>para revisar", "<b>3</b>aprovados"))

    # --- Atualizar, tarefa e interromper ---

    def test_06_atualizar_tarefas(self):
        self.get("/atualizar", nome="GET /atualizar primeira vez", contem=(
            "<h1>Atualizar andamentos</h1>", "Primeira atualização deste relatório", "name='desde'", "name='tipo' value='rodada'",
            "name='tipo' value='conferencia'", "Conferência"), ausente=("Interromper",))
        comum.save_json(comum.ESTADO_FILE, {"x": 1})
        self.get("/atualizar", nome="GET /atualizar já atualizado", ausente=("Primeira atualização",))
        self.post("/tarefa", {"tipo": "invalido"}, nome="POST /tarefa tipo inválido")
        self.assertEqual(REGISTRO["POST /tarefa tipo inválido"]["status"], 400)
        self.assertEqual(ProcFalso.chamadas, [])

        ProcFalso.codigo = None
        self.post("/tarefa", {"tipo": "rodada", "desde": "2026-09-15"}, nome="POST /tarefa rodada")
        self.get("/atualizar", nome="GET /atualizar tarefa em andamento", contem=(
            "Atualização no jus.br", "em andamento", "action='/interromper'", "location.reload", "disabled"))
        self.get("/", nome="GET / com tarefa em andamento", contem=("tarefa em andamento: Atualização no jus.br",))
        self.seguir(self.post("/tarefa", {"tipo": "conferencia"}, nome="POST /tarefa com outra em andamento"),
                    "mensagem /tarefa duplicada", contem=("Já há uma tarefa em andamento. Aguarde ou interrompa.",))
        with mock.patch("os.killpg") as killpg, mock.patch("os.getpgid", lambda pid: pid + 1):
            self.seguir(self.post("/interromper", nome="POST /interromper"), "mensagem /interromper",
                        contem=("Tarefa interrompida. O que já foi coletado fica salvo.",))
            REGISTRO["killpg"] = [list(map(str, c.args)) for c in killpg.call_args_list]
        self.assertEqual(len(killpg.call_args_list), 1)

        ProcFalso.codigo = 0
        self.get("/atualizar", nome="GET /atualizar tarefa concluída", contem=("concluída (código 0)", "Ir para a revisão"),
                 ausente=("action='/interromper'",))
        with mock.patch("os.killpg") as killpg:
            self.post("/interromper", nome="POST /interromper sem tarefa rodando")
            self.assertEqual(killpg.call_args_list, [])
        self.post("/tarefa", {"tipo": "conferencia", "dias": "ab12"}, nome="POST /tarefa conferência")
        ProcFalso.codigo = 1
        self.get("/atualizar", nome="GET /atualizar conferência concluída", contem=("Conferência (12 dias)", "concluída (código 1)"))
        self.post("/tarefa", {"tipo": "teste_acesso"}, nome="POST /tarefa teste de acesso")
        ProcFalso.codigo = None
        self.get("/acesso", nome="GET /acesso com teste em andamento", contem=("<pre class='log'>", "setTimeout"))
        ProcFalso.codigo = 0
        self.get("/acesso", nome="GET /acesso com teste concluído", contem=("<pre class='log'>",), ausente=("setTimeout",))
        # o que foi pedido ao sistema operacional (comandos, pasta e ambiente, sem rodar nada)
        REGISTRO["processos pedidos"] = [
            {"args": [_normalizar(a, self.TOKEN) for a in args], "cwd": _normalizar(str(kw.get("cwd")), self.TOKEN),
             "relatorio_env": sorted(k for k in kw["env"] if k.startswith("RELATORIO_")),
             "projeto_env": kw["env"].get("RELATORIO_PROJETO"), "unbuffered": kw["env"].get("PYTHONUNBUFFERED"),
             "stderr": kw.get("stderr"), "start_new_session": kw.get("start_new_session")}
            for args, kw in ProcFalso.chamadas]
        self.assertEqual(len(ProcFalso.chamadas), 3)
        _tarefa().clear()
        ProcFalso.codigo = None

    # --- Planilha do mês e relatório ---

    def _modelo_xlsx(self, caminho):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Processos"
        ws["A1"], ws["P1"] = "Número", "Andamentos"
        ws["A2"], ws["P2"] = PROC, "Até 29/09/2026 sem atualizações."
        ws["A3"] = PROC_2
        wb.save(caminho)

    def test_07_planilha_e_relatorio(self):
        self.get("/planilha", nome="GET /planilha vazia", contem=(
            "<h1>Planilha do mês</h1>", "<b>3</b> andamento(s) aprovado(s)", "Referência atual: nenhuma",
            "enctype='multipart/form-data'", "Nenhuma ainda."))
        self.seguir(self.post("/planilha", nome="POST /planilha sem modelo"), "mensagem /planilha sem modelo",
                    contem=("Envie a planilha do relatório anterior",))
        import io as _io
        self.seguir(self.post("/planilha", {"modelo": (_io.BytesIO(b"x"), "lista.txt")}, nome="POST /planilha arquivo errado",
                              content_type="multipart/form-data"), "mensagem /planilha arquivo errado",
                    contem=("Envie uma planilha .xlsx.",))
        modelo = TMP / "modelo-entrada.xlsx"
        self._modelo_xlsx(modelo)
        r = self.post("/planilha", {"modelo": (io.BytesIO(modelo.read_bytes()), "2026-09-01 - Relatório Exemplo - atualizada.xlsx"),
                                    "sem_novidade": "1"}, nome="POST /planilha gera", content_type="multipart/form-data")
        self.seguir(r, "mensagem /planilha gerada", contem=("Planilha gerada: 1 processo(s) atualizado(s).",
                                                            "Planilhas geradas", "Mostrar no Finder", "/relatorio?p=planilhas/"))
        gerados = sorted((comum.RELATORIOS_DIR / "planilhas").glob("*.xlsx"))
        self.assertEqual(len(gerados), 1)
        REGISTRO["planilha gerada"] = re.sub(r"^\d{4}-\d{2}-\d{2}", "«HOJE»", gerados[0].name)
        proj = comum.projeto()
        REGISTRO["projeto depois da planilha"] = {"modelo": _normalizar(proj["planilha_modelo"], self.TOKEN),
                                                  "ultimo": _normalizar(proj["ultimo_relatorio"], self.TOKEN)}
        self.get("/planilha", nome="GET /planilha com arquivo", contem=("Referência atual:", "Relatório Exemplo"))
        # sem enviar arquivo: usa a referência atual (a que acabou de ser gerada)
        self.seguir(self.post("/planilha", {}, nome="POST /planilha com a referência atual"),
                    "mensagem /planilha referência atual", contem=("Planilha gerada:",))
        rel = "planilhas/" + gerados[0].name
        r = self.get(f"/relatorio?p={rel}", nome="GET /relatorio xlsx")
        self.assertIn("attachment", r.headers["Content-Disposition"])
        (comum.RELATORIOS_DIR / "Relatório teste.html").write_text("<p>relatório fictício</p>", encoding="utf-8")
        self.get("/relatorio?p=Relatório teste.html", nome="GET /relatorio html", contem=("relatório fictício",))
        self.get("/relatorio?p=../../../etc/passwd", nome="GET /relatorio fora da pasta")
        self.get("/relatorio?p=planilhas", nome="GET /relatorio pasta")
        self.get("/relatorio", nome="GET /relatorio sem p")
        with mock.patch("subprocess.run") as run:
            self.seguir(self.post("/mostrar", {"p": rel}, nome="POST /mostrar"), "tela depois de /mostrar")
            REGISTRO["mostrar chamou"] = [[_normalizar(str(a), self.TOKEN) for a in c.args[0]] for c in run.call_args_list]
            self.post("/mostrar", {"p": "../x"}, nome="POST /mostrar fora da pasta")
            self.post("/mostrar", {"p": "nao-existe.xlsx"}, nome="POST /mostrar inexistente")
            self.assertEqual(len(run.call_args_list), 1)
        # modelo referido que sumiu
        proj["planilha_modelo"] = str(TMP / "sumiu.xlsx")
        comum.salvar_projeto(proj)
        self.get("/planilha", nome="GET /planilha referência sumida", contem=("(arquivo não encontrado)",))

    # --- Configuração ---

    def test_08_configuracao(self):
        self.get("/config", nome="GET /config", contem=("<h1>Configuração do relatório</h1>", "name='planilha_modelo'",
                                                        "name='ultimo_relatorio'", "Pasta deste relatório: «TMP»"))
        self.seguir(self.post("/config", {"nome": "Grupo Exemplo", "planilha_modelo": "/nao/existe.xlsx"},
                              nome="POST /config planilha inexistente"), "mensagem /config planilha inexistente",
                    contem=("Planilha não encontrada: /nao/existe.xlsx",))
        modelo = TMP / "modelo-entrada.xlsx"
        self.seguir(self.post("/config", {"nome": "Grupo Exemplo Renomeado", "planilha_modelo": str(modelo),
                                          "ultimo_relatorio": "2026-09-30"}, nome="POST /config salva"),
                    "mensagem /config salva", contem=("Configuração salva.", "Grupo Exemplo Renomeado", "2026-09-30"))
        self.seguir(self.post("/config", {"nome": "", "planilha_modelo": "", "ultimo_relatorio": ""},
                              nome="POST /config nome vazio mantém o nome"), "mensagem /config nome vazio",
                    contem=("Configuração salva.", "Grupo Exemplo Renomeado"))
        self.assertEqual(comum.projeto()["planilha_modelo"], "")

    # --- Acesso e escritório ---

    def test_09_acesso(self):
        self.get("/acesso", nome="GET /acesso nada configurado", contem=(
            "<h1>Acesso e escritório</h1>", "não configurada", "name='cert_senha'", "name='totp_secret'",
            "name='identificadores'", "name='revisor'", "(configurar)"), ausente=("✓", "Falta configurar o acesso"))
        self.get("/", nome="GET / com aviso de acesso", contem=("Falta configurar o acesso", "<a href='/acesso'"))
        self.seguir(self.post("/acesso", {"totp_secret": "1!1"}, nome="POST /acesso segredo inválido"),
                    "mensagem /acesso segredo inválido", contem=("O segredo do autenticador não é válido",))
        self.assertEqual(self.cofre, {})
        self.assertFalse(comum.CONFIG_FILE.exists())
        self.seguir(self.post("/acesso", {"cert_senha": "senha-ficticia", "totp_secret": "jbsw y3dp ehpk 3pxp",
                                          "identificadores": "Fulano Teste OAB 12.345\n\n  Beltrano Teste  \n",
                                          "revisor": " Revisora Teste "}, nome="POST /acesso salva tudo"),
                    "mensagem /acesso salva", contem=("Senha do certificado guardada no cofre do sistema.",
                                                      "Segredo do autenticador guardado no cofre do sistema.",
                                                      "Dados do escritório salvos.", "configurada ✓", "Fulano Teste OAB 12.345",
                                                      "Revisora Teste", "Código de agora", "Acesso e escritório ✓"),
                    ausente=("Falta configurar o acesso",))
        REGISTRO["cofre"] = sorted(self.cofre)
        self.assertEqual(self.cofre["totp_secret"], "JBSWY3DPEHPK3PXP")
        cfg = comum.load_json(comum.CONFIG_FILE, {})
        REGISTRO["config depois do acesso"] = {k: cfg.get(k) for k in ("identificadores_escritorio", "revisor")}
        self.seguir(self.post("/acesso", {"identificadores": "", "revisor": ""}, nome="POST /acesso só dados do escritório"),
                    "mensagem /acesso só dados", contem=("Dados do escritório salvos.",), ausente=("Senha do certificado guardada",))
        self.assertEqual(sorted(self.cofre), ["cert_senha", "totp_secret"], "em branco, mantém o que está guardado")

    # --- Clientes e processos (cadastro.py, já separado) ---

    def test_10_cadastro(self):
        self.get("/cadastro", nome="GET /cadastro vazio", contem=("<h1>Clientes e processos</h1>", "action='/cadastro/cliente'"))
        self.seguir(self.post("/cadastro/descobrir", {"dias": "x"}, nome="POST /cadastro/descobrir sem clientes"),
                    "mensagem /cadastro descobrir", contem=("Cadastre ao menos um cliente antes de procurar no DJEN.",))
        self.seguir(self.post("/cadastro/cliente", {"acao": "salvar", "indice": "novo", "nome": CLIENTE,
                                                    "variacoes": "EMPRESA TESTE, TESTE LTDA", "contato": "a@exemplo.test",
                                                    "responsavel": "Fulano"}, nome="POST /cadastro/cliente novo"),
                    "mensagem /cadastro cliente salvo", contem=("Cliente salvo: " + CLIENTE, "EMPRESA TESTE, TESTE LTDA"))
        self.seguir(self.post("/cadastro/cliente", {"acao": "salvar", "indice": "novo", "nome": CLIENTE.lower()},
                              nome="POST /cadastro/cliente repetido"), "mensagem /cadastro cliente repetido",
                    contem=("já está cadastrado.",))
        self.seguir(self.post("/cadastro/importar", nome="POST /cadastro/importar vazio"), "mensagem /cadastro importar vazio",
                    contem=("Nada para importar",))
        with contextlib.redirect_stdout(io.StringIO()):
            self.seguir(self.post("/cadastro/importar", {"texto": f"{PROC}\n{PROC_2}\n1234567-07.2026.8.06.0001\n",
                                                         "cliente": CLIENTE, "polo": "passivo", "parte_contraria": "Fulano de Tal",
                                                         "responsavel": "Beltrano"}, nome="POST /cadastro/importar texto"),
                        "mensagem /cadastro importar", contem=("2 número(s) lido(s), 1 novo(s)", "Recusados"))
        self.get("/cadastro", nome="GET /cadastro com dados", contem=(PROC, PROC_2, "action='/cadastro/processo'"))
        self.seguir(self.post("/cadastro/processo", {"acao": "salvar", "numero": PROC, "cliente": CLIENTE, "polo": "ativo",
                                                     "parte_contraria": "Outro", "responsavel": "Sicrano", "ativo": "1"},
                              nome="POST /cadastro/processo salvar"), "mensagem /cadastro processo salvo",
                    contem=("Processo salvo: " + PROC,))
        self.seguir(self.post("/cadastro/cliente", {"acao": "remover", "indice": "0"}, nome="POST /cadastro/cliente remover em uso"),
                    "mensagem /cadastro remover em uso", contem=("tem 2 processo(s) na carteira",))
        with mock.patch("carteira.completar_com_djen", lambda apenas=None: (0, 0)):
            self.seguir(self.post("/cadastro/completar", nome="POST /cadastro/completar"), "mensagem /cadastro completar",
                        contem=("DJEN: 0 processo(s) completado(s)",))
        import djen
        comum.save_json(djen.DESCOBERTA_FILE, {"pendentes": [
            {"numero": "1234569-95.2026.5.07.0001", "tribunal": "TRT7", "cliente": CLIENTE, "polo_cliente": "ativo",
             "outras_partes": "Fulano", "ultima": "Notificação"},
            {"numero": "1234570-71.2025.8.06.0001", "tribunal": "TJCE", "cliente": CLIENTE, "polo_cliente": "passivo",
             "outras_partes": "Beltrano", "ultima": "Intimação"}]})
        self.get("/cadastro", nome="GET /cadastro com pendentes", contem=("1234569-95.2026.5.07.0001", "action='/cadastro/pendente'"))
        with contextlib.redirect_stdout(io.StringIO()):
            self.seguir(self.post("/cadastro/pendente", {"numero": "1234569-95.2026.5.07.0001", "acao": "incluir"},
                                  nome="POST /cadastro/pendente incluir"), "mensagem /cadastro pendente incluído",
                        contem=("incluído na carteira.",))
            self.seguir(self.post("/cadastro/pendente", {"numero": "1234570-71.2025.8.06.0001", "acao": "ignorar"},
                                  nome="POST /cadastro/pendente ignorar"), "mensagem /cadastro pendente ignorado",
                        contem=("ignorado.",))
        self.seguir(self.post("/cadastro/processo", {"acao": "remover", "numero": PROC_2}, nome="POST /cadastro/processo remover"),
                    "mensagem /cadastro processo removido", contem=("Processo removido da carteira: " + PROC_2,))
        REGISTRO["carteira depois"] = sorted((p["numero"], p.get("polo_cliente"), p.get("ativo", True))
                                             for p in comum.load_json(comum.CARTEIRA_FILE, []))

    # --- comparação com o instantâneo ---

    def test_99_equivalencia_com_o_instantaneo(self):
        if "carteira depois" not in REGISTRO:
            self.skipTest("rode a classe inteira (os passos anteriores alimentam o registro)")
        atual = json.loads(json.dumps(REGISTRO, ensure_ascii=False, sort_keys=True))
        if os.environ.get("PAINEL_GRAVAR"):
            GOLDEN.parent.mkdir(parents=True, exist_ok=True)
            GOLDEN.write_text(json.dumps(atual, ensure_ascii=False, indent=1, sort_keys=True) + "\n", encoding="utf-8")
            self.skipTest(f"instantâneo gravado em {GOLDEN} ({len(atual)} itens)")
        esperado = json.loads(GOLDEN.read_text(encoding="utf-8"))
        self.assertEqual(sorted(atual), sorted(esperado), "passos a mais ou a menos que no instantâneo")
        for chave in sorted(esperado):
            with self.subTest(passo=chave):
                self.assertEqual(atual[chave], esperado[chave])


if __name__ == "__main__":
    unittest.main()
