"""Testes do kit de pedidos das iniciais (WS-16): esquema, prompt, validador, pacote, gravação, planilha e tela.

Tudo com dados fictícios e sem rede: números de processo vêm de `ficticio.numero_ficticio` (ou do exemplo.json,
que só usa os números sintéticos permitidos), PDFs são gerados por `simulado.pdf_minimo`.

    python3 -m unittest tests/test_pedidos.py -v
"""
import copy
import io
import json
import re
import sys
import unittest
import zipfile
from decimal import Decimal
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent))
from isolamento import TMP  # noqa: E402,F401  (antes de tudo)
import ficticio  # noqa: E402
import simulado  # noqa: E402
import comum  # noqa: E402
import ficha as fch  # noqa: E402
import pedidos  # noqa: E402
import taxonomia  # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent


# ---------------------------------------------------------------- fabricação de resultados fictícios

def cadastro_de(**extra):
    c = {"reclamante": "Pessoa Fictícia 0001", "funcao": "Auxiliar de produção", "categoria_funcao": "Operacional",
         "empresa_principal": "Empresa Fictícia 001 Ltda", "outras_empresas": [], "terceiros": [],
         "municipio": "Cidade Exemplo", "uf": "CE", "vara": "1ª Vara do Trabalho de Cidade Exemplo",
         "tribunal": "TRT7", "data_ajuizamento": "2025-03-10", "valor_causa": 100000.0,
         "criterio_valor_causa": "soma dos pedidos", "advogado_reclamante": None,
         "tipo_acao": "Reclamação trabalhista", "terceirizado": "Não"}
    c.update(extra)
    return c


def pedido(materia="Horas extras e reflexos", valor=60000.0, situacao="atribuído", entra=True, pagina=3, texto=None, **extra):
    p = {"materia": materia, "pedido_como_formulado": texto or f"Pedido de {materia.lower()}", "valor_atribuido": valor,
         "situacao_valor": situacao, "pagina_pdf": pagina, "entra_nos_totais": entra}
    p.update(extra)
    return p


def pedidos_padrao():
    """Soma 100.000,00, igual ao valor da causa padrão."""
    return [pedido("Horas extras e reflexos", 60000.0, pagina=3), pedido("Dano moral", 25000.0, pagina=5),
            pedido("Adicional de insalubridade", 15000.0, pagina=6)]


def processo(numero, cadastro=None, peds=None, achados=None):
    return {"numero": numero, "cadastro": cadastro if cadastro is not None else cadastro_de(),
            "pedidos": peds if peds is not None else pedidos_padrao(), "achados": achados or []}


def resultado(*procs, **extra):
    return {"versao_prompt": pedidos.prompt_versao(), "processos": list(procs), **extra}


def validar(dados, carteira, **kw):
    texto = dados if isinstance(dados, str) else json.dumps(dados, ensure_ascii=False)
    return pedidos.validar(texto, carteira=carteira, **kw)


def codigos(res, nivel=None):
    return [a["codigo"] for a in res["erros"] + res["avisos"] if nivel is None or a["nivel"] == nivel]


# ---------------------------------------------------------------- modelos versionados

class Modelos(unittest.TestCase):
    def test_exemplo_confere_com_o_esquema_e_so_usa_numeros_sinteticos(self):
        ex = pedidos.exemplo()
        self.assertEqual(pedidos.verificar_esquema(ex, pedidos.esquema()), [])
        for p in ex["processos"]:
            self.assertTrue(ficticio.dv_confere(p["numero"]), p["numero"])
            self.assertRegex(p["numero"], r"^123456[7-9]-|^1234570-")

    def test_materias_do_exemplo_estao_no_vocabulario(self):
        for p in pedidos.exemplo()["processos"]:
            for ped in p["pedidos"]:
                self.assertIn(ped["materia"], taxonomia.MATERIA)

    def test_prompt_versionado_e_marcado_como_nao_validado(self):
        bruto = (RAIZ / "src" / "modelos" / "pedidos" / "prompt.md").read_text(encoding="utf-8")
        self.assertRegex(bruto.splitlines()[0], r"versao: \d+\.\d+\.\d+")
        self.assertIn("NÃO VALIDADO com petições reais", bruto.splitlines()[0])
        self.assertEqual(pedidos.prompt_versao(), "1.0.0")
        self.assertFalse(pedidos.prompt_validado())

    def test_prompt_pronto_traz_vocabulario_esquema_exemplo_e_regras(self):
        texto = pedidos.prompt_pronto()
        self.assertNotIn("{{", texto)
        self.assertNotIn("<!--", texto)
        for materia in taxonomia.MATERIA:
            self.assertIn(f"- {materia}", texto)
        self.assertIn('"situacao_valor"', texto)           # esquema embutido
        self.assertIn("Pessoa Fictícia 0001", texto)       # exemplo embutido
        for regra in ("exatamente como o reclamante atribuiu", "arts. 467 e 477", "sem valor atribuído",
                      "fora do valor da causa", "encargos embutidos na causa", "somente com JSON", "página ilegível",
                      "Versão do prompt: 1.0.0"):
            self.assertIn(regra.lower(), texto.lower(), regra)

    def test_prompt_acompanha_o_vocabulario_atual(self):
        novo = dict(taxonomia.MATERIA, **{"Matéria Nova De Teste": ("Tema X", "Mérito", True)})
        with mock.patch.object(taxonomia, "MATERIA", novo):
            self.assertIn("- Matéria Nova De Teste", pedidos.prompt_pronto())

    @unittest.skipUnless(__import__("importlib").util.find_spec("jsonschema"), "jsonschema não instalado")
    def test_verificador_minimo_equivale_ao_jsonschema(self):
        import jsonschema
        sch = pedidos.esquema()
        base = pedidos.exemplo()
        variantes = [base]
        for mexer in (lambda d: d["processos"][0].pop("numero"),
                      lambda d: d["processos"][0]["pedidos"][0].__setitem__("entra_nos_totais", "sim"),
                      lambda d: d["processos"][0]["pedidos"][0].__setitem__("situacao_valor", "outra"),
                      lambda d: d["processos"][0]["cadastro"].__setitem__("uf", "ce"),
                      lambda d: d["processos"][0]["cadastro"].__setitem__("data_ajuizamento", "ontem"),
                      lambda d: d["processos"][0]["pedidos"][0].__setitem__("pagina_pdf", 0),
                      lambda d: d["processos"][0]["pedidos"][0].__setitem__("pagina_pdf", 1.5),
                      lambda d: d["processos"][0]["pedidos"][0].__setitem__("valor_atribuido", True),
                      lambda d: d["processos"][0]["cadastro"].__setitem__("terceirizado", "Talvez"),
                      lambda d: d["processos"][0]["cadastro"].__setitem__("outras_empresas", "uma"),
                      lambda d: d.__setitem__("processos", []),
                      lambda d: d["processos"][0]["achados"][0].pop("descricao"),
                      lambda d: d["processos"][0]["cadastro"].pop("valor_causa")):
            v = copy.deepcopy(base)
            mexer(v)
            variantes.append(v)
        for v in variantes:
            esperado = bool(list(jsonschema.Draft202012Validator(sch).iter_errors(v)))
            obtido = bool(pedidos.verificar_esquema(v, sch))
            self.assertEqual(obtido, esperado, json.dumps(v)[:200])


# ---------------------------------------------------------------- validador (cada tipo de erro)

class Validador(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.n = [ficticio.numero_ficticio(i, 2025, 5, 7, i + 1) for i in range(6)]
        cls.carteira = cls.n

    def test_resultado_correto_nao_tem_erro_nem_atencao(self):
        res = validar(resultado(processo(self.n[0]), processo(self.n[1])), self.carteira)
        self.assertTrue(res["ok"], res["erros"])
        self.assertEqual(codigos(res, "atencao"), [])
        self.assertEqual([p["numero"] for p in res["dados"]["processos"]], self.n[:2])
        r = res["dados"]["processos"][0]["resumo"]
        self.assertEqual((r["valor_causa"], r["soma_pedidos"], r["diferenca"], r["n_pedidos"]),
                         ("100000.00", "100000.00", "0.00", 3))
        self.assertEqual(r["materia_maior_peso"], "Horas extras e reflexos")

    def test_dados_saem_normalizados(self):
        cad = cadastro_de(valor_causa="R$ 100.000,00", data_ajuizamento="10/03/2025", empresa_principal="  Empresa   Fictícia 001 Ltda ")
        peds = [pedido(valor="R$ 60.000,00"), pedido("Dano moral", "25000.00"), pedido("Adicional de insalubridade", 15000)]
        res = validar(resultado(processo(self.n[0].replace(".", "").replace("-", ""), cad, peds)), self.carteira)
        self.assertTrue(res["ok"], res["erros"])
        p = res["dados"]["processos"][0]
        self.assertEqual(p["numero"], self.n[0])                      # com máscara
        self.assertEqual(p["cadastro"]["data_ajuizamento"], "2025-03-10")
        self.assertEqual(p["cadastro"]["valor_causa"], "100000.00")
        self.assertEqual([x["valor_atribuido"] for x in p["pedidos"]], ["60000.00", "25000.00", "15000.00"])
        self.assertEqual(p["cadastro"]["empresa_principal"], "Empresa Fictícia 001 Ltda")

    def test_json_invalido_e_cortado(self):
        texto = json.dumps(resultado(processo(self.n[0])), ensure_ascii=False)
        res = pedidos.validar(texto[:-30], carteira=self.carteira)
        self.assertFalse(res["ok"])
        self.assertIsNone(res["dados"])
        self.assertEqual(codigos(res), ["json_invalido"])
        self.assertIn("cortada", res["erros"][0]["mensagem"])
        self.assertEqual(codigos(pedidos.validar("isto nao e json", carteira=self.carteira)), ["json_invalido"])
        self.assertEqual(codigos(pedidos.validar("   ", carteira=self.carteira)), ["json_vazio"])

    def test_json_dentro_de_cerca_ou_com_texto_em_volta(self):
        texto = json.dumps(resultado(processo(self.n[0])), ensure_ascii=False)
        for embrulhado in (f"Aqui está o resultado:\n```json\n{texto}\n```\nQualquer dúvida, avise.",
                           f"Segue o JSON {{conforme pedido}}: {texto} Fim."):
            res = pedidos.validar(embrulhado, carteira=self.carteira)
            self.assertTrue(res["ok"], res["erros"])
            self.assertIn("resposta_com_texto_extra", codigos(res, "info"))

    def test_lista_ou_processo_solto_sem_raiz_e_aceito_com_aviso(self):
        res = validar([processo(self.n[0])], self.carteira)
        self.assertTrue(res["ok"])
        self.assertIn("lista_sem_raiz", codigos(res, "info"))
        res = validar(processo(self.n[0]), self.carteira)
        self.assertTrue(res["ok"])
        self.assertIn("processo_unico_sem_raiz", codigos(res, "info"))

    def test_esquema_invalido_na_raiz_e_por_processo(self):
        res = validar({"processos": []}, self.carteira)
        self.assertEqual((res["ok"], res["dados"], codigos(res)), (False, None, ["esquema_invalido"]))
        res = validar({"outra": 1}, self.carteira)
        self.assertFalse(res["ok"])
        sem_numero = processo(self.n[0])
        del sem_numero["numero"]
        tipo_errado = processo(self.n[1], peds=[pedido(entra="sim")])
        res = validar(resultado(sem_numero, tipo_errado, processo(self.n[2])), self.carteira)
        self.assertFalse(res["ok"])
        self.assertEqual(codigos(res, "erro"), ["esquema_invalido", "esquema_invalido"])
        self.assertEqual([p["numero"] for p in res["dados"]["processos"]], [self.n[2]])   # o bom não se perde
        self.assertEqual([r["posicao"] for r in res["recusados"]], [1, 2])
        self.assertTrue(any("entra_nos_totais" in a["mensagem"] for a in res["erros"]))

    def test_processo_que_nao_e_objeto(self):
        res = validar({"processos": ["texto solto", processo(self.n[0])]}, self.carteira)
        self.assertEqual(codigos(res, "erro"), ["esquema_invalido"])
        self.assertEqual(len(res["dados"]["processos"]), 1)

    def test_numero_com_digito_errado_ou_ilegivel(self):
        errado = ficticio.numero_com_dv_errado(self.n[0])
        res = validar(resultado(processo(errado), processo("SEM NÚMERO - arquivo.pdf"), processo(self.n[1])), self.carteira)
        self.assertEqual(codigos(res, "erro"), ["numero_invalido", "numero_invalido"])
        self.assertIn("dígito verificador", res["erros"][0]["mensagem"])
        self.assertEqual([p["numero"] for p in res["dados"]["processos"]], [self.n[1]])

    def test_processo_desconhecido_da_carteira(self):
        fora = ficticio.numero_ficticio(900, 2025, 5, 7, 9)
        res = validar(resultado(processo(fora), processo(self.n[0])), self.carteira)
        self.assertEqual(codigos(res, "erro"), ["processo_desconhecido"])
        self.assertEqual(res["recusados"][0]["numero"], fora)
        self.assertEqual(len(res["dados"]["processos"]), 1)

    def test_processo_vinculado_aponta_o_principal(self):
        f = fch.nova_ficha(self.n[0])
        vinc = ficticio.numero_ficticio(901, 2025, 5, 7, 0)
        fch.vincular(f, vinc, "agravo")
        with ficticio.projeto_de_teste([f]):
            res = pedidos.validar(json.dumps(resultado(processo(vinc))))
        self.assertEqual(codigos(res, "erro"), ["processo_vinculado"])
        self.assertEqual(res["erros"][0]["candidatos"], [self.n[0]])

    def test_processo_duplicado_no_resultado(self):
        res = validar(resultado(processo(self.n[0]), processo(self.n[1]), processo(self.n[0])), self.carteira)
        self.assertEqual(codigos(res, "erro"), ["processo_duplicado", "processo_duplicado"])
        self.assertEqual([p["numero"] for p in res["dados"]["processos"]], [self.n[1]])
        self.assertIn("1, 3", res["erros"][0]["mensagem"])

    def test_valor_nao_numerico_e_negativo(self):
        cad = cadastro_de(valor_causa="a apurar")
        peds = [pedido(valor="a apurar"), pedido("Dano moral", "entre 1000 e 2000"), pedido("Adicional noturno", -5),
                pedido("Tempo de espera", "R$ 100,00 (cem reais)")]
        res = validar(resultado(processo(self.n[0], cad, peds)), self.carteira)
        self.assertFalse(res["ok"])
        self.assertEqual(sorted(codigos(res, "erro")), sorted(["valor_negativo"] + ["valor_nao_numerico"] * 4))
        self.assertEqual(res["dados"]["processos"], [])

    def test_pedido_atribuido_sem_valor_e_erro(self):
        res = validar(resultado(processo(self.n[0], peds=[pedido(valor=None)])), self.carteira)
        self.assertEqual(codigos(res, "erro"), ["valor_ausente"])

    def test_processo_sem_pedidos_e_erro(self):
        res = validar(resultado(processo(self.n[0], peds=[])), self.carteira)
        self.assertEqual(codigos(res, "erro"), ["processo_sem_pedidos"])

    def test_materia_fora_do_vocabulario_vira_outros_com_sugestoes(self):
        peds = [pedido("Horas extras e reflexos", 60000.0), pedido("Dno morl", 25000.0),
                pedido("Adicional de periculosidade", 15000.0)]
        res = validar(resultado(processo(self.n[0], peds=peds)), self.carteira)
        self.assertTrue(res["ok"])
        aviso = next(a for a in res["avisos"] if a["codigo"] == "materia_fora_do_vocabulario")
        self.assertEqual(aviso["nivel"], "atencao")
        self.assertIn("Dano moral", aviso["candidatos"])
        self.assertIn("pedido 2", aviso["onde"])
        ped = res["dados"]["processos"][0]["pedidos"][1]
        self.assertEqual((ped["materia"], ped["materia_original"], ped["materia_sugerida"]), ("Outros", "Dno morl", "Dno morl"))

    def test_materia_totalmente_estranha_nao_e_adivinhada(self):
        res = validar(resultado(processo(self.n[0], peds=[pedido("Xyzzy plugh", 100000.0)])), self.carteira)
        aviso = next(a for a in res["avisos"] if a["codigo"] == "materia_fora_do_vocabulario")
        self.assertEqual(aviso["candidatos"], [])
        self.assertEqual(res["dados"]["processos"][0]["pedidos"][0]["materia"], "Outros")

    def test_materia_por_sinonimo_e_normalizada_com_info(self):
        peds = [pedido("horas extras", 60000.0), pedido("Dano moral", 25000.0), pedido("Insalubridade", 15000.0)]
        res = validar(resultado(processo(self.n[0], peds=peds)), self.carteira)
        self.assertEqual(codigos(res, "atencao"), [])
        self.assertEqual(codigos(res, "info").count("materia_normalizada"), 2)
        self.assertEqual([p["materia"] for p in res["dados"]["processos"][0]["pedidos"]],
                         ["Horas extras e reflexos", "Dano moral", "Adicional de insalubridade"])

    def test_outros_com_sugestao_vira_info(self):
        peds = pedidos_padrao()[:2] + [pedido("Outros", 15000.0, materia_sugerida="Multas dos arts. 467 e 477 da CLT")]
        res = validar(resultado(processo(self.n[0], peds=peds)), self.carteira)
        self.assertIn("materia_sugerida", codigos(res, "info"))
        self.assertEqual(res["dados"]["processos"][0]["pedidos"][2]["materia_sugerida"], "Multas dos arts. 467 e 477 da CLT")

    def test_soma_diferente_do_valor_da_causa_avisa_com_o_valor(self):
        cad = cadastro_de(valor_causa=120000.0)
        res = validar(resultado(processo(self.n[0], cad)), self.carteira)
        self.assertTrue(res["ok"])
        aviso = next(a for a in res["avisos"] if a["codigo"] == "soma_diverge_do_valor_da_causa")
        self.assertEqual(aviso["nivel"], "atencao")
        for trecho in ("R$ 100.000,00", "R$ 120.000,00", "R$ 20.000,00", "a menos"):
            self.assertIn(trecho, aviso["mensagem"])
        self.assertEqual(res["dados"]["processos"][0]["resumo"]["diferenca"], "20000.00")
        a_mais = validar(resultado(processo(self.n[0], cadastro_de(valor_causa=90000.0))), self.carteira)
        self.assertIn("a mais", next(a for a in a_mais["avisos"] if a["codigo"] == "soma_diverge_do_valor_da_causa")["mensagem"])

    def test_tolerancia_da_soma_e_configuravel(self):
        cad = cadastro_de(valor_causa=100000.5)
        self.assertEqual(codigos(validar(resultado(processo(self.n[0], cad)), self.carteira), "atencao"), [])
        cad = cadastro_de(valor_causa=100010.0)
        self.assertIn("soma_diverge_do_valor_da_causa", codigos(validar(resultado(processo(self.n[0], cad)), self.carteira)))
        self.assertEqual(codigos(validar(resultado(processo(self.n[0], cad)), self.carteira, tolerancia="50"), "atencao"), [])

    def test_criterio_que_nao_e_soma_deixa_a_divergencia_so_informativa(self):
        for criterio in ("valor arbitrado pelo reclamante", "valor da alçada", "sem soma dos pedidos"):
            cad = cadastro_de(valor_causa=50000.0, criterio_valor_causa=criterio)
            res = validar(resultado(processo(self.n[0], cad)), self.carteira)
            self.assertEqual(codigos(res, "atencao"), [], criterio)
            self.assertIn("soma_diverge_do_valor_da_causa", codigos(res, "info"))

    def test_encargos_embutidos_compoem_o_valor_da_causa(self):
        peds = [pedido("Horas extras e reflexos", 90000.0),
                pedido("Encargos (INSS/IR/custas) no valor da causa", 10000.0, "encargos embutidos na causa", entra=False)]
        res = validar(resultado(processo(self.n[0], peds=peds)), self.carteira)
        self.assertEqual(codigos(res, "atencao"), [])
        r = res["dados"]["processos"][0]["resumo"]
        self.assertEqual((r["soma_pedidos"], r["encargos_embutidos"], r["diferenca"]), ("90000.00", "10000.00", "0.00"))

    def test_pedidos_sem_valor_e_fora_da_causa_nao_entram_na_soma(self):
        peds = pedidos_padrao() + [pedido("Honorários advocatícios", None, "sem valor atribuído", entra=False),
                                   pedido("Tempo de espera", 7000.0, "fora do valor da causa", entra=False)]
        res = validar(resultado(processo(self.n[0], peds=peds)), self.carteira)
        self.assertEqual(codigos(res, "atencao"), [])
        r = res["dados"]["processos"][0]["resumo"]
        self.assertEqual((r["soma_pedidos"], r["n_pedidos"], r["pedidos_sem_valor"]), ("100000.00", 5, 1))

    def test_valor_da_causa_ausente_avisa(self):
        res = validar(resultado(processo(self.n[0], cadastro_de(valor_causa=None))), self.carteira)
        self.assertIn("valor_causa_ausente", codigos(res, "atencao"))
        self.assertIsNone(res["dados"]["processos"][0]["resumo"]["diferenca"])

    def test_situacao_do_valor_incoerente(self):
        peds = pedidos_padrao() + [pedido("Tempo de espera", 3000.0, "sem valor atribuído", entra=False),
                                   pedido("Controle de jornada", None, "fora do valor da causa", entra=False),
                                   pedido("Dano estético", None, "sem valor atribuído", entra=True)]
        res = validar(resultado(processo(self.n[0], peds=peds)), self.carteira)
        self.assertEqual(codigos(res, "atencao").count("situacao_valor_incoerente"), 2)
        self.assertEqual(codigos(res, "atencao").count("entra_nos_totais_incoerente"), 1)

    def test_situacao_do_valor_com_acento_ou_sinonimo(self):
        peds = pedidos_padrao()[:2] + [pedido("Adicional de insalubridade", 15000.0, "Atribuido")]
        peds.append(pedido("Honorários advocatícios", None, "Sem valor", entra=False))
        res = validar(resultado(processo(self.n[0], peds=peds)), self.carteira)
        self.assertTrue(res["ok"], res["erros"])
        self.assertIn("situacao_valor_normalizada", codigos(res, "info"))
        res = validar(resultado(processo(self.n[0], peds=[pedido(situacao="quase atribuído")])), self.carteira)
        self.assertEqual(codigos(res, "erro"), ["esquema_invalido"])

    def test_data_e_uf_invalidas_avisam_e_ficam_vazias(self):
        cad = cadastro_de(data_ajuizamento="2025-13-45", uf="XX")
        res = validar(resultado(processo(self.n[0], cad)), self.carteira)
        self.assertTrue(res["ok"])
        self.assertEqual(sorted(codigos(res, "atencao")), ["data_invalida", "uf_invalida"])
        c = res["dados"]["processos"][0]["cadastro"]
        self.assertEqual((c["data_ajuizamento"], c["uf"]), (None, None))

    def test_campo_desconhecido_no_cadastro_e_cadastro_incompleto(self):
        cad = cadastro_de(cor_favorita="azul", vara=None, funcao=None)
        res = validar(resultado(processo(self.n[0], cad)), self.carteira)
        self.assertIn("campo_desconhecido", codigos(res, "atencao"))
        self.assertNotIn("cor_favorita", res["dados"]["processos"][0]["cadastro"])
        incompleto = next(a for a in res["avisos"] if a["codigo"] == "cadastro_incompleto")
        self.assertIn("funcao", incompleto["mensagem"])
        self.assertIn("vara", incompleto["mensagem"])

    def test_pedido_repetido_avisa(self):
        peds = pedidos_padrao() + [copy.deepcopy(pedidos_padrao()[0])]
        peds[-1]["valor_atribuido"] = 0.0
        res = validar(resultado(processo(self.n[0], peds=peds)), self.carteira)
        self.assertIn("pedido_duplicado", codigos(res, "atencao"))

    def test_processo_ja_extraido_avisa_que_sera_substituido(self):
        with ficticio.projeto_de_teste([fch.nova_ficha(n) for n in self.n[:2]]):
            dados = pedidos.validar(json.dumps(resultado(processo(self.n[0]))))["dados"]
            pedidos.gravar(None, dados)
            res = pedidos.validar(json.dumps(resultado(processo(self.n[0]), processo(self.n[1]))))
            ja = [a for a in res["avisos"] if a["codigo"] == "processo_ja_extraido"]
            self.assertEqual(len(ja), 1)
            self.assertIn(self.n[0], ja[0]["onde"])

    def test_versao_do_prompt_diferente_e_info(self):
        res = validar(resultado(processo(self.n[0]), versao_prompt="0.1"), self.carteira)
        self.assertIn("versao_do_prompt_diferente", codigos(res, "info"))

    def test_texto_gigante_e_recusado(self):
        res = pedidos.validar("x" * (pedidos.LIMITE_TEXTO + 1), carteira=self.carteira)
        self.assertEqual(codigos(res), ["texto_grande_demais"])

    def test_avisos_seguem_o_formato_estruturado(self):
        res = validar(resultado(processo(self.n[0], cadastro_de(valor_causa=1.0), [pedido("Xyzzy", 5.0)])), self.carteira)
        for a in res["erros"] + res["avisos"]:
            self.assertEqual(set(a), {"nivel", "codigo", "onde", "mensagem", "candidatos"})
            self.assertIn(a["nivel"], ("info", "atencao", "erro"))
            self.assertRegex(a["codigo"], r"^[a-z_]+$")


# ---------------------------------------------------------------- pacote, gravação, planilha

def inicial_pdf(pasta, nome="inicial.pdf", texto="Petição inicial fictícia"):
    pasta.mkdir(parents=True, exist_ok=True)
    arq = pasta / nome
    arq.write_bytes(simulado.pdf_minimo(texto))
    return arq


class ComProjeto(unittest.TestCase):
    """Relatório temporário com fichas fictícias."""

    def setUp(self):
        self.fichas = ficticio.gerar_carteira(6, clientes=2, semente=5)
        ctx = ficticio.projeto_de_teste(self.fichas)
        self.proj = ctx.__enter__()
        self.addCleanup(ctx.__exit__, None, None, None)
        self.numeros = [f["numero"] for f in self.fichas]
        self.data = comum.DATA

    def evento_inicial(self, numero, pdf, tipo="Petição Inicial", descricao="Petição Inicial", data="2025-03-10"):
        eventos = comum.load_json(comum.EVENTOS_FILE, [])
        eventos.append({"id": f"{numero}:{len(eventos)}", "tipo_evento": "documento", "numero": numero, "tipo": tipo,
                        "descricao": descricao, "data": data, "arquivo": str(pdf), "status": "coletado"})
        comum.save_json(comum.EVENTOS_FILE, eventos)

    def texto_resultado(self, indices, **kw):
        procs = []
        for i in indices:
            f = self.fichas[i]
            procs.append(processo(f["numero"], cadastro_de(reclamante=f"Pessoa Fictícia {i + 1:04d}", **kw)))
        return json.dumps(resultado(*procs), ensure_ascii=False)


class Pacote(ComProjeto):
    def test_localiza_inicial_por_evento_pasta_manual_e_pasta_de_documentos(self):
        pdf = inicial_pdf(self.data / "baixados", "x.pdf")
        self.evento_inicial(self.numeros[0], pdf)
        manual = inicial_pdf(self.data / "iniciais", f"{self.numeros[1]}.pdf")
        doc = inicial_pdf(comum.DOCS_DIR / "cliente-x" / comum.slug(self.numeros[2]), "190000001 - Petição Inicial - Petição Inicial.pdf")
        inicial_pdf(comum.DOCS_DIR / "cliente-x" / comum.slug(self.numeros[4]), "190000002 - Sentença - Sentença.pdf")
        self.assertEqual(pedidos.localizar_inicial(self.numeros[0]), pdf)
        self.assertEqual(pedidos.localizar_inicial(self.numeros[1]), manual)
        self.assertEqual(pedidos.localizar_inicial(self.numeros[2]), doc)
        self.assertIsNone(pedidos.localizar_inicial(self.numeros[4]))      # só tem sentença
        self.assertIsNone(pedidos.localizar_inicial(self.numeros[5]))

    def test_evento_de_outro_tipo_ou_sem_arquivo_nao_conta(self):
        pdf = inicial_pdf(self.data / "baixados", "y.pdf")
        self.evento_inicial(self.numeros[0], pdf, "Sentença", "Sentença")
        self.evento_inicial(self.numeros[1], self.data / "nao-existe.pdf")
        self.assertIsNone(pedidos.localizar_inicial(self.numeros[0]))
        self.assertIsNone(pedidos.localizar_inicial(self.numeros[1]))

    def test_inicial_exata_vence_emenda(self):
        emenda = inicial_pdf(self.data / "baixados", "emenda.pdf")
        exata = inicial_pdf(self.data / "baixados", "exata.pdf")
        self.evento_inicial(self.numeros[0], emenda, "Petição", "Emenda à petição inicial", "2025-03-01")
        self.evento_inicial(self.numeros[0], exata, "Petição Inicial", "Petição Inicial", "2025-03-10")
        self.assertEqual(pedidos.localizar_inicial(self.numeros[0]), exata)

    def test_pacote_traz_so_quem_tem_inicial_e_nao_foi_extraido(self):
        self.assertFalse(self.fichas[3]["ativo"])          # processo encerrado também tem pedidos a extrair
        for i in (0, 1, 3):
            inicial_pdf(self.data / "iniciais", f"{self.numeros[i]}.pdf")
        lista = pedidos.pacote()
        self.assertEqual([c["numero"] for c in lista], [self.numeros[0], self.numeros[1], self.numeros[3]])   # inclui encerrado
        self.assertTrue(all(Path(c["pdf_inicial"]).is_file() and not c["ja_extraido"] for c in lista))
        pedidos.gravar(None, pedidos.validar(self.texto_resultado([1]))["dados"])
        self.assertEqual([c["numero"] for c in pedidos.pacote()], [self.numeros[0], self.numeros[3]])
        todos = pedidos.pacote(incluir_extraidos=True)
        self.assertEqual([(c["numero"], c["ja_extraido"]) for c in todos],
                         [(self.numeros[0], False), (self.numeros[1], True), (self.numeros[3], False)])

    def test_situacao_lista_todos_os_processos(self):
        inicial_pdf(self.data / "iniciais", f"{self.numeros[0]}.pdf")
        linhas = pedidos.situacao()
        self.assertEqual(len(linhas), 6)
        self.assertEqual(sum(1 for s in linhas if s["pdf_inicial"]), 1)
        self.assertEqual(linhas[0]["cliente"], fch.obter(self.fichas[0], "cliente"))

    def test_projeto_por_pasta_ou_slug(self):
        inicial_pdf(self.data / "iniciais", f"{self.numeros[0]}.pdf")
        for ref in (self.proj["pasta"], self.proj["slug"], self.proj):
            self.assertEqual([c["numero"] for c in pedidos.pacote(ref)], [self.numeros[0]])

    def test_exportar_pacote_para_pasta_com_lotes_prompt_e_leia_me(self):
        for i in range(5):
            inicial_pdf(self.data / "iniciais", f"{self.numeros[i]}.pdf", texto=f"Inicial fictícia {i}")
        destino = TMP / "saida-pacote"
        r = pedidos.exportar_pacote(None, destino, por_lote=2)
        self.assertEqual((len(r["arquivos"]), r["lotes"], r["sem_inicial"]), (5, 3, []))
        self.assertTrue((destino / "lote-01" / f"{self.numeros[0]}.pdf").is_file())
        self.assertTrue((destino / "lote-03" / f"{self.numeros[4]}.pdf").is_file())
        self.assertEqual((destino / "prompt.md").read_text(encoding="utf-8"), pedidos.prompt_pronto())
        self.assertIn("NÃO VALIDADO", (destino / "LEIA-ME.txt").read_text(encoding="utf-8"))
        self.assertIn("fora do seu computador", (destino / "LEIA-ME.txt").read_text(encoding="utf-8"))

    def test_exportar_pacote_de_numeros_escolhidos_e_sem_inicial(self):
        inicial_pdf(self.data / "iniciais", f"{self.numeros[0]}.pdf")
        r = pedidos.exportar_pacote(None, TMP / "saida-2", numeros=[self.numeros[0], self.numeros[5]])
        self.assertEqual([p.name for p in r["arquivos"]], [f"{self.numeros[0]}.pdf"])
        self.assertEqual(r["sem_inicial"], [self.numeros[5]])

    def test_pacote_em_zip(self):
        for i in (0, 2):
            inicial_pdf(self.data / "iniciais", f"{self.numeros[i]}.pdf")
        with pedidos.pacote_zip(None) as arq:
            z = zipfile.ZipFile(arq)
            self.assertEqual(sorted(z.namelist()), sorted([f"{self.numeros[0]}.pdf", f"{self.numeros[2]}.pdf",
                                                           "prompt.md", "LEIA-ME.txt"]))
            self.assertTrue(z.read(f"{self.numeros[0]}.pdf").startswith(b"%PDF"))
            self.assertIn("Versão do prompt", z.read("prompt.md").decode("utf-8"))


class Gravacao(ComProjeto):
    def test_ida_e_volta_do_exemplo(self):
        ex = pedidos.exemplo()
        numeros = [p["numero"] for p in ex["processos"]]
        with ficticio.projeto_de_teste([fch.nova_ficha(n, cliente="Cliente Exemplo 01 Ltda") for n in numeros]):
            texto = (RAIZ / "src" / "modelos" / "pedidos" / "exemplo.json").read_text(encoding="utf-8")
            res = pedidos.validar(texto)
            self.assertTrue(res["ok"], res["erros"])
            self.assertEqual(codigos(res, "atencao"), [])   # o exemplo não deve gerar nenhum aviso de atenção
            saida = pedidos.gravar(None, res["dados"])
            self.assertEqual(saida["gravados"], numeros)
            guardado = pedidos.carregar()["processos"]
            self.assertEqual(list(guardado), numeros)
            for original, p in zip(ex["processos"], res["dados"]["processos"]):
                reg = guardado[p["numero"]]
                self.assertEqual(reg["cadastro"], p["cadastro"])
                self.assertEqual(len(reg["pedidos"]), len(original["pedidos"]))
                self.assertEqual([x["pedido_como_formulado"] for x in reg["pedidos"]],
                                 [x["pedido_como_formulado"] for x in original["pedidos"]])
                self.assertEqual([x["valor_atribuido"] for x in reg["pedidos"]],
                                 [None if x["valor_atribuido"] is None else f"{x['valor_atribuido']:.2f}" for x in original["pedidos"]])
                self.assertEqual(reg["achados"][0]["descricao"], original["achados"][0]["descricao"])
                self.assertEqual(reg["prompt_versao"], "1.0.0")
            # revalidar o que foi gravado (reexportado como resultado) devolve os mesmos dados
            de_volta = {"versao_prompt": "1.0.0", "processos": [
                {"numero": n, "cadastro": r["cadastro"], "pedidos": r["pedidos"], "achados": r["achados"]}
                for n, r in guardado.items()]}
            res2 = pedidos.validar(json.dumps(de_volta))
            self.assertTrue(res2["ok"], res2["erros"])
            for a, b in zip(res["dados"]["processos"], res2["dados"]["processos"]):
                self.assertEqual(a["pedidos"], b["pedidos"])
                self.assertEqual(a["cadastro"], b["cadastro"])

    def test_gravar_substitui_o_mesmo_processo_e_mantem_os_outros(self):
        pedidos.gravar(None, pedidos.validar(self.texto_resultado([0, 1]))["dados"])
        novo = resultado(processo(self.numeros[1], peds=[pedido("Dano moral", 100000.0, pagina=1)]))
        saida = pedidos.gravar(None, pedidos.validar(json.dumps(novo))["dados"])
        self.assertEqual((saida["gravados"], saida["substituidos"]), ([self.numeros[1]], [self.numeros[1]]))
        guardado = pedidos.carregar()["processos"]
        self.assertEqual(set(guardado), {self.numeros[0], self.numeros[1]})
        self.assertEqual(len(guardado[self.numeros[1]]["pedidos"]), 1)
        self.assertEqual(len(guardado[self.numeros[0]]["pedidos"]), 3)

    def test_gravar_so_aceita_dados_do_validador(self):
        with self.assertRaises(ValueError):
            pedidos.gravar(None, None)
        with self.assertRaises(ValueError):
            pedidos.gravar(None, {"processos": "x"})

    def test_gravar_sem_pedir_nao_toca_nas_fichas(self):
        antes = comum.load_json(comum.CARTEIRA_FILE, [])
        pedidos.gravar(None, pedidos.validar(self.texto_resultado([0, 1]))["dados"])
        self.assertEqual(comum.load_json(comum.CARTEIRA_FILE, []), antes)

    def test_atualizar_fichas_preenche_so_o_vazio_com_origem_humano(self):
        f0 = self.fichas[0]
        ficha_vazia = fch.nova_ficha(f0["numero"], cliente="Cliente Exemplo 01 Ltda")
        fch.definir(ficha_vazia, "vara", "Vara Antiga Coletada", "coletado")
        fch.definir(ficha_vazia, "uf", "SP", "humano")
        comum.save_json(comum.CARTEIRA_FILE, [ficha_vazia])
        dados = pedidos.validar(self.texto_resultado([0]))["dados"]
        previa = {l["campo"]: l for l in pedidos.conferir_cadastro(None, dados)}
        self.assertEqual(previa["vara"]["acao"], "divergente")
        self.assertEqual(previa["uf"]["acao"], "divergente")
        self.assertEqual(previa["municipio"]["acao"], "preencher")
        saida = pedidos.gravar(None, dados, atualizar_fichas=True)
        f = fch.carregar(todas=True)[0]
        self.assertEqual(fch.obter(f, "municipio"), "Cidade Exemplo")
        self.assertEqual(fch.origem(f, "municipio"), "humano")
        self.assertEqual(fch.obter(f, "valor_causa"), "100000.00")
        self.assertEqual(fch.obter(f, "autores"), "Pessoa Fictícia 0001")
        self.assertEqual(fch.obter(f, "reus"), "Empresa Fictícia 001 Ltda")
        self.assertEqual(fch.obter(f, "data_ajuizamento"), "2025-03-10")
        self.assertEqual((fch.obter(f, "vara"), fch.obter(f, "uf")), ("Vara Antiga Coletada", "SP"))   # divergentes intactos
        self.assertEqual({d["campo"] for d in saida["divergencias"]}, {"vara", "uf"})
        self.assertIn("municipio", {a["campo"] for a in saida["fichas_atualizadas"]})

    def test_sobrescrever_so_com_pedido_explicito(self):
        ficha_x = fch.nova_ficha(self.numeros[0], cliente="Cliente Exemplo 01 Ltda")
        fch.definir(ficha_x, "vara", "Vara Antiga Coletada", "coletado")
        comum.save_json(comum.CARTEIRA_FILE, [ficha_x])
        dados = pedidos.validar(self.texto_resultado([0]))["dados"]
        saida = pedidos.gravar(None, dados, atualizar_fichas=True, sobrescrever=True)
        f = fch.carregar(todas=True)[0]
        self.assertEqual(fch.obter(f, "vara"), "1ª Vara do Trabalho de Cidade Exemplo")
        self.assertEqual(fch.origem(f, "vara"), "humano")
        self.assertEqual(saida["divergencias"], [])

    def test_campos_so_de_pedidos_nao_vao_para_a_ficha(self):
        pedidos.gravar(None, pedidos.validar(self.texto_resultado([0]))["dados"], atualizar_fichas=True)
        f = fch.carregar(todas=True)[0]
        for campo in f.get("campos", {}):
            self.assertIn(campo, fch.CAMPOS)
        self.assertEqual(fch.validar(f), [])
        self.assertNotIn("funcao", f["campos"])


class Planilha(ComProjeto):
    def setUp(self):
        super().setUp()
        texto = self.texto_resultado([0, 1, 2])
        pedidos.gravar(None, pedidos.validar(texto)["dados"])
        # um com achados, encargos e pedido sem valor
        peds = [pedido("Horas extras e reflexos", 80000.0), pedido("Honorários advocatícios", None, "sem valor atribuído", entra=False),
                pedido("Encargos (INSS/IR/custas) no valor da causa", 20000.0, "encargos embutidos na causa", entra=False)]
        extra = resultado(processo(self.numeros[3], peds=peds, achados=[{"tipo": "pedido_sem_valor", "descricao": "Honorários sem valor.", "impacto": "baixo"}]))
        pedidos.gravar(None, pedidos.validar(json.dumps(extra, ensure_ascii=False))["dados"])

    def test_abas_e_colunas_do_modelo(self):
        abas = pedidos.linhas_para_planilha()
        self.assertEqual(list(abas), ["cadastro", "pedidos", "resumo", "parametros"])
        self.assertEqual(len(abas["cadastro"]["colunas"]), 26)
        self.assertEqual(abas["pedidos"]["colunas"][5], "Pedido (como formulado na inicial)")
        self.assertEqual(abas["resumo"]["colunas"][2], "Soma dos pedidos quantificados")
        self.assertEqual(abas["parametros"]["colunas"], ["Matéria", "Tema", "Classe", "Conta nos rankings", "Pode ser causa geradora"])
        for aba in abas.values():
            for linha in aba["linhas"]:
                self.assertEqual(len(linha), len(aba["colunas"]))
        self.assertEqual(len(abas["cadastro"]["linhas"]), 4)
        self.assertEqual(len(abas["pedidos"]["linhas"]), 3 * 3 + 3)
        self.assertEqual(len(abas["parametros"]["linhas"]), len(taxonomia.MATERIA))

    def test_percentual_e_resumo(self):
        abas = pedidos.linhas_para_planilha()
        primeiro = abas["pedidos"]["linhas"][0]
        self.assertAlmostEqual(primeiro[8], 0.6)                       # 60.000 / 100.000
        resumo = {l[0]: l for l in abas["resumo"]["linhas"]}[self.numeros[3]]
        self.assertEqual(resumo[1:4], [100000.0, 80000.0, 0.0])        # causa, soma, diferença (já desconta encargos)
        self.assertEqual((resumo[5], resumo[6], resumo[7]), (3, 1, "Horas extras e reflexos"))
        self.assertIn("Honorários sem valor.", resumo[8])
        self.assertIn("encargos embutidos", resumo[8])

    def test_cadastro_traz_o_que_a_ficha_sabe_do_julgamento(self):
        linha = {l[0]: l for l in pedidos.linhas_para_planilha()["cadastro"]["linhas"]}[self.numeros[0]]
        col = pedidos.COLUNAS["cadastro"]
        self.assertEqual(linha[col.index("Reclamante")], "Pessoa Fictícia 0001")
        self.assertEqual(linha[col.index("Situação")], fch.obter(self.fichas[0], "situacao"))
        self.assertEqual(linha[col.index("Data do ajuizamento")], "10/03/2025")

    def test_exportar_xlsx_ida_e_volta(self):
        from openpyxl import load_workbook
        destino = pedidos.exportar_xlsx(None, TMP / "pedidos-saida.xlsx")
        wb = load_workbook(destino)
        self.assertEqual(wb.sheetnames, ["Cadastro", "Pedidos", "Resumo", "Parâmetros por matéria"])
        abas = pedidos.linhas_para_planilha()
        for chave, nome in pedidos.ABAS_XLSX.items():
            ws = wb[nome]
            linhas = [list(r) for r in ws.iter_rows(values_only=True)]
            self.assertEqual(linhas[0], abas[chave]["colunas"])
            self.assertEqual(len(linhas) - 1, len(abas[chave]["linhas"]))
        pedidos_ws = [list(r) for r in wb["Pedidos"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual([r[4] for r in pedidos_ws], [l[4] for l in abas["pedidos"]["linhas"]])
        self.assertEqual([r[6] for r in pedidos_ws], [l[6] for l in abas["pedidos"]["linhas"]])
        with self.assertRaises(FileExistsError):                          # nunca sobrescreve
            pedidos.exportar_xlsx(None, destino)


class Provedor(unittest.TestCase):
    def test_indisponivel_por_padrao(self):
        self.assertEqual(pedidos.provedor_disponivel(), (False, pedidos.MENSAGEM_INDISPONIVEL))
        with self.assertRaises(ValueError) as e:
            pedidos.enviar_pelo_provedor(None, ["x"])
        self.assertIn("indisponível", str(e.exception))

    def test_ponto_de_extensao(self):
        with mock.patch.object(pedidos, "PROVEDOR", lambda projeto, numeros: f"resposta para {numeros}"):
            self.assertTrue(pedidos.provedor_disponivel()[0])
            self.assertEqual(pedidos.enviar_pelo_provedor(None, ["a", "b"]), "resposta para ['a', 'b']")


# ---------------------------------------------------------------- tela

class Tela(ComProjeto):
    @classmethod
    def setUpClass(cls):
        import acesso
        cls.cofre = {}
        cls.patches = [mock.patch.object(acesso, "obter", lambda chave: cls.cofre.get(chave)),
                       mock.patch.object(acesso, "guardar", lambda chave, valor: cls.cofre.__setitem__(chave, valor))]
        for p in cls.patches:
            p.start()
        import revisao
        cls.app, cls.TOKEN = revisao.app, revisao.TOKEN
        cls.app.config["TESTING"] = True

    @classmethod
    def tearDownClass(cls):
        for p in cls.patches:
            p.stop()

    def setUp(self):
        super().setUp()
        self.c = self.app.test_client()

    def post(self, caminho, dados=None, token=True):
        corpo = dict(dados or {})
        if token:
            corpo["token"] = self.TOKEN
        return self.c.post(caminho, data=corpo)

    def test_pagina_inicial_com_avisos_e_lista(self):
        for i in (0, 1):
            inicial_pdf(self.data / "iniciais", f"{self.numeros[i]}.pdf")
        pedidos.gravar(None, pedidos.validar(self.texto_resultado([1]))["dados"])
        r = self.c.get("/pedidos")
        corpo = r.get_data(as_text=True)
        self.assertEqual(r.status_code, 200)
        for trecho in ("Confidencialidade", "sai do seu computador", "NÃO VALIDADO com petições reais", "Prompt versão 1.0.0",
                       "Colar", "Cole aqui a resposta inteira da IA", "Baixar pacote", "Inicial não localizada",
                       "Pedidos extraídos em", "Enviar pelo provedor (indisponível)", self.numeros[0]):
            self.assertIn(trecho.lower(), corpo.lower(), trecho)
        self.assertRegex(corpo, rf"value='{re.escape(self.numeros[0])}' checked")
        self.assertNotRegex(corpo, rf"value='{re.escape(self.numeros[1])}' checked")     # já extraído: desmarcado

    def test_todo_post_exige_token(self):
        for rota in ("/pedidos/validar", "/pedidos/gravar", "/pedidos/pacote", "/pedidos/planilha", "/pedidos/enviar"):
            self.assertEqual(self.post(rota, {"texto": "x"}, token=False).status_code, 403, rota)

    def test_baixar_o_prompt(self):
        r = self.c.get("/pedidos/prompt")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_data(as_text=True), pedidos.prompt_pronto())
        self.assertIn("attachment", r.headers["Content-Disposition"])

    def test_colar_resultado_mostra_conferencia_antes_de_gravar(self):
        texto = self.texto_resultado([0, 1], valor_causa=120000.0)
        r = self.post("/pedidos/validar", {"texto": texto})
        corpo = r.get_data(as_text=True)
        self.assertEqual(r.status_code, 200)
        self.assertIn("Confira antes de gravar", corpo)
        self.assertIn("soma_diverge_do_valor_da_causa", corpo)
        self.assertIn("R$ 20.000,00", corpo)
        self.assertIn("Gravar 2 processo(s)", corpo)
        self.assertIn("Achados da IA", corpo)
        self.assertFalse(pedidos.arquivo_de_pedidos().exists())            # nada gravado só por conferir

    def test_erro_aparece_e_nao_oferece_gravar_o_que_esta_ruim(self):
        r = self.post("/pedidos/validar", {"texto": "isto não é json"})
        corpo = r.get_data(as_text=True)
        self.assertIn("json_invalido", corpo)
        self.assertIn("Nada a gravar", corpo)
        self.assertNotIn("Gravar 0", corpo)

    def test_texto_colado_e_escapado(self):
        r = self.post("/pedidos/validar", {"texto": "<script>alert(1)</script>"})
        corpo = r.get_data(as_text=True)
        self.assertNotIn("<script>alert(1)</script>", corpo)
        self.assertIn("&lt;script&gt;", corpo)

    def test_gravar_exige_confirmacao_dos_achados(self):
        texto = self.texto_resultado([0])
        r = self.post("/pedidos/gravar", {"texto": texto})
        self.assertEqual(r.status_code, 200)
        self.assertIn("Marque que conferiu", r.get_data(as_text=True))
        self.assertFalse(pedidos.arquivo_de_pedidos().exists())

    def test_gravar_depois_de_conferir(self):
        texto = self.texto_resultado([0, 1])
        r = self.post("/pedidos/gravar", {"texto": texto, "confirmo": "1"})
        self.assertEqual(r.status_code, 302)
        self.assertIn("/pedidos?msg=", r.headers["Location"])
        self.assertEqual(set(pedidos.carregar()["processos"]), set(self.numeros[:2]))

    def test_gravar_so_os_processos_sem_erro(self):
        bom, ruim = self.numeros[0], ficticio.numero_com_dv_errado(self.numeros[1])
        texto = json.dumps(resultado(processo(bom), processo(ruim)), ensure_ascii=False)
        r = self.post("/pedidos/gravar", {"texto": texto, "confirmo": "1"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(list(pedidos.carregar()["processos"]), [bom])
        self.assertIn("ficaram%20de%20fora", r.headers["Location"])

    def test_gravar_nao_grava_se_nao_ha_processo_valido(self):
        r = self.post("/pedidos/gravar", {"texto": "{}", "confirmo": "1"})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(pedidos.arquivo_de_pedidos().exists())

    def test_gravar_com_atualizacao_da_ficha(self):
        comum.save_json(comum.CARTEIRA_FILE, [fch.nova_ficha(self.numeros[0], cliente="Cliente Exemplo 01 Ltda")])
        texto = self.texto_resultado([0])
        self.post("/pedidos/gravar", {"texto": texto, "confirmo": "1", "atualizar_fichas": "1"})
        f = next(x for x in fch.carregar(todas=True) if x["numero"] == self.numeros[0])
        self.assertEqual(fch.origem(f, "municipio"), "humano")
        self.assertEqual(fch.obter(f, "municipio"), "Cidade Exemplo")

    def test_baixar_pacote_em_zip(self):
        for i in (0, 2):
            inicial_pdf(self.data / "iniciais", f"{self.numeros[i]}.pdf")
        r = self.post("/pedidos/pacote", {"numero": [self.numeros[0]], "por_lote": "0"})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.mimetype, "application/zip")
        z = zipfile.ZipFile(io.BytesIO(r.data))
        self.assertEqual(sorted(z.namelist()), sorted([f"{self.numeros[0]}.pdf", "prompt.md", "LEIA-ME.txt"]))
        r.close()
        r = self.post("/pedidos/pacote", {"por_lote": "1"})           # nada marcado: todos os sem pedidos
        z = zipfile.ZipFile(io.BytesIO(r.data))
        self.assertIn(f"lote-02/{self.numeros[2]}.pdf", z.namelist())
        r.close()

    def test_pacote_sem_nenhuma_inicial_avisa(self):
        r = self.post("/pedidos/pacote", {})
        self.assertEqual(r.status_code, 302)
        self.assertIn("Nenhuma", r.headers["Location"].replace("%20", " "))

    def test_planilha_baixada(self):
        r = self.post("/pedidos/planilha", {})
        self.assertEqual(r.status_code, 302)                           # ainda sem pedidos gravados
        pedidos.gravar(None, pedidos.validar(self.texto_resultado([0, 1]))["dados"])
        r = self.post("/pedidos/planilha", {})
        self.assertEqual(r.status_code, 200)
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.data))
        self.assertEqual(len(list(wb["Cadastro"].iter_rows(min_row=2))), 2)

    def test_enviar_pelo_provedor_indisponivel_e_extensao(self):
        r = self.post("/pedidos/enviar", {})
        self.assertEqual(r.status_code, 302)
        self.assertIn("indispon", r.headers["Location"].replace("%C3%AD", "í"))
        with mock.patch.object(pedidos, "PROVEDOR", lambda projeto, numeros: self.texto_resultado([0])):
            corpo = self.c.get("/pedidos").get_data(as_text=True)
            self.assertNotIn("(indisponível)", corpo)
            r = self.post("/pedidos/enviar", {})
            self.assertIn("Confira antes de gravar", r.get_data(as_text=True))


# ---------------------------------------------------------------- confidencialidade dos arquivos do kit

class Confidencialidade(unittest.TestCase):
    ARQUIVOS = ("src/pedidos.py", "src/painel/pedidos.py", "src/modelos/pedidos/prompt.md", "src/modelos/pedidos/esquema.json",
                "src/modelos/pedidos/exemplo.json", "docs/pedidos-iniciais.md", "tests/test_pedidos.py")
    CNJ = re.compile(r"\b[0-9]{7}-[0-9]{2}\.[0-9]{4}\.[0-9]\.[0-9]{2}\.[0-9]{4}\b")
    PERMITIDO = re.compile(r"0000000-00\.0000\.0\.00\.0000|9999999-99\.9999\.9\.99\.9999|123456[7-9]-|1234570-")

    def test_nenhum_numero_de_processo_real(self):
        for nome in self.ARQUIVOS:
            texto = (RAIZ / nome).read_text(encoding="utf-8")
            for m in self.CNJ.finditer(texto):
                self.assertRegex(m[0], self.PERMITIDO, f"{nome}: {m[0]}")

    def test_exemplo_so_tem_nomes_claramente_ficticios(self):
        texto = (RAIZ / "src/modelos/pedidos/exemplo.json").read_text(encoding="utf-8")
        for nome in re.findall(r'"(?:reclamante|empresa_principal|advogado_reclamante)": "([^"]+)"', texto):
            self.assertRegex(nome, r"Fictíci[ao]", nome)


if __name__ == "__main__":
    unittest.main()
