"""Testes do escritor DOCX do modelo A (src/escritores/docx_a.py). Só dados fictícios, sem rede.

    python3 -m unittest tests/test_docx_a.py -v

Os arquivos de entrada "tipo Google Docs" (runs fragmentados, rPr repetido, parágrafos vazios entre tabelas) vêm do
gerador do spike S2 (`spikes/s2_docx/gerar_modelo.py`, só dev: o pacote de distribuição não leva `spikes/`); os
números de processo são calculados em tempo de execução (`ficticio.numero_ficticio`). Os testes do LibreOffice são
pulados se o `soffice` não existir. NADA aqui prova que o Word ou o Google Docs abrem os arquivos: só python-docx, lxml
e LibreOffice foram usados (ver docs/fase2/conferencia-docx.md).
"""
import copy
import importlib.util
import re
import shutil
import subprocess
import sys
import tempfile
import time
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import isolamento  # noqa: F401  (antes de qualquer módulo da ferramenta)
import ficticio  # noqa: E402  (põe src/ no sys.path)
from lxml import etree  # noqa: E402

import ficha  # noqa: E402
from escritores import docx_a as da  # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent
_spec = importlib.util.spec_from_file_location("s2_gerar_modelo", RAIZ / "spikes" / "s2_docx" / "gerar_modelo.py")
gm = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(gm)

W = da.W_NS
w = da.w
DATA_BASE_NOVA = "07/10/2026"
SOFFICE = shutil.which("soffice")
TMP = Path(tempfile.mkdtemp(prefix="docx-a-"))
NUM_NOVO = ficticio.numero_ficticio(900)
NUM_NOVO_2 = ficticio.numero_ficticio(901)
NUM_NOVO_3 = ficticio.numero_ficticio(902)
NUM_AUSENTE = "9999999-99.9999.9.99.9999"       # número de exemplo permitido pelo empacotar.sh


def tearDownModule():
    shutil.rmtree(TMP, ignore_errors=True)


PROCS = gm.processos()


def num(i):
    return PROCS[i]["numeros"][0]


def doc_xml(caminho):
    with zipfile.ZipFile(caminho) as z:
        return etree.fromstring(z.read("word/document.xml"))


def c14n(el):
    return etree.tostring(el, method="c14n")


def corpo(caminho):
    return doc_xml(caminho).find(w("body"))


def tabelas(caminho):
    return list(corpo(caminho).findall(w("tbl")))


def texto_el(el):
    return "".join(t.text or "" for t in el.iter(w("t")))


def runs_do_par_andamentos(caminho, i):
    """Runs do parágrafo de 'Andamentos' do processo i (tabela i+1: a 0 é o resumo)."""
    tbl = tabelas(caminho)[i + 1]
    ultima = tbl.findall(w("tr"))[-1]
    par = [p for p in ultima.findall(w("tc"))[-1].findall(w("p"))][-1]
    return par, par.findall(w("r"))


def so_digitos_e_letras(s):
    return re.sub(r"\s+", "", s)


def escrever_doc(origem, destino, fn, parte="word/document.xml"):
    """Reescreve uma parte XML de `origem` aplicando fn(root); o resto do pacote fica igual."""
    with zipfile.ZipFile(origem) as zin, zipfile.ZipFile(destino, "w") as zout:
        for info in zin.infolist():
            dados = zin.read(info.filename)
            if info.filename == parte:
                raiz = etree.fromstring(dados)
                fn(raiz)
                dados = etree.tostring(raiz, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(info, dados)


def pdf_texto(docx_path):
    """Converte com o LibreOffice e devolve o texto do PDF (sem espaços/quebras, para comparar)."""
    from pypdf import PdfReader
    saida = Path(tempfile.mkdtemp(dir=TMP))
    perfil = saida / "perfil"
    r = subprocess.run([SOFFICE, f"-env:UserInstallation=file://{perfil}", "--headless", "--convert-to", "pdf",
                        "--outdir", str(saida), str(docx_path)], capture_output=True, text=True, timeout=240)
    pdf = saida / (Path(docx_path).stem + ".pdf")
    if not pdf.exists():
        raise AssertionError(f"soffice falhou: {r.stdout} {r.stderr}")
    leitor = PdfReader(str(pdf))
    return so_digitos_e_letras("".join(p.extract_text() for p in leitor.pages)), len(leitor.pages)


def proc_novo(numeros=None, tipos=None, **kw):
    p = {"numeros": numeros or [NUM_NOVO], "tipos": tipos or [], "assunto": "Ação de teste nova",
         "autores": "Fulana Exemplo de Tal", "reus": gm.CLIENTE, "ajuizamento": "2026-09-30", "valor_causa": "12345.60",
         "data_citacao": None, "juizo": "1ª Vara Fictícia da Comarca Modelo", "area": "Cível", "materia": "Contratos",
         "momento_atual": "aguardando citação",
         "andamentos": [{"data": "01/10/2026", "texto": "foi distribuída a ação."},
                        {"data": "03/10/2026", "texto": "foi determinada a citação da parte ré."}]}
    p.update(kw)
    return p


class Base(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.molde = TMP / "molde.docx"
        if not cls.molde.exists():
            gm.gerar(cls.molde, semente=7)

    def destino(self, nome):
        p = TMP / f"{self.id().split('.')[-1]}-{nome}.docx"
        if p.exists():
            p.unlink()
        return p

    def atualizacoes(self):
        return [
            {"numero": num(0), "momento_atual": "Cumprimento de sentença",
             "andamentos": [{"data": "2026-10-02", "texto": "O juiz proferiu sentença de procedência."},
                            {"data": "05/10/2026", "texto": "a ré interpôs recurso de apelação."}]},
            {"numeros": [PROCS[1]["numeros"][1]], "andamentos": [{"data": "01/10/2026", "texto": "foi julgado o agravo."}]},
            {"numero": num(2)},   # sem novidade: só renova o fecho
        ]

    def aplicar(self, atualizacoes=None, nome="saida", molde=None, **kw):
        destino = self.destino(nome)
        res = da.atualizar(molde or self.molde, destino, atualizacoes if atualizacoes is not None else self.atualizacoes(),
                           DATA_BASE_NOVA, **kw)
        return destino, res

    def codigos(self, res):
        return [a["codigo"] for a in res["avisos"]]


class TestLeitura(Base):
    def test_estrutura_do_fixture(self):
        est = da.ler_estrutura(self.molde)
        self.assertEqual(est["data_base"], gm.DATA_BASE)
        self.assertEqual(est["cliente"], gm.CLIENTE)
        self.assertEqual(len(est["resumo"]), 12)
        self.assertEqual(len(est["processos"]), 12)
        self.assertEqual(est["avisos"], [])
        tres = est["processos"][1]
        self.assertEqual(len(tres["numeros"]), 3)
        self.assertEqual(est["resumo"][1]["numeros"], tres["numeros"])
        self.assertEqual(est["processos"][3]["valor_causa"], "")          # sem valor da causa
        self.assertEqual(est["processos"][4]["data_citacao"], "-")
        self.assertEqual(est["processos"][0]["momento_atual"], "AGUARDANDO SENTENÇA")
        self.assertEqual(est["processos"][0]["fecho"]["data"], gm.DATA_BASE)
        self.assertIsNone(est["processos"][8]["fecho"])                    # sem fecho
        self.assertTrue(all(a["data_em_negrito"] for p in est["processos"] for a in p["andamentos"]))
        self.assertEqual(da.verificar_coerencia(self.molde), [])

    def test_vinculados_lidos_do_titulo(self):
        est = da.ler_estrutura(self.molde)
        self.assertEqual(est["processos"][1]["vinculados"],
                         [{"numero": PROCS[1]["numeros"][1], "tipo": "agravo"},
                          {"numero": PROCS[1]["numeros"][2], "tipo": "apenso"}])
        self.assertEqual(est["processos"][0]["vinculados"], [])

    def test_runs_fragmentados_no_fixture(self):
        """O fixture precisa ser mesmo 'sujo': muitos runs por parágrafo e datas quebradas."""
        _, runs = runs_do_par_andamentos(self.molde, 0)
        self.assertGreater(len(runs), 12)
        quebradas = [r for r in runs if re.fullmatch(r"[\d/]{1,9}", texto_el(r)) and len(texto_el(r)) < 10 and "/" in texto_el(r)
                     and len(texto_el(r)) != 10]
        self.assertTrue(quebradas, "esperava ao menos uma data partida em runs")

    def test_separar_momento(self):
        self.assertEqual(da.separar_momento("CUMPRIMENTO DE SENTENÇA (HONORÁRIOS SUSPENSOS)"),
                         ("CUMPRIMENTO DE SENTENÇA", "HONORÁRIOS SUSPENSOS"))
        self.assertEqual(da.separar_momento("PROCESSO ARQUIVADO (DECISÃO FAVORÁVEL)"),
                         ("PROCESSO ARQUIVADO", "DECISÃO FAVORÁVEL"))
        self.assertEqual(da.separar_momento("AGUARDANDO SENTENÇA"), ("AGUARDANDO SENTENÇA", None))
        self.assertEqual(da.separar_momento(None), ("", None))


class TestAtualizacao(Base):
    def test_resumo_titulo_data_base_andamentos_e_fecho(self):
        destino, res = self.aplicar()
        est = da.ler_estrutura(destino)
        self.assertEqual(est["data_base"], DATA_BASE_NOVA)
        p0 = est["processos"][0]
        self.assertEqual(p0["momento_atual"], "CUMPRIMENTO DE SENTENÇA")                    # título
        self.assertEqual(est["resumo"][0]["momento_atual"], "CUMPRIMENTO DE SENTENÇA")      # resumo (negrito, caixa alta)
        self.assertEqual(est["resumo"][0]["ultimo_andamento"], "05/10/2026")
        self.assertEqual([a["data"] for a in p0["andamentos"]][-2:], ["02/10/2026", "05/10/2026"])
        self.assertTrue(p0["andamentos"][-1]["texto"].startswith("a ré interpôs"))
        self.assertTrue(p0["andamentos"][-2]["texto"].startswith("o juiz proferiu"))        # caixa ajustada
        # padrão: processo COM novidade termina na novidade (sem fecho); o fecho antigo sai
        self.assertIsNone(p0["fecho"])
        self.assertEqual(p0["andamentos_texto"].count("sem atualizações"), 0)
        self.assertTrue(p0["andamentos_texto"].endswith("Em 05/10/2026, a ré interpôs recurso de apelação."))
        # agravo reconhecido por UM dos três números; resumo atualizado
        self.assertEqual(est["resumo"][1]["ultimo_andamento"], "01/10/2026")
        # sem novidade: texto igual, só o fecho (a data) muda
        antes = da.ler_estrutura(self.molde)["processos"][2]["andamentos_texto"]
        depois = est["processos"][2]["andamentos_texto"]
        self.assertEqual(depois, antes.replace(gm.DATA_BASE, DATA_BASE_NOVA))
        self.assertEqual(est["processos"][2]["fecho"]["data"], DATA_BASE_NOVA)
        self.assertEqual({m["campo"] for m in res["mudancas"] if m["numero"] == num(2)}, {"andamentos_fecho"})
        self.assertEqual(set(res["processos_atualizados"]), {num(0), PROCS[1]["numeros"][0], num(2)})

    def test_resultado_tem_o_formato_do_contrato(self):
        destino, res = self.aplicar()
        for chave in ("destino", "processos_atualizados", "processos_novos", "ignorados", "mudancas", "avisos",
                      "textos_gravados", "campos_gravados", "nao_encontrados", "gravado"):
            self.assertIn(chave, res)
        self.assertTrue(res["gravado"])
        self.assertEqual(Path(res["destino"]), destino)
        self.assertTrue(all(set(m) == {"numero", "campo", "antes", "depois"} for m in res["mudancas"]))
        self.assertTrue(all({"nivel", "codigo", "onde", "mensagem", "candidatos"} <= set(a) for a in res["avisos"]))
        self.assertTrue(all(a["nivel"] in ("info", "atencao", "erro") and a["codigo"] for a in res["avisos"]))
        est = da.ler_estrutura(destino)
        # textos_gravados: o texto de andamentos que ficou no arquivo, de todos os processos
        self.assertEqual(len(res["textos_gravados"]), 12)
        self.assertEqual(res["textos_gravados"][num(0)], est["processos"][0]["andamentos_texto"])
        self.assertEqual(res["campos_gravados"][num(0)], {"momento_atual": "CUMPRIMENTO DE SENTENÇA",
                                                          "ultimo_andamento": "05/10/2026"})

    def test_fecho_apos_novidade_verdadeiro(self):
        destino, _ = self.aplicar([{"numero": num(0), "andamentos": [{"data": "02/10/2026", "texto": "foi proferida sentença."}]}],
                                  fecho_apos_novidade=True)
        p = da.ler_estrutura(destino)["processos"][0]
        self.assertEqual(p["fecho"]["data"], DATA_BASE_NOVA)
        self.assertTrue(p["andamentos_texto"].endswith(f"Em 02/10/2026, foi proferida sentença. Em {DATA_BASE_NOVA}, sem atualizações."))
        self.assertEqual(da.verificar_coerencia(destino, fechos=False), [])

    def test_datas_novas_em_negrito_no_xml(self):
        destino, _ = self.aplicar()
        par, runs = runs_do_par_andamentos(destino, 0)
        for data in ("02/10/2026", "05/10/2026"):
            ach = [r for r in runs if texto_el(r) == data]
            self.assertEqual(len(ach), 1, data)
            self.assertTrue(da._negrito(ach[0].find(w("rPr"))), data)
        # e o texto novo NÃO está em negrito
        for r in runs:
            if texto_el(r).startswith(", o juiz proferiu"):
                self.assertFalse(da._negrito(r.find(w("rPr"))))
        # processo sem novidade: a data do fecho está em negrito
        _, runs2 = runs_do_par_andamentos(destino, 2)
        ach = [r for r in runs2 if texto_el(r) == DATA_BASE_NOVA]
        self.assertEqual(len(ach), 1)
        self.assertTrue(da._negrito(ach[0].find(w("rPr"))))

    def test_runs_novos_copiam_formatacao_do_vizinho(self):
        destino, _ = self.aplicar()
        _, antes = runs_do_par_andamentos(self.molde, 0)
        _, depois = runs_do_par_andamentos(destino, 0)
        modelo = next(r for r in reversed(antes) if texto_el(r).strip() and not da._negrito(r.find(w("rPr"))))
        novo = next(r for r in depois if texto_el(r).startswith(", o juiz proferiu"))

        def sem_b(rpr):
            rpr = copy.deepcopy(rpr)
            for b in rpr.findall(w("b")) + rpr.findall(w("bCs")):
                rpr.remove(b)
            return c14n(rpr)
        self.assertEqual(sem_b(novo.find(w("rPr"))), sem_b(modelo.find(w("rPr"))))
        # fonte/tamanho do negrito também são os do texto (Arial 11)
        data = next(r for r in depois if texto_el(r) == "02/10/2026")
        rpr = data.find(w("rPr"))
        self.assertEqual(rpr.find(w("sz")).get(w("val")), "22")
        self.assertEqual(rpr.find(w("rFonts")).get(w("ascii")), "Arial")

    def test_preservacao_do_que_nao_deveria_mudar(self):
        destino, _ = self.aplicar()
        t0, t1 = tabelas(self.molde), tabelas(destino)
        self.assertEqual(len(t0), len(t1))
        atualizados = {0, 1, 2}               # índices de processos tocados (tabela i+1)
        for i in range(1, len(t0)):
            if i - 1 not in atualizados:
                self.assertEqual(c14n(t0[i]), c14n(t1[i]), f"bloco {i - 1} mudou sem ter sido atualizado")
        # quadro-resumo: linhas dos processos não tocados idênticas; células intocadas das tocadas também
        r0, r1 = t0[0].findall(w("tr")), t1[0].findall(w("tr"))
        self.assertEqual(c14n(r0[0]), c14n(r1[0]))
        for k in range(1, len(r0)):
            if k - 1 not in atualizados:
                self.assertEqual(c14n(r0[k]), c14n(r1[k]))
        for k in atualizados:
            c0, c1 = r0[k + 1].findall(w("tc")), r1[k + 1].findall(w("tc"))
            for j in (0, 1):     # nº e assunto
                self.assertEqual(c14n(c0[j]), c14n(c1[j]))
            for j in range(4):   # larguras e bordas (tcPr) nunca mudam
                self.assertEqual(c14n(c0[j].find(w("tcPr"))), c14n(c1[j].find(w("tcPr"))))
        # blocos tocados: tudo igual, menos título (momento) e célula de andamentos
        for i in atualizados:
            a, b = t0[i + 1], t1[i + 1]
            self.assertEqual(c14n(a.find(w("tblPr"))), c14n(b.find(w("tblPr"))))
            self.assertEqual(c14n(a.find(w("tblGrid"))), c14n(b.find(w("tblGrid"))))
            la, lb = a.findall(w("tr")), b.findall(w("tr"))
            for k in range(1, len(la) - 1):
                self.assertEqual(c14n(la[k]), c14n(lb[k]))
            self.assertEqual(c14n(la[-1].findall(w("tc"))[0]), c14n(lb[-1].findall(w("tc"))[0]))     # rótulo
            self.assertEqual(c14n(la[-1].findall(w("tc"))[-1].find(w("tcPr"))), c14n(lb[-1].findall(w("tc"))[-1].find(w("tcPr"))))
            self.assertEqual(c14n(la[0].find(w("trPr"))), c14n(lb[0].find(w("trPr"))))
        # parágrafos fora das tabelas: só o da Data-Base muda
        p0 = [c14n(p) for p in corpo(self.molde).findall(w("p"))]
        p1 = [c14n(p) for p in corpo(destino).findall(w("p"))]
        self.assertEqual(len(p0), len(p1))
        diferentes = [k for k in range(len(p0)) if p0[k] != p1[k]]
        self.assertEqual(len(diferentes), 1)
        self.assertIn("Data-Base", texto_el(corpo(destino).findall(w("p"))[diferentes[0]]))
        # outras partes do pacote: byte a byte
        with zipfile.ZipFile(self.molde) as z0, zipfile.ZipFile(destino) as z1:
            self.assertEqual(z0.namelist(), z1.namelist())
            for n in z0.namelist():
                if n != "word/document.xml":
                    self.assertEqual(z0.read(n), z1.read(n), n)

    def test_so_acrescenta_texto_do_advogado_e_runs_antigos_intactos(self):
        destino, _ = self.aplicar()
        # processo 1 tem anotação do advogado em destaque (itálico + marca-texto) no fim do texto
        antes = da.ler_estrutura(self.molde)["processos"][1]["andamentos_texto"]
        depois = da.ler_estrutura(destino)["processos"][1]["andamentos_texto"]
        corpo_antigo = antes[:antes.rindex(" Em 18/09/2026")]
        self.assertTrue(depois.startswith(corpo_antigo), "texto antigo foi reescrito")
        self.assertIn("(Anotação do advogado: conferir o valor com o cliente.)", depois)
        par, runs = runs_do_par_andamentos(destino, 1)
        nota = [r for r in runs if "Anotação do advogado" in texto_el(r)]
        self.assertEqual(len(nota), 1)
        self.assertIsNotNone(nota[0].find(w("rPr")).find(w("highlight")))                      # destaque do advogado preservado
        for r in runs:                                                                          # ...e NÃO herdado pelo texto novo
            if "foi julgado o agravo" in texto_el(r):
                self.assertIsNone(r.find(w("rPr")).find(w("highlight")))
                self.assertIsNone(r.find(w("rPr")).find(w("i")))
        # todo run antigo que termina antes do fecho (menos o espaço separador) continua idêntico, byte a byte
        _, runs0 = runs_do_par_andamentos(self.molde, 1)
        texto0 = "".join(texto_el(r) for r in runs0)
        limite = texto0.rindex("Em 18/09/2026") - 1
        pos, verificados = 0, 0
        for k, r in enumerate(runs0):
            fim = pos + len(texto_el(r))
            if fim <= limite:
                self.assertEqual(c14n(r), c14n(runs[k]), f"run {k} mudou")
                verificados += 1
            pos = fim
        self.assertGreater(verificados, 10)

    def test_recarrega_com_python_docx(self):
        import docx
        destino, _ = self.aplicar(self.atualizacoes() + [{"novo": proc_novo()}], nome="pydocx")
        d = docx.Document(str(destino))
        self.assertEqual(len(d.tables), 14)                      # resumo + 13 blocos
        self.assertEqual(len(d.tables[0].rows), 14)              # cabeçalho + 13 processos
        self.assertIn("02/10/2026", d.tables[1].rows[-1].cells[-1].text)
        self.assertEqual(d.tables[0].rows[1].cells[2].text, "CUMPRIMENTO DE SENTENÇA")
        runs = [r for r in d.tables[1].rows[-1].cells[-1].paragraphs[0].runs if r.text == "02/10/2026"]
        self.assertTrue(runs and runs[0].bold)
        self.assertEqual(runs[0].font.name, "Arial")
        self.assertEqual(runs[0].font.size.pt, 11)
        d.save(str(TMP / "reabre.docx"))                          # e o python-docx consegue regravar

    def test_idempotencia(self):
        upds = self.atualizacoes() + [{"novo": proc_novo()}]
        d1, r1 = self.aplicar(upds, nome="um")
        self.assertTrue(r1["mudancas"])
        d2 = self.destino("dois")
        r2 = da.atualizar(d1, d2, upds, DATA_BASE_NOVA)
        self.assertEqual(r2["mudancas"], [])
        self.assertEqual(r2["processos_atualizados"], [])
        self.assertEqual(r2["processos_novos"], [])
        with zipfile.ZipFile(d1) as z1, zipfile.ZipFile(d2) as z2:
            self.assertEqual(z1.read("word/document.xml"), z2.read("word/document.xml"))
        est = da.ler_estrutura(d2)
        t = est["processos"][0]["andamentos_texto"]
        self.assertEqual(t.count("02/10/2026"), 1)
        self.assertEqual(t.count("sem atualizações"), 0)          # novidade já gravada: a reaplicação não põe fecho
        self.assertEqual(len(est["processos"]), 13)
        self.assertTrue(any(i["motivo"] == "andamento já presente" for i in r2["ignorados"]))
        self.assertTrue(any(i["motivo"] == "processo já existe" for i in r2["ignorados"]))
        # reaplicação com 'texto_conhecido' antigo não acusa edição manual
        upds_c = copy.deepcopy(upds)
        upds_c[0]["texto_conhecido"] = da.ler_estrutura(self.molde)["processos"][0]["andamentos_texto"]
        r3 = da.atualizar(d1, self.destino("tres"), upds_c, DATA_BASE_NOVA)
        self.assertFalse([a for a in r3["avisos"] if a["codigo"] == "edicao_manual"])

    def test_nao_duplica_andamento_ja_presente_mesmo_com_variacao(self):
        existente = "a ré juntou documentos e rol de testemunhas."
        upds = [{"numero": num(2), "andamentos": [
            {"data": "01/09/2026", "texto": "A ré juntou documentos e rol de testemunhas"},     # caixa e ponto diferentes
            {"data": "01/09/2026", "texto": "foi juntada nova procuração aos autos."},          # outra coisa, mesma data
        ]}]
        destino, res = self.aplicar(upds)
        t = da.ler_estrutura(destino)["processos"][2]["andamentos_texto"]
        self.assertEqual(t.count(existente), 1)
        self.assertIn("nova procuração", t)
        self.assertEqual(sum(1 for i in res["ignorados"] if i["motivo"] == "andamento já presente"), 1)

    def test_edicao_manual_recente_so_acrescenta_e_avisa(self):
        conhecido = da.ler_estrutura(self.molde)["processos"][0]["andamentos_texto"]

        def mexer(raiz):    # o advogado acrescentou uma frase à mão depois do último ciclo
            tbl = raiz.find(w("body")).findall(w("tbl"))[1]
            par = tbl.findall(w("tr"))[-1].findall(w("tc"))[-1].findall(w("p"))[-1]
            fecho = [r for r in par.findall(w("r"))]
            novo = da._novo_run(" Em ", da._rpr_modelo(fecho[0].find(w("rPr")), False))
            par.insert(len(par) - 4, novo)       # antes dos runs do fecho: edição no meio do final
            par.insert(len(par) - 4, da._novo_run("(nota do advogado: cliente ligou)", da._rpr_modelo(fecho[0].find(w("rPr")), False)))
        editado = TMP / "editado.docx"
        escrever_doc(self.molde, editado, mexer)
        upds = [{"numero": num(0), "texto_conhecido": conhecido,
                 "andamentos": [{"data": "02/10/2026", "texto": "o juiz proferiu sentença."}]}]
        destino, res = self.aplicar(upds, molde=editado)
        avisos = [a for a in res["avisos"] if a["codigo"] == "edicao_manual"]
        self.assertEqual(len(avisos), 1)
        t = da.ler_estrutura(destino)["processos"][0]["andamentos_texto"]
        self.assertIn("(nota do advogado: cliente ligou)", t)       # nada apagado
        self.assertIn("o juiz proferiu sentença.", t)
        # sem divergência, sem aviso
        destino2, res2 = self.aplicar([{"numero": num(0), "texto_conhecido": conhecido,
                                        "andamentos": [{"data": "02/10/2026", "texto": "o juiz proferiu sentença."}]}], nome="limpo")
        self.assertFalse([a for a in res2["avisos"] if a["codigo"] == "edicao_manual"])

    def test_edicao_manual_e_detectada_mesmo_com_andamentos_antigos_na_lista(self):
        """Na vida real a lista traz os eventos antigos (já presentes) junto com o novo: a checagem não pode ficar cega."""
        conhecido = da.ler_estrutura(self.molde)["processos"][0]["andamentos_texto"]

        def mexer(raiz):
            tbl = raiz.find(w("body")).findall(w("tbl"))[1]
            par = tbl.findall(w("tr"))[-1].findall(w("tc"))[-1].findall(w("p"))[-1]
            par.insert(len(par) - 4, da._novo_run("(nota solta do advogado)", da._rpr_modelo(par.findall(w("r"))[0].find(w("rPr")), False)))
        editado = TMP / "editado2.docx"
        escrever_doc(self.molde, editado, mexer)
        antigos = [{"data": d, "texto": t} for d, t in PROCS[0]["andamentos"]]
        destino, res = self.aplicar([{"numero": num(0), "texto_conhecido": conhecido,
                                      "andamentos": antigos + [{"data": "02/10/2026", "texto": "foi proferida sentença."}]}],
                                    molde=editado)
        self.assertEqual(len([a for a in res["avisos"] if a["codigo"] == "edicao_manual"]), 1)
        self.assertEqual(len([i for i in res["ignorados"] if i["motivo"] == "andamento já presente"]), 4)

    def test_duplicata_provavelmente_digitada_a_mao_nao_e_repetida(self):
        # o andamento de 30/07 já está no texto (digitado pelo sistema ou à mão), com palavras um pouco diferentes
        upds = [{"numero": num(0), "andamentos": [{"data": "30/07/2026", "texto": "os autos foram conclusos ao juiz para sentença."}]}]
        destino, res = self.aplicar(upds)
        t = da.ler_estrutura(destino)["processos"][0]["andamentos_texto"]
        self.assertEqual(t.count("30/07/2026"), 1)
        self.assertEqual(len(res["ignorados"]), 1)

    def test_fecho_variantes(self):
        sem_novidade = [{"numero": num(i)} for i in (7, 9, 8)]
        destino, res = self.aplicar(sem_novidade)
        est = da.ler_estrutura(destino)
        # 7: fecho todo em negrito -> continua em negrito; só a data muda
        t7 = est["processos"][7]["andamentos_texto"]
        self.assertTrue(t7.endswith(f"Em {DATA_BASE_NOVA}, sem atualizações."))
        _, runs = runs_do_par_andamentos(destino, 7)
        self.assertTrue(all(da._negrito(r.find(w("rPr"))) for r in runs if texto_el(r) in (DATA_BASE_NOVA,)))
        # 9: 'sem atualização' (singular) mantém a redação original
        self.assertTrue(est["processos"][9]["andamentos_texto"].endswith(f"Em {DATA_BASE_NOVA}, sem atualização."))
        # 8: não tinha fecho -> acrescentado com aviso informativo
        self.assertTrue(est["processos"][8]["andamentos_texto"].endswith(f"Em {DATA_BASE_NOVA}, sem atualizações."))
        self.assertTrue(est["processos"][8]["fecho"]["data_em_negrito"])
        self.assertTrue([a for a in res["avisos"] if a["codigo"] == "fecho_acrescentado"])

    def test_renovar_fecho_dos_demais(self):
        destino, res = self.aplicar([{"numero": num(0)}], renovar_fecho_dos_demais=True)
        est = da.ler_estrutura(destino)
        for p in est["processos"]:
            self.assertEqual((p["fecho"] or {}).get("data"), DATA_BASE_NOVA, p["numeros"][0])
        self.assertEqual(da.verificar_coerencia(destino), [])

    def test_processos_fora_da_lista_ficam_intactos_com_aviso(self):
        destino, res = self.aplicar([{"numero": num(0)}])
        self.assertIn("processos_sem_atualizacao", self.codigos(res))
        t0, t1 = tabelas(self.molde), tabelas(destino)
        for i in range(2, len(t0)):
            self.assertEqual(c14n(t0[i]), c14n(t1[i]))

    def test_numero_nao_encontrado(self):
        destino, res = self.aplicar([{"numero": NUM_AUSENTE, "andamentos": [{"data": "01/10/2026", "texto": "x."}]}])
        self.assertEqual(res["nao_encontrados"], [NUM_AUSENTE])
        self.assertTrue([a for a in res["avisos"] if a["codigo"] == "processo_nao_encontrado"])

    def test_original_nunca_e_sobrescrito(self):
        antes = self.molde.read_bytes()
        with self.assertRaises(ValueError):
            da.atualizar(self.molde, self.molde, [], DATA_BASE_NOVA)
        link = TMP / "link.docx"
        if link.exists() or link.is_symlink():
            link.unlink()
        link.symlink_to(self.molde)
        with self.assertRaises(ValueError):
            da.atualizar(self.molde, link, [], DATA_BASE_NOVA)
        existente = self.destino("ja_existe")
        existente.write_bytes(b"x")
        with self.assertRaises(FileExistsError):
            da.atualizar(self.molde, existente, [], DATA_BASE_NOVA)
        self.aplicar()
        self.assertEqual(self.molde.read_bytes(), antes)

    def test_mesclagem_vertical_preservada(self):
        destino, _ = self.aplicar([{"numero": num(6), "momento_atual": "Trânsito em julgado",
                                    "andamentos": [{"data": "01/10/2026", "texto": "a apelação foi julgada."}]}])
        a, b = tabelas(self.molde)[7], tabelas(destino)[7]
        self.assertEqual(len(list(a.iter(w("vMerge")))), 2)
        self.assertEqual(len(list(b.iter(w("vMerge")))), 2)
        ra, rb = a.findall(w("tr")), b.findall(w("tr"))
        for k in range(2, len(ra) - 1):
            self.assertEqual(c14n(ra[k]), c14n(rb[k]))

    def test_sem_mudancas_o_document_xml_fica_identico(self):
        d2 = self.destino("mesma")
        res2 = da.atualizar(self.molde, d2, [], gm.DATA_BASE)
        self.assertEqual(res2["mudancas"], [])
        with zipfile.ZipFile(self.molde) as z0, zipfile.ZipFile(d2) as z1:
            self.assertEqual(z0.read("word/document.xml"), z1.read("word/document.xml"))

    def test_documento_sem_quadro_resumo_ainda_atualiza_o_titulo(self):
        molde = TMP / "sem_resumo.docx"

        def mexer(raiz):
            corpo_ = raiz.find(w("body"))
            corpo_.remove(corpo_.findall(w("tbl"))[0])
        escrever_doc(self.molde, molde, mexer)
        destino, res = self.aplicar([{"numero": num(0), "momento_atual": "Cumprimento de sentença",
                                      "andamentos": [{"data": "02/10/2026", "texto": "foi proferida sentença."}]}], molde=molde)
        self.assertIn("sem_resumo", self.codigos(res))
        p = da.ler_estrutura(destino)["processos"][0]
        self.assertEqual(p["momento_atual"], "CUMPRIMENTO DE SENTENÇA")
        self.assertIn("Em 02/10/2026, foi proferida sentença.", p["andamentos_texto"])

    def test_opcao_desconhecida_e_erro_de_programacao(self):
        with self.assertRaises(TypeError):
            da.atualizar(self.molde, self.destino("x"), [], DATA_BASE_NOVA, fecho_apos_novidades=True)


class TestProcessoNovo(Base):
    def proc(self, **kw):
        return proc_novo(**kw)

    def test_clona_bloco_e_linha_do_resumo(self):
        destino, res = self.aplicar([{"novo": self.proc()}])
        self.assertEqual(res["processos_novos"], [self.proc()["numeros"][0]])
        est = da.ler_estrutura(destino)
        self.assertEqual(len(est["processos"]), 13)
        self.assertEqual(len(est["resumo"]), 13)
        novo = est["processos"][-1]
        self.assertEqual(novo["numeros"], self.proc()["numeros"])
        self.assertEqual(novo["momento_atual"], "AGUARDANDO CITAÇÃO")
        self.assertEqual(novo["valor_causa"], "R$ 12.345,60")
        self.assertEqual(novo["data_citacao"], "-")
        self.assertEqual(novo["ajuizamento"], "30/09/2026")
        self.assertEqual([a["data"] for a in novo["andamentos"]], ["01/10/2026", "03/10/2026"])
        self.assertTrue(all(a["data_em_negrito"] for a in novo["andamentos"]))
        self.assertIsNone(novo["fecho"])                            # tem novidade: termina nela
        self.assertEqual(est["resumo"][-1], {"numeros": novo["numeros"], "assunto": "Ação de teste nova",
                                             "momento_atual": "AGUARDANDO CITAÇÃO", "ultimo_andamento": "03/10/2026"})
        self.assertEqual(da.verificar_coerencia(destino, fechos=False), [])
        # nada de texto do bloco-modelo sobrou no clone
        modelo_texto = da.ler_estrutura(self.molde)["processos"][-1]
        for campo in ("assunto", "autores", "juizo", "materia"):
            self.assertNotEqual(novo[campo], modelo_texto[campo])
        self.assertNotIn(modelo_texto["andamentos"][0]["texto"], novo["andamentos_texto"])
        # estrutura (larguras, mesclagens, bordas) igual à do modelo clonado
        tb = tabelas(destino)
        self.assertEqual(c14n(tb[-1].find(w("tblGrid"))), c14n(tb[-2].find(w("tblGrid"))))
        self.assertEqual(c14n(tb[-1].find(w("tblPr"))), c14n(tb[-2].find(w("tblPr"))))
        spans = lambda t: [[tc.find(w("tcPr")).find(w("gridSpan")) is not None for tc in tr.findall(w("tc"))] for tr in t.findall(w("tr"))]
        self.assertEqual(spans(tb[-1]), spans(tb[-2]))
        # parágrafo vazio separa as tabelas (senão o Word as funde)
        corpo_ = corpo(destino)
        filhos = list(corpo_)
        i = filhos.index(corpo_.findall(w("tbl"))[-1])
        self.assertEqual(filhos[i - 1].tag, w("p"))
        self.assertFalse(texto_el(filhos[i - 1]).strip())

    def test_processo_novo_sem_andamento_tem_fecho(self):
        destino, _ = self.aplicar([{"novo": self.proc(andamentos=[], ultimo_andamento="2026-09-30")}])
        novo = da.ler_estrutura(destino)["processos"][-1]
        self.assertEqual(novo["andamentos_texto"], f"Em {DATA_BASE_NOVA}, sem atualizações.")
        self.assertEqual(novo["fecho"]["data"], DATA_BASE_NOVA)
        self.assertTrue(novo["fecho"]["data_em_negrito"])
        self.assertEqual(da.ler_estrutura(destino)["resumo"][-1]["ultimo_andamento"], "30/09/2026")

    def test_processo_novo_com_tres_numeros(self):
        p = self.proc(numeros=[NUM_NOVO, NUM_NOVO_2, NUM_NOVO_3], tipos=["agravo", "apenso"])
        destino, _ = self.aplicar([{"novo": p}])
        est = da.ler_estrutura(destino)
        self.assertEqual(est["processos"][-1]["numeros"], p["numeros"])
        self.assertEqual(est["resumo"][-1]["numeros"], p["numeros"])
        self.assertIn("AGRAVO Nº", est["processos"][-1]["titulo"])
        self.assertEqual([v["tipo"] for v in est["processos"][-1]["vinculados"]], ["agravo", "apenso"])
        self.assertEqual(da.verificar_coerencia(destino, fechos=False), [])

    def test_novo_nao_duplica_se_numero_ja_existe(self):
        p = self.proc(numeros=[num(0)])
        destino, res = self.aplicar([{"novo": p}])
        self.assertEqual(res["processos_novos"], [])
        self.assertEqual(len(da.ler_estrutura(destino)["processos"]), 12)
        self.assertIn("ja_existe", self.codigos(res))

    def test_modelo_ruim_nao_e_clonado(self):
        """Com os 7 primeiros processos, o último bloco tem mesclagem vertical: o clone deve vir de um bloco limpo."""
        molde = TMP / "molde7.docx"
        gm.gerar(molde, n=7)
        destino, res = self.aplicar([{"novo": self.proc()}], molde=molde, nome="m7")
        self.assertEqual(len(list(tabelas(destino)[-1].iter(w("vMerge")))), 0)
        self.assertFalse([a for a in res["avisos"] if a["codigo"] == "modelo_imperfeito"])
        self.assertEqual(da.ler_estrutura(destino)["processos"][-1]["reus"], gm.CLIENTE)

    def test_bloco_com_imagem_nao_e_clonado(self):
        def mexer(raiz):    # um desenho dentro do último bloco (logotipo, por exemplo)
            tbl = raiz.find(w("body")).findall(w("tbl"))[-1]
            par = tbl.findall(w("tr"))[1].findall(w("tc"))[1].findall(w("p"))[0]
            run = etree.SubElement(par, w("r"))
            etree.SubElement(run, w("drawing"))
        molde = TMP / "com_imagem.docx"
        escrever_doc(self.molde, molde, mexer)
        destino, res = self.aplicar([{"novo": self.proc()}], molde=molde, nome="img")
        self.assertEqual(len(list(tabelas(destino)[-1].iter(w("drawing")))), 0)
        self.assertEqual(len(list(tabelas(destino)[12].iter(w("drawing")))), 1)       # o original continua com o desenho

    def test_so_ha_blocos_com_imagem_clona_e_avisa(self):
        molde1 = TMP / "um_bloco.docx"
        gm.gerar(molde1, n=1)

        def mexer(raiz):
            tbl = raiz.find(w("body")).findall(w("tbl"))[-1]
            par = tbl.findall(w("tr"))[1].findall(w("tc"))[1].findall(w("p"))[0]
            etree.SubElement(etree.SubElement(par, w("r")), w("drawing"))
        molde = TMP / "um_bloco_img.docx"
        escrever_doc(molde1, molde, mexer)
        destino, res = self.aplicar([{"novo": self.proc()}], molde=molde, nome="img2")
        self.assertIn("modelo_imperfeito", self.codigos(res))

    def test_celula_solta_do_modelo_nao_vaza_para_o_clone(self):
        """Se o único bloco tem texto fora do padrão, o clone esvazia a célula e avisa."""
        molde = TMP / "molde1.docx"
        gm.gerar(molde, n=1)

        def mexer(raiz):    # texto solto numa linha extra do bloco
            tbl = raiz.find(w("body")).findall(w("tbl"))[1]
            extra = copy.deepcopy(tbl.findall(w("tr"))[-1])
            for t in extra.iter(w("t")):
                t.text = "Observação manual do advogado sobre este processo"
            tbl.append(extra)
        sujo = TMP / "sujo.docx"
        escrever_doc(molde, sujo, mexer)
        destino, res = self.aplicar([{"novo": self.proc()}], molde=sujo, nome="sujo")
        self.assertNotIn("Observação manual", texto_el(tabelas(destino)[-1]))
        self.assertTrue([a for a in res["avisos"] if a["codigo"] in ("celula_limpa", "modelo_imperfeito")])
        self.assertIn("Observação manual", texto_el(tabelas(destino)[1]))        # o original do advogado ficou

    def test_historico_migrado_entra_no_processo_novo_com_datas_em_negrito(self):
        historico = ("Em 10/03/2026 foi distribuída a ação. Em 20/04/2026 foi apresentada contestação. "
                     "Em 18/09/2026, sem atualizações.")                 # o fecho do relatório antigo não vai junto
        p = self.proc(historico=historico, andamentos=[
            {"data": "20/04/2026", "texto": "Foi apresentada contestação."},      # já consta no histórico: não repete
            {"data": "03/10/2026", "texto": "foi determinada a citação da parte ré."}])
        destino, res = self.aplicar([{"novo": p}])
        novo = da.ler_estrutura(destino)["processos"][-1]
        self.assertEqual([a["data"] for a in novo["andamentos"]], ["10/03/2026", "20/04/2026", "03/10/2026"])
        self.assertTrue(all(a["data_em_negrito"] for a in novo["andamentos"]))
        self.assertEqual(novo["andamentos_texto"].count("contestação"), 1)
        self.assertNotIn("18/09/2026", novo["andamentos_texto"])
        self.assertIsNone(novo["fecho"])
        self.assertEqual(da.ler_estrutura(destino)["resumo"][-1]["ultimo_andamento"], "03/10/2026")
        self.assertTrue(any(i["motivo"] == "andamento já presente" for i in res["ignorados"]))


class TestEscala(Base):
    def test_200_processos(self):
        procs = []
        for i in range(200):
            p = PROCS[i % 12]
            procs.append({"numeros": [gm.cnj(i, 2020 + i % 6, 8, 6, 100 + i)], "assunto": p["assunto"], "autores": p["autores"],
                          "reus": p["reus"], "ajuizamento": p["ajuizamento"], "valor_causa": p["valor_causa"] or None,
                          "data_citacao": p["citacao"], "juizo": p["juizo"], "area": p["area"], "materia": p["materia"],
                          "momento_atual": p["momento"], "andamentos": [{"data": d, "texto": t} for d, t in p["andamentos"]]})
        base = self.destino("g200")
        da.gerar(base, {"cliente": gm.CLIENTE, "data_base": gm.DATA_BASE, "processos": procs})
        upds = [{"numero": p["numeros"][0], "momento_atual": "Cumprimento de sentença",
                 "andamentos": [{"data": "02/10/2026", "texto": "foi proferida sentença."}]} for p in procs]
        t0 = time.time()
        destino, res = self.aplicar(upds, molde=base, nome="a200")
        self.assertLess(time.time() - t0, 30)
        self.assertEqual(len(res["processos_atualizados"]), 200)
        self.assertEqual(da.verificar_coerencia(destino), [])
        self.assertEqual(len(da.ler_estrutura(destino)["processos"]), 200)


class TestVariantes(Base):
    def test_texto_de_andamentos_na_linha_de_baixo(self):
        molde = TMP / "abaixo.docx"
        gm.gerar(molde, variante="andamentos_abaixo")
        est = da.ler_estrutura(molde)
        self.assertEqual(len(est["processos"][0]["andamentos"]), 4)
        destino, _ = self.aplicar([{"numero": num(0), "andamentos": [{"data": "02/10/2026", "texto": "foi proferida sentença."}]},
                                   {"numero": num(2)}], molde=molde, nome="abaixo")
        est = da.ler_estrutura(destino)
        self.assertEqual(est["processos"][0]["andamentos"][-1]["data"], "02/10/2026")
        self.assertEqual(est["processos"][2]["fecho"]["data"], DATA_BASE_NOVA)

    def test_varias_fragmentacoes(self):
        """O resultado não pode depender de onde o Google quebrou os runs."""
        esperado = None
        for semente in range(1, 13):
            molde = TMP / f"frag{semente}.docx"
            gm.gerar(molde, semente=semente)
            destino, res = self.aplicar(self.atualizacoes(), molde=molde, nome=f"frag{semente}")
            est = da.ler_estrutura(destino)
            textos = [p["andamentos_texto"] for p in est["processos"]]
            if esperado is None:
                esperado = textos
            self.assertEqual(textos, esperado, f"semente {semente}")
            self.assertEqual(da.verificar_coerencia(destino, fechos=False), [], f"semente {semente}")
            self.assertTrue(all(a["data_em_negrito"] for p in est["processos"] for a in p["andamentos"]))

    def test_espaco_nao_separavel_no_fecho(self):
        def mexer(raiz):
            tbl = raiz.find(w("body")).findall(w("tbl"))[1]
            par = tbl.findall(w("tr"))[-1].findall(w("tc"))[-1].findall(w("p"))[-1]
            for t in par.iter(w("t")):
                t.text = (t.text or "").replace(", sem atualizações", ", sem atualizações")
        nbsp = TMP / "nbsp.docx"
        escrever_doc(self.molde, nbsp, mexer)
        self.assertEqual(da.ler_estrutura(nbsp)["processos"][0]["fecho"]["data"], gm.DATA_BASE)
        destino, res = self.aplicar([{"numero": num(0)}], molde=nbsp, nome="nbsp")
        self.assertEqual(da.ler_estrutura(destino)["processos"][0]["fecho"]["data"], DATA_BASE_NOVA)
        self.assertEqual(da.ler_estrutura(destino)["processos"][0]["andamentos_texto"].count("sem"), 1)

    def test_advogado_escreveu_outro_paragrafo_depois_do_fecho(self):
        def mexer(raiz):
            tbl = raiz.find(w("body")).findall(w("tbl"))[1]
            tc = tbl.findall(w("tr"))[-1].findall(w("tc"))[-1]
            par = tc.findall(w("p"))[-1]
            novo = copy.deepcopy(par)
            for r in novo.findall(w("r"))[1:]:
                novo.remove(r)
            primeiro = novo.findall(w("r"))[0]
            for t in primeiro.iter(w("t")):
                t.text = "Observação do advogado: aguardar retorno do cliente."
            par.addnext(novo)
        molde = TMP / "dois_par.docx"
        escrever_doc(self.molde, molde, mexer)
        # sem novidade: só a data do fecho muda, ali mesmo
        d1, r1 = self.aplicar([{"numero": num(0)}], molde=molde, nome="sem_novidade")
        pars = [texto_el(p) for p in tabelas(d1)[1].findall(w("tr"))[-1].findall(w("tc"))[-1].findall(w("p"))]
        self.assertTrue(pars[0].endswith(f"Em {DATA_BASE_NOVA}, sem atualizações."))
        self.assertEqual(pars[1], "Observação do advogado: aguardar retorno do cliente.")
        self.assertTrue([a for a in r1["avisos"] if a["codigo"] == "fecho_fora_do_fim"])
        # com novidade: texto novo vai no último parágrafo; o do advogado e o fecho antigo continuam
        d2, _ = self.aplicar([{"numero": num(0), "andamentos": [{"data": "02/10/2026", "texto": "foi proferida sentença."}]}],
                             molde=molde, nome="com_novidade")
        pars = [texto_el(p) for p in tabelas(d2)[1].findall(w("tr"))[-1].findall(w("tc"))[-1].findall(w("p"))]
        self.assertTrue(pars[0].endswith(f"Em {gm.DATA_BASE}, sem atualizações."))
        self.assertTrue(pars[1].startswith("Observação do advogado") and "Em 02/10/2026, foi proferida sentença." in pars[1])

    def test_controle_de_alteracoes_e_comentario_no_trecho_final(self):
        """Limite documentado: com revisão/comentário no parágrafo, nada é removido; só acrescenta e avisa."""
        def mexer(raiz):
            tbl = raiz.find(w("body")).findall(w("tbl"))[1]
            par = tbl.findall(w("tr"))[-1].findall(w("tc"))[-1].findall(w("p"))[-1]
            runs = par.findall(w("r"))
            ins = etree.Element(w("ins"))
            ins.set(w("id"), "900"), ins.set(w("author"), "Advogada Exemplo"), ins.set(w("date"), "2026-09-20T10:00:00Z")
            runs[-1].addprevious(ins)
            for r in runs[-3:]:
                ins.append(r)
            ini = etree.Element(w("commentRangeStart"))
            ini.set(w("id"), "5")
            par.insert(1, ini)
        sujo = TMP / "revisao.docx"
        escrever_doc(self.molde, sujo, mexer)
        destino, res = self.aplicar([{"numero": num(0), "andamentos": [{"data": "02/10/2026", "texto": "foi proferida sentença."}]}],
                                    molde=sujo, nome="revisao")
        self.assertTrue([a for a in res["avisos"] if a["codigo"] == "revisao_no_trecho"])
        r = doc_xml(destino)
        self.assertEqual(len(list(r.iter(w("ins")))), 1)
        self.assertEqual(len(list(r.iter(w("commentRangeStart")))), 1)
        t = da.ler_estrutura(destino)["processos"][0]["andamentos_texto"]
        self.assertIn(f"Em {gm.DATA_BASE}, sem atualizações.", t)       # fecho antigo preservado dentro da revisão
        self.assertTrue(t.endswith("Em 02/10/2026, foi proferida sentença."))

    def test_controle_de_alteracoes_ligado_nas_configuracoes_avisa(self):
        def mexer(raiz):
            etree.SubElement(raiz, w("trackRevisions"))
        molde = TMP / "track.docx"
        escrever_doc(self.molde, molde, mexer, parte="word/settings.xml")
        destino, res = self.aplicar([{"numero": num(2)}], molde=molde, nome="track")
        self.assertIn("controle_alteracoes_ligado", self.codigos(res))

    def test_tem_comentarios_so_informa(self):
        molde = TMP / "comentarios.docx"
        with zipfile.ZipFile(self.molde) as zin, zipfile.ZipFile(molde, "w") as zout:
            for info in zin.infolist():
                zout.writestr(info, zin.read(info.filename))
            zout.writestr("word/comments.xml", f'<?xml version="1.0" encoding="UTF-8"?><w:comments xmlns:w="{W}"/>')
        destino, res = self.aplicar([{"numero": num(2)}], molde=molde, nome="coment")
        self.assertIn("tem_comentarios", self.codigos(res))
        with zipfile.ZipFile(destino) as z:
            self.assertIn("word/comments.xml", z.namelist())


class TestDeteccaoDeDuplicata(Base):
    """Marcadores de data além de 'Em DD/MM/AAAA': 'No dia', 'Em DD/MM', 'Em DD/MM/AA', data por extenso."""
    HISTORICO = ("No dia 18/06/2026, foi proferida sentença de procedência. Em 20/06 o réu foi intimado da sentença. "
                 "Em 22 de junho de 2026, foi interposto recurso de apelação. Em 25/06/26 foram apresentadas as razões. "
                 "Em 1º de julho foi determinada a remessa dos autos ao tribunal.")

    def molde_com_historico(self, historico=None):
        destino = TMP / "historico.docx"
        proc = proc_novo(historico=historico or self.HISTORICO, andamentos=[], numeros=[NUM_NOVO])
        da.gerar(destino, {"cliente": "Cliente Exemplo", "data_base": "2026-09-18", "processos": [proc]}, sobrescrever_destino=True)
        return destino

    def test_leitura_normaliza_as_datas(self):
        est = da.ler_estrutura(self.molde_com_historico())
        p = est["processos"][0]
        self.assertEqual([a["data"] for a in p["andamentos"]],
                         ["18/06/2026", "20/06/2026", "22/06/2026", "25/06/2026", "01/07/2026"])
        self.assertTrue(all(a["data_em_negrito"] for a in p["andamentos"]))

    def test_cada_forma_e_reconhecida_como_ja_presente(self):
        molde = self.molde_com_historico()
        novos = [("18/06/2026", "Foi proferida sentença de procedência."), ("20/06/2026", "o réu foi intimado da sentença."),
                 ("22/06/2026", "foi interposto recurso de apelação."), ("25/06/2026", "foram apresentadas as razões."),
                 ("01/07/2026", "foi determinada a remessa dos autos ao tribunal.")]
        destino, res = self.aplicar([{"numero": NUM_NOVO, "andamentos": [{"data": d, "texto": t} for d, t in novos]}],
                                    molde=molde, nome="formas")
        self.assertEqual(len([i for i in res["ignorados"] if i["motivo"] == "andamento já presente"]), 5)
        self.assertEqual(res["processos_atualizados"], [NUM_NOVO])        # só o fecho (não houve novidade)
        self.assertEqual({m["campo"] for m in res["mudancas"] if m["numero"]}, {"andamentos_fecho"})
        texto = da.ler_estrutura(destino)["processos"][0]["andamentos_texto"]
        for trecho in ("sentença de procedência", "intimado da sentença", "recurso de apelação", "razões", "remessa"):
            self.assertEqual(texto.count(trecho), 1, trecho)

    def test_andamento_diferente_na_mesma_data_entra(self):
        molde = self.molde_com_historico()
        destino, res = self.aplicar([{"numero": NUM_NOVO, "andamentos": [
            {"data": "22/06/2026", "texto": "foi certificada a tempestividade e deferido o processamento."}]}],
            molde=molde, nome="outra")
        self.assertEqual(res["ignorados"], [])
        self.assertIn("mesma_data", self.codigos(res))
        self.assertIn("tempestividade", da.ler_estrutura(destino)["processos"][0]["andamentos_texto"])

    def test_limiares_sao_configuraveis(self):
        molde = self.molde_com_historico()
        variacao = [{"numero": NUM_NOVO, "andamentos": [
            {"data": "18/06/2026", "texto": "foi proferida sentença de procedência pelo juiz da causa."}]}]
        _, padrao = self.aplicar(variacao, molde=molde, nome="padrao")
        self.assertEqual(len(padrao["ignorados"]), 1)                        # 0,9 >= 0,7: duplicata
        _, estrito = self.aplicar(variacao, molde=molde, nome="estrito", limiar_duplicata=0.95, limiar_parecido=0.95)
        self.assertEqual(estrito["ignorados"], [])                           # com limiares altos, entra
        _, meio = self.aplicar(variacao, molde=molde, nome="meio", limiar_duplicata=0.95, limiar_parecido=0.5)
        self.assertEqual(len(meio["ignorados"]), 1)
        self.assertIn("possivel_duplicata_manual", self.codigos(meio))       # parecido: não entra, mas avisa

    def test_forma_curta_no_meio_da_frase_nao_e_marcador(self):
        molde = self.molde_com_historico("Em 18/06/2026, foi proferida decisão que fixou 5/6 do valor e designou prazo. "
                                         "Em 22/06/2026, foi interposto agravo.")
        p = da.ler_estrutura(molde)["processos"][0]
        self.assertEqual([a["data"] for a in p["andamentos"]], ["18/06/2026", "22/06/2026"])

    def test_data_em_minuscula_no_meio_da_frase_nao_e_marcador(self):
        p = da.ler_estrutura(self.molde)["processos"][10]       # 'em 14/10/2026' e 'até 30/09/2026' no meio da frase
        self.assertEqual([a["data"] for a in p["andamentos"]], ["22/07/2026", "04/09/2026"])

    def test_ano_das_formas_curtas_vem_da_frase_anterior_ou_da_data_base(self):
        a = da._parse_andamentos("Em 20/06 foi feito algo. Em 05/01 outra coisa.", "18/09/2026")
        self.assertEqual([x["data"] for x in a], ["20/06/2026", "05/01/2026"])
        b = da._parse_andamentos("Em 20/12 foi feito algo.", "18/09/2026")      # depois da data-base: ano anterior
        self.assertEqual([x["data"] for x in b], ["20/12/2025"])
        self.assertEqual(da._parse_andamentos("Em 31/02/2026 foi algo.", "18/09/2026"), [])   # data inválida não é marcador


class TestCamposMecanicos(Base):
    """Exceções mecânicas ao 'só acrescenta' e o aviso edicao_manual_sobrescrita."""

    def editar_texto(self, de, para, indice=0, no_resumo=False):
        """Troca `de` por `para` na tabela do processo `indice` (ou no quadro-resumo), mesmo com o texto partido em runs."""
        def mexer(raiz):
            tbl = raiz.find(w("body")).findall(w("tbl"))[0 if no_resumo else indice + 1]
            for p in tbl.iter(w("p")):
                texto = da._texto_par(p)
                i = texto.find(de)
                if i >= 0:
                    da._substituir(p, i, i + len(de), para)
                    return
            raise AssertionError(f"{de!r} não encontrado")
        return mexer

    def test_titulo_e_resumo_divergentes_avisam_ao_trocar(self):
        molde = TMP / "titulo_manual.docx"
        escrever_doc(self.molde, molde, self.editar_texto("AGUARDANDO SENTENÇA", "AGUARDANDO ALGO"))
        destino, res = self.aplicar([{"numero": num(0), "momento_atual": "Cumprimento de sentença"}], molde=molde)
        avisos = [a for a in res["avisos"] if a["codigo"] == "edicao_manual_sobrescrita"]
        self.assertEqual(len(avisos), 1)
        self.assertEqual(da.ler_estrutura(destino)["processos"][0]["momento_atual"], "CUMPRIMENTO DE SENTENÇA")
        self.assertEqual(da.verificar_coerencia(destino, fechos=False), [])

    def test_momento_alterado_a_mao_nos_dois_lugares_so_avisa_se_o_ultimo_gravado_for_conhecido(self):
        a = TMP / "momento_manual.docx"
        molde = TMP / "momento_manual2.docx"
        escrever_doc(self.molde, a, self.editar_texto("AGUARDANDO SENTENÇA", "MOMENTO ESCRITO A MAO"))
        escrever_doc(a, molde, self.editar_texto("AGUARDANDO SENTENÇA", "MOMENTO ESCRITO A MAO", no_resumo=True))
        upd = {"numero": num(0), "momento_atual": "Cumprimento de sentença"}
        _, sem_conhecido = self.aplicar([upd], molde=molde, nome="sem")
        self.assertNotIn("edicao_manual_sobrescrita", self.codigos(sem_conhecido))
        _, com = self.aplicar([upd], molde=molde, nome="com",
                              campos_conhecidos={num(0): {"momento_atual": "AGUARDANDO SENTENÇA"}})
        self.assertEqual(self.codigos(com).count("edicao_manual_sobrescrita"), 1)
        # se o arquivo ainda tem o que o sistema gravou, não há o que avisar
        _, igual = self.aplicar([upd], molde=self.molde, nome="igual",
                                campos_conhecidos={num(0): {"momento_atual": "AGUARDANDO SENTENÇA"}})
        self.assertNotIn("edicao_manual_sobrescrita", self.codigos(igual))

    def test_ultimo_andamento_alterado_a_mao(self):
        molde = TMP / "ultimo_manual.docx"
        escrever_doc(self.molde, molde, self.editar_texto("30/07/2026", "31/07/2026", no_resumo=True))
        destino, res = self.aplicar([{"numero": num(0), "ultimo_andamento": "2026-10-02"}], molde=molde, nome="u",
                                    campos_conhecidos={num(0): {"ultimo_andamento": "30/07/2026"}})
        self.assertEqual(self.codigos(res).count("edicao_manual_sobrescrita"), 1)
        self.assertEqual(da.ler_estrutura(destino)["resumo"][0]["ultimo_andamento"], "02/10/2026")

    def test_ultimo_andamento_nunca_recua(self):
        destino, res = self.aplicar([{"numero": num(0), "ultimo_andamento": "2026-01-02"}])
        self.assertEqual(da.ler_estrutura(destino)["resumo"][0]["ultimo_andamento"], "30/07/2026")

    def test_fecho_com_data_mexida_a_mao(self):
        conhecido = da.ler_estrutura(self.molde)["processos"][2]["andamentos_texto"]
        molde = TMP / "fecho_manual.docx"
        escrever_doc(self.molde, molde, self.editar_texto(gm.DATA_BASE, "01/09/2026", indice=2))
        _, res = self.aplicar([{"numero": num(2), "texto_conhecido": conhecido}], molde=molde, nome="a")
        self.assertEqual(self.codigos(res).count("edicao_manual_sobrescrita"), 1)
        _, sem = self.aplicar([{"numero": num(2), "texto_conhecido": conhecido}], molde=self.molde, nome="b")
        self.assertNotIn("edicao_manual_sobrescrita", self.codigos(sem))

    def test_data_base_mexida_a_mao(self):
        molde = TMP / "db_manual.docx"

        def mexer(raiz):
            for t in raiz.iter(w("t")):
                if (t.text or "") == gm.DATA_BASE:
                    t.text = "19/09/2026"
                    break
        escrever_doc(self.molde, molde, mexer)
        self.assertEqual(da.ler_estrutura(molde)["data_base"], "19/09/2026")
        _, res = self.aplicar([], molde=molde, nome="a", data_base_conhecida=gm.DATA_BASE)
        self.assertEqual(self.codigos(res).count("edicao_manual_sobrescrita"), 1)

    def test_qualificador_do_momento_vai_no_titulo_e_no_resumo(self):
        destino, res = self.aplicar([{"numero": num(0), "momento_atual": "CUMPRIMENTO DE SENTENÇA",
                                      "momento_qualificador": "honorários suspensos"}])
        est = da.ler_estrutura(destino)
        self.assertEqual(est["processos"][0]["momento_atual"], "CUMPRIMENTO DE SENTENÇA (HONORÁRIOS SUSPENSOS)")
        self.assertEqual(est["processos"][0]["momento"], "CUMPRIMENTO DE SENTENÇA")
        self.assertEqual(est["processos"][0]["qualificador"], "HONORÁRIOS SUSPENSOS")
        self.assertEqual(est["resumo"][0]["momento_atual"], "CUMPRIMENTO DE SENTENÇA (HONORÁRIOS SUSPENSOS)")
        self.assertEqual(da.verificar_coerencia(destino, fechos=False), [])
        # o mesmo pedido de novo não muda nada; e a forma 'X (Y)' num campo só vale igual
        d2 = self.destino("de_novo")
        r2 = da.atualizar(destino, d2, [{"numero": num(0), "momento_atual": "CUMPRIMENTO DE SENTENÇA (HONORÁRIOS SUSPENSOS)"}],
                          DATA_BASE_NOVA)
        self.assertEqual(r2["mudancas"], [])

    def test_qualificador_escrito_pelo_advogado_e_preservado_se_o_momento_nao_mudou(self):
        d1, _ = self.aplicar([{"numero": num(0), "momento_atual": "AGUARDANDO SENTENÇA", "momento_qualificador": "prazo em dobro"}])
        self.assertEqual(da.ler_estrutura(d1)["processos"][0]["momento_atual"], "AGUARDANDO SENTENÇA (PRAZO EM DOBRO)")
        # a ficha só diz o momento canônico: o parêntese do advogado fica
        d2 = self.destino("preserva")
        da.atualizar(d1, d2, [{"numero": num(0), "momento_atual": "AGUARDANDO SENTENÇA"}], DATA_BASE_NOVA)
        self.assertEqual(da.ler_estrutura(d2)["processos"][0]["momento_atual"], "AGUARDANDO SENTENÇA (PRAZO EM DOBRO)")
        # momento diferente: troca tudo (o qualificador antigo não vale para o novo momento)
        d3 = self.destino("troca")
        da.atualizar(d1, d3, [{"numero": num(0), "momento_atual": "TRÂNSITO EM JULGADO"}], DATA_BASE_NOVA)
        self.assertEqual(da.ler_estrutura(d3)["processos"][0]["momento_atual"], "TRÂNSITO EM JULGADO")

    def test_processo_novo_com_qualificador_e_varios_numeros(self):
        p = proc_novo(numeros=[NUM_NOVO, NUM_NOVO_2], tipos=["reajuizamento"],
                      momento_atual="processo arquivado", momento_qualificador="decisão favorável")
        destino, _ = self.aplicar([{"novo": p}])
        est = da.ler_estrutura(destino)
        self.assertEqual(est["processos"][-1]["momento_atual"], "PROCESSO ARQUIVADO (DECISÃO FAVORÁVEL)")
        self.assertEqual(est["resumo"][-1]["momento_atual"], "PROCESSO ARQUIVADO (DECISÃO FAVORÁVEL)")
        self.assertEqual(est["resumo"][-1]["numeros"], [NUM_NOVO, NUM_NOVO_2])
        self.assertEqual(est["processos"][-1]["vinculados"], [{"numero": NUM_NOVO_2, "tipo": "reajuizamento"}])


class TestGerarDoZero(Base):
    def relatorio(self, n=5):
        procs = []
        for i, p in enumerate(PROCS[:n]):
            procs.append({"numeros": p["numeros"], "tipos": p["tipos"], "assunto": p["assunto"], "autores": p["autores"],
                          "reus": p["reus"], "ajuizamento": p["ajuizamento"], "valor_causa": p["valor_causa"] or None,
                          "data_citacao": p["citacao"], "juizo": p["juizo"], "area": p["area"], "materia": p["materia"],
                          "momento_atual": p["momento"], "andamentos": [{"data": d, "texto": t} for d, t in p["andamentos"]]})
        procs.append(dict(procs[0], numeros=[NUM_NOVO], andamentos=[]))
        return {"cliente": gm.CLIENTE, "data_base": "2026-10-07", "processos": procs}

    def test_gerar_recarrega_e_e_coerente(self):
        for estilo in da.ESTILOS:
            destino = self.destino(estilo)
            res = da.gerar(destino, self.relatorio(), estilo=estilo)
            self.assertEqual(len(res["processos_novos"]), 6)
            self.assertTrue(res["gravado"])
            import docx
            d = docx.Document(str(destino))
            self.assertEqual(len(d.tables), 7, estilo)            # resumo + 6 blocos (o bloco e a linha de exemplo saíram)
            est = da.ler_estrutura(destino)
            self.assertEqual(est["cliente"], gm.CLIENTE)
            self.assertEqual(est["data_base"], DATA_BASE_NOVA)
            self.assertEqual(len(est["resumo"]), 6)
            self.assertEqual(est["processos"][1]["numeros"], PROCS[1]["numeros"])
            self.assertEqual(est["processos"][3]["valor_causa"], "-")
            self.assertEqual(est["processos"][5]["andamentos_texto"], f"Em {DATA_BASE_NOVA}, sem atualizações.")
            self.assertTrue(all(a["data_em_negrito"] for p in est["processos"] for a in p["andamentos"]))
            self.assertEqual(da.verificar_coerencia(destino), [])
            self.assertEqual(res["textos_gravados"][num(0)], est["processos"][0]["andamentos_texto"])
            self.assertNotIn("0000000-00.0000.0.00.0000", texto_el(corpo(destino)))     # nada do modelo sobrou
            self.assertNotIn("NOME DO CLIENTE", texto_el(corpo(destino)))
            self.assertNotIn("01/01/2000", texto_el(corpo(destino)))

    def test_gerar_nao_tem_fecho_depois_de_novidade(self):
        destino = self.destino("zero")
        da.gerar(destino, self.relatorio())
        p = da.ler_estrutura(destino)["processos"][0]
        self.assertIsNone(p["fecho"])
        destino2 = self.destino("zero_fecho")
        da.gerar(destino2, self.relatorio(), fecho_apos_novidade=True)
        self.assertEqual(da.ler_estrutura(destino2)["processos"][0]["fecho"]["data"], DATA_BASE_NOVA)

    def test_gerado_do_zero_pode_ser_atualizado(self):
        for estilo in da.ESTILOS:
            base = self.destino(f"zero2-{estilo}")
            da.gerar(base, self.relatorio(), estilo=estilo)
            destino, res = self.aplicar([{"numero": num(0), "andamentos": [{"data": "05/10/2026", "texto": "foi proferida sentença."}],
                                          "momento_atual": "Aguardando julgamento da apelação"},
                                         {"novo": proc_novo(numeros=[NUM_NOVO_2])}], molde=base, nome=f"zero3-{estilo}")
            est = da.ler_estrutura(destino)
            self.assertEqual(len(est["processos"]), 7)
            self.assertEqual(est["processos"][0]["momento_atual"], "AGUARDANDO JULGAMENTO DA APELAÇÃO")
            self.assertEqual(da.verificar_coerencia(destino, fechos=False), [])

    def test_gerar_e_deterministico(self):
        a, b = self.destino("det1"), self.destino("det2")
        da.gerar(a, self.relatorio())
        da.gerar(b, self.relatorio())
        with zipfile.ZipFile(a) as za, zipfile.ZipFile(b) as zb:
            self.assertEqual(za.read("word/document.xml"), zb.read("word/document.xml"))

    def test_gerar_nao_sobrescreve(self):
        destino = self.destino("zero4")
        da.gerar(destino, self.relatorio())
        with self.assertRaises(FileExistsError):
            da.gerar(destino, self.relatorio())

    def test_estilo_desconhecido(self):
        with self.assertRaises(ValueError):
            da.gerar(self.destino("x"), self.relatorio(), estilo="barroco")

    def test_compacto_le_igual_ao_estilo_a(self):
        a, c = self.destino("a"), self.destino("c")
        da.gerar(a, self.relatorio(), estilo="a")
        da.gerar(c, self.relatorio(), estilo="compacto")
        ea, ec = da.ler_estrutura(a), da.ler_estrutura(c)
        for chave in ("numeros", "assunto", "autores", "reus", "ajuizamento", "valor_causa", "data_citacao", "juizo", "area",
                      "materia", "andamentos_texto", "momento_atual"):
            self.assertEqual([p[chave] for p in ea["processos"]], [p[chave] for p in ec["processos"]], chave)
        self.assertEqual(ea["resumo"], ec["resumo"])
        # layout diferente: o compacto usa seis colunas de grade e fonte menor
        self.assertEqual(len(tabelas(c)[1].find(w("tblGrid")).findall(w("gridCol"))), 6)
        self.assertEqual(len(tabelas(a)[1].find(w("tblGrid")).findall(w("gridCol"))), 4)
        self.assertEqual(next(tabelas(c)[1].iter(w("sz"))).get(w("val")), "18")


class TestModelosSanitizados(Base):
    PASTA = RAIZ / "src" / "modelos" / "docx_a"

    def test_modelos_existem_e_sao_regeneraveis(self):
        for estilo in da.ESTILOS:
            versionado = self.PASTA / f"modelo_{estilo}.docx"
            self.assertTrue(versionado.exists(), versionado)
            novo = TMP / f"recriado_{estilo}.docx"
            da.criar_modelo(estilo, novo)
            with zipfile.ZipFile(versionado) as z0, zipfile.ZipFile(novo) as z1:
                self.assertEqual(z0.read("word/document.xml"), z1.read("word/document.xml"), estilo)

    def test_criar_modelo_e_deterministico(self):
        a, b = TMP / "det_a.docx", TMP / "det_b.docx"
        da.criar_modelo("compacto", a)
        da.criar_modelo("compacto", b)
        self.assertEqual(a.read_bytes(), b.read_bytes())

    def test_modelo_so_tem_marcadores(self):
        for estilo in da.ESTILOS:
            caminho = self.PASTA / f"modelo_{estilo}.docx"
            with zipfile.ZipFile(caminho) as z:
                todo = " ".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith(".xml") or n.endswith(".rels"))
            self.assertEqual(set(da.CNJ.findall(todo)), {"0000000-00.0000.0.00.0000"}, estilo)
            texto = texto_el(corpo(caminho))
            for esperado in ("NOME DO CLIENTE", "(assunto)", "01/01/2000", "MOMENTO ATUAL"):
                self.assertIn(esperado, texto)
            est = da.ler_estrutura(caminho)
            self.assertEqual(len(est["processos"]), 1)
            self.assertEqual(len(est["resumo"]), 1)
            self.assertEqual(est["avisos"], [])
            with zipfile.ZipFile(caminho) as z:
                core = z.read("docProps/core.xml").decode()
            self.assertNotIn("python-docx", core)

    def test_modelo_tem_um_bloco_limpo_para_clonar(self):
        for estilo in da.ESTILOS:
            doc = da._abrir(self.PASTA / f"modelo_{estilo}.docx")
            modelo, limpo = da._escolher_modelo(doc.blocos())
            self.assertTrue(limpo, estilo)

    def test_modelo_ausente_da_erro_claro(self):
        original = da.PASTA_MODELOS
        try:
            da.PASTA_MODELOS = TMP / "nao_existe"
            with self.assertRaises(FileNotFoundError):
                da.gerar(self.destino("x"), TestGerarDoZero.relatorio(self))
        finally:
            da.PASTA_MODELOS = original


class TestEnsaio(Base):
    def test_ensaio_da_conferencia_manual(self):
        d1 = self.destino("ensaio1")
        res = da.ensaio(self.molde, d1, DATA_BASE_NOVA)
        self.assertTrue(res["gravado"])
        self.assertEqual(res["processos_novos"], ["0000000-00.0000.0.00.0000"])
        est = da.ler_estrutura(d1)
        self.assertEqual(len(est["processos"]), 13)
        self.assertIn("ENSAIO", est["processos"][0]["andamentos_texto"])
        self.assertIsNone(est["processos"][0]["fecho"])
        self.assertEqual(est["processos"][5]["fecho"]["data"], DATA_BASE_NOVA)
        self.assertEqual(da.verificar_coerencia(d1, fechos=False), [])
        # o mesmo ensaio sobre a própria saída não muda nada
        d2 = self.destino("ensaio2")
        res2 = da.ensaio(d1, d2, DATA_BASE_NOVA)
        self.assertEqual(res2["mudancas"], [])
        # a linha de comando roda e informa
        import contextlib
        import io
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            da._cli(["docx_a.py", "ensaio", str(self.molde), str(self.destino("ensaio3")), DATA_BASE_NOVA])
        self.assertIn("Gravado", buf.getvalue())
        # e roda como script, de qualquer pasta (é como o usuário a chama na conferência manual)
        r = subprocess.run([sys.executable, str(RAIZ / "src" / "escritores" / "docx_a.py"), "ler", str(self.molde)],
                           capture_output=True, text=True, cwd=str(TMP), timeout=120)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn(gm.CLIENTE, r.stdout)


class TestMolde(Base):
    def test_molde_ilegivel_vira_aviso_e_nao_grava(self):
        lixo = TMP / "lixo.docx"
        lixo.write_bytes(b"isto nao e um docx")
        destino = self.destino("lixo")
        res = da.atualizar(lixo, destino, [], DATA_BASE_NOVA)
        self.assertFalse(res["gravado"])
        self.assertFalse(destino.exists())
        self.assertEqual([a["codigo"] for a in res["avisos"]], ["molde_ilegivel"])
        self.assertEqual(res["avisos"][0]["nivel"], "erro")

    def test_zip_sem_document_xml(self):
        vazio = TMP / "vazio.docx"
        with zipfile.ZipFile(vazio, "w") as z:
            z.writestr("outra.txt", "x")
        res = da.atualizar(vazio, self.destino("v"), [], DATA_BASE_NOVA)
        self.assertEqual(res["avisos"][0]["codigo"], "molde_ilegivel")

    def test_molde_inexistente(self):
        res = da.atualizar(TMP / "nao_existe.docx", self.destino("n"), [], DATA_BASE_NOVA)
        self.assertEqual(res["avisos"][0]["codigo"], "molde_ilegivel")


# ------------------------------------------------------------------ contrato gravar(molde, estado, destino)

def ficha_de(i, **campos):
    p = PROCS[i]
    f = {"numero": p["numeros"][0], "vinculados": [{"numero": n, "tipo": t} for n, t in zip(p["numeros"][1:], p["tipos"])],
         "campos": {"assunto": {"valor": p["assunto"], "origem": "migrado"},
                    "autores": {"valor": p["autores"], "origem": "migrado"}, "reus": {"valor": p["reus"], "origem": "migrado"},
                    "momento_atual": {"valor": p["momento"], "origem": "migrado"}}}
    for k, v in campos.items():
        f["campos"][k] = {"valor": v, "origem": "coletado"}
    return f


class TestContratoGravar(Base):
    def estado(self):
        return {"cliente": gm.CLIENTE, "data_base": "2026-10-07", "perfil": {}, "parametros": {},
                "fichas": [ficha_de(0, momento_atual="CUMPRIMENTO DE SENTENÇA", ultimo_andamento="2026-10-02"),
                           ficha_de(1),
                           dict(ficha_de(2), numero=NUM_NOVO, vinculados=[], campos={
                               "assunto": {"valor": "Ação nova", "origem": "coletado"},
                               "valor_causa": {"valor": "1500.00", "origem": "coletado"},
                               "data_ajuizamento": {"valor": "2026-09-30", "origem": "coletado"},
                               "vara": {"valor": "1ª Vara Fictícia", "origem": "coletado"},
                               "momento_atual": {"valor": "AGUARDANDO CITAÇÃO", "origem": "coletado"}})],
                "eventos": [{"numero": num(0), "data": "02/10/2026", "frase": "Foi proferida sentença de procedência.",
                             "conteudo": "", "status": "aprovado"},
                            {"numero": PROCS[1]["numeros"][1], "data": "01/10/2026", "frase": "Foi julgado o agravo de instrumento.",
                             "grau": "2º grau", "status": "aprovado"},
                            {"numero": num(0), "data": "03/10/2026", "frase": "Foi proferido despacho.", "status": "rascunho"}]}

    def test_gravar_com_molde(self):
        destino = self.destino("g1")
        res = da.gravar(self.molde, self.estado(), destino)
        for chave in ("destino", "processos_atualizados", "processos_novos", "mudancas", "avisos", "ignorados", "textos_gravados"):
            self.assertIn(chave, res)
        self.assertEqual(Path(res["destino"]), destino)
        self.assertEqual(res["processos_novos"], [NUM_NOVO])
        self.assertTrue(all(set(m) == {"numero", "campo", "antes", "depois"} for m in res["mudancas"]))
        self.assertTrue(all({"nivel", "codigo", "onde", "mensagem", "candidatos"} <= set(a) for a in res["avisos"]))
        est = da.ler_estrutura(destino)
        self.assertEqual(est["processos"][0]["momento_atual"], "CUMPRIMENTO DE SENTENÇA")
        self.assertIn("Em 02/10/2026, foi proferida sentença de procedência.", est["processos"][0]["andamentos_texto"])
        self.assertNotIn("03/10/2026", est["processos"][0]["andamentos_texto"])         # rascunho não vai para o cliente
        self.assertIn("Em 01/10/2026, no 2º grau, foi julgado o agravo de instrumento.", est["processos"][1]["andamentos_texto"])
        self.assertEqual(est["processos"][-1]["valor_causa"], "R$ 1.500,00")
        self.assertEqual(est["processos"][-1]["juizo"], "1ª Vara Fictícia")
        # textos_gravados vem com a chave da ficha (principal), não a do título
        self.assertEqual(res["textos_gravados"][PROCS[1]["numeros"][0]], est["processos"][1]["andamentos_texto"])
        # idempotente
        res3 = da.gravar(destino, self.estado(), self.destino("g3"))
        self.assertEqual(res3["mudancas"], [])
        with self.assertRaises(ValueError):
            da.gravar(self.molde, self.estado(), self.molde)

    def test_gravar_sem_molde_cria_do_zero(self):
        destino = self.destino("g4")
        res = da.gravar(None, self.estado(), destino)
        self.assertEqual(len(res["processos_novos"]), 3)
        self.assertEqual(da.verificar_coerencia(destino), [])
        self.assertEqual(len(da.ler_estrutura(destino)["processos"]), 3)

    def test_estilo_vem_do_perfil(self):
        estado = self.estado()
        estado["perfil"] = {"estilo_texto": "b"}
        destino = self.destino("g5")
        da.gravar(None, estado, destino)
        self.assertEqual(len(tabelas(destino)[1].find(w("tblGrid")).findall(w("gridCol"))), 6)        # compacto
        destino2 = self.destino("g6")
        da.gravar(None, self.estado(), destino2, estilo="a")
        self.assertEqual(len(tabelas(destino2)[1].find(w("tblGrid")).findall(w("gridCol"))), 4)

    def test_qualificador_do_momento_vem_da_ficha(self):
        estado = self.estado()
        # três jeitos de a ficha dizer o qualificador: campo próprio, chave plana e parênteses no próprio valor
        estado["fichas"][0]["campos"]["momento_qualificador"] = {"valor": "honorários suspensos", "origem": "humano"}
        estado["fichas"][1]["momento_qualificador"] = "decisão favorável"
        estado["fichas"][1]["campos"]["momento_atual"]["valor"] = "PROCESSO ARQUIVADO"
        estado["fichas"][2]["campos"]["momento_atual"]["valor"] = "SUSPENSO (AGUARDANDO O TEMA)"
        destino = self.destino("q1")
        da.gravar(self.molde, estado, destino)
        est = da.ler_estrutura(destino)
        self.assertEqual(est["processos"][0]["momento_atual"], "CUMPRIMENTO DE SENTENÇA (HONORÁRIOS SUSPENSOS)")
        self.assertEqual(est["processos"][1]["momento_atual"], "PROCESSO ARQUIVADO (DECISÃO FAVORÁVEL)")
        self.assertEqual(est["processos"][-1]["momento_atual"], "SUSPENSO (AGUARDANDO O TEMA)")
        self.assertEqual(est["resumo"][0]["momento_atual"], est["processos"][0]["momento_atual"])
        self.assertEqual(da.verificar_coerencia(destino, fechos=False), [])

    def test_molde_ilegivel_nao_levanta_excecao(self):
        lixo = TMP / "lixo2.docx"
        lixo.write_bytes(b"nao e docx")
        res = da.gravar(lixo, self.estado(), self.destino("lixo"))
        self.assertFalse(res["gravado"])
        self.assertEqual(res["avisos"][0]["codigo"], "molde_ilegivel")

    def test_historico_da_linha_de_base_entra_ao_gerar(self):
        estado = self.estado()
        estado["fichas"][2]["linha_de_base"] = {"data_base": "2026-09-18", "arquivo": "x.xlsx", "ultimo_andamento": "2026-09-10",
                                                "andamentos_texto": "Em 10/09/2026 foi distribuída a ação. Até 18/09/2026 sem andamentos."}
        destino = self.destino("g7")
        da.gravar(None, estado, destino)
        p = [x for x in da.ler_estrutura(destino)["processos"] if x["numeros"] == [NUM_NOVO]][0]
        self.assertEqual(p["andamentos_texto"], f"Em 10/09/2026 foi distribuída a ação. Em {DATA_BASE_NOVA}, sem atualizações.")
        self.assertTrue(p["andamentos"][0]["data_em_negrito"])

    def test_ultimo_texto_gravado_alimenta_a_deteccao_de_edicao_manual(self):
        d1 = self.destino("c1")
        r1 = da.gravar(self.molde, self.estado(), d1)
        e2 = self.estado()
        for f in e2["fichas"]:
            f["ultimo_texto_gravado"] = {"data_base": e2["data_base"], "texto": r1["textos_gravados"][f["numero"]],
                                         "arquivo": "x.docx", "campos": r1["campos_gravados"][f["numero"]]}

        def mexer(raiz):    # o advogado mexe no momento (título e resumo) e acrescenta frase no fim do texto do 1º processo
            for t in raiz.iter(w("t")):
                if t.text == "CUMPRIMENTO DE SENTENÇA":
                    t.text = "MOMENTO DO ADVOGADO"
            tbl = raiz.find(w("body")).findall(w("tbl"))[1]
            par = tbl.findall(w("tr"))[-1].findall(w("tc"))[-1].findall(w("p"))[-1]
            par.append(da._novo_run(" (nota do advogado)", da._rpr_modelo(par.findall(w("r"))[0].find(w("rPr")), False)))
        editado = TMP / "c1_editado.docx"
        escrever_doc(d1, editado, mexer)
        e2["data_base"] = "2026-11-05"
        e2["eventos"].append({"numero": num(0), "data": "20/10/2026", "frase": "Foi interposto recurso.", "status": "aprovado"})
        r2 = da.gravar(editado, e2, self.destino("c2"))
        cods = self.codigos(r2)
        self.assertEqual(cods.count("edicao_manual_sobrescrita"), 1)           # momento
        self.assertEqual(cods.count("edicao_manual"), 1)                        # frase solta no fim do texto
        t = da.ler_estrutura(r2["destino"])["processos"][0]["andamentos_texto"]
        self.assertIn("(nota do advogado)", t)                                  # nunca apagada
        self.assertIn("Em 20/10/2026, foi interposto recurso.", t)


class TestVariosCiclos(Base):
    """200 processos em ciclos mensais seguidos: nada duplica, nada se perde, fecho só sem novidade, idempotente."""

    def setUp(self):
        self.fichas = ficticio.gerar_carteira(200, clientes=1, semente=3)
        for f in self.fichas:
            ficticio.anexar_linha_de_base(f, data_base="2026-09-18")
        self.eventos = []

    def novos_eventos(self, ciclo, data_iso, passo):
        """Um evento aprovado a cada `passo` processos, com data própria do ciclo."""
        for i, f in enumerate(self.fichas):
            if i % passo == ciclo % passo:
                self.eventos.append({"numero": f["numero"], "data": ficha.data_br(data_iso), "status": "aprovado",
                                     "frase": f"Foi proferido despacho de ordem {ciclo} no processo.", "conteudo": ""})

    def estado(self, data_base):
        return {"cliente": "Cliente Exemplo 01 Ltda", "data_base": data_base, "perfil": {}, "parametros": {},
                "fichas": self.fichas, "eventos": copy.deepcopy(self.eventos)}

    def guardar(self, estado, res):
        for f in self.fichas:
            f["ultimo_texto_gravado"] = {"data_base": estado["data_base"], "texto": res["textos_gravados"][f["numero"]],
                                         "arquivo": "x.docx", "campos": res["campos_gravados"][f["numero"]]}

    def test_tres_ciclos_sem_duplicar_e_idempotentes(self):
        # ciclo 0: relatório inicial (do zero), sem eventos novos
        e0 = self.estado("2026-09-18")
        d0 = self.destino("c0")
        r0 = da.gravar(None, e0, d0)
        self.assertEqual(len(r0["processos_novos"]), 200)
        self.guardar(e0, r0)
        est0 = da.ler_estrutura(d0)
        self.assertEqual(len(est0["processos"]), 200)
        # o histórico migrado entrou, com as datas em negrito, e o fecho (sem novidade) vem com a data-base
        for p in est0["processos"]:
            self.assertEqual(p["fecho"]["data"], "18/09/2026")
            self.assertTrue(all(a["data_em_negrito"] for a in p["andamentos"]))
        anterior, anterior_res = d0, r0
        datas = {1: ("2026-10-02", "07/10/2026", "2026-10-07"), 2: ("2026-11-03", "05/11/2026", "2026-11-05"),
                 3: ("2026-12-02", "04/12/2026", "2026-12-04")}
        for ciclo, (data_evento, data_br, data_base) in datas.items():
            for e in self.eventos:                       # o que já foi para o relatório volta como "relatado"
                e["status"] = "relatado"
            self.novos_eventos(ciclo, data_evento, 3)
            estado = self.estado(data_base)
            destino = self.destino(f"c{ciclo}")
            t0 = time.time()
            res = da.gravar(anterior, estado, destino)
            self.assertLess(time.time() - t0, 30)
            est = da.ler_estrutura(destino)
            self.assertEqual(len(est["processos"]), 200)
            self.assertEqual(res["processos_novos"], [])
            self.assertEqual(res["ignorados"], [])        # eventos já relatados que constam no texto não poluem a lista
            self.assertEqual(da.verificar_coerencia(destino, fechos=False, ultimo=False), [])
            self.assertNotIn("edicao_manual", self.codigos(res))
            self.assertNotIn("edicao_manual_sobrescrita", self.codigos(res))
            self.assertNotIn("incoerencia_documento", self.codigos(res))
            self.assertEqual([a for a in res["avisos"] if a["nivel"] == "erro"], [])
            novos_no_ciclo = {e["numero"] for e in self.eventos if e["data"] == ficha.data_br(data_evento)}
            for p in est["processos"]:
                numero = p["numeros"][0]
                # cada evento (de qualquer ciclo) aparece exatamente uma vez
                for e in self.eventos:
                    if e["numero"] == numero:
                        self.assertEqual(p["andamentos_texto"].count(f"Em {e['data']}, foi proferido despacho de ordem"), 1,
                                         (ciclo, numero, e["data"]))
                # fecho só quando o processo não teve andamento novo NESTE ciclo
                if numero in novos_no_ciclo:
                    self.assertIsNone(p["fecho"], (ciclo, numero))
                    self.assertTrue(p["andamentos_texto"].endswith(
                        f"Em {ficha.data_br(data_evento)}, foi proferido despacho de ordem {ciclo} no processo."), (ciclo, numero))
                else:
                    self.assertEqual(p["fecho"]["data"], data_br, (ciclo, numero))
                    self.assertTrue(p["andamentos_texto"].endswith(f"Em {data_br}, sem atualizações."))
            # o histórico migrado nunca é reescrito: o texto do ciclo anterior (sem o fecho) é prefixo do novo
            for f in self.fichas:
                antigo = da._sem_fecho(anterior_res["textos_gravados"][f["numero"]])
                self.assertTrue(res["textos_gravados"][f["numero"]].startswith(antigo), (ciclo, f["numero"]))
            # idempotência: reaplicar o mesmo estado sobre a própria saída não muda nada, byte a byte
            self.guardar(estado, res)
            res_b = da.gravar(destino, self.estado(data_base), self.destino(f"c{ciclo}b"))
            self.assertEqual(res_b["mudancas"], [], ciclo)
            with zipfile.ZipFile(destino) as z1, zipfile.ZipFile(res_b["destino"]) as z2:
                self.assertEqual(z1.read("word/document.xml"), z2.read("word/document.xml"))
            anterior, anterior_res = destino, res

    def test_ciclo_com_processo_novo_no_meio(self):
        e0 = self.estado("2026-09-18")
        d0 = self.destino("n0")
        da.gravar(None, {**e0, "fichas": self.fichas[:190]}, d0)
        self.novos_eventos(1, "2026-10-02", 3)
        res = da.gravar(d0, self.estado("2026-10-07"), self.destino("n1"))
        self.assertEqual(len(res["processos_novos"]), 10)
        est = da.ler_estrutura(res["destino"])
        self.assertEqual(len(est["processos"]), 200)
        self.assertEqual(len(est["resumo"]), 200)
        self.assertEqual(da.verificar_coerencia(res["destino"], fechos=False, ultimo=False), [])
        novo = [p for p in est["processos"] if self.fichas[195]["numero"] in p["numeros"]][0]
        self.assertTrue(novo["andamentos_texto"])                       # o histórico migrado da ficha entrou

    def test_momentos_e_numeros_dos_vinculados_das_fichas(self):
        destino = self.destino("v")
        da.gravar(None, self.estado("2026-10-07"), destino)
        est = da.ler_estrutura(destino)
        com_vinculado = [f for f in self.fichas if f["vinculados"]]
        self.assertTrue(com_vinculado, "a carteira desta semente deveria ter processos vinculados")
        for f in com_vinculado:
            p = [x for x in est["processos"] if f["numero"] == x["numeros"][0]][0]
            self.assertEqual(p["numeros"], ficha.todos_os_numeros(f))
            self.assertEqual([v["tipo"] for v in p["vinculados"]], [v["tipo"] for v in f["vinculados"]])
        for f in self.fichas:
            p = [x for x in est["processos"] if f["numero"] == x["numeros"][0]][0]
            self.assertEqual(p["momento_atual"], ficha.obter(f, "momento_atual"))
            self.assertEqual(p["vinculados"] == [], not f["vinculados"])


@unittest.skipUnless(SOFFICE, "soffice não instalado")
class TestLibreOffice(Base):
    def test_atualizado_abre_e_texto_novo_aparece_no_pdf(self):
        upds = self.atualizacoes() + [{"novo": proc_novo()}]
        destino, _ = self.aplicar(upds, nome="pdf")
        texto, paginas = pdf_texto(destino)
        self.assertGreater(paginas, 2)
        for trecho in ("Data-Base:07/10/2026", "02/10/2026", "ojuizproferiusentençadeprocedência", "05/10/2026",
                       "Em07/10/2026,sematualizações", NUM_NOVO, "Açãodetestenova", "CUMPRIMENTODESENTENÇA",
                       "foidistribuídaaação"):
            self.assertIn(so_digitos_e_letras(trecho), texto, trecho)
        self.assertNotIn("18/09/2026,sematualizações" + "Em", texto)
        # o que não foi tocado continua lá
        self.assertIn(so_digitos_e_letras("Anotação do advogado: conferir o valor com o cliente."), texto)

    def test_original_e_gerado_do_zero_tambem_abrem(self):
        texto, paginas = pdf_texto(self.molde)
        self.assertIn("COMERCIALEXEMPLOALFALTDA.", texto)
        self.assertGreater(paginas, 2)
        for estilo in da.ESTILOS:
            base = self.destino(f"zero_pdf_{estilo}")
            da.gerar(base, TestGerarDoZero.relatorio(self), estilo=estilo)
            texto2, _ = pdf_texto(base)
            self.assertIn("PROCESSONº" + so_digitos_e_letras(PROCS[0]["numeros"][0]), texto2)
            self.assertIn(so_digitos_e_letras("Data-Base: 07/10/2026"), texto2)
            self.assertNotIn("NOMEDOCLIENTE", texto2)

    def test_resumo_e_blocos_batem_no_pdf(self):
        destino, _ = self.aplicar()
        texto, _ = pdf_texto(destino)
        # momento novo aparece no quadro-resumo e no título do bloco: pelo menos 2 vezes
        self.assertGreaterEqual(texto.count("CUMPRIMENTODESENTENÇA"), 2)

    def test_200_processos_gerados_abrem_no_libreoffice(self):
        fichas = ficticio.gerar_carteira(200, clientes=1, semente=5)
        estado = {"cliente": "Cliente Exemplo 01 Ltda", "data_base": "2026-10-07", "perfil": {"estilo_texto": "b"},
                  "parametros": {}, "fichas": fichas, "eventos": []}
        destino = self.destino("g200")
        da.gravar(None, estado, destino)
        texto, paginas = pdf_texto(destino)
        self.assertGreater(paginas, 10)
        self.assertIn(so_digitos_e_letras(fichas[-1]["numero"]), texto)


class TestConfidencialidade(unittest.TestCase):
    """Nenhum número de processo real nos arquivos deste workstream (mesmo padrão do empacotar.sh)."""
    PADRAO = re.compile(r"\b[0-9]{7}-[0-9]{2}\.[0-9]{4}\.[0-9]\.[0-9]{2}\.[0-9]{4}\b")
    PERMITIDOS = re.compile(r"0000000-00\.0000\.0\.00\.0000|9999999-99\.9999\.9\.99\.9999|123456[7-9]-|1234570-")

    def test_nenhum_numero_real(self):
        arquivos = [RAIZ / "src" / "escritores" / "__init__.py", RAIZ / "src" / "escritores" / "docx_a.py",
                    RAIZ / "tests" / "test_docx_a.py", RAIZ / "docs" / "fase2" / "conferencia-docx.md"]
        arquivos += list((RAIZ / "src" / "modelos" / "docx_a").glob("*.py"))
        arquivos += list(RAIZ.glob("docs/fase2/RFC-docx*.md"))
        for caminho in arquivos:
            if not caminho.exists():
                continue
            for linha in caminho.read_text(encoding="utf-8").splitlines():
                for achado in self.PADRAO.findall(linha):
                    self.assertTrue(self.PERMITIDOS.search(achado), f"{caminho.name}: {achado}")
        for modelo in (RAIZ / "src" / "modelos" / "docx_a").glob("*.docx"):
            with zipfile.ZipFile(modelo) as z:
                for nome in z.namelist():
                    for achado in self.PADRAO.findall(z.read(nome).decode("utf-8", "ignore")):
                        self.assertTrue(self.PERMITIDOS.search(achado), f"{modelo.name}/{nome}: {achado}")


if __name__ == "__main__":
    unittest.main()
