"""Testes do escritor .xlsx (modelo B): `src/escritores/xlsx_b.py`, `src/planilha.py` e o modelo padrão.

    python3 -m unittest tests/test_xlsx_b.py -v

Só dados fictícios (números de processo gerados em tempo de execução por `ficticio`), sem rede.
Duas partes:

  1. MOTOR CIRÚRGICO: a suíte do spike S1 migrada (fixtures do spike em spikes/s1_xlsx/gerar_modelo.py: um
     modelo B "de cliente" com tabela, gráficos, tabela dinâmica, validação e formatação condicional; as
     classes que dependem dele são puladas se a pasta spikes/ não existir) e ampliada com as recusas novas.
  2. ESCRITOR DO CONTRATO (`gravar`): modelo padrão, atualização do arquivo do cliente, política das colunas,
     idempotência, avisos, cache de fórmulas, indicadores conferidos contra um cálculo independente.

Os testes que abrem o LibreOffice (/usr/bin/soffice) são pulados se ele não existir. Eles provam que o
LibreOffice abre e recalcula; NÃO provam nada sobre Excel ou Google Sheets (ver docs/fase2/conferencia-xlsx.md).
"""
import csv
import datetime
import hashlib
import importlib.util
import io
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
import warnings
import zipfile
import xml.etree.ElementTree as ET
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import isolamento  # noqa: F401,E402  (antes de qualquer módulo da ferramenta)
import ficticio  # noqa: E402  (já põe src/ no sys.path)
import ficha as fi  # noqa: E402
from escritores import xlsx_b as x  # noqa: E402

import openpyxl  # noqa: E402

warnings.filterwarnings("ignore", module="openpyxl")
RAIZ = Path(__file__).resolve().parent.parent
SPIKE = RAIZ / "spikes" / "s1_xlsx" / "gerar_modelo.py"
if SPIKE.exists():
    _spec = importlib.util.spec_from_file_location("gerar_modelo_s1", SPIKE)
    g = importlib.util.module_from_spec(_spec)
    _spec.loader.exec_module(g)
else:      # pacote sem spikes/: as classes que usam o modelo do spike são puladas
    g = None
requer_spike = unittest.skipUnless(g is not None, "spikes/s1_xlsx ausente (fixtures do spike)")
SOFFICE = shutil.which("soffice") or "/usr/bin/soffice"
TEM_LO = Path(SOFFICE).exists()
M = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def partes(caminho):
    with zipfile.ZipFile(caminho) as z:
        return {i.filename: z.read(i.filename) for i in z.infolist()}


def sha(b):
    return hashlib.sha256(b).hexdigest()


class BaseLO(unittest.TestCase):
    """Pasta temporária e utilitários do LibreOffice."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp(prefix="xlsx_b-"))
        cls.perfil_lo = cls.tmp / "perfil-lo"

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def saida(self, nome):
        return self.tmp / nome

    def lo(self, arquivo, formato, filtro=None):
        """Converte com o LibreOffice; devolve a pasta de saída."""
        pasta = self.tmp / ("lo-" + Path(arquivo).stem + "-" + formato.split(":")[0])
        shutil.rmtree(pasta, ignore_errors=True)
        pasta.mkdir()
        p = subprocess.run(
            [SOFFICE, f"-env:UserInstallation=file://{self.perfil_lo}", "--headless", "--convert-to",
             filtro or formato, "--outdir", str(pasta), str(arquivo)],
            capture_output=True, text=True, timeout=240)
        self.assertEqual(p.returncode, 0, p.stderr)
        return pasta

    def lo_valores(self, arquivo):
        """Abre no LibreOffice, deixa ele salvar de novo como xlsx e lê os valores calculados."""
        pasta = self.lo(arquivo, "xlsx")
        return openpyxl.load_workbook(pasta / Path(arquivo).name, data_only=True)

    @staticmethod
    def linhas_novas(a, b):
        return [g.linha_ficticia(i) for i in range(a, b)]


@requer_spike
class Base(BaseLO):
    """Os três modelos do spike: completo (10 linhas), sujo e pré-formatado."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.completo = cls.tmp / "completo.xlsx"
        cls.sujo = cls.tmp / "sujo.xlsx"
        cls.pre = cls.tmp / "pre.xlsx"
        g.gerar(cls.completo, "completo")
        g.gerar(cls.sujo, "sujo")
        g.gerar(cls.pre, "preformatado")


class TestInserir200(Base):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.novas = cls.linhas_novas(10, 210)
        cls.dest = cls.tmp / "200.xlsx"
        cls.res = x.inserir_linhas(cls.completo, cls.dest, "Processos", cls.novas)

    def test_a_resultado(self):
        r = self.res
        self.assertEqual((r.linhas_inseridas, r.primeira_linha, r.ultima_linha), (200, 12, 211))
        self.assertEqual(r.avisos, [])

    def test_b_partes_nao_editadas_identicas(self):
        antes, depois = partes(self.completo), partes(self.dest)
        removidas = set(antes) - set(depois)
        self.assertEqual(set(depois) - set(antes), set())          # nenhuma parte nova
        self.assertEqual(removidas, {"xl/calcChain.xml"})          # única remoção: calcChain
        alteradas = {n for n in depois if antes[n] != depois[n]}
        # cada parte que mudou, e por quê:
        justificativa = {
            "xl/worksheets/sheet1.xml": "aba Processos: linhas novas, dimension, CF/DV, cache das fórmulas",
            "xl/tables/table1.xml": "ref e autoFilter da Tabela",
            "xl/workbook.xml": "intervalos nomeados e fullCalcOnLoad",
            "xl/_rels/workbook.xml.rels": "relação do calcChain removida",
            "[Content_Types].xml": "tipo do calcChain removido",
            "xl/charts/chart2.xml": "série do gráfico que aponta para a coluna da tabela",
            "xl/worksheets/sheet3.xml": "Indicadores: fórmulas A1 que cobrem a coluna inteira + cache",
        }
        self.assertEqual(alteradas, set(justificativa), f"mudaram: {sorted(alteradas)}")
        self.assertEqual(sorted(alteradas), self.res.partes_alteradas)
        # tudo o mais é idêntico byte a byte: estilos, textos, tema, outras abas, gráfico 1,
        # desenhos, tabela dinâmica (cache e definição), propriedades
        for n in set(antes) - alteradas - removidas:
            self.assertEqual(antes[n], depois[n], n)
        for essencial in ("xl/styles.xml", "xl/sharedStrings.xml", "xl/theme/theme1.xml", "xl/charts/chart1.xml",
                          "xl/drawings/drawing1.xml", "xl/pivotTables/pivotTable1.xml",
                          "xl/pivotCache/pivotCacheDefinition1.xml", "xl/worksheets/sheet2.xml",
                          "xl/worksheets/sheet4.xml", "xl/worksheets/sheet5.xml", "docProps/core.xml"):
            self.assertIn(essencial, depois)
            self.assertEqual(antes[essencial], depois[essencial], essencial)

    def test_b2_linhas_antigas_do_xml_intactas(self):
        """A linha 1 (cabeçalho) e as células que não são fórmula das linhas antigas voltam idênticas."""
        a = partes(self.completo)["xl/worksheets/sheet1.xml"].decode()
        d = partes(self.dest)["xl/worksheets/sheet1.xml"].decode()
        self.assertEqual(re.search(r'<row r="1".*?</row>', a, re.S).group(0),
                         re.search(r'<row r="1".*?</row>', d, re.S).group(0))
        for n in range(2, 12):
            ca = re.findall(r'<c r="[A-Z]+%d"[^>]*?(?<!<f)>(?:(?!<f>).)*?</c>' % n,
                            re.search(rf'<row r="{n}".*?</row>', a, re.S).group(0))
            for celula in ca:
                if "<f>" not in celula:
                    self.assertIn(celula, d)

    def test_c_openpyxl_carrega(self):
        wb = openpyxl.load_workbook(self.dest)
        ws = wb["Processos"]
        self.assertEqual(ws.max_row, 211)
        self.assertEqual(ws.tables["tblProcessos"].ref, "A1:AC211")
        self.assertEqual(ws["A211"].value, g.cnj_ficticio(210))
        self.assertEqual(ws["G12"].number_format, "dd/mm/yyyy")           # estilo copiado da linha-modelo
        self.assertEqual(ws["K12"].number_format, "#,##0.00")
        self.assertEqual(ws["G12"].value.date(), self.novas[0]["Data do Ajuizamento"])
        self.assertEqual(ws["W12"].value[:3], "=IF")                      # fórmula propagada
        self.assertEqual(ws["Y12"].value, '=IF(X12="","",X12-G12)')       # A1 deslocada
        self.assertEqual(ws["N211"].value, '=IF(OR(M211="Encerrado",M211="Arquivado"),"Não","Sim")')
        self.assertEqual(wb.sheetnames, ["Processos", "Parâmetros", "Indicadores", "Dashboard", "Dinâmica"])

    def test_d_estruturas_acompanham(self):
        d = partes(self.dest)
        s = d["xl/worksheets/sheet1.xml"].decode()
        self.assertIn('<dimension ref="A1:AC211"/>', s)
        sq = re.findall(r'sqref="([^"]*)"', s)
        for esperado in ("A2:AC211", "M2:N211", "K2:K211", "O2:O211 Q2:Q211", "P2:P211", "M2:M211", "AA2:AA211"):
            self.assertIn(esperado, sq)
        wb = d["xl/workbook.xml"].decode()
        self.assertIn("Processos!$A$1:$AC$211", wb)
        self.assertIn("Processos!$K$2:$K$211", wb)
        self.assertIn("'Parâmetros'!$B$3", wb)                       # nome que não aponta para a tabela: igual
        self.assertIn('fullCalcOnLoad="1"', wb)
        self.assertNotIn("calcChain", d["xl/_rels/workbook.xml.rels"].decode())
        self.assertNotIn("calcChain", d["[Content_Types].xml"].decode())
        self.assertIn("$K$2:$K$211", d["xl/charts/chart2.xml"].decode())
        self.assertIn("$B$12:$B$18", d["xl/charts/chart1.xml"].decode())   # gráfico da outra aba: intacto
        ind = d["xl/worksheets/sheet3.xml"].decode()
        self.assertIn("Processos!$K$2:$K$211,Processos!$P$2:$P$211", ind)
        self.assertIn("COUNTIFS(Processos!$M$2:$M$211,A12)", ind)
        self.assertIn("tblProcessos[Valor da Causa]", ind)           # referência estruturada: não se mexe

    def test_e_xml_coerente(self):
        self.assertEqual(x.validar(self.dest), [])
        d = partes(self.dest)
        for nome in ("xl/worksheets/sheet1.xml", "xl/worksheets/sheet3.xml", "xl/tables/table1.xml",
                     "xl/workbook.xml", "xl/charts/chart2.xml", "[Content_Types].xml", "xl/_rels/workbook.xml.rels"):
            ET.fromstring(d[nome])                                    # bem-formado
        raiz = ET.fromstring(d["xl/worksheets/sheet1.xml"])
        linhas = raiz.find(f"{M}sheetData")
        numeros = [int(r.get("r")) for r in linhas]
        self.assertEqual(numeros, list(range(1, 212)))
        refs = [c.get("r") for r in linhas for c in r]
        self.assertEqual(len(refs), len(set(refs)))                  # refs de célula únicas
        tab = ET.fromstring(d["xl/tables/table1.xml"])
        self.assertEqual(tab.get("ref"), f"A1:AC{len(numeros)}")
        self.assertEqual(tab.find(f"{M}autoFilter").get("ref"), tab.get("ref"))
        ids = [c.get("id") for c in tab.find(f"{M}tableColumns")]
        self.assertEqual(len(ids), len(set(ids)))
        self.assertEqual(len(ids), 29)

    @unittest.skipUnless(TEM_LO, "LibreOffice ausente")
    def test_f_libreoffice_pdf_e_csv(self):
        pasta = self.lo(self.dest, "pdf")
        pdf = (pasta / "200.pdf").read_bytes()
        self.assertTrue(pdf.startswith(b"%PDF") and len(pdf) > 5000)
        pasta = self.lo(self.dest, "csv", "csv:Text - txt - csv (StarCalc):44,34,76,1,,0,false,true,false,false,false,-1")
        arquivos = {p.name.split("-", 1)[1]: p for p in pasta.glob("*.csv")}
        self.assertEqual(len(arquivos), 5)                            # uma por aba
        linhas = list(csv.reader(io.StringIO(next(p for n, p in arquivos.items() if n.startswith("Processos")
                                                  ).read_text("utf-8"))))
        self.assertEqual(len(linhas), 211)
        self.assertEqual(linhas[-1][0], g.cnj_ficticio(210))

    @unittest.skipUnless(TEM_LO, "LibreOffice ausente")
    def test_g_libreoffice_recalcula_formulas(self):
        wb = self.lo_valores(self.dest)
        todas = self.linhas_novas(0, 210)
        exp = g.indicadores_esperados(todas)
        ind = wb["Indicadores"]
        v = {r: ind.cell(r, 2).value for r in range(2, 19)}
        self.assertEqual(v[2], exp["total"])
        self.assertEqual(v[3], exp["ativos"])
        self.assertAlmostEqual(v[4], exp["valor_causa_ativos"], 2)
        self.assertAlmostEqual(v[5], exp["prazo_medio"], 6)
        self.assertAlmostEqual(v[6], exp["economia"], 2)
        self.assertAlmostEqual(v[7], exp["causa_provavel"], 2)
        self.assertAlmostEqual(v[8], exp["total"] / 350 * 1000, 6)
        self.assertEqual([v[12 + i] for i in range(7)], exp["por_situacao"])
        p = wb["Processos"]
        for i in (0, 9, 10, 11, 150, 209):                            # antigas e novas
            d = g.derivados(todas[i])
            r = 2 + i
            self.assertEqual(p[f"N{r}"].value, d["Ativo"])
            for col, nome in (("W", "Valor Economizado"), ("Y", "Taxa de resolução (em dias)")):
                esperado = d[nome]
                atual = p[f"{col}{r}"].value
                if esperado == "":
                    self.assertIn(atual, (None, ""))
                else:
                    self.assertAlmostEqual(float(atual), float(esperado), 4)

    @unittest.skipUnless(TEM_LO, "LibreOffice ausente")
    def test_h_libreoffice_tabela_dinamica_acompanha(self):
        pasta = self.lo(self.dest, "csv", "csv:Text - txt - csv (StarCalc):44,34,76,1,,0,false,true,false,false,false,-1")
        din = next(p for p in pasta.glob("*.csv") if "-Din" in p.name).read_text("utf-8")
        self.assertIn(",210,", din)                                   # a dinâmica reflete as 210 linhas
        pasta = self.lo(self.dest, "ods")
        with zipfile.ZipFile(pasta / "200.ods") as z:
            self.assertIn("data-pilot-table", z.read("content.xml").decode())


    @unittest.skipUnless(TEM_LO, "LibreOffice ausente")
    def test_i_fullcalconload_sozinho_nao_basta_no_libreoffice(self):
        """Achado do spike: o LibreOffice ignora fullCalcOnLoad e confia no cache das fórmulas. Por isso o
        protótipo também apaga o cache (invalidar_cache=True, padrão). Este teste documenta o comportamento:
        se algum dia o LibreOffice passar a honrar fullCalcOnLoad, ele falha e a decisão pode ser revista."""
        dest = self.tmp / "sem-invalidar.xlsx"
        ed = x.Edicao(self.completo, "Processos", invalidar_cache=False)
        ed.inserir_linhas(self.novas)
        ed.salvar(dest)
        self.assertIn('fullCalcOnLoad="1"', partes(dest)["xl/workbook.xml"].decode())
        wb = self.lo_valores(dest)
        self.assertEqual(wb["Indicadores"]["B2"].value, 10)           # velho: 10 em vez de 210
        self.assertEqual(wb["Processos"]["W211"].value, 524604.5)     # as fórmulas das linhas novas calculam


class TestSujo(Base):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.dest = cls.tmp / "sujo200.xlsx"
        cls.res = x.inserir_linhas(cls.sujo, cls.dest, "Processos", cls.linhas_novas(10, 110))

    def test_sujo_tem_as_sujeiras(self):
        p = partes(self.sujo)
        sheet = p["xl/worksheets/sheet1.xml"].decode()
        self.assertIn('t="shared"', sheet)
        self.assertNotIn("tableStyleInfo", p["xl/tables/table1.xml"].decode())
        self.assertGreater(p["xl/sharedStrings.xml"].count(b"<si>"), 6000)
        self.assertIn('<row r="30"', sheet)                           # linhas vazias formatadas no fim

    def test_insercao_e_validacao(self):
        self.assertEqual(self.res.linhas_inseridas, 100)
        self.assertEqual(x.validar(self.dest), [])
        antes, depois = partes(self.sujo), partes(self.dest)
        self.assertEqual(antes["xl/sharedStrings.xml"], depois["xl/sharedStrings.xml"])   # 6.000 textos: nem tocados
        self.assertEqual(antes["xl/worksheets/sheet6.xml"], depois["xl/worksheets/sheet6.xml"])
        self.assertNotIn("tableStyleInfo", depois["xl/tables/table1.xml"].decode())
        self.assertIn('ref="A1:AC111"', depois["xl/tables/table1.xml"].decode())

    def test_linhas_vazias_do_fim_sao_reaproveitadas(self):
        s = partes(self.dest)["xl/worksheets/sheet1.xml"].decode()
        self.assertEqual(s.count('<row r="12"'), 1)
        self.assertEqual(s.count('<row r="30"'), 1)
        numeros = [int(n) for n in re.findall(r'<row r="(\d+)"', s)]
        self.assertEqual(numeros, sorted(set(numeros)))
        self.assertEqual(max(numeros), 111)

    def test_celulas_sem_estilo_e_formula_compartilhada(self):
        wb = openpyxl.load_workbook(self.dest)
        ws = wb["Processos"]
        self.assertEqual(ws["Y12"].value, '=IF(X12="","",X12-G12)')   # veio da fórmula compartilhada (mestre Y2)
        self.assertEqual(ws["Y111"].value, '=IF(X111="","",X111-G111)')
        s = partes(self.dest)["xl/worksheets/sheet1.xml"].decode()
        # linha 12 já existia (vazia e formatada): célula mantém o estilo que tinha; linha 41 é nova e a
        # linha-modelo (última de dados) tem a coluna AC sem estilo: a célula nova também fica sem
        self.assertIn(' s="', re.search(r'<c r="AC12"[^>]*?/?>', s).group(0))
        m41 = re.search(r'<c r="AC41"[^>]*?/?>', s)
        self.assertTrue(m41 is None or ' s="' not in m41.group(0))
        self.assertIn(' s="', re.search(r'<c r="AB41"[^>]*?>', s).group(0))   # AB tem estilo na linha-modelo
        self.assertIn(' s="', re.search(r'<c r="A12"[^>]*?>', s).group(0))

    def test_cf_e_dv_do_sujo(self):
        s = partes(self.dest)["xl/worksheets/sheet1.xml"].decode()
        self.assertIn('sqref="P2:P111"', s)
        self.assertIn('sqref="O2:O111 Q2:Q111"', s)

    @unittest.skipUnless(TEM_LO, "LibreOffice ausente")
    def test_libreoffice(self):
        wb = self.lo_valores(self.dest)
        exp = g.indicadores_esperados(self.linhas_novas(0, 110))
        ind = wb["Indicadores"]
        self.assertEqual(ind["B2"].value, exp["total"])
        self.assertEqual(ind["B3"].value, exp["ativos"])
        self.assertAlmostEqual(ind["B6"].value, exp["economia"], 2)
        self.assertEqual(wb["Processos"]["Y2"].value in (None, ""), True)   # fórmula compartilhada recalcula


@unittest.skipUnless(TEM_LO, "LibreOffice ausente")
class TestArquivoRessalvoPeloLibreOffice(Base):
    """Fixture 'estrangeira': o modelo re-salvo pelo LibreOffice tem outra cara (atributos aca/customFormat,
    aspas como &quot;, tabela dinâmica por intervalo COM cache de registros, extLst x14 da barra de dados,
    sem calculatedColumnFormula). Serve de segunda opinião sobre o que o protótipo assume do XML."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        pasta = cls.tmp / "estrangeiro"
        pasta.mkdir()
        p = subprocess.run([SOFFICE, f"-env:UserInstallation=file://{cls.perfil_lo}", "--headless",
                            "--convert-to", "xlsx", "--outdir", str(pasta), str(cls.completo)],
                           capture_output=True, text=True, timeout=240)
        assert p.returncode == 0, p.stderr
        cls.estr = pasta / "completo.xlsx"

    def test_insercao_em_arquivo_estrangeiro(self):
        self.assertEqual(x.validar(self.estr), [])
        self.assertIn("xl/pivotCache/pivotCacheRecords1.xml", partes(self.estr))
        dest = self.tmp / "estr100.xlsx"
        r = x.inserir_linhas(self.estr, dest, "Processos", self.linhas_novas(10, 110))
        self.assertEqual(r.linhas_inseridas, 100)
        self.assertEqual(x.validar(dest), [])
        d = partes(dest)
        self.assertIn("<xm:sqref>O2:O111 Q2:Q111</xm:sqref>", d["xl/worksheets/sheet1.xml"].decode())   # x14
        cache = d["xl/pivotCache/pivotCacheDefinition1.xml"].decode()
        self.assertIn('ref="A1:AC111"', cache)
        self.assertIn('refreshOnLoad="1"', cache)
        antes = partes(self.estr)
        self.assertEqual(antes["xl/pivotCache/pivotCacheRecords1.xml"], d["xl/pivotCache/pivotCacheRecords1.xml"])
        wb = self.lo_valores(dest)
        exp = g.indicadores_esperados(self.linhas_novas(0, 110))
        self.assertEqual(wb["Indicadores"]["B2"].value, exp["total"])
        self.assertAlmostEqual(wb["Indicadores"]["B6"].value, exp["economia"], 2)
        self.assertEqual([wb["Indicadores"].cell(12 + i, 2).value for i in range(7)], exp["por_situacao"])

    def test_coluna_nova_com_cache_de_dinamica_salvo_e_recusada(self):
        with self.assertRaises(x.NaoSuportado):
            x.acrescentar_coluna(self.estr, self.tmp / "estr_col.xlsx", "Processos", "Nova")


class TestEscreverCelulas(Base):
    def test_humanos_nunca_sobrescritos(self):
        dest = self.saida("humanos.xlsx")
        humanos = ["Probabilidade", "Valor Estimado", "Garantias Processuais"]
        r = x.escrever_celulas(self.completo, dest, "Processos", {
            (2, "Probabilidade"): "Provável",        # humano com valor: ignorado
            (2, "Garantias Processuais"): 1500.0,    # humano vazio: grava
            (2, "Objeto"): "Objeto novo",            # não humano: grava por cima
            (2, "Ativo"): "Sim",                     # fórmula: ignorada
            (3, "Valor Estimado"): 1.0,              # humano com valor: ignorado
            (3, "AA"): 0.55,                         # coluna por letra
        }, humanos=humanos)
        self.assertEqual(r.celulas_escritas, 3)
        motivos = dict(r.ignoradas)
        self.assertEqual(motivos["P2"], "campo humano já preenchido")
        self.assertEqual(motivos["Q3"], "campo humano já preenchido")
        self.assertEqual(motivos["N2"], "tem fórmula")
        self.assertEqual(x.validar(dest), [])
        ws = openpyxl.load_workbook(dest)["Processos"]
        orig = openpyxl.load_workbook(self.completo)["Processos"]
        self.assertEqual(ws["P2"].value, orig["P2"].value)
        self.assertEqual(ws["Q3"].value, orig["Q3"].value)
        self.assertEqual(ws["U2"].value, 1500.0)
        self.assertEqual(ws["U2"].number_format, "#,##0.00")          # estilo da célula preservado
        self.assertEqual(ws["J2"].value, "Objeto novo")
        self.assertEqual(ws["AA3"].value, 0.55)
        self.assertEqual(ws["AA3"].number_format, orig["AA3"].number_format)
        self.assertEqual(ws["N2"].value, orig["N2"].value)
        # partes não editadas: iguais
        a, d = partes(self.completo), partes(dest)
        for n in a:
            if n not in ("xl/worksheets/sheet1.xml", "xl/worksheets/sheet3.xml", "xl/workbook.xml",
                         "xl/_rels/workbook.xml.rels", "[Content_Types].xml", "xl/calcChain.xml"):
                self.assertEqual(a[n], d[n], n)

    def test_sem_preservar_sobrescreve(self):
        dest = self.saida("humanos2.xlsx")
        r = x.escrever_celulas(self.completo, dest, "Processos", {(2, "Probabilidade"): "Remota"},
                               humanos=["Probabilidade"], preservar_humanos=False)
        self.assertEqual(r.ignoradas, [])
        self.assertEqual(openpyxl.load_workbook(dest)["Processos"]["P2"].value, "Remota")

    def test_limpar_e_data(self):
        dest = self.saida("limpar.xlsx")
        r = x.escrever_celulas(self.completo, dest, "Processos", {
            (2, "Objeto"): x.LIMPAR, (2, "Data do Ajuizamento"): datetime.date(2020, 2, 29)})
        ws = openpyxl.load_workbook(dest)["Processos"]
        self.assertIsNone(ws["J2"].value)
        self.assertEqual(ws["G2"].value.date() if hasattr(ws["G2"].value, "date") else ws["G2"].value,
                         datetime.date(2020, 2, 29))
        self.assertEqual(r.avisos, [])

    def test_data_em_celula_sem_formato_de_data_avisa(self):
        r = x.escrever_celulas(self.completo, self.saida("d.xlsx"), "Processos",
                               {(2, "Município"): datetime.date(2020, 1, 1)})
        self.assertTrue(any("sem formato de data" in a for a in r.avisos))

    def test_texto_com_caracteres_especiais(self):
        dest = self.saida("esp.xlsx")
        texto = "linha1\nlinha2 & <b>\"aspas\"</b> \x01ilegal  espaços  "
        x.escrever_celulas(self.completo, dest, "Processos", {(4, "Andamentos"): texto})
        v = openpyxl.load_workbook(dest)["Processos"]["L4"].value
        self.assertEqual(v, texto.replace("\x01", ""))
        self.assertEqual(x.validar(dest), [])


class TestAcrescentarColuna(Base):
    def test_coluna_calculada_estruturada(self):
        dest = self.saida("col.xlsx")
        f = f'IF({g.TABELA}[[#This Row],[Valor da Causa]]>20000,"Alto","Baixo")'
        r = x.acrescentar_coluna(self.completo, dest, "Processos", "Faixa de valor", formula=f)
        self.assertEqual(x.validar(dest), [])
        d = partes(dest)
        tab = d["xl/tables/table1.xml"].decode()
        self.assertIn('ref="A1:AD11"', tab)
        self.assertIn('<tableColumns count="30">', tab)
        self.assertIn('name="Faixa de valor"><calculatedColumnFormula>', tab)
        self.assertIn("A1:AD11", re.search(r'<autoFilter[^>]*>', tab).group(0))
        self.assertIn("Processos!$A$1:$AD$11", d["xl/workbook.xml"].decode())          # nome da tabela inteira
        self.assertIn("Processos!$K$2:$K$11", d["xl/workbook.xml"].decode())           # coluna só: igual
        self.assertIn('<dimension ref="A1:AD11"/>', d["xl/worksheets/sheet1.xml"].decode())
        self.assertIn('count="30"', d["xl/pivotCache/pivotCacheDefinition1.xml"].decode())   # cache + dinâmica
        self.assertIn('<pivotFields count="30">', d["xl/pivotTables/pivotTable1.xml"].decode())
        wb = openpyxl.load_workbook(dest)
        ws = wb["Processos"]
        self.assertEqual(ws["AD1"].value, "Faixa de valor")
        self.assertEqual(ws["AD1"].font.b, ws["AC1"].font.b)           # estilo do cabeçalho copiado
        self.assertEqual(ws["AD11"].value[:3], "=IF")
        self.assertEqual(ws.tables["tblProcessos"].ref, "A1:AD11")
        self.assertEqual(r.avisos, [])

    def test_valores_e_a1(self):
        dest = self.saida("col2.xlsx")
        r = x.acrescentar_coluna(self.completo, dest, "Processos", "Observação interna",
                                 valores=[f"obs {i}" for i in range(10)])
        r2 = x.acrescentar_coluna(self.completo, self.saida("col3.xlsx"), "Processos", "Dobro",
                                  formula="K2*2")
        self.assertTrue(any("referência A1" in a for a in r2.avisos))
        self.assertEqual(x.validar(self.saida("col3.xlsx")), [])
        ws = openpyxl.load_workbook(dest)["Processos"]
        self.assertEqual([ws.cell(r, 30).value for r in (2, 11)], ["obs 0", "obs 9"])
        ws3 = openpyxl.load_workbook(self.saida("col3.xlsx"))["Processos"]
        self.assertEqual(ws3["AD7"].value, "=K7*2")                    # A1 relativa desloca por linha

    def test_idempotente(self):
        d1, d2 = self.saida("i1.xlsx"), self.saida("i2.xlsx")
        x.acrescentar_coluna(self.completo, d1, "Processos", "Observação interna", valores=["a"] * 10)
        r = x.acrescentar_coluna(d1, d2, "Processos", "Observação interna", valores=["b"] * 10)
        ws = openpyxl.load_workbook(d2)["Processos"]
        self.assertEqual(ws.max_column, 30)
        self.assertEqual(ws["AD2"].value, "a")                         # já tinha valor: não sobrescreve
        self.assertTrue(any("já existe" in a for a in r.avisos))

    @unittest.skipUnless(TEM_LO, "LibreOffice ausente")
    def test_libreoffice_coluna_calculada_e_linhas_depois(self):
        d1, d2 = self.saida("c1.xlsx"), self.saida("c2.xlsx")
        f = f'IF({g.TABELA}[[#This Row],[Valor da Causa]]>20000,"Alto","Baixo")'
        x.acrescentar_coluna(self.completo, d1, "Processos", "Faixa de valor", formula=f)
        # depois, 50 linhas novas: a coluna calculada nova deve propagar (vem da Tabela)
        x.inserir_linhas(d1, d2, "Processos", self.linhas_novas(10, 60))
        self.assertEqual(x.validar(d2), [])
        wb = self.lo_valores(d2)
        ws = wb["Processos"]
        todas = self.linhas_novas(0, 60)
        for i in (0, 5, 10, 59):
            esperado = "Alto" if todas[i]["Valor da Causa"] > 20000 else "Baixo"
            self.assertEqual(ws.cell(2 + i, 30).value, esperado)
        self.assertEqual(ws.max_row, 61)

    def test_dinamica_com_cache_salvo_nao_suportada(self):
        com_registros = self.saida("reg.xlsx")
        p = partes(self.completo)
        p["xl/pivotCache/pivotCacheDefinition1.xml"] = p["xl/pivotCache/pivotCacheDefinition1.xml"].replace(
            b"<cacheSource", b'<cacheSource', 1)
        rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.'
                'openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.'
                'openxmlformats.org/officeDocument/2006/relationships/pivotCacheRecords" '
                'Target="pivotCacheRecords1.xml"/></Relationships>').encode()
        p["xl/pivotCache/_rels/pivotCacheDefinition1.xml.rels"] = rels
        p["xl/pivotCache/pivotCacheRecords1.xml"] = b'<pivotCacheRecords xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="0"/>'
        with zipfile.ZipFile(com_registros, "w") as z:
            for n, d in p.items():
                z.writestr(n, d)
        with self.assertRaises(x.NaoSuportado):
            x.acrescentar_coluna(com_registros, self.saida("reg2.xlsx"), "Processos", "Nova")
        self.assertFalse(self.saida("reg2.xlsx").exists())


class TestSegurancaEErros(Base):
    def test_coluna_desconhecida_nao_descarta_dado(self):
        dest = self.saida("e1.xlsx")
        with self.assertRaises(x.ColunaDesconhecida):
            x.inserir_linhas(self.completo, dest, "Processos", [{"Número do Processo": "x", "Coluna Fantasma": 1}])
        self.assertFalse(dest.exists())

    def test_mapeamento_de_colunas_e_nomes_tolerantes(self):
        dest = self.saida("e2.xlsx")
        x.inserir_linhas(self.completo, dest, "Processos",
                         [{"numero": "1234567-00.2020.8.06.0001", "reus": "Ré Teste", "recurso": "Sim"}],
                         colunas={"numero": "Número do Processo", "reus": "reu(s)",
                                  "recurso": "Houve recurso da empresa"})
        ws = openpyxl.load_workbook(dest)["Processos"]
        self.assertEqual((ws["A12"].value, ws["C12"].value, ws["Z12"].value),
                         ("1234567-00.2020.8.06.0001", "Ré Teste", "Sim"))

    def test_conteudo_abaixo_da_tabela(self):
        com = self.saida("abaixo.xlsx")
        x.escrever_celulas(self.completo, com, "Processos", {(13, "A"): "nota do usuário"})
        with self.assertRaises(x.ConflitoDeConteudo):
            x.inserir_linhas(com, self.saida("abaixo2.xlsx"), "Processos", self.linhas_novas(10, 20))

    def test_linha_de_totais_nao_suportada(self):
        p = partes(self.completo)
        p["xl/tables/table1.xml"] = p["xl/tables/table1.xml"].replace(b"<table ", b'<table totalsRowCount="1" ', 1)
        arq = self.saida("tot.xlsx")
        with zipfile.ZipFile(arq, "w") as z:
            for n, d in p.items():
                z.writestr(n, d)
        with self.assertRaises(x.NaoSuportado):
            x.inserir_linhas(arq, self.saida("tot2.xlsx"), "Processos", self.linhas_novas(10, 11))

    def test_original_nunca_alterado(self):
        antes = sha(self.completo.read_bytes())
        x.inserir_linhas(self.completo, self.saida("o.xlsx"), "Processos", self.linhas_novas(10, 15))
        self.assertEqual(sha(self.completo.read_bytes()), antes)
        with self.assertRaises(x.ErroXlsx):
            x.inserir_linhas(self.completo, self.completo, "Processos", self.linhas_novas(10, 11))
        self.assertEqual(sha(self.completo.read_bytes()), antes)

    def test_texto_acima_do_limite_do_excel(self):
        with self.assertRaises(x.ErroXlsx):
            x.escrever_celulas(self.completo, self.saida("big.xlsx"), "Processos", {(2, "Andamentos"): "a" * 32768})
        x.escrever_celulas(self.completo, self.saida("big.xlsx"), "Processos", {(2, "Andamentos"): "a" * 32767})

    def test_aba_inexistente(self):
        with self.assertRaises(x.ErroXlsx):
            x.inserir_linhas(self.completo, self.saida("n.xlsx"), "Nao existe", [{"A": 1}])

    def test_arquivo_que_nao_e_xlsx(self):
        ruim = self.saida("ruim.xlsx")
        ruim.write_bytes(b"isto nao e um zip")
        with self.assertRaises(x.ErroXlsx):
            x.inserir_linhas(ruim, self.saida("r.xlsx"), "Processos", [{"A": 1}])

    def test_saida_deterministica(self):
        a, b = self.saida("det1.xlsx"), self.saida("det2.xlsx")
        x.inserir_linhas(self.completo, a, "Processos", self.linhas_novas(10, 60))
        x.inserir_linhas(self.completo, b, "Processos", self.linhas_novas(10, 60))
        self.assertEqual(a.read_bytes(), b.read_bytes())

    def test_gerador_deterministico(self):
        a, b = self.saida("g1.xlsx"), self.saida("g2.xlsx")
        g.gerar(a)
        g.gerar(b)
        self.assertEqual(a.read_bytes(), b.read_bytes())
        self.assertEqual(a.read_bytes(), self.completo.read_bytes())

    def test_composicao_100_mais_100_igual_a_200(self):
        d1, d2, d3 = self.saida("k1.xlsx"), self.saida("k2.xlsx"), self.saida("k3.xlsx")
        x.inserir_linhas(self.completo, d1, "Processos", self.linhas_novas(10, 110))
        x.inserir_linhas(d1, d2, "Processos", self.linhas_novas(110, 210))
        x.inserir_linhas(self.completo, d3, "Processos", self.linhas_novas(10, 210))
        a, b = partes(d2), partes(d3)
        self.assertEqual(set(a), set(b))
        for n in a:
            self.assertEqual(a[n], b[n], n)

    def test_dinamica_por_intervalo_e_refresh(self):
        p = partes(self.completo)
        cache = p["xl/pivotCache/pivotCacheDefinition1.xml"].decode()
        cache = cache.replace('<worksheetSource name="tblProcessos"/>',
                              '<worksheetSource ref="A1:AC11" sheet="Processos"/>').replace(' refreshOnLoad="1"', "")
        self.assertNotIn("refreshOnLoad", cache)
        p["xl/pivotCache/pivotCacheDefinition1.xml"] = cache.encode()
        arq = self.saida("pivref.xlsx")
        with zipfile.ZipFile(arq, "w") as z:
            for n, d in p.items():
                z.writestr(n, d)
        dest = self.saida("pivref2.xlsx")
        x.inserir_linhas(arq, dest, "Processos", self.linhas_novas(10, 30))
        novo = partes(dest)["xl/pivotCache/pivotCacheDefinition1.xml"].decode()
        self.assertIn('ref="A1:AC31"', novo)
        self.assertIn('refreshOnLoad="1"', novo)
        self.assertEqual(x.validar(dest), [])

    def test_validar_pega_problemas(self):
        p = partes(self.completo)
        p["xl/tables/table1.xml"] = p["xl/tables/table1.xml"].replace(b'ref="A1:AC11"', b'ref="A1:AC50"', 1)
        arq = self.saida("ruim_tab.xlsx")
        with zipfile.ZipFile(arq, "w") as z:
            for n, d in p.items():
                z.writestr(n, d)
        probs = x.validar(arq)
        self.assertTrue(any("autoFilter" in s for s in probs), probs)


class TestAlternativaPreformatada(Base):
    """Alternativa 'template com linhas pré-formatadas': a tabela já cobre 300 linhas e só se preenchem células."""

    def test_preencher_sem_mudar_estrutura(self):
        dest = self.saida("pre_cheio.xlsx")
        novas = self.linhas_novas(10, 210)
        celulas = {}
        for k, l in enumerate(novas):
            for cab, v in l.items():
                if cab not in g.FORMULAS and v is not None:
                    celulas[(12 + k, cab)] = v
        r = x.escrever_celulas(self.pre, dest, "Processos", celulas)
        self.assertEqual(x.validar(dest), [])
        a, d = partes(self.pre), partes(dest)
        alteradas = {n for n in d if a[n] != d[n]}
        # estrutura intacta: tabela, nomes, gráficos, CF/DV não mudam; só as abas com cache de fórmula
        self.assertEqual(alteradas, {"xl/worksheets/sheet1.xml", "xl/worksheets/sheet3.xml", "xl/workbook.xml",
                                     "xl/_rels/workbook.xml.rels", "[Content_Types].xml"})
        self.assertEqual(r.celulas_escritas, len(celulas))

    @unittest.skipUnless(TEM_LO, "LibreOffice ausente")
    def test_efeito_colateral_das_linhas_em_branco(self):
        """O custo da alternativa: as fórmulas das linhas vazias contam como processo 'Sim' (ativo)."""
        dest = self.saida("pre_cheio2.xlsx")
        celulas = {(12 + k, cab): v for k, l in enumerate(self.linhas_novas(10, 60))
                   for cab, v in l.items() if cab not in g.FORMULAS and v is not None}
        x.escrever_celulas(self.pre, dest, "Processos", celulas)
        wb = self.lo_valores(dest)
        real = g.indicadores_esperados(self.linhas_novas(0, 60))
        ind = wb["Indicadores"]
        self.assertEqual(ind["B2"].value, real["total"])               # COUNTA da coluna de número: certo
        self.assertEqual(ind["B3"].value, real["ativos"] + 240)       # 240 linhas em branco viram "ativo": errado
        self.assertGreater(ind["B3"].value, real["ativos"])

# =====================================================================================================
# 2. Escritor do contrato: modelo padrão, gravar, planilha.py
# =====================================================================================================

MODELO = x.MODELO_PADRAO
GERADOR = RAIZ / "src" / "modelos" / "xlsx_b" / "gerar_modelo.py"
DATA_BASE = "2026-10-07"


def evento(numero, data_br, frase, conteudo="", **extra):
    """Evento aprovado fictício no formato de comum.py."""
    ev = {"numero": numero, "status": "aprovado", "data": data_br, "frase": frase, "conteudo": conteudo,
          "detectado_em": f"{data_br[6:]}-{data_br[3:5]}-{data_br[:2]}T10:00:00", "grau": "1º grau"}
    ev.update(extra)
    return ev


def estado_de(fichas, data_base=DATA_BASE, eventos=(), perfil=None, parametros=None, **extra):
    e = {"cliente": "Cliente Exemplo", "data_base": data_base, "fichas": list(fichas), "eventos": list(eventos),
         "perfil": perfil or {}, "parametros": parametros or {}}
    e.update(extra)
    return e


def enriquecer(fichas, semente=7):
    """Completa campos de julgamento em parte das fichas, para os indicadores terem o que contar."""
    import random
    for i, f in enumerate(fichas):
        rng = random.Random(f"{semente}:{i}")
        if not f["ativo"]:
            if rng.random() < 0.7:
                fi.definir(f, "valor_arbitrado", f"{rng.randint(1000, 90000)}.{rng.randint(0, 99):02d}", "humano")
            if rng.random() < 0.4:
                fi.definir(f, "valor_execucao", f"{rng.randint(1000, 90000)}.{rng.randint(0, 99):02d}", "humano")
            if rng.random() < 0.7:
                fi.definir(f, "data_transito", fi.obter(f, "ultimo_andamento"), "humano")
        if rng.random() < 0.5:
            fi.definir(f, "valor_estimado", f"{rng.randint(500, 80000)}.00", "humano")
        if rng.random() < 0.3:
            fi.definir(f, "houve_recurso", rng.choice(["Sim", "Não"]), "derivado")
    return fichas


def _edate_menos_12(d):
    try:
        return d.replace(year=d.year - 1)
    except ValueError:               # 29/02
        return d.replace(year=d.year - 1, day=28)


def esperado(fichas, data_ref, headcount, fator=1):
    """Os indicadores do modelo padrão, calculados em Python direto das fichas (independente da planilha)."""
    D = fi.dinheiro
    ativos = [f for f in fichas if f["ativo"]]
    enc = [f for f in fichas if not f["ativo"]]
    cont = lambda r: sum(1 for f in fichas if fi.obter(f, "resultado") == r)   # noqa: E731
    improc, arq, acordo = cont("Improcedente"), cont("Arquivado / desistência"), cont("Acordo")
    parcial, proc = cont("Parcialmente procedente"), cont("Procedente")
    decididos = improc + arq + acordo + parcial + proc
    dias = []
    for f in fichas:
        tr, aj = fi.data(fi.obter(f, "data_transito")), fi.data(fi.obter(f, "data_ajuizamento"))
        if tr and aj:
            dias.append((tr - aj).days)
    dias.sort()
    mediana = (dias[len(dias) // 2] if len(dias) % 2 else (dias[len(dias) // 2 - 1] + dias[len(dias) // 2]) / 2) if dias else ""
    ref = fi.data(data_ref)
    novos = sum(1 for f in fichas if fi.data(fi.obter(f, "data_ajuizamento"))
                and _edate_menos_12(ref) < fi.data(fi.obter(f, "data_ajuizamento")) <= ref)
    recursos = sum(1 for f in fichas if fi.obter(f, "houve_recurso") == "Sim")
    arb = lambda f: D(fi.obter(f, "valor_arbitrado"))   # noqa: E731
    exe = lambda f: D(fi.obter(f, "valor_execucao"))     # noqa: E731
    causa = lambda f: D(fi.obter(f, "valor_causa")) or Decimal(0)   # noqa: E731
    lancado = [f for f in enc if arb(f) is not None or exe(f) is not None]
    realizado = sum((arb(f) if arb(f) is not None else exe(f) for f in lancado), Decimal(0))
    causa_lancado = sum((causa(f) for f in lancado), Decimal(0))
    total_causa = sum((causa(f) for f in fichas), Decimal(0))
    return {
        "Total de processos": len(fichas), "Processos ativos": len(ativos), "Processos encerrados": len(enc),
        "Litígios por 100 funcionários": len(fichas) / headcount * 100 if headcount else "",
        "Novos ajuizamentos em 12 meses": novos,
        "Tempo médio de resolução (dias)": sum(dias) / len(dias) if dias else "",
        "Mediana do tempo de resolução (dias)": mediana,
        "Processos com recurso da empresa": recursos, "Processos decididos": decididos,
        "Taxa de recurso": recursos / decididos if decididos else "",
        "Improcedentes": improc, "Arquivadas / desistência": arq, "Acordos": acordo,
        "Parcialmente procedentes": parcial, "Procedentes": proc,
        "Taxa de êxito": (improc + arq) / decididos if decididos else "",
        "% por acordo": acordo / decididos if decididos else "",
        "Valor da causa total": total_causa, "Valor da causa total corrigido": total_causa * Decimal(str(fator)),
        "Valor da causa dos encerrados": sum((causa(f) for f in enc), Decimal(0)),
        "Encerrados com valor realizado lançado": len(lancado),
        "Encerrados sem valor lançado (fora da economia)": len(enc) - len(lancado),
        "Valor da causa dos encerrados com valor lançado": causa_lancado,
        "Valor realizado dos encerrados": realizado, "Economia efetiva": causa_lancado - realizado,
    }


def indicadores_da_planilha(wb):
    """{rótulo: valor} da aba Indicadores (só as linhas de indicador, que têm valor na coluna B)."""
    ws = wb["Indicadores"]
    saida = {}
    for r in range(2, ws.max_row + 1):
        rotulo = ws.cell(r, 1).value
        if rotulo and ws.cell(r, 3).value:           # a coluna C só tem nota nos indicadores
            saida[rotulo] = ws.cell(r, 2).value
    return saida


def linhas_da_aba(caminho, aba="Processos"):
    """[{cabeçalho: valor}] da aba (fórmulas aparecem como texto '=...'; sem cache)."""
    ws = openpyxl.load_workbook(caminho)[aba]
    cab = [c.value for c in ws[1]]
    return [dict(zip(cab, [c.value for c in linha])) for linha in ws.iter_rows(min_row=2)
            if any(c.value is not None for c in linha)]


def _carregar_gerador():
    spec = importlib.util.spec_from_file_location("gerar_modelo_b", GERADOR)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


class TestModeloPadrao(BaseLO):
    def test_estrutura(self):
        self.assertTrue(MODELO.exists(), "gere com src/modelos/xlsx_b/gerar_modelo.py")
        self.assertEqual(x.validar(MODELO), [])
        wb = openpyxl.load_workbook(MODELO)
        self.assertEqual(wb.sheetnames, ["Processos", "Parâmetros", "Indicadores", "Dashboard", "Histórico",
                                         "Campos não migrados"])
        ws = wb["Processos"]
        cab = [c.value for c in ws[1]]
        self.assertEqual(cab[:29], [x.CAMPOS_B[c] for c in x.COLUNAS_PADRAO])
        self.assertEqual(len(cab), 37)
        self.assertEqual(len(x.COLUNAS_PADRAO), 29)
        self.assertIn("tblProcessos", ws.tables)
        self.assertEqual(ws.tables["tblProcessos"].ref, "A1:AK2")
        self.assertEqual(sorted(wb["Histórico"].tables), ["tblHistorico"])
        self.assertEqual(sorted(wb["Campos não migrados"].tables), ["tblNaoMigrados"])
        # colunas extras ocultas; as 29 visíveis
        ocultas = {c for c, d in ws.column_dimensions.items() if d.hidden}
        extras = {x.col_letra(cab.index(x.CAMPOS_B[c]) + 1) for c in x.COLUNAS_EXTRAS}
        self.assertEqual(ocultas, extras)
        # colunas calculadas declaradas na Tabela (propagam para as linhas novas)
        tab = partes(MODELO)["xl/tables/table1.xml"].decode()
        self.assertEqual(tab.count("<calculatedColumnFormula>"), 2)
        self.assertGreaterEqual(len(ws.data_validations.dataValidation), 8)
        self.assertGreaterEqual(len(ws.conditional_formatting), 2)

    def test_formulas_sem_funcoes_que_quebram_no_excel(self):
        p = partes(MODELO)
        todas = "".join(p[n].decode() for n in p if n.startswith("xl/worksheets/sheet"))
        for proibida in ("FILTER(", "UNIQUE(", "XLOOKUP(", "LET(", "SORT(", "LAMBDA(", "_xlfn", "_xlws", "SEQUENCE("):
            self.assertNotIn(proibida, todas)
        for necessaria in ("COUNTIFS(", "SUMIFS(", "AVERAGE(", "MEDIAN(", "COUNTA("):
            self.assertIn(necessaria, todas)

    def test_sanitizado_sem_numero_de_processo_nem_nome(self):
        p = partes(MODELO)
        texto = "".join(d.decode("utf-8", "replace") for n, d in p.items() if n.endswith((".xml", ".rels")))
        self.assertEqual(re.findall(r"\d{7}-\d{2}\.\d{4}\.\d\.\d{2}\.\d{4}", texto), [])
        self.assertNotRegex(texto, r"Ltda|S\.A\.")

    def test_modelo_commitado_igual_ao_gerado(self):
        """Regerar com o script dá a mesma estrutura e o mesmo conteúdo (bytes podem variar com a versão do openpyxl)."""
        novo = self.saida("regerado.xlsx")
        _carregar_gerador().gerar(novo)
        a, b = partes(MODELO), partes(novo)
        self.assertEqual(sorted(a), sorted(b))
        for n in a:
            if n.endswith(".xml") and not n.startswith("docProps"):
                self.assertEqual(ET.canonicalize(a[n].decode()), ET.canonicalize(b[n].decode()), n)

    def test_gerador_deterministico(self):
        mod = _carregar_gerador()
        mod.gerar(self.saida("m1.xlsx"))
        mod.gerar(self.saida("m2.xlsx"))
        self.assertEqual(self.saida("m1.xlsx").read_bytes(), self.saida("m2.xlsx").read_bytes())

    @unittest.skipUnless(TEM_LO, "LibreOffice ausente")
    def test_abre_no_libreoffice_vazio(self):
        wb = self.lo_valores(MODELO)
        ind = indicadores_da_planilha(wb)
        self.assertEqual(ind["Total de processos"], 0)
        self.assertEqual(ind["Valor da causa total"], 0)
        self.assertIn(ind["Taxa de êxito"], (None, ""))           # sem decididos: vazio, não erro


class TestGravarModelo200(BaseLO):
    """gravar(None, ...) com 200 fichas fictícias: criação do zero a partir do modelo padrão."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.fichas = enriquecer(ficticio.gerar_carteira(200, clientes=5, semente=1))
        cls.parametros = {"headcount": 350, "empresas_do_grupo": ["Empresa Exemplo A Ltda", "Empresa Exemplo B Ltda"],
                          "fator_correcao": 1.05}
        cls.estado = estado_de(cls.fichas, parametros=cls.parametros)
        cls.dest = cls.tmp / "do-zero.xlsx"
        cls.res = x.gravar(None, cls.estado, cls.dest)

    def erros(self, res):
        return [a for a in res["avisos"] if a["nivel"] == "erro"]

    def test_a_resultado_do_contrato(self):
        r = self.res
        self.assertEqual(self.erros(r), [])
        self.assertEqual(r["destino"], self.dest)
        for chave in ("destino", "processos_atualizados", "processos_novos", "ignorados", "mudancas", "avisos",
                      "textos_gravados"):
            self.assertIn(chave, r)
        self.assertEqual(r["processos_novos"], [f["numero"] for f in self.fichas])
        self.assertEqual(r["processos_atualizados"], [])
        self.assertEqual(len(r["textos_gravados"]), 200)
        self.assertEqual(r["cache_formulas"], "invalidado")
        self.assertTrue(any(a["codigo"] == "cache_de_formulas_invalidado" for a in r["avisos"]))
        for a in r["avisos"]:
            self.assertEqual(set(a), {"nivel", "codigo", "onde", "mensagem", "candidatos"})

    def test_b_200_linhas_inseridas_e_estrutura_coerente(self):
        self.assertEqual(x.validar(self.dest), [])
        wb = openpyxl.load_workbook(self.dest)
        ws = wb["Processos"]
        self.assertEqual(ws.tables["tblProcessos"].ref, "A1:AK201")
        numeros = [ws.cell(r, 1).value for r in range(2, 202)]
        self.assertEqual(numeros, [f["numero"] if not f["vinculados"] else "; ".join(fi.todos_os_numeros(f))
                                   for f in self.fichas])
        self.assertIsNone(ws.cell(202, 1).value)
        # estilo (formato de data e de dinheiro) herdado na última linha; colunas calculadas propagadas
        cab = {c.value: c.column for c in ws[1]}
        self.assertEqual(ws.cell(201, cab["Data do Ajuizamento"]).number_format, "dd/mm/yyyy")
        self.assertEqual(ws.cell(201, cab["Valor da Causa"]).number_format, "#,##0.00")
        self.assertTrue(str(ws.cell(201, cab["Valor Economizado"]).value).startswith("=IF("))
        self.assertTrue(str(ws.cell(201, cab["Taxa de resolução (em dias)"]).value).startswith("=IF("))
        s = partes(self.dest)["xl/worksheets/sheet1.xml"].decode()
        for esperado_ in ("A2:AK201", "AA2:AA201", "P2:P201"):
            self.assertTrue(any(esperado_ in sq for sq in re.findall(r'sqref="([^"]*)"', s)), esperado_)
        self.assertIn('<dimension ref="A1:AK201"/>', s)

    def test_c_partes_nao_editadas_identicas(self):
        antes, depois = partes(MODELO), partes(self.dest)
        self.assertEqual(set(antes), set(depois))
        alteradas = {n for n in depois if antes[n] != depois[n]}
        for intacta in ("xl/styles.xml", "xl/theme/theme1.xml", "xl/sharedStrings.xml", "xl/workbook.xml"):
            if intacta in antes and intacta != "xl/workbook.xml":
                self.assertNotIn(intacta, alteradas, intacta)
        for n in antes:
            if n.startswith(("xl/charts/", "xl/drawings/")):
                self.assertEqual(antes[n], depois[n], n)
        self.assertEqual(self.res["partes_alteradas"], sorted(alteradas))
        self.assertIn("xl/worksheets/sheet1.xml", alteradas)
        self.assertIn('fullCalcOnLoad="1"', depois["xl/workbook.xml"].decode())

    def test_d_colunas_extras_ocultas_e_perfil_ativa(self):
        ws = openpyxl.load_workbook(self.dest)["Processos"]
        cab = [c.value for c in ws[1]]
        for campo in x.COLUNAS_EXTRAS:
            letra = x.col_letra(cab.index(x.CAMPOS_B[campo]) + 1)
            self.assertTrue(ws.column_dimensions[letra].hidden, campo)
        for campo in x.COLUNAS_PADRAO:
            letra = x.col_letra(cab.index(x.CAMPOS_B[campo]) + 1)
            self.assertFalse(ws.column_dimensions[letra].hidden, campo)
        # perfil: ativa momento atual e esconde garantias; as colunas dos indicadores nunca somem
        perfil = {"colunas_ativas": ["numero", "autores", "reus", "momento_atual", "garantias", "outras_partes"]}
        dest = self.saida("perfil.xlsx")
        r = x.gravar(None, estado_de(self.fichas[:5], perfil=perfil), dest)
        self.assertEqual(self.erros(r), [])
        ws = openpyxl.load_workbook(dest)["Processos"]
        oculto = lambda campo: ws.column_dimensions[x.col_letra(cab.index(x.CAMPOS_B[campo]) + 1)].hidden  # noqa: E731
        self.assertFalse(oculto("momento_atual"))
        self.assertFalse(oculto("garantias"))
        self.assertTrue(oculto("vara"))                    # fora do perfil
        self.assertFalse(oculto("situacao"))               # usada pelos indicadores: sempre visível
        self.assertTrue(oculto("ultimo_andamento"))
        # valor do campo ativo foi gravado
        linhas = linhas_da_aba(dest)
        self.assertEqual(linhas[0]["Momento Atual"], fi.obter(self.fichas[0], "momento_atual"))
        self.assertIsNone(linhas[0]["Vara"])

    def test_e_andamentos_e_campos_gravados(self):
        linhas = linhas_da_aba(self.dest)
        for f, linha in zip(self.fichas, linhas):
            self.assertEqual(linha["Andamentos"], f"Até {fi.data_br(DATA_BASE)} sem atualizações.")
            self.assertEqual(linha["Réu(s)"], fi.obter(f, "reus"))
            self.assertEqual(linha["Situação"], fi.obter(f, "situacao"))
            self.assertEqual(linha["Ativo"], "Sim" if f["ativo"] else "Não")
            if fi.obter(f, "valor_causa"):
                self.assertAlmostEqual(linha["Valor da Causa"], float(fi.obter(f, "valor_causa")), 2)
            if fi.obter(f, "data_ajuizamento"):
                self.assertEqual(linha["Data do Ajuizamento"].date().isoformat(), fi.obter(f, "data_ajuizamento"))
            self.assertEqual(linha["Resultado"], fi.obter(f, "resultado"))

    @unittest.skipUnless(TEM_LO, "LibreOffice ausente")
    def test_f_indicadores_conferem_com_calculo_independente(self):
        wb = self.lo_valores(self.dest)
        obtido = indicadores_da_planilha(wb)
        previsto = esperado(self.fichas, DATA_BASE, 350, 1.05)
        self.assertEqual(set(obtido), set(previsto))
        for rotulo, valor in previsto.items():
            if valor == "":
                self.assertIn(obtido[rotulo], (None, ""), rotulo)
            else:
                self.assertAlmostEqual(float(obtido[rotulo]), float(valor), 4, rotulo)
        # os indicadores não são todos zero: a conferência está exercitando alguma coisa
        self.assertGreater(previsto["Economia efetiva"], 0)
        self.assertGreater(previsto["Encerrados com valor realizado lançado"], 5)
        self.assertGreater(previsto["Processos decididos"], 20)
        # colunas calculadas conferem linha a linha
        ws = wb["Processos"]
        cab = {c.value: c.column for c in ws[1]}
        for r, f in zip(range(2, 202), self.fichas):
            causa = fi.dinheiro(fi.obter(f, "valor_causa"))
            est = fi.dinheiro(fi.obter(f, "valor_estimado"))
            econ = ws.cell(r, cab["Valor Economizado"]).value
            if not f["ativo"] and causa is not None and est is not None:
                self.assertAlmostEqual(float(econ), float(causa - est), 2)
            else:
                self.assertIn(econ, (None, ""))
            tr, aj = fi.data(fi.obter(f, "data_transito")), fi.data(fi.obter(f, "data_ajuizamento"))
            taxa = ws.cell(r, cab["Taxa de resolução (em dias)"]).value
            if tr and aj:
                self.assertEqual(int(taxa), (tr - aj).days)
            else:
                self.assertIn(taxa, (None, ""))

    @unittest.skipUnless(TEM_LO, "LibreOffice ausente")
    def test_g_quadros_dos_graficos_e_historico(self):
        wb = self.lo_valores(self.dest)
        ws = wb["Indicadores"]
        quadro = {}
        for r in range(1, ws.max_row + 1):
            if ws.cell(r, 1).value and ws.cell(r, 3).value is None and ws.cell(r, 2).value is not None:
                quadro[ws.cell(r, 1).value] = ws.cell(r, 2).value
        for situacao in ("Ativo", "Encerrado", "Suspenso"):
            self.assertEqual(quadro[situacao], sum(1 for f in self.fichas if fi.obter(f, "situacao") == situacao))
        self.assertEqual(quadro["Trabalhista"], sum(1 for f in self.fichas if fi.obter(f, "area") == "Trabalhista"))
        h = wb["Histórico"]
        self.assertEqual([c.value for c in h[2]][1:4], [200, sum(1 for f in self.fichas if f["ativo"]),
                                                       sum(1 for f in self.fichas if not f["ativo"])])
        self.assertEqual(h["A2"].value.date().isoformat(), DATA_BASE)
        p = wb["Parâmetros"]
        self.assertEqual(p["B2"].value, 350)
        self.assertEqual(p["B3"].value.date().isoformat(), DATA_BASE)
        self.assertEqual(p["B5"].value, "Empresa Exemplo A Ltda; Empresa Exemplo B Ltda")
        d = wb["Dashboard"]
        self.assertEqual(d["A4"].value, 200)

    @unittest.skipUnless(TEM_LO, "LibreOffice ausente")
    def test_h_pdf_pelo_libreoffice(self):
        pasta = self.lo(self.dest, "pdf")
        pdf = (pasta / "do-zero.pdf").read_bytes()
        self.assertTrue(pdf.startswith(b"%PDF") and len(pdf) > 5000)

    def test_i_deterministico_e_original_do_modelo_intacto(self):
        antes = sha(MODELO.read_bytes())
        outro = self.saida("do-zero-2.xlsx")
        r2 = x.gravar(None, self.estado, outro)
        self.assertEqual(self.erros(r2), [])
        self.assertEqual(outro.read_bytes(), self.dest.read_bytes())
        self.assertEqual(sha(MODELO.read_bytes()), antes)

    def test_j_gravar_de_novo_sobre_o_resultado_e_idempotente(self):
        """Usar a saída como molde com o mesmo estado: nenhuma linha duplicada, nenhum texto repetido."""
        dest2 = self.saida("segunda.xlsx")
        r = x.gravar(self.dest, self.estado, dest2)
        self.assertEqual(self.erros(r), [])
        self.assertEqual(r["processos_novos"], [])
        self.assertEqual(linhas_da_aba(dest2), linhas_da_aba(self.dest))
        self.assertEqual(r["textos_gravados"], self.res["textos_gravados"])
        self.assertEqual(x.validar(dest2), [])
        # e o histórico não ganhou linha repetida
        self.assertEqual(len(linhas_da_aba(dest2, "Histórico")), 1)

def construir_cliente(caminho, abas, campos=None, ocultas=(), tabela=True, cabecalhos=None, extra_colunas=()):
    """Planilha 'de cliente' fictícia feita com openpyxl: {nome_da_aba: [fichas]}; cada aba com tabela do Excel.
    `campos`: campos da ficha, na ordem das colunas; `cabecalhos`: {campo: texto do cabeçalho} para variar nomes."""
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table, TableStyleInfo
    campos = list(campos or x.COLUNAS_PADRAO)
    cabecalhos = cabecalhos or {}
    wb = Workbook()
    wb.remove(wb.active)
    for k, (nome, fichas) in enumerate(abas.items()):
        ws = wb.create_sheet(nome)
        titulos = [cabecalhos.get(c, x.CAMPOS_B[c]) for c in campos] + list(extra_colunas)
        ws.append(titulos)
        for f in fichas:
            linha = []
            for c in campos:
                v = x._valor_da_ficha(f, c)
                if c == "andamentos":
                    v = f"Até 30/09/2026 sem atualizações."
                linha.append(v)
            ws.append(linha + [None] * len(extra_colunas))
        for i, c in enumerate(campos, start=1):
            tipo = fi.CAMPOS.get(c, ("", "", "texto", None))[2]
            for r in range(2, len(fichas) + 2):
                if tipo == "data":
                    ws.cell(r, i).number_format = "dd/mm/yyyy"
                elif tipo == "dinheiro":
                    ws.cell(r, i).number_format = "#,##0.00"
        for c in ocultas:
            ws.column_dimensions[x.col_letra(campos.index(c) + 1)].hidden = True
        if tabela:
            ws.add_table(Table(displayName=f"tbl{k}", ref=f"A1:{x.col_letra(len(titulos))}{max(2, len(fichas) + 1)}",
                               tableStyleInfo=TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)))
    wb.save(caminho)
    return caminho


def numero_da_linha(caminho, numero, aba="Processos"):
    ws = openpyxl.load_workbook(caminho)[aba]
    for r in range(2, ws.max_row + 1):
        if numero in str(ws.cell(r, 1).value or ""):
            return r
    return None


def texto_andamentos(caminho, numero, aba="Processos"):
    ws = openpyxl.load_workbook(caminho)[aba]
    cab = {c.value: c.column for c in ws[1]}
    return ws.cell(numero_da_linha(caminho, numero, aba), cab["Andamentos"]).value


def valor_da_celula(caminho, numero, cabecalho, aba="Processos"):
    ws = openpyxl.load_workbook(caminho)[aba]
    cab = {c.value: c.column for c in ws[1]}
    return ws.cell(numero_da_linha(caminho, numero, aba), cab[cabecalho]).value


def codigos(res, nivel=None):
    return [a["codigo"] for a in res["avisos"] if nivel is None or a["nivel"] == nivel]


class TestAtualizarArquivoDoCliente(BaseLO):
    """gravar(molde, ...): a planilha do cliente é o molde (política de colunas, andamentos, idempotência)."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        todas = enriquecer(ficticio.gerar_carteira(34, clientes=2, semente=2))
        cls.base, cls.novas = todas[:30], todas[30:]
        cls.A = cls.tmp / "A.xlsx"
        cls.res_a = x.gravar(None, estado_de(cls.base, data_base="2026-09-30"), cls.A)
        # quem não tem resultado/probabilidade lançados (para testar preenchimento e conflito)
        cls.sem_resultado = [f for f in cls.base if not fi.obter(f, "resultado") and f["ativo"]]
        cls.sem_prob = [f for f in cls.base if not fi.obter(f, "probabilidade")]

    def setUp(self):
        self.fichas = [self.copia(f) for f in self.base]
        self.por_numero = {f["numero"]: f for f in self.fichas}

    @staticmethod
    def copia(f):
        import copy
        return copy.deepcopy(f)

    def gravar(self, molde, fichas, data_base="2026-10-31", nome="saida.xlsx", eventos=(), **kw):
        dest = self.saida(nome)
        if dest.exists():
            dest.unlink()
        perfil = kw.pop("perfil", None)
        res = x.gravar(molde, estado_de(fichas, data_base=data_base, eventos=eventos, perfil=perfil), dest, **kw)
        self.assertEqual(codigos(res, "erro"), [], res["avisos"])
        return res, dest

    def test_a_andamentos_so_acrescenta_e_fecho_so_sem_novidade(self):
        import planilha
        f3, f4, f5 = self.fichas[3], self.fichas[4], self.fichas[5]
        ev3 = evento(f3["numero"], "05/10/2026", "Foi proferida decisão", "determinando a citação do réu")
        ev4a = evento(f4["numero"], "08/10/2026", "Foi designada audiência", "para o dia 20/11/2026")
        ev4b = evento(f4["numero"], "08/10/2026", "Foi juntada petição", "da parte autora")
        res, dest = self.gravar(self.A, self.fichas, eventos=[ev3, ev4b, ev4a])
        self.assertEqual(res["processos_novos"], [])
        # com novidade: o fecho antigo sai e o andamento entra, sem fecho novo
        self.assertEqual(texto_andamentos(dest, f3["numero"]), planilha.frase_planilha(ev3))
        self.assertEqual(res["textos_gravados"][f3["numero"]], planilha.frase_planilha(ev3))
        # dois andamentos no mesmo dia: os dois entram, em ordem
        t4 = texto_andamentos(dest, f4["numero"])
        self.assertEqual(t4, f"{planilha.frase_planilha(ev4b)} {planilha.frase_planilha(ev4a)}")
        self.assertNotIn("sem atualizações", t4)
        # sem novidade: o fecho é renovado com a data-base
        self.assertEqual(texto_andamentos(dest, f5["numero"]), "Até 31/10/2026 sem atualizações.")
        self.assertIn(f3["numero"], res["processos_atualizados"])
        self.assertIn(f5["numero"], res["processos_atualizados"])

    def test_b_gravar_duas_vezes_nao_duplica(self):
        f3 = self.fichas[3]
        ev = evento(f3["numero"], "05/10/2026", "Foi proferida sentença", "julgando procedente o pedido")
        r1, B = self.gravar(self.A, self.fichas, eventos=[ev], nome="B.xlsx")
        r2, C = self.gravar(B, self.fichas, eventos=[ev], nome="C.xlsx")
        self.assertEqual(linhas_da_aba(C), linhas_da_aba(B))
        self.assertEqual(r2["processos_novos"], [])
        self.assertEqual(r2["processos_atualizados"], [])
        self.assertEqual([i["motivo"] for i in r2["ignorados"] if i["numero"] == f3["numero"]],
                         ["andamento_ja_presente"])
        self.assertIn("andamento_ja_presente", codigos(r2))      # aviso agregado, com os processos
        self.assertEqual(r2["textos_gravados"], r1["textos_gravados"])
        # ciclo seguinte, sem novidade: o fecho vem DEPOIS do andamento e depois só se renova
        r3, D = self.gravar(C, self.fichas, data_base="2026-11-30", nome="D.xlsx")
        t = texto_andamentos(D, f3["numero"])
        self.assertTrue(t.endswith("Até 30/11/2026 sem atualizações."), t)
        self.assertEqual(t.count("Em 05/10/2026"), 1)
        r4, E = self.gravar(D, self.fichas, data_base="2026-12-31", nome="E.xlsx")
        t = texto_andamentos(E, f3["numero"])
        self.assertTrue(t.endswith("Até 31/12/2026 sem atualizações."), t)
        self.assertEqual(t.count("Até "), 1)

    def test_c_edicao_manual_do_texto_e_preservada_e_avisada(self):
        f4 = self.fichas[4]
        linha = numero_da_linha(self.A, f4["numero"])
        editado = self.saida("editado.xlsx")
        x.escrever_celulas(self.A, editado, "Processos", {(linha, "Andamentos"): "Anotação do advogado sobre o caso."})
        f4["ultimo_texto_gravado"] = {"data_base": "2026-09-30", "texto": "Até 30/09/2026 sem atualizações.",
                                      "arquivo": "A.xlsx"}
        ev = evento(f4["numero"], "08/10/2026", "Foi proferida decisão", "determinando a perícia")
        res, dest = self.gravar(editado, self.fichas, eventos=[ev])
        t = texto_andamentos(dest, f4["numero"])
        self.assertTrue(t.startswith("Anotação do advogado sobre o caso. Em 08/10/2026"), t)
        self.assertIn("andamentos_editados_a_mao", codigos(res))
        # sem a memória do último texto gravado, a edição passa em silêncio (e o texto continua preservado)
        f4.pop("ultimo_texto_gravado")
        res2, dest2 = self.gravar(editado, self.fichas, eventos=[ev], nome="s2.xlsx")
        self.assertNotIn("andamentos_editados_a_mao", codigos(res2))
        self.assertEqual(texto_andamentos(dest2, f4["numero"]), t)

    def test_d_fecho_editado_a_mao_avisa_edicao_manual_sobrescrita(self):
        f7 = self.fichas[7]
        linha = numero_da_linha(self.A, f7["numero"])
        editado = self.saida("fecho-editado.xlsx")
        x.escrever_celulas(self.A, editado, "Processos", {(linha, "Andamentos"): "Até 15/09/2026 sem atualizações."})
        f7["ultimo_texto_gravado"] = {"data_base": "2026-09-30", "texto": "Até 30/09/2026 sem atualizações.",
                                      "arquivo": "A.xlsx"}
        res, dest = self.gravar(editado, self.fichas)
        self.assertEqual(texto_andamentos(dest, f7["numero"]), "Até 31/10/2026 sem atualizações.")
        avisos = [a for a in res["avisos"] if a["codigo"] == "edicao_manual_sobrescrita"]
        self.assertEqual([a["onde"] for a in avisos], [f7["numero"]])

    def test_e_colunas_mecanicas_trocam_objetivas_e_julgamento_preservam(self):
        f0 = self.por_numero[self.sem_resultado[0]["numero"]]
        f0["ativo"] = False
        fi.definir(f0, "situacao", "Encerrado", "derivado", forcar=True)
        fi.definir(f0, "resultado", "Improcedente", "sugerido")
        f1 = self.fichas[1]
        antigo = valor_da_celula(self.A, f1["numero"], "Valor da Causa")
        fi.definir(f1, "valor_causa", "99999.00", "coletado", forcar=True)
        f2 = self.por_numero[next(f for f in self.sem_prob if f["numero"] not in (f0["numero"], f1["numero"]))["numero"]]
        humano = self.saida("humano.xlsx")
        x.escrever_celulas(self.A, humano, "Processos", {(numero_da_linha(self.A, f2["numero"]), "Probabilidade"): "Remota"})
        fi.definir(f2, "probabilidade", "Provável", "sugerido")
        res, dest = self.gravar(humano, self.fichas)
        # mecânicas: trocadas
        self.assertEqual(valor_da_celula(dest, f0["numero"], "Situação"), "Encerrado")
        self.assertEqual(valor_da_celula(dest, f0["numero"], "Ativo"), "Não")
        # julgamento com a célula vazia: preenchido (origem sugerido)
        self.assertEqual(valor_da_celula(dest, f0["numero"], "Resultado"), "Improcedente")
        self.assertIn(f0["numero"], res["processos_atualizados"])
        # objetiva já preenchida e diferente: preservada, com aviso agregado
        self.assertEqual(valor_da_celula(dest, f1["numero"], "Valor da Causa"), antigo)
        av = [a for a in res["avisos"] if a["codigo"] == "valor_divergente" and "Valor da Causa" in a["onde"]]
        self.assertEqual(len(av), 1)
        self.assertIn(f1["numero"], av[0]["candidatos"])
        # julgamento preenchido por humano: nunca sobrescrito, com aviso
        self.assertEqual(valor_da_celula(dest, f2["numero"], "Probabilidade"), "Remota")
        av = [a for a in res["avisos"] if a["codigo"] == "campo_divergente" and "Probabilidade" in a["onde"]]
        self.assertEqual(len(av), 1)
        self.assertIn(f2["numero"], av[0]["candidatos"])
        # mudanças registradas com antes e depois
        m = {(c["numero"], c["campo"]): c for c in res["mudancas"]}
        self.assertEqual(m[(f0["numero"], "Situação")]["antes"], "Ativo")
        self.assertEqual(m[(f0["numero"], "Situação")]["depois"], "Encerrado")
        self.assertIsNone(m[(f0["numero"], "Resultado")]["antes"])

    def test_f_formulas_nunca_sao_sobrescritas(self):
        f = next(f for f in self.fichas if not f["ativo"])
        fi.definir(f, "valor_economizado", "123456.00", "humano")
        fi.definir(f, "taxa_resolucao_dias", 999, "humano")
        res, dest = self.gravar(self.A, self.fichas)
        ws = openpyxl.load_workbook(dest)["Processos"]
        cab = {c.value: c.column for c in ws[1]}
        r = numero_da_linha(dest, f["numero"])
        self.assertTrue(str(ws.cell(r, cab["Valor Economizado"]).value).startswith("=IF("))
        self.assertTrue(str(ws.cell(r, cab["Taxa de resolução (em dias)"]).value).startswith("=IF("))
        self.assertEqual(x.validar(dest), [])

    def test_g_mecanica_alterada_a_mao_avisa_quando_ha_memoria(self):
        f = self.fichas[9]
        linha = numero_da_linha(self.A, f["numero"])
        editado = self.saida("situacao-manual.xlsx")
        x.escrever_celulas(self.A, editado, "Processos", {(linha, "Situação"): "Suspenso"})
        f["ultimos_valores_gravados"] = {"situacao": fi.obter(f, "situacao")}      # o que o sistema gravou
        res, dest = self.gravar(editado, self.fichas)
        self.assertEqual(valor_da_celula(dest, f["numero"], "Situação"), fi.obter(f, "situacao"))
        av = [a for a in res["avisos"] if a["codigo"] == "edicao_manual_sobrescrita" and a["onde"] == f["numero"]]
        self.assertEqual(len(av), 1)
        # sem memória: troca sem aviso
        f.pop("ultimos_valores_gravados")
        res2, _ = self.gravar(editado, self.fichas, nome="sem-memoria.xlsx")
        self.assertEqual([a for a in res2["avisos"] if a["codigo"] == "edicao_manual_sobrescrita"], [])

    def test_h_processo_novo_entra_ao_fim_com_estilo_e_formulas(self):
        novo = self.novas[0]
        res, dest = self.gravar(self.A, self.fichas + [self.copia(novo)])
        self.assertEqual(res["processos_novos"], [novo["numero"]])
        self.assertEqual(x.validar(dest), [])
        ws = openpyxl.load_workbook(dest)["Processos"]
        self.assertEqual(ws.max_row, 32)
        self.assertEqual(ws.tables["tblProcessos"].ref, "A1:AK32")
        self.assertEqual(ws["A32"].value, novo["numero"])
        self.assertTrue(str(ws["W32"].value).startswith("=IF("))
        self.assertEqual(ws["G32"].number_format, "dd/mm/yyyy")
        self.assertEqual(res["textos_gravados"][novo["numero"]], "Até 31/10/2026 sem atualizações.")
        # partes não editadas idênticas (gráficos, desenhos, estilos)
        a, d = partes(self.A), partes(dest)
        for n in a:
            if n.startswith("xl/charts/") and b"Hist" in a[n]:
                continue                  # o gráfico do histórico cresce com a linha nova do retrato mensal
            if n.startswith(("xl/charts/", "xl/drawings/")) or n in ("xl/styles.xml", "xl/theme/theme1.xml"):
                self.assertEqual(a[n], d[n], n)

    def test_i_linha_de_base_entra_no_texto_de_processo_novo_e_de_celula_vazia(self):
        novo = self.copia(self.novas[1])
        ficticio.anexar_linha_de_base(novo, data_base="2026-09-30")
        base = novo["linha_de_base"]["andamentos_texto"]
        self.assertTrue(base)
        ev = evento(novo["numero"], "05/10/2026", "Foi proferida decisão", "determinando a citação do réu")
        res, dest = self.gravar(self.A, self.fichas + [novo], eventos=[ev])
        t = texto_andamentos(dest, novo["numero"])
        self.assertTrue(t.startswith(base.rstrip()), t[:80])
        self.assertTrue(t.endswith("determinando a citação do réu."), t[-80:])
        # célula de Andamentos vazia numa linha existente: parte do histórico migrado
        f = self.fichas[11]
        ficticio.anexar_linha_de_base(f, data_base="2026-09-30")
        vazio = self.saida("vazio.xlsx")
        x.escrever_celulas(self.A, vazio, "Processos", {(numero_da_linha(self.A, f["numero"]), "Andamentos"): x.LIMPAR})
        res2, dest2 = self.gravar(vazio, self.fichas, nome="vazio-out.xlsx")
        self.assertTrue(texto_andamentos(dest2, f["numero"]).startswith(f["linha_de_base"]["andamentos_texto"].rstrip()))

    def test_j_texto_acima_do_limite_do_excel_so_afeta_aquele_processo(self):
        f, g2 = self.fichas[12], self.fichas[13]
        grande = self.saida("grande.xlsx")
        x.escrever_celulas(self.A, grande, "Processos",
                           {(numero_da_linha(self.A, f["numero"]), "Andamentos"): "a" * 32760})
        ev = evento(f["numero"], "05/10/2026", "Foi proferida decisão", "determinando a citação do réu")
        ev2 = evento(g2["numero"], "05/10/2026", "Foi proferida decisão", "determinando a citação do réu")
        dest = self.saida("grande-out.xlsx")
        res = x.gravar(grande, estado_de(self.fichas, data_base="2026-10-31", eventos=[ev, ev2]), dest)
        self.assertIsNotNone(res["destino"])
        erros = [a for a in res["avisos"] if a["nivel"] == "erro"]
        self.assertEqual([(a["codigo"], a["onde"]) for a in erros], [("texto_acima_do_limite", f["numero"])])
        self.assertIn({"numero": f["numero"], "motivo": "texto_acima_do_limite"}, res["ignorados"])
        self.assertEqual(texto_andamentos(dest, f["numero"]), "a" * 32760)       # intacto, não truncado
        self.assertTrue(texto_andamentos(dest, g2["numero"]).startswith("Em 05/10/2026"))
        self.assertEqual(x.validar(dest), [])

    def test_k_perfil_restringe_colunas_gravadas(self):
        f0 = self.por_numero[self.sem_resultado[0]["numero"]]
        fi.definir(f0, "resultado", "Improcedente", "sugerido")
        res, dest = self.gravar(self.A, self.fichas, perfil={"colunas_ativas": ["numero", "andamentos", "situacao"]})
        self.assertIsNone(valor_da_celula(dest, f0["numero"], "Resultado"))       # coluna fora do perfil
        res2, dest2 = self.gravar(self.A, self.fichas, nome="s2.xlsx",
                                  perfil={"colunas_ativas": ["numero", "andamentos", "resultado", "campo_que_nao_existe"]})
        self.assertEqual(valor_da_celula(dest2, f0["numero"], "Resultado"), "Improcedente")
        self.assertIn("coluna_desconhecida_no_perfil", codigos(res2))

    def test_l_data_de_referencia_e_historico_do_segundo_ciclo(self):
        res, dest = self.gravar(self.A, self.fichas, data_base="2026-10-31")
        wb = openpyxl.load_workbook(dest)
        self.assertEqual(wb["Parâmetros"]["B3"].value.date(), datetime.date(2026, 10, 31))
        hist = linhas_da_aba(dest, "Histórico")
        self.assertEqual([h["Data-base"].date().isoformat() for h in hist], ["2026-09-30", "2026-10-31"])
        self.assertEqual(hist[1]["Processos"], 30)
        # o gráfico do histórico acompanhou a linha nova
        grafico = [d.decode() for n, d in partes(dest).items() if n.startswith("xl/charts/chart") and b"Hist" in d]
        self.assertTrue(any("$A$2:$A$3" in c or "$B$1:$B$3" in c or "$B$2:$B$3" in c for c in grafico), grafico)
        self.assertEqual(x.validar(dest), [])
        # mesma data-base de novo: a linha é atualizada, não duplicada
        res2, dest2 = self.gravar(dest, self.fichas, data_base="2026-10-31", nome="h2.xlsx")
        self.assertEqual(len(linhas_da_aba(dest2, "Histórico")), 2)
        # retratos anteriores informados pelo coordenador entram uma vez, em ordem
        antes = {"data_base": "2026-08-31", "totais": {"processos": 25, "ativos": 20, "encerrados": 5,
                                                        "valor_causa": "1000.00", "valor_estimado": "0.00",
                                                        "valor_economizado": "0.00"}}
        dest3 = self.saida("h3.xlsx")
        r3 = x.gravar(self.A, estado_de(self.fichas, data_base="2026-10-31", historico=[antes, antes]), dest3)
        self.assertEqual(codigos(r3, "erro"), [])
        datas = [h["Data-base"].date().isoformat() for h in linhas_da_aba(dest3, "Histórico")]
        self.assertEqual(datas, ["2026-09-30", "2026-08-31", "2026-10-31"])   # só se acrescenta ao fim

    def test_m_campos_nao_migrados_idempotente(self):
        nm = [{"coluna": "Observação do advogado", "amostra": ["a", "b"],
               "valores": {self.fichas[0]["numero"]: "ligar na segunda", self.fichas[1]["numero"]: "ok"}},
              {"coluna": "Cor", "amostra": ["azul"]}]
        dest = self.saida("nm.xlsx")
        r = x.gravar(None, estado_de(self.fichas[:3], campos_nao_migrados=nm), dest)
        self.assertEqual(codigos(r, "erro"), [])
        linhas = linhas_da_aba(dest, "Campos não migrados")
        self.assertEqual([(l["Coluna de origem"], l["Processo"], l["Valor"]) for l in linhas],
                         [("Observação do advogado", self.fichas[0]["numero"], "ligar na segunda"),
                          ("Observação do advogado", self.fichas[1]["numero"], "ok"), ("Cor", None, "azul")])
        dest2 = self.saida("nm2.xlsx")
        x.gravar(dest, estado_de(self.fichas[:3], campos_nao_migrados=nm), dest2)
        self.assertEqual(len(linhas_da_aba(dest2, "Campos não migrados")), 3)
        # molde sem a aba: aviso, sem erro
        r3 = x.gravar(construir_cliente(self.saida("sem-aba.xlsx"), {"Processos": self.fichas[:2]}),
                      estado_de(self.fichas[:2], campos_nao_migrados=nm), self.saida("sem-aba-out.xlsx"))
        self.assertIn("campos_nao_migrados_sem_aba", codigos(r3))

    def test_n_numero_duplicado_na_planilha_nao_e_atualizado(self):
        f = self.fichas[2]
        duplicada = self.saida("dup.xlsx")
        outra = numero_da_linha(self.A, self.fichas[20]["numero"])
        x.escrever_celulas(self.A, duplicada, "Processos", {(outra, "Número do Processo"): f["numero"]})
        res, dest = self.gravar(duplicada, self.fichas)
        self.assertIn("numero_duplicado_na_planilha", codigos(res))
        self.assertIn({"numero": f["numero"], "motivo": "numero_duplicado_na_planilha"}, res["ignorados"])

    def test_o_vinculado_na_celula_acha_a_linha_do_principal(self):
        f = next(f for f in self.fichas if f["vinculados"])
        principal = f["numero"]
        r = numero_da_linha(self.A, principal)
        self.assertIn(f["vinculados"][0]["numero"], str(valor_da_celula(self.A, principal, "Número do Processo")))
        # o principal sumiu da célula e só o vinculado ficou: ainda acha a linha, sem criar processo novo
        so_vinculado = self.saida("so-vinculado.xlsx")
        x.escrever_celulas(self.A, so_vinculado, "Processos", {(r, "Número do Processo"): f["vinculados"][0]["numero"]})
        res, dest = self.gravar(so_vinculado, self.fichas)
        self.assertEqual(res["processos_novos"], [])

    def test_p_evento_de_processo_fora_das_fichas_vira_aviso(self):
        ev = evento(ficticio.numero_ficticio(4999), "05/10/2026", "Foi proferida decisão", "determinando a citação")
        res, dest = self.gravar(self.A, self.fichas, eventos=[ev])
        self.assertIn("evento_sem_ficha", codigos(res))

    def test_q_evento_nao_aprovado_nao_entra(self):
        f = self.fichas[3]
        ev = evento(f["numero"], "05/10/2026", "Foi proferida decisão", "determinando a citação", status="rascunho")
        res, dest = self.gravar(self.A, self.fichas, eventos=[ev])
        self.assertEqual(texto_andamentos(dest, f["numero"]), "Até 31/10/2026 sem atualizações.")

    def test_r_fecho_apos_novidade_opcional(self):
        f = self.fichas[3]
        ev = evento(f["numero"], "05/10/2026", "Foi proferida decisão", "determinando a citação do réu")
        res, dest = self.gravar(self.A, self.fichas, eventos=[ev], fecho_apos_novidade=True)
        self.assertTrue(texto_andamentos(dest, f["numero"]).endswith("determinando a citação do réu. Até 31/10/2026 sem atualizações."))

    def test_s_original_nunca_alterado_e_destino_igual_ao_molde(self):
        antes = sha(self.A.read_bytes())
        res, dest = self.gravar(self.A, self.fichas, eventos=[])
        self.assertEqual(sha(self.A.read_bytes()), antes)
        r = x.gravar(self.A, estado_de(self.fichas), self.A)
        self.assertIsNone(r["destino"])
        self.assertEqual(codigos(r, "erro"), ["destino_igual_ao_molde"])
        self.assertEqual(sha(self.A.read_bytes()), antes)


if __name__ == "__main__":
    unittest.main(verbosity=2)
