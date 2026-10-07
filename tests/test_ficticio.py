"""A base sintética em si: a carteira fictícia é determinística, válida e variada; os defeitos intencionais são
detectáveis; as listas brutas são lidas por carteira.ler_lista; nenhum número real entra no repositório.

    python3 -m unittest tests/test_ficticio.py -v
"""
import re
import sys
import tempfile
import unittest
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import isolamento  # noqa: E402,F401  (antes de tudo)
import carteira  # noqa: E402
import comum  # noqa: E402
import ficha  # noqa: E402
import ficticio  # noqa: E402
import taxonomia  # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent
CNJ = re.compile(r"\b[0-9]{7}-[0-9]{2}\.[0-9]{4}\.[0-9]\.[0-9]{2}\.[0-9]{4}\b")
PERMITIDOS = re.compile(r"0000000-00\.0000\.0\.00\.0000|9999999-99\.9999\.9\.99\.9999|123456[7-9]-|1234570-")


class TestNumero(unittest.TestCase):
    def test_numero_ficticio_tem_dv_correto_e_e_reconhecido(self):
        for n in range(50):
            numero = ficticio.numero_ficticio(n, 2020 + n % 6, [8, 5, 4][n % 3], 1 + n % 6, n % 20)
            self.assertTrue(carteira.dv_correto(carteira.CNJ.fullmatch(numero).groups()), numero)
            self.assertTrue(ficticio.dv_confere(numero))
            validos, invalidos = carteira.numeros_no_texto(numero)
            self.assertEqual((len(validos), invalidos), (1, []))

    def test_dv_errado_e_recusado(self):
        errado = ficticio.numero_com_dv_errado(ficticio.numero_ficticio(3))
        self.assertFalse(ficticio.dv_confere(errado))
        self.assertEqual(carteira.numeros_no_texto(errado), ([], [errado]))


class TestCarteira(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.fichas = ficticio.gerar_carteira(200)

    def test_determinismo(self):
        self.assertEqual(self.fichas, ficticio.gerar_carteira(200))
        self.assertNotEqual(self.fichas, ficticio.gerar_carteira(200, semente=2))
        self.assertEqual(ficticio.gerar_carteira(30), self.fichas[:30], "n menor é o começo da carteira maior")

    def test_tamanho_validade_e_numeros(self):
        self.assertEqual(len(self.fichas), 200)
        self.assertEqual([(f["numero"], ficha.validar(f)) for f in self.fichas if ficha.validar(f)], [])
        todos = [n for f in self.fichas for n in ficha.todos_os_numeros(f)]
        self.assertEqual(len(todos), len(set(todos)), "número repetido na carteira limpa")
        for n in todos:
            m = carteira.CNJ.fullmatch(n)
            self.assertTrue(m and ficticio.dv_confere(n), n)
            self.assertTrue(ficticio.SEQ_BASE <= int(m[1]) < ficticio.SEQ_LIMITE, f"fora da faixa sintética: {n}")
        for f in self.fichas:
            self.assertEqual(f["tribunal"], carteira.numeros_no_texto(f["numero"])[0][0][1])

    def test_distribuicao_realista(self):
        ativos = sum(f["ativo"] for f in self.fichas) / len(self.fichas)
        self.assertTrue(0.45 <= ativos <= 0.65, ativos)
        tribunais = Counter(re.match(r"TRT|TRF|TJ", f["tribunal"])[0] for f in self.fichas)
        self.assertTrue({"TJ", "TRT", "TRF"} <= set(tribunais), tribunais)
        self.assertGreaterEqual(len({f["tribunal"] for f in self.fichas}), 10)
        self.assertEqual(len({ficha.obter(f, "cliente") for f in self.fichas}), 5)
        self.assertGreaterEqual(sum(1 for f in self.fichas if f["vinculados"]), 10)
        humanos = [f for f in self.fichas if any(ficha.origem(f, c) == "humano" for c in ficha.CAMPOS_DE_JULGAMENTO)]
        self.assertTrue(30 <= len(humanos) <= 90, len(humanos))
        self.assertGreaterEqual(len({ficha.obter(f, "momento_atual") for f in self.fichas}), 15)
        self.assertEqual({ficha.obter(f, "area") for f in self.fichas} - set(taxonomia.AREA), set())
        self.assertTrue(all(ficha.obter(f, "momento_atual") in taxonomia.MOMENTO_ATUAL for f in self.fichas))
        self.assertEqual(Counter(ficha.obter(f, "resultado") for f in self.fichas).keys() - set(taxonomia.RESULTADO) - {None}, set())

    def test_coerencia_e_nomes_ficticios(self):
        for f in self.fichas:
            momento = ficha.obter(f, "momento_atual")
            self.assertEqual(f["ativo"], taxonomia.momento_ativo(momento))
            if not f["ativo"]:
                self.assertTrue(ficha.obter(f, "resultado"), f"encerrado sem resultado: {f['numero']}")
            if ficha.obter(f, "resultado") == "Acordo":
                self.assertTrue(ficha.obter(f, "valor_acordo"), "acordo sem valor na carteira limpa")
            self.assertTrue(re.fullmatch(r"Cliente Exemplo \d\d .+", ficha.obter(f, "cliente")))
            for campo in ("autores", "reus", "parte_contraria"):
                self.assertRegex(ficha.obter(f, campo), r"^(Pessoa Fictícia \d{4}|Empresa Fictícia \d{3} Ltda|Ente Público Exemplo|Cliente Exemplo \d\d .+)$")
            self.assertTrue(ficha.data(ficha.obter(f, "data_ajuizamento")) <= ficha.data(ficha.obter(f, "ultimo_andamento")))
            self.assertTrue(ficha.data(ficha.obter(f, "ultimo_andamento")) <= ficha.data(ficticio.HOJE))
        self.assertEqual(ficticio.detectar_defeitos(self.fichas), {k: [] for k in ficticio.detectar_defeitos(self.fichas)})

    def test_vinculados_validos(self):
        com = [f for f in self.fichas if f["vinculados"]]
        for f in com:
            for v in f["vinculados"]:
                self.assertIn(v["tipo"], taxonomia.TIPO_VINCULO)
                self.assertTrue(ficticio.dv_confere(v["numero"]))
        self.assertTrue(com)

    def test_ida_e_volta_no_arquivo(self):
        with ficticio.projeto_de_teste(self.fichas[:20]) as proj:
            de_novo = ficha.carregar(todas=True)
            self.assertEqual(de_novo, self.fichas[:20])
            self.assertTrue(str(proj["pasta"]).startswith(str(isolamento.TMP)))
            self.assertEqual(comum.PROJETO, proj["slug"])
        self.assertIsNone(comum.PROJETO, "comum volta ao estado anterior")
        self.assertTrue(str(comum.CARTEIRA_FILE).startswith(str(isolamento.TMP)))


class TestDefeitos(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.fichas = ficticio.gerar_carteira(200, com_defeitos=True)

    def test_mesmo_tamanho_e_determinismo(self):
        self.assertEqual(len(self.fichas), 200)
        self.assertEqual(self.fichas, ficticio.gerar_carteira(200, com_defeitos=True))
        self.assertEqual(self.fichas.defeitos, ficticio.gerar_carteira(200, com_defeitos=True).defeitos)
        self.assertEqual(ficticio.gerar_carteira(200).defeitos, {})

    def test_cada_defeito_existe_e_e_detectavel(self):
        previstos = self.fichas.defeitos
        achados = ficticio.detectar_defeitos(self.fichas)
        for tipo in ("numero_duplicado", "materia_dois_rotulos", "acordo_sem_valor", "encerrado_sem_resultado",
                     "ativo_em_conflito", "dv_errado", "grafias_cliente"):
            self.assertTrue(previstos[tipo], f"sem defeito do tipo {tipo}")
            self.assertTrue(achados[tipo], f"defeito não detectado: {tipo}")
        self.assertEqual(sorted(achados["numero_duplicado"]), sorted(previstos["numero_duplicado"]))
        self.assertEqual(sorted(achados["materia_dois_rotulos"]), sorted(previstos["materia_dois_rotulos"]))
        self.assertEqual(sorted(achados["acordo_sem_valor"]), sorted(previstos["acordo_sem_valor"]))
        self.assertEqual(sorted(achados["encerrado_sem_resultado"]), sorted(previstos["encerrado_sem_resultado"]))
        self.assertEqual(sorted(achados["ativo_em_conflito"]), sorted(previstos["ativo_em_conflito"]))
        self.assertEqual(sorted(achados["dv_errado"]), sorted(previstos["dv_errado"]))
        self.assertEqual(len(achados["grafias_cliente"]), len(previstos["grafias_cliente"]))

    def test_ficha_validar_pega_dv_e_conflito_e_nada_mais(self):
        problemas = {f["numero"]: ficha.validar(f) for f in self.fichas if ficha.validar(f)}
        esperados = set(self.fichas.defeitos["dv_errado"]) | set(self.fichas.defeitos["ativo_em_conflito"])
        self.assertEqual(set(problemas), esperados)
        for numero in self.fichas.defeitos["dv_errado"]:
            self.assertTrue(any("Número CNJ inválido" in p for p in problemas[numero]))
        for numero in self.fichas.defeitos["ativo_em_conflito"]:
            self.assertTrue(any("conflita" in p for p in problemas[numero]))

    def test_dv_errado_e_recusado_pela_carteira(self):
        errado = self.fichas.defeitos["dv_errado"][0]
        self.assertEqual(carteira.numeros_no_texto(errado), ([], [errado]))


class TestListasBrutas(unittest.TestCase):
    def test_cada_arquivo_e_lido_com_a_contagem_esperada(self):
        fichas = ficticio.gerar_carteira(60)
        with tempfile.TemporaryDirectory() as d:
            arquivos = ficticio.gerar_lista_bruta(d, fichas)
            self.assertEqual(set(arquivos), {"txt", "csv", "xlsx", "xlsx_bagunca"})
            for nome, esperado in arquivos.items():
                self.assertTrue(esperado["arquivo"].exists(), nome)
                registros, invalidos = carteira.ler_lista(str(esperado["arquivo"]), "Cliente Exemplo 01 Ltda")
                self.assertEqual(sorted(r["numero"] for r in registros), sorted(esperado["validos"]), nome)
                self.assertEqual(invalidos, esperado["invalidos"], nome)
                self.assertGreaterEqual(len(esperado["validos"]), 3)
                self.assertEqual(len(esperado["invalidos"]), 1)
            # o csv e o xlsx simples trazem o que mais a lista tiver
            registros, _ = carteira.ler_lista(str(arquivos["csv"]["arquivo"]))
            self.assertTrue(all(r["polo_cliente"] in ("ativo", "passivo") and r["cliente"] for r in registros))
            registros, _ = carteira.ler_lista(str(arquivos["xlsx"]["arquivo"]))
            self.assertTrue(all(r["cliente"].startswith("Cliente Exemplo") for r in registros))

    def test_xlsx_bagunca_e_de_fato_bagunçado(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as d:
            arq = ficticio.gerar_lista_bruta(d, ficticio.gerar_carteira(40))["xlsx_bagunca"]["arquivo"]
            linhas = list(load_workbook(arq).active.iter_rows(values_only=True))
            self.assertEqual(linhas[2][0], "Seq", "cabeçalho na linha 3")
            self.assertTrue(any(all(c is None for c in l) for l in linhas), "tem linha vazia")
            self.assertTrue(any(isinstance(l[1], str) and re.search(r"[A-Za-z]", l[1]) and CNJ.search(l[1]) for l in linhas[3:]),
                            "número misturado com texto")

    def test_blocos_distintos_e_determinismo(self):
        fichas = ficticio.gerar_carteira(60)
        with tempfile.TemporaryDirectory() as a, tempfile.TemporaryDirectory() as b:
            x, y = ficticio.gerar_lista_bruta(a, fichas), ficticio.gerar_lista_bruta(b, fichas)
            for nome in x:
                self.assertEqual(x[nome]["arquivo"].read_bytes() if nome in ("txt", "csv") else x[nome]["validos"],
                                 y[nome]["arquivo"].read_bytes() if nome in ("txt", "csv") else y[nome]["validos"])
            todos = [n for e in x.values() for n in e["validos"]]
            self.assertEqual(len(todos), len(set(todos)), "cada arquivo usa um bloco diferente da carteira")


class TestHistorico(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.fichas = ficticio.gerar_carteira(200)

    def test_texto_em_estilo_de_relatorio_e_coerente(self):
        por_resultado = {"Procedente": "julgando procedente", "Parcialmente procedente": "julgando parcialmente procedentes",
                         "Improcedente": "julgando improcedente", "Acordo": "homologado acordo",
                         "Extinto sem resolução de mérito": "extinguindo o processo sem resolução do mérito"}
        vistos = set()
        for f in self.fichas:
            texto = ficticio.gerar_historico(f)
            frases = ficticio.gerar_textos_de_andamento(f)
            self.assertTrue(texto.startswith("Em ") and texto == " ".join(frases))
            self.assertTrue(all(re.match(r"Em \d{2}/\d{2}/\d{4} ", fr) for fr in frases))
            datas = [ficha.parse_data(fr[3:13]) for fr in frases]
            self.assertEqual(datas, sorted(datas), "ordem cronológica")
            self.assertEqual(datas[-1], ficha.obter(f, "ultimo_andamento"), "o último fato é o último andamento")
            resultado = ficha.obter(f, "resultado")
            if resultado in por_resultado and ficha.obter(f, "momento_atual") not in ("PROCESSO ARQUIVADO",):
                self.assertIn(por_resultado[resultado], texto, f"{resultado}: {texto}")
                vistos.add(resultado)
            self.assertNotRegex(texto, r"\d{7}-\d{2}\.\d{4}")  # sem número de processo no texto
        self.assertEqual(vistos, set(por_resultado))

    def test_ate_corta_na_data_base(self):
        f = next(x for x in self.fichas if ficha.obter(x, "momento_atual") == "TRÂNSITO EM JULGADO")
        completo = ficticio.gerar_textos_de_andamento(f)
        meio = ficha.parse_data(completo[2][3:13])
        cortado = ficticio.gerar_textos_de_andamento(f, ate=meio)
        self.assertEqual(cortado, completo[:3])

    def test_linha_de_base_no_formato_do_contrato(self):
        f = ficticio.anexar_linha_de_base(ficticio.gerar_carteira(5)[0], data_base="2026-09-18")
        base = f["linha_de_base"]
        self.assertEqual(set(base), {"data_base", "andamentos_texto", "arquivo", "ultimo_andamento"})
        self.assertLessEqual(base["ultimo_andamento"], "2026-09-18")
        com = ficticio.gerar_carteira(5, com_linha_de_base=True)
        self.assertTrue(all(x["linha_de_base"]["andamentos_texto"] for x in com))


class TestHigiene(unittest.TestCase):
    def test_nenhum_numero_literal_fora_do_permitido(self):
        """Mesmo padrão do empacotar.sh, sobre tests/, src/ e docs/fase2/."""
        achados = []
        for pasta in ("tests", "src", "docs/fase2"):
            for arq in sorted((RAIZ / pasta).rglob("*")):
                if not arq.is_file() or arq.suffix in (".pyc", ".xlsx", ".pdf") or "__pycache__" in arq.parts:
                    continue
                try:
                    texto = arq.read_text(encoding="utf-8")
                except UnicodeDecodeError:
                    continue
                for n, linha in enumerate(texto.splitlines(), 1):
                    for m in CNJ.finditer(linha):
                        if not PERMITIDOS.search(m.group(0)):
                            achados.append(f"{arq.relative_to(RAIZ)}:{n}: {m.group(0)}")
        self.assertEqual(achados, [])

    def test_nomes_sao_genericos(self):
        permitidos = re.compile(r"^(Cliente Exemplo \d\d( .+)?|Pessoa Fictícia \d{4}|Empresa Fictícia \d{3} Ltda|Ente Público Exemplo|"
                                r"Responsável Exemplo [A-C]|Caso \d{3}|CLIENTE EXEMPLO \d\d .+)$")
        for f in ficticio.gerar_carteira(200, com_defeitos=True):
            for campo in ("cliente", "autores", "reus", "parte_contraria", "responsavel", "apelido"):
                valor = ficha.obter(f, campo)
                if valor:
                    self.assertRegex(valor, permitidos)


if __name__ == "__main__":
    unittest.main(verbosity=2)
