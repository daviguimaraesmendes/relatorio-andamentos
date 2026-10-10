"""Planilha FICTÍCIA no formato "desformatado" de contingências jurídicas, para testar a migração (leitor, lacunas, escritor, painel).

Reproduz a ESTRUTURA de um relatório real de passivos contingentes, sem nenhum dado dele: nomes, CNPJs e números de processo
são inventados (números por `ficticio.numero_ficticio`). O que a planilha tem de difícil:

- aba de processos ativos com 3 linhas de título (1 a 3), linha vazia, cabeçalho na linha 5 (congelado em A6) e 14 colunas:
  AUTOR/RECLAMANTE, RÉU/RECLAMADO, CNPJ ... PROCESSADO, PROCESSO Nº., NATUREZA DA AÇÃO, ATIVO POTENCIAL, PASSIVO POTENCIAL,
  DEPÓSITO JUDICIAL REALIZADO?, VALOR DO DEPÓSITO JUDICIAL, POSSIBILIDADE DE PERDA, "%", PROVISÃO CONSTITUÍDA,
  BREVE RESUMO DO CASO e OBSERVAÇÃO (o histórico);
- valores misturados: `R$ 260.123,33`, `R$41.000,00`, número puro 124229.25, texto com espaço no fim, "-", negativo;
- percentual como fração com formato de porcentagem (0,5) e célula vazia;
- grau de perda colado à justificativa ("REMOTA - Processo extinto ..."), com e sem acento, em maiúsculas e minúsculas;
- uma linha com número de processo inválido no meio dos dados (vira erro `numero_invalido`);
- depois dos dados: linha de total, quadro de critérios (legenda com texto e percentuais) e linhas vazias;
- aba "Arquivados" (cabeçalho na linha 3), com data de ajuizamento, passivo atualizado, pagamento realizado e a coluna
  VALOR ECONOMIZADO em FÓRMULA (`=G4-K4`), sem valor guardado (o openpyxl não grava cache);
- linha "VALOR TOTAL ECONOMIZADO" com SOMA no fim da aba de arquivados.

`GABARITO` resume o que um leitor correto deve achar (os testes comparam com ele).
"""
import datetime
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import ficticio  # noqa: E402

ABA_ATIVOS = "Relatório Empresa Exemplo"
ABA_ARQUIVADOS = "Arquivados - Empresa Exemplo"
EMPRESA = "EMPRESA EXEMPLO COMERCIAL LTDA"
REU = "Empresa Exemplo Comercial Ltda"
CABECALHO_ATIVOS = ["AUTOR/RECLAMANTE", "RÉU/RECLAMADO", "CNPJ EMPRESA EXEMPLO PROCESSADA", "PROCESSO Nº.", "NATUREZA DA AÇÃO",
                    "ATIVO POTENCIAL", "PASSIVO POTENCIAL ", "DEPÓSITO JUDICIAL REALIZADO?", "VALOR DO DEPÓSITO JUDICIAL",
                    "POSSIBILIDADE DE PERDA ", "%", "PROVISÃO CONSTITUÍDA", "BREVE RESUMO DO CASO", "OBSERVAÇÃO"]
CABECALHO_ARQUIVADOS = ["AUTOR/RECLAMANTE", "RÉU/RECLAMADO", "PROCESSO Nº.", "NATUREZA DA AÇÃO", "DATA DE AJUIZAMENTO",
                        "PASSIVO POTENCIAL ", "PASSIVO POTENCIAL ATUALIZADO (CORREÇÃO INPC E JUROS 1% AO MÊS)",
                        "DEPÓSITO JUDICIAL REALIZADO?", "VALOR DO DEPÓSITO JUDICIAL", "POSSIBILIDADE DE PERDA ",
                        "PAGAMENTO REALIZADO", "VALOR ECONOMIZADO"]


def cnpj_ficticio(n):
    """CNPJ inventado e claramente falso (a varredura de confidencialidade só aceita repetições de dígito)."""
    d = str(n % 10)
    return f"{d * 2}.{d * 3}.{d * 3}/{d * 4}-{d * 2}"


def historico(n):
    return (f"Em 10/05/2026, apresentamos contestação. Em 15/06/2026, foi designada audiência de instrução para o caso {n}. "
            "Em 20/07/2026, realizada a audiência de instrução com a oitiva das testemunhas arroladas pelas partes; os autos seguem conclusos para sentença, aguardando decisão do juízo.")


# Cada linha de ativos: (autor, natureza, ativo_potencial, passivo, depósito Sim/Não, valor do depósito, perda, % , provisão)
LINHAS_ATIVOS = [
    ("Pessoa Fictícia 0001", "INDENIZAÇÃO POR DANOS MORAIS E MATERIAIS", None, "R$ 260.123,33", "NÃO", None, "POSSÍVEL", 0.5, "R$ 130.061,66"),
    ("Pessoa Fictícia 0002", "OBRIGAÇÃO DE FAZER E INDENIZAÇÃO POR DANOS MORAIS", None, "R$41.000,00", "NÃO", None, "PROVÁVEL", 1.0, "R$41.000,00"),
    ("Pessoa Fictícia 0003", "TRABALHISTA", None, 124229.25, "SIM", "R$ 70.115,85", "POSSÍVEL", 0.5, 62114.63),
    ("Pessoa Fictícia 0004", "INDENIZAÇÃO - ACIDENTE DE TRÂNSITO", None, 25361.4, "SIM", 144458.62, "PROVÁVEL", 1.0, 25361.4),
    ("Pessoa Fictícia 0005", "INDENIZAÇÃO POR DANOS MORAIS", "R$ 53.839,36", "R$ 175.913,36 ", "NÃO", None, "REMOTA", None, None),
    ("Pessoa Fictícia 0006", "COBRANÇA", None, "-", "NÃO", None, "possível", None, "-"),
    ("Pessoa Fictícia 0007", "INDENIZAÇÃO POR DANOS MATERIAIS", "-R$ 500,00", "R$ 50.000,00", "NÃO", None, "POSSIVEL - Sem decisão", 0.5, "R$ 25.000,00"),
    ("Pessoa Fictícia 0008", "OBRIGAÇÃO DE FAZER", None, "R$ 9.000,00", None, "R$ 9.000,00", "PROVÁVEL", 1.0, "R$ 9.000,00"),
    ("Pessoa Fictícia 0009", "INDENIZAÇÃO POR DANOS MORAIS", None, "R$ 10.000,00", "NÃO", None, "PROVÁVEL", None, "R$ 2.500,00"),
]
# Arquivados: (autor, natureza, ajuizamento, passivo, atualizado, depósito, valor do depósito, perda, pagamento)
LINHAS_ARQUIVADOS = [
    ("Pessoa Fictícia 0101", "INDENIZAÇÃO POR DANOS MORAIS E MATERIAIS", datetime.datetime(2023, 3, 14), "R$22.921,00", 41149.37, "NÃO", None,
     "REMOTA - Processo extinto sem resolução do mérito, com trânsito em julgado", 0.0),
    ("Pessoa Fictícia 0102", "TRABALHISTA ", datetime.datetime(2022, 11, 2), 43872.47, "R$ 48.665,84", "NÃO", None,
     "PROVÁVEL - Celebração de acordo", 6099.83),
    ("Pessoa Fictícia 0103", "INDENIZAÇÃO POR DANOS MORAIS", datetime.datetime(2024, 1, 9), "R$ 35.233,00", "R$ 7.352,00", "SIM", "R$ 7.352,00",
     "REMOTA - Realizado acordo entre as partes, sem encargo", None),
]
N_ATIVOS = len(LINHAS_ATIVOS)
N_ARQUIVADOS = len(LINHAS_ARQUIVADOS)
# número do processo de cada linha: ativos 1..N, arquivados 101..
NUMEROS_ATIVOS = [ficticio.numero_ficticio(i + 1) for i in range(N_ATIVOS)]
NUMEROS_ARQUIVADOS = [ficticio.numero_ficticio(101 + i) for i in range(N_ARQUIVADOS)]

GABARITO = {
    "processos": N_ATIVOS + N_ARQUIVADOS,
    "ativos": N_ATIVOS,
    "arquivados": N_ARQUIVADOS,
    # contagens de campo preenchido (depois da leitura): ativos + arquivados
    "passivo_potencial": N_ATIVOS - 1 + N_ARQUIVADOS,            # um "-" nos ativos
    "passivo_atualizado": N_ARQUIVADOS,
    "ativo_potencial": 2,                                          # R$ 53.839,36 e -R$ 500,00
    "provisao": 7,                                                 # nove ativos, dois sem provisão ("-" e vazio)
    "percentual_provisao": 7,                                      # 6 lidos da coluna "%"; o do processo 9 é derivado (provisão ÷ passivo)
    "deposito_judicial": N_ATIVOS - 1 + N_ARQUIVADOS + 1,        # um vazio informado, mas derivado do valor (processo 8)
    "depositos_recursais": 4,                                      # processos 3, 4, 8 e o arquivado 103
    "probabilidade": N_ATIVOS + N_ARQUIVADOS,
    "justificativa_probabilidade": 4,                              # processo 7 e os três arquivados
    "pagamento_realizado": 2,                                      # 0,00 e 6.099,83 (o terceiro está vazio)
    "valor_economizado": N_ARQUIVADOS,                             # fórmulas =G-K calculadas pelo programa
    "cnpj_processado": N_ATIVOS,
    "total_passivo_ativos": "695627.34",                           # soma dos passivos lidos dos ativos (8 valores)
}


def gravar(caminho):
    """Grava a planilha fictícia em `caminho` (Path) e a devolve."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = ABA_ATIVOS
    ws["A1"], ws["A2"], ws["A3"] = EMPRESA, "Acompanhamento de Passivos Contingentes", "Mês: Setembro- 2026"
    for j, c in enumerate(CABECALHO_ATIVOS, start=1):
        ws.cell(5, j, c)
    ws.freeze_panes = "A6"
    linha = 6
    for i, (autor, natureza, ativo_pot, passivo, dep, dep_valor, perda, pct, provisao) in enumerate(LINHAS_ATIVOS):
        valores = [autor, REU, cnpj_ficticio(i + 1), NUMEROS_ATIVOS[i], natureza, ativo_pot, passivo, dep,
                   dep_valor, perda, pct, provisao, f"Ação fictícia número {i + 1}, movida contra a empresa.", historico(i + 1)]
        for j, v in enumerate(valores, start=1):
            ws.cell(linha, j, v)
        if pct is not None:
            ws.cell(linha, 11).number_format = "0%"
        for j in (6, 7, 9, 12):
            ws.cell(linha, j).number_format = "[$R$ -416]#,##0.00"
        linha += 1
        if i == 1:                                     # linha com número inválido no meio dos dados
            ws.cell(linha, 1, "Pessoa Fictícia 0999")
            ws.cell(linha, 2, EMPRESA)
            ws.cell(linha, 4, "SEM NÚMERO AINDA")
            ws.cell(linha, 5, "COBRANÇA")
            ws.cell(linha, 7, "R$ 1,00")
            linha += 1
    linha += 1
    ws.cell(linha, 1, "TOTAL DE VALORES A RECEBER E A PAGAR")
    linha += 4
    ws.cell(linha, 1, "CRITÉRIOS NA COMPOSIÇÃO DAS CONTIGÊNCIAS")
    ws.cell(linha, 2, "CLASSIFICAÇÃO")
    ws.cell(linha, 4, "VALOR CONTINGENCIADO ")
    for criterio, grau, valor in (("Processo em fase de instrução", "POSSÍVEL ", "50% do valor da causa + juízo de proporcionalidade"),
                                  ("Processo com condenação em primeiro grau", "PROVÁVEL", "100% - Registro do valor da condenação"),
                                  ("Processo julgado improcedente", "REMOTA", 0.0)):
        linha += 1
        ws.cell(linha, 1, criterio)
        ws.cell(linha, 2, grau)
        ws.cell(linha, 4, valor)
    wa = wb.create_sheet(ABA_ARQUIVADOS)
    wa["A2"] = "AÇÕES FINALIZADAS "
    for j, c in enumerate(CABECALHO_ARQUIVADOS, start=1):
        wa.cell(3, j, c)
    linha = 4
    for i, (autor, natureza, ajuizamento, passivo, atualizado, dep, dep_valor, perda, pagamento) in enumerate(LINHAS_ARQUIVADOS):
        valores = [autor, EMPRESA, NUMEROS_ARQUIVADOS[i], natureza, ajuizamento, passivo, atualizado, dep, dep_valor, perda, pagamento,
                   f"=G{linha}-K{linha}"]
        for j, v in enumerate(valores, start=1):
            wa.cell(linha, j, v)
        wa.cell(linha, 5).number_format = "dd/mm/yyyy"
        linha += 1
    wa.cell(linha, 1, "VALOR TOTAL ECONOMIZADO")
    wa.cell(linha, 12, f"=SUM(L4:L{linha - 1})")
    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)
    return caminho
