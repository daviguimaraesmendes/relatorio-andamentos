"""Testes dos fluxos ponta a ponta (src/fluxos.py, WS-14): migrar, converter, inicial, atualizar e entregar, com o
coletor simulado, 200 processos e 5 clientes, sem rede e sem IA real.

Tudo é dado fictício e gerado em tempo de execução (tests/ficticio.py). O que NÃO é exercitado aqui: o ColetorReal
(certificado, jus.br, TRT), provedores de IA reais, Word, Excel e Google (só python-docx, openpyxl, LibreOffice e
Chromium offline).

    python3 -m unittest tests/test_fluxos.py -v
"""
import datetime
import json
import os
import shutil
import sys
import unittest
from decimal import Decimal
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent))
from isolamento import TMP  # noqa: E402  (antes de tudo)
import ficticio  # noqa: E402
import simulado  # noqa: E402
import comum  # noqa: E402
import ficha  # noqa: E402
import fluxos  # noqa: E402
import historico  # noqa: E402

D = Decimal
DATA_1, DATA_2, DATA_3 = "2026-07-31", "2026-08-31", "2026-09-30"


# ------------------------------------------------------------------ apoio

class Relogio:
    """Relógio falso da fila: `dormir` só adianta o relógio (nenhum teste espera de verdade)."""

    def __init__(self, inicio=datetime.datetime(2026, 10, 7, 22, 0)):
        self.agora = inicio

    def __call__(self):
        return self.agora

    def dormir(self, segundos):
        self.agora += datetime.timedelta(seconds=segundos)


def opcoes_da_fila(relogio=None, **extra):
    relogio = relogio or Relogio()
    return {"relogio": relogio, "dormir": relogio.dormir, "pausa_s": (0, 0), "janelas": ["20:00-06:00"],
            "espera_base_s": 60, "tentativas_max": 3, **extra}


class ColetorAte:
    """Coletor simulado que só enxerga o que existia até uma data: faz o tempo passar entre os ciclos mudando `ate`."""

    def __init__(self, base, ate):
        self.base, self.ate = base, ate
        self.completos = []            # números cujo coletar() terminou sem erro (cada um conta uma coleta de verdade)

    def coletar(self, processo, profundidade, desde):
        r = self.base.coletar(processo, profundidade, desde)
        if r.get("erro"):
            return r
        r = dict(r)
        r["movimentos"] = [m for m in r["movimentos"] if m["data"] <= self.ate]
        r["documentos"] = [d for d in r["documentos"] if d["data"] <= self.ate]
        self.completos.append(processo["numero"])
        return r

    def __getattr__(self, nome):
        return getattr(self.base, nome)


class ProvedorFalso:
    """Provedor de IA falso (CONTRATOS §8): conta chamadas e devolve um resumo que cita o início do documento."""

    def __init__(self, motor="local:falso"):
        self.motor, self.chamadas, self.clientes = motor, 0, []

    def gerar(self, sistema, usuario, *, esquema=None, cliente="", partes=()):
        self.chamadas += 1
        self.clientes.append(cliente)
        texto = usuario.split("<<<\n", 1)[1].rsplit("\n>>>", 1)[0] if "<<<\n" in usuario else ""
        trecho = " ".join(texto.split())[:90]
        dados = {"conteudo": "determinando o cumprimento do que foi decidido", "trecho_origem": trecho, "prazo": None,
                 "audiencia": None, "efeito": "neutro"}
        return {"texto": json.dumps(dados), "json": dados, "motor": self.motor, "avisos": []}


def magra(f):
    """A ficha 'magra' de uma lista de números: só número, cliente, polo, parte contrária e responsável."""
    m = ficha.nova_ficha(f["numero"])
    for campo in ("cliente", "polo_cliente", "parte_contraria", "responsavel"):
        v = ficha.obter(f, campo)
        if v:
            ficha.definir(m, campo, v, "migrado")
    return m


def aprovar_tudo(projeto_slug=None, so_numeros=None):
    """A pessoa que revisa: todo rascunho vira aprovado."""
    lista = comum.eventos()
    n = 0
    for e in lista:
        if e["status"] == "rascunho" and (so_numeros is None or e["numero"] in so_numeros):
            e["status"] = "aprovado"
            n += 1
    comum.salvar_eventos(lista)
    return n


def dec(v):
    return ficha.dinheiro(v) or D(0)


class Base(unittest.TestCase):
    N, CLIENTES = 200, 5

    def abrir_projeto(self, magras=True, n=None, clientes=None, semente=1, **kw):
        n = n or self.N
        self.verdade = ficticio.gerar_carteira(n, clientes=clientes or self.CLIENTES, semente=semente)
        fichas = [magra(f) for f in self.verdade] if magras else self.verdade
        shutil.rmtree(TMP / "projetos-de-teste", ignore_errors=True)
        self.proj = ficticio.criar_projeto_de_teste(fichas, nome="Relatório de Fluxos")
        self.addCleanup(ficticio.restaurar_comum)
        self.slug = self.proj["slug"]
        self.relogio = Relogio()
        self.simulado = simulado.ColetorSimulado(self.verdade, semente=1, taxa_falha=kw.pop("taxa_falha", 0.0),
                                                 pasta=self.proj["pasta"] / "docs-sim", **kw)
        self.coletor = ColetorAte(self.simulado, DATA_1)
        return self.proj

    def fila(self, **extra):
        return opcoes_da_fila(self.relogio, **extra)

    def inicial(self, **kw):
        kw.setdefault("profundidade", "padrao")
        kw.setdefault("modo", "imediato")
        kw.setdefault("data_base", DATA_1)
        kw.setdefault("coletor", self.coletor)
        kw.setdefault("fila_opcoes", self.fila())
        return fluxos.inicial(self.slug, **kw)

    def atualizar(self, arquivos=None, **kw):
        kw.setdefault("modo", "imediato")
        kw.setdefault("coletor", self.coletor)
        kw.setdefault("fila_opcoes", self.fila())
        return fluxos.atualizar(self.slug, arquivos, **kw)

    def fichas(self):
        with fluxos._em(self.slug):
            return ficha.carregar(todas=True)

    def eventos(self):
        with fluxos._em(self.slug):
            return comum.eventos()

    def codigos(self, r):
        return {a["codigo"] for a in r["avisos"]}


# ------------------------------------------------------------------ inicial

class Inicial(Base):
    N, CLIENTES = 12, 2

    def test_para_na_revisao_e_depois_entrega(self):
        self.abrir_projeto()
        r = self.inicial()
        self.assertTrue(r["ok"], r["resumo"])
        self.assertEqual(r["etapa"], "revisao")
        self.assertEqual(r["arquivos"], [])
        self.assertGreater(r["pendentes_de_revisao"], 0)
        eventos = self.eventos()
        self.assertEqual({e["status"] for e in eventos} - {"rascunho", "descartado"}, set())
        # nenhum evento nasce aprovado e todo rascunho traz motor e profundidade
        for e in eventos:
            self.assertEqual(e["profundidade"], "padrao")
            if e["status"] == "rascunho":
                self.assertIn("motor", e)
        # a capa foi gravada nas fichas, com origem coletado, e polo/parte contrária continuam os da lista
        for f in self.fichas():
            self.assertEqual(ficha.origem(f, "valor_causa"), "coletado")
            self.assertEqual(ficha.origem(f, "polo_cliente"), "migrado")
        r2 = self.inicial()                        # sem aprovar nada: continua parado na revisão, sem recoletar
        self.assertEqual(r2["etapa"], "revisao")
        self.assertEqual(len(self.coletor.completos), len(set(self.coletor.completos)), "ninguém foi recoletado")
        aprovar_tudo()
        r3 = self.inicial()
        self.assertEqual(r3["etapa"], "entregue", r3["resumo"])
        nomes = sorted(Path(a).name for a in r3["arquivos"])
        self.assertEqual(sum(n.endswith(".docx") for n in nomes), 2)
        self.assertTrue({"qualidade.html", "O que mudou neste ciclo.txt"} <= set(nomes))
        self.assertTrue(any(n.startswith("Planilha") for n in nomes) and any(n.startswith("Painel") for n in nomes))

    def test_tres_profundidades(self):
        quantidade = {}
        for prof in ("rapido", "padrao", "completo"):
            self.abrir_projeto(semente=2)
            r = self.inicial(profundidade=prof)
            self.assertEqual(r["etapa"], "revisao", (prof, r["resumo"]))
            docs = [e for e in self.eventos() if e["tipo_evento"] == "documento"]
            quantidade[prof] = len(docs)
            for e in self.eventos():
                self.assertEqual(e["profundidade"], prof)
            if prof == "rapido":
                self.assertEqual(docs, [], "no nível rápido nenhum documento é aberto")
            ficticio.restaurar_comum()
        self.assertEqual(quantidade["rapido"], 0)
        self.assertGreater(quantidade["padrao"], 0)
        self.assertGreaterEqual(quantidade["completo"], quantidade["padrao"])

    def test_momento_atual_ativo_e_ultimo_andamento(self):
        self.abrir_projeto()
        self.inicial()
        com_momento = [f for f in self.fichas() if ficha.obter(f, "momento_atual")]
        self.assertGreater(len(com_momento), 0)
        for f in com_momento:
            self.assertEqual(ficha.origem(f, "momento_atual"), "coletado")
            import taxonomia
            ativo = taxonomia.momento_ativo(ficha.obter(f, "momento_atual"))
            if ativo is not None:
                self.assertEqual(f["ativo"], ativo)
            self.assertEqual(ficha.validar(f), [])
        import traduzir
        for f in self.fichas():
            evs = [e for e in self.eventos() if e["numero"] in ficha.todos_os_numeros(f)]
            relevantes = [e for e in evs if e["tipo_evento"] == "documento" or traduzir.traduzir_movimento(e["titulo"])[1]]
            if evs:             # cálculo independente: a maior data entre documentos e movimentos que não são de rotina
                datas = [ficha.parse_data(e["data"]) for e in (relevantes or evs)]
                self.assertEqual(ficha.obter(f, "ultimo_andamento"), max(datas), f["numero"])

    def test_julgamento_so_sugerido_e_so_depois_da_revisao(self):
        self.abrir_projeto()
        self.inicial()
        for f in self.fichas():          # antes da revisão nenhum campo de julgamento foi tocado
            for c in ficha.CAMPOS_DE_JULGAMENTO:
                self.assertIsNone(ficha.obter(f, c))
        humano = self.fichas()[0]
        humano_numero = humano["numero"]
        with fluxos._em(self.slug):
            fs = ficha.carregar(todas=True)
            ficha.definir(fs[0], "probabilidade", "Remota", "humano")
            ficha.definir(fs[0], "resultado", "Improcedente", "humano")
            ficha.salvar(fs)
        aprovar_tudo()
        r = self.inicial()
        self.assertEqual(r["etapa"], "entregue", r["resumo"])
        f0 = next(f for f in self.fichas() if f["numero"] == humano_numero)
        self.assertEqual((ficha.origem(f0, "probabilidade"), ficha.obter(f0, "probabilidade")), ("humano", "Remota"))
        self.assertEqual((ficha.origem(f0, "resultado"), ficha.obter(f0, "resultado")), ("humano", "Improcedente"))
        origens = {ficha.origem(f, c) for f in self.fichas() for c in ficha.CAMPOS_DE_JULGAMENTO} - {None}
        self.assertTrue(origens <= {"sugerido", "humano"}, origens)

    def test_so_quem_nao_tem_relatorio_e_coletado(self):
        self.abrir_projeto()
        with fluxos._em(self.slug):
            fs = ficha.carregar(todas=True)
            for f in fs[:4]:                                   # estes já têm relatório anterior
                f["linha_de_base"] = {"data_base": "2026-06-30", "andamentos_texto": "Em 01/06/2026 foi feito algo.",
                                      "arquivo": "x.docx", "ultimo_andamento": "2026-06-01"}
            ficha.salvar(fs)
        r = self.inicial()
        self.assertEqual(r["processos"], self.N - 4)
        self.assertEqual(len(set(self.coletor.completos)), self.N - 4)
        r = self.inicial(todos=True, data_base="2026-07-31")      # todos os que seguem ativos (encerrado não é acompanhado)
        self.assertEqual(r["processos"], sum(1 for f in self.fichas() if f.get("ativo", True)))

    def test_conferir_manualmente_nao_trava_o_resto(self):
        self.abrir_projeto(falhar_em={})
        alvo = [f["numero"] for f in self.verdade]
        self.simulado.falhar_em = {alvo[0]: "captcha", alvo[1]: "segredo", alvo[2]: ("timeout", 1), alvo[3]: ("timeout", 9)}
        r = self.inicial()
        manual = {c["numero"]: c["codigo"] for c in r["conferir_manualmente"]}
        self.assertEqual(manual.get(alvo[0]), "captcha")
        self.assertEqual(manual.get(alvo[1]), "segredo")
        self.assertIn(alvo[3], manual)                            # timeout que não passa vira conferência depois das tentativas
        self.assertNotIn(alvo[2], manual)                          # timeout passageiro: repetiu e deu certo
        self.assertEqual(r["etapa"], "revisao")
        self.assertGreater(r["fila"]["coletado"], self.N - 6)
        aprovar_tudo()
        r = self.inicial()
        self.assertEqual(r["etapa"], "entregue")
        guardado = fluxos.ultimo_ciclo(self.slug)
        self.assertEqual({c["numero"] for c in guardado["conferir_manualmente"]}, set(manual))



# ------------------------------------------------------------------ entregas A, B e C com 200 processos

def _abrir_navegador():
    """(playwright, navegador) ou (None, None) se o Chromium não estiver disponível."""
    try:
        from playwright.sync_api import sync_playwright
        pw = sync_playwright().start()
        return pw, pw.chromium.launch()
    except Exception:  # noqa: BLE001
        return None, None


class Entregas200(Base):
    N, CLIENTES = 200, 5

    def test_a_b_e_c_abrem_e_batem_com_calculo_independente(self):
        import openpyxl
        import docx
        from escritores import docx_a
        self.abrir_projeto()
        self.assertEqual(self.inicial()["etapa"], "revisao")
        aprovar_tudo()
        antes = self.eventos()
        r = self.inicial(entregas=["docx_a", "xlsx_b", "dashboard"])
        self.assertEqual(r["etapa"], "entregue", r["resumo"])
        self.assertEqual([a for a in r["avisos"] if a["nivel"] == "erro"], [])
        fichas = self.fichas()
        por_numero = {f["numero"]: f for f in fichas}
        self.assertEqual(len(fichas), self.N)
        arquivos = [Path(a) for a in r["arquivos"]]
        docxs = sorted(a for a in arquivos if a.suffix == ".docx")
        xlsx = next(a for a in arquivos if a.suffix == ".xlsx")
        html = next(a for a in arquivos if a.suffix == ".html" and a.name.startswith("Painel"))
        self.assertEqual(len(docxs), self.CLIENTES)
        # --- cálculo independente, a partir das fichas gravadas (não dos escritores)
        total = len(fichas)
        ativos = sum(1 for f in fichas if f.get("ativo", True))
        soma_causa = sum((dec(ficha.obter(f, "valor_causa")) for f in fichas), D(0))
        soma_causa_ativos = sum((dec(ficha.obter(f, "valor_causa")) for f in fichas if f.get("ativo", True)), D(0))
        # --- A: um bloco por processo, com o momento no título e as datas dos eventos aprovados no texto
        vistos = set()
        for caminho in docxs:
            doc = docx.Document(caminho)                                  # abre em python-docx
            est = docx_a.ler_estrutura(caminho)
            cliente = est["cliente"]
            esperados = {f["numero"] for f in fichas if ficha.obter(f, "cliente") == cliente}
            achados = {b["numeros"][0] for b in est["processos"]}
            self.assertEqual(achados, esperados, cliente)
            self.assertEqual(len(doc.tables), 1 + len(esperados))        # quadro-resumo + um bloco por processo
            vistos |= achados
            for b in est["processos"]:
                f = por_numero[b["numeros"][0]]
                momento = ficha.obter(f, "momento_atual")
                if momento:
                    self.assertEqual(b["momento"], momento, f["numero"])
                for e in antes:
                    if e["numero"] in ficha.todos_os_numeros(f) and e["status"] == "aprovado":
                        self.assertIn(e["data"], b["andamentos_texto"], (f["numero"], e["titulo"]))
        self.assertEqual(vistos, set(por_numero))
        # --- B: uma linha por processo; soma de Valor da Causa igual à das fichas; texto com as datas
        wb = openpyxl.load_workbook(xlsx, data_only=True)
        ws = wb["Processos"]
        cab = [c.value for c in ws[1]]
        col = {nome: i for i, nome in enumerate(cab)}
        linhas = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[col["Número do Processo"]]]
        self.assertEqual(len(linhas), total)
        self.assertEqual({l[col["Número do Processo"]] for l in linhas}, set(por_numero))
        self.assertEqual(sum((D(str(l[col["Valor da Causa"]] or 0)) for l in linhas), D(0)).quantize(D("0.01")), soma_causa)
        for l in linhas:
            f = por_numero[l[col["Número do Processo"]]]
            texto = l[col["Andamentos"]] or ""
            for e in antes:
                if e["numero"] in ficha.todos_os_numeros(f) and e["status"] == "aprovado":
                    self.assertIn(e["data"], texto, (f["numero"], e["titulo"]))
        # --- C: abre no navegador, sem rede, e os números do painel são os mesmos
        pw, navegador = _abrir_navegador()
        if pw is None:
            self.skipTest("Chromium indisponível neste ambiente")
        self.addCleanup(pw.stop)
        ctx = navegador.new_context()
        externas = []
        ctx.route("**/*", lambda rota: rota.continue_() if rota.request.url.startswith(("file:", "data:", "blob:", "about:"))
                  else (externas.append(rota.request.url), rota.abort()))
        pagina = ctx.new_page()
        erros = []
        pagina.on("pageerror", lambda e: erros.append(str(e)))
        pagina.goto(html.resolve().as_uri())
        pagina.wait_for_function("document.documentElement.dataset.pronto === '1'", timeout=60000)
        ind = pagina.evaluate("JSON.parse(JSON.stringify(DASH.estado.ind))")
        ctx.close()
        navegador.close()
        self.assertEqual((externas, erros), ([], []))
        self.assertEqual((ind["total"], ind["ativos"], ind["encerrados"]), (total, ativos, total - ativos))
        self.assertAlmostEqual(ind["valorCausaAtivos"], float(soma_causa_ativos), delta=0.01)

    def test_qualidade_o_que_mudou_e_retrato_mensal(self):
        self.abrir_projeto()
        self.inicial()
        aprovar_tudo()
        r = self.inicial()
        self.assertEqual(r["etapa"], "entregue", r["resumo"])
        self.assertIsNotNone(r["qualidade"])
        self.assertIn("O que mudou", Path(next(a for a in r["arquivos"] if "mudou" in str(a))).read_text(encoding="utf-8"))
        depois = r["o_que_mudou"]["totais"]["depois"]
        self.assertEqual((depois["processos"], depois["ativos"]), (self.N, sum(1 for f in self.fichas() if f.get("ativo", True))))
        retrato = Path(r["retrato"])
        dados = json.loads(retrato.read_text(encoding="utf-8"))
        fichas = self.fichas()
        self.assertEqual(dados["data_base"], DATA_1)
        self.assertEqual(dados["totais"]["processos"], len(fichas))
        self.assertEqual(D(dados["totais"]["valor_causa"]), sum((dec(ficha.obter(f, "valor_causa")) for f in fichas), D(0)))
        self.assertEqual(len(dados["por_processo"]), len(fichas))
        with fluxos._em(self.slug):
            self.assertEqual([x["data_base"] for x in historico.carregar(self.slug)], [DATA_1])

if __name__ == "__main__":
    unittest.main()
