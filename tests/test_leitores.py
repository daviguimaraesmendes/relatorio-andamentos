"""Leitores de relatórios existentes (WS-2): detecção de formato, modelo A (.docx), modelo B (.xlsx), listas e tabelas
fora do padrão. Só dados fictícios, sem rede; números de processo vêm de `ficticio.numero_ficticio` (nada literal).

    python3 -m unittest tests/test_leitores.py -v

Usa os protótipos dos spikes S1/S2 (`spikes/`) para gerar fixtures e para a ida e volta (ler -> regravar -> ler); esses
testes são pulados se `spikes/` ou o python-docx não existirem. O Excel, o Word e o Google Docs/Sheets reais NÃO são
exercitados aqui: só openpyxl, python-docx, lxml e (quando há) o LibreOffice.
"""
import copy
import datetime
import importlib.util
import json
import random
import re
import shutil
import sys
import tempfile
import time
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import isolamento  # noqa: F401,E402
import ficticio  # noqa: E402  (já põe src/ no sys.path)

import carteira as cart  # noqa: E402
import ficha  # noqa: E402
import leitores  # noqa: E402
from leitores import base, grade  # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent
SPIKES = RAIZ / "spikes"
TMP = Path(tempfile.mkdtemp(prefix="leitores-"))
try:
    import docx as python_docx  # noqa: E402
except ImportError:      # pragma: no cover
    python_docx = None
try:
    from lxml import etree  # noqa: E402
except ImportError:      # pragma: no cover
    etree = None


def tearDownModule():
    shutil.rmtree(TMP, ignore_errors=True)


def _carregar_spike(nome, caminho):
    spec = importlib.util.spec_from_file_location(nome, caminho)
    modulo = importlib.util.module_from_spec(spec)
    sys.modules[nome] = modulo
    spec.loader.exec_module(modulo)
    return modulo


if (SPIKES / "s2_docx").exists() and (SPIKES / "s1_xlsx").exists():
    s2_gm = _carregar_spike("s2_gerar_modelo", SPIKES / "s2_docx" / "gerar_modelo.py")
    s2_da = _carregar_spike("s2_docx_atualizador", SPIKES / "s2_docx" / "docx_atualizador.py")
    s1_gm = _carregar_spike("s1_gerar_modelo", SPIKES / "s1_xlsx" / "gerar_modelo.py")
    s1_xc = _carregar_spike("s1_xlsx_cirurgico", SPIKES / "s1_xlsx" / "xlsx_cirurgico.py")
else:      # pragma: no cover
    s2_gm = s2_da = s1_gm = s1_xc = None

COM_SPIKES = s2_da is not None and python_docx is not None
SOFFICE = shutil.which("soffice")

_CACHE = {}


def fichas_200():
    if "fichas" not in _CACHE:
        _CACHE["fichas"] = ficticio.gerar_carteira(200, clientes=5, semente=1, com_linha_de_base=True)
    return _CACHE["fichas"]


def pasta(nome):
    p = TMP / nome
    p.mkdir(parents=True, exist_ok=True)
    return p


def codigos(rel, nivel=None):
    return [a["codigo"] for a in rel["avisos"] if nivel is None or a["nivel"] == nivel]


def avisos_de(rel, codigo):
    return [a for a in rel["avisos"] if a["codigo"] == codigo]


def valor(p, campo):
    c = p["campos"].get(campo)
    return c["valor"] if c else None


def por_numero(rel):
    return {p["numero"]: p for p in rel["processos"]}


def norm_espacos(texto):
    return " ".join(str(texto or "").split())


def norm_andamento(texto):
    """Espaços colapsados e sem a vírgula logo depois da data ('Em 18/06/2026, foi ...' == 'Em 18/06/2026 foi ...')."""
    return re.sub(r"(\d{2}/\d{2}/\d{4}),", r"\1", norm_espacos(texto))


# ---------------------------------------------------------------- fabricação de documentos (modelo A)

def montar_docx(destino, cliente, data_base, blocos, resumo=None, embrulhar=False):
    """.docx do modelo A com python-docx. blocos: [{"titulo", "campos": [(rótulo, valor)], "andamentos": [(texto, negrito)]}]
    resumo: [(numeros, assunto, momento, ultimo)] ou None (sem quadro-resumo). embrulhar: cada bloco dentro de uma
    tabela de 1 célula (tabela aninhada)."""
    d = python_docx.Document()
    if cliente:
        d.add_paragraph(cliente)
    if data_base:
        p = d.add_paragraph()
        p.add_run("Data-Base: ").bold = True
        p.add_run(data_base)
    d.add_paragraph("")
    if resumo is not None:
        t = d.add_table(rows=1, cols=4)
        for c, rotulo in zip(t.rows[0].cells, ["Nº DO PROCESSO", "ASSUNTO", "MOMENTO ATUAL DO PROCESSO", "ÚLTIMO ANDAMENTO"]):
            c.text = rotulo
        for numeros, assunto, momento, ultimo in resumo:
            cels = t.add_row().cells
            cels[0].text = "\n".join(numeros) if isinstance(numeros, (list, tuple)) else numeros
            cels[1].text, cels[2].text, cels[3].text = assunto, momento, ultimo
        d.add_paragraph("")
    for b in blocos:
        alvo = d
        if embrulhar:
            envolta = d.add_table(rows=1, cols=1)
            alvo = envolta.rows[0].cells[0]
        t = alvo.add_table(rows=0, cols=2) if embrulhar else d.add_table(rows=0, cols=2)
        linha = t.add_row()
        titulo = linha.cells[0].merge(linha.cells[1])
        titulo.text = b["titulo"]
        for rotulo, v in b["campos"]:
            cels = t.add_row().cells
            cels[0].text, cels[1].text = rotulo, v
        if b.get("andamentos") is not None:
            cels = t.add_row().cells
            cels[0].text = "Andamentos:"
            p = cels[1].paragraphs[0]
            for texto, negrito in b["andamentos"]:
                p.add_run(texto).bold = negrito
        if not embrulhar:
            d.add_paragraph("")
        else:
            alvo.add_paragraph("")
    destino = Path(destino)
    d.save(str(destino))
    return destino


def andamento_runs(*itens):
    """[(data_iso, texto)] -> runs 'Em ' + data em negrito + ', texto.' com espaço entre as frases."""
    runs = []
    for i, (data, texto) in enumerate(itens):
        runs += [(("" if i == 0 else " ") + "Em ", False), (ficha.data_br(data), True), (f", {texto}", False)]
    return runs


def bloco_basico(numeros_titulo, momento="AGUARDANDO SENTENÇA", assunto="Cobrança", extra_campos=(), andamentos=None, fecho=None):
    runs = andamento_runs(*(andamentos or [("2026-03-02", "foi proferido despacho inicial."), ("2026-04-20", "foi apresentada contestação.")]))
    if fecho:
        runs += [(" Em ", False), (ficha.data_br(fecho), True), (", sem atualizações.", False)]
    return {"titulo": f"{numeros_titulo} [ {momento} ]" if momento else numeros_titulo,
            "campos": [("Assunto", assunto), ("Autor(es)", "Pessoa Fictícia 0001"), ("Réu(s)", "Cliente Exemplo 01 Ltda"),
                       ("Ajuizamento", "15/02/2026"), ("Valor da Causa", "R$ 25.000,00"), ("Data de citação", "02/03/2026"),
                       ("Juízo", "1ª Vara Cível da Comarca de Cidade Modelo"), ("Área do Direito", "Cível"),
                       ("Matéria Principal", "Responsabilidade civil"), *extra_campos],
            "andamentos": runs}


def reescrever_docx(origem, destino, fn):
    """Reescreve word/document.xml aplicando fn(texto_xml) -> texto_xml; o resto do pacote fica igual."""
    with zipfile.ZipFile(origem) as zin, zipfile.ZipFile(destino, "w") as zout:
        for info in zin.infolist():
            dados = zin.read(info.filename)
            if info.filename == "word/document.xml":
                dados = fn(dados.decode("utf-8")).encode("utf-8")
            zout.writestr(info, dados)
    return destino


# ---------------------------------------------------------------- fabricação de planilhas (modelo B)

CABECALHO_B = list(grade.CABECALHOS_B)
CAMPO_DA_COLUNA_B = {
    "Autor(es)": "autores", "Réu(s)": "reus", "Vara": "vara", "Município": "municipio", "Data do Ajuizamento": "data_ajuizamento",
    "Área do Direito": "area", "Matéria Principal": "materia_principal", "Objeto": "objeto", "Valor da Causa": "valor_causa",
    "Situação": "situacao", "Valor Arbitrado em Juízo": "valor_arbitrado", "Probabilidade": "probabilidade",
    "Valor Estimado": "valor_estimado", "Valor da Execução": "valor_execucao", "Custas Processuais": "custas",
    "Depósitos Recursais": "depositos_recursais", "Garantias Processuais": "garantias", "Resultado": "resultado",
    "Data do trânsito em julgado": "data_transito", "Houve recurso da empresa?": "houve_recurso",
    "Reclamante terceirizado?": "terceirizado", "Outra(s) Parte(s)": "outras_partes"}
ROTULO_DO_TIPO = {"agravo": "Agravo", "apenso": "Apenso", "recurso": "Recurso"}


def _celula_numeros(f):
    partes = [f["numero"]] + [f"{ROTULO_DO_TIPO[v['tipo']]}: {v['numero']}" for v in f["vinculados"]]
    return "\n".join(partes)


def _formatar_data(i, iso):
    d = datetime.date.fromisoformat(iso)
    return [d, d.strftime("%d/%m/%Y"), (d - datetime.date(1899, 12, 30)).days][i % 3]


def _formatar_dinheiro(i, texto):
    v = float(texto)
    br = f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return [v, f"R$ {br}", br][i % 3]


def gerar_xlsx_b(destino, fichas, abas=("Ativos", "Arquivados"), titulo_antes=True, renomear=None, extras=False):
    """Planilha do modelo B fictícia, com formatos misturados (data como data/texto/série, dinheiro como número/R$/BR),
    várias abas, título e linha em branco acima do cabeçalho, linha-marcador e linha de total, aba de parâmetros.
    `renomear`: {coluna do modelo B: outro nome} para simular cabeçalhos diferentes; `extras`: duas colunas sem destino."""
    from openpyxl import Workbook
    cabecalho = CABECALHO_B
    renomear = renomear or {}
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Parâmetros"
    ws0.append(["Parâmetro", "Valor"])
    ws0.append(["Número de funcionários", 1200])
    ws0.append(["Data de referência", datetime.date(2026, 9, 30)])
    ws0.append(["Empresas do grupo", "Empresa Fictícia Alfa Ltda; Empresa Fictícia Beta S.A."])
    ws0.append(["Fator de correção", "1,0523"])
    ws0.append(["Observação livre", "texto qualquer"])
    ativos = [f for f in fichas if f.get("ativo", True)]
    encerrados = [f for f in fichas if not f.get("ativo", True)]
    for nome, grupo in zip(abas, (ativos, encerrados)):
        ws = wb.create_sheet(nome)
        if titulo_antes:
            ws.append([f"Relatório de contencioso - {nome}"])
            ws.append([])
        ws.append([renomear.get(c, c) for c in cabecalho] + (["Observação interna", "Código interno"] if extras else []))
        for i, f in enumerate(grupo):
            linha = []
            for col in cabecalho:
                if col == "Número do Processo":
                    linha.append(_celula_numeros(f))
                elif col == "Tribunal":
                    linha.append(f["tribunal"])
                elif col == "Ativo":
                    linha.append("Sim" if f.get("ativo", True) else "Não")
                elif col == "Andamentos":
                    texto = (f.get("linha_de_base") or {}).get("andamentos_texto", "")
                    linha.append(texto + (" Até 30/09/2026 sem atualizações." if i % 3 == 0 else ""))
                elif col in CAMPO_DA_COLUNA_B:
                    campo = CAMPO_DA_COLUNA_B[col]
                    v = ficha.obter(f, campo)
                    if v is None:
                        linha.append(None)
                    elif ficha.CAMPOS[campo][2] == "data":
                        linha.append(_formatar_data(i, v))
                    elif ficha.CAMPOS[campo][2] == "dinheiro":
                        linha.append(_formatar_dinheiro(i, v))
                    elif ficha.CAMPOS[campo][2] == "sim_nao":
                        linha.append({"Sim": ["Sim", "S", "sim"], "Não": ["Não", "N", "nao"]}[v][i % 3])
                    else:
                        linha.append(v)
                else:
                    linha.append(None)       # Valor Economizado, Taxa de resolução, Percentual: vazios
            if extras:
                linha += [f"obs {i}", f"C-{i:04d}"]
            ws.append(linha)
        ws.append([])
        ws.append(["--- fim da lista ---"])
        ws.append(["Total", None, None, f"=COUNTA(A1:A{ws.max_row})"])
    wb.save(str(destino))
    return destino


CAMPOS_COMPARAVEIS_B = tuple(CAMPO_DA_COLUNA_B.values())


def esperado_do_ficha(f, campos):
    """{campo: valor da ficha} dos campos pedidos que a ficha tem preenchidos."""
    return {c: ficha.obter(f, c) for c in campos if ficha.obter(f, c) not in (None, "")}


def lido_do_processo(p, campos):
    return {c: p["campos"][c]["valor"] for c in campos if c in p["campos"]}


# ================================================================ unidades: conversão de valores e mapeador

class TestConversoes(unittest.TestCase):
    def test_datas(self):
        casos = {"18/09/2026": "2026-09-18", "18/9/2026": "2026-09-18", "18.09.2026": "2026-09-18", "2026-09-18": "2026-09-18",
                 "18/09/26": "2026-09-18", "18 de setembro de 2026": "2026-09-18", "distribuído em 18/09/2026": "2026-09-18",
                 datetime.date(2026, 9, 18): "2026-09-18", datetime.datetime(2026, 9, 18, 10, 30): "2026-09-18",
                 (datetime.date(2026, 9, 18) - datetime.date(1899, 12, 30)).days: "2026-09-18", "-": None, "": None}
        for entrada, esperado in casos.items():
            self.assertEqual(base.converter_data(entrada)[0], esperado, entrada)
        self.assertEqual(base.converter_data("31/02/2026")[1], "invalida")
        self.assertEqual(base.converter_data("em breve")[1], "invalida")
        self.assertEqual(base.converter_data("a definir"), (None, None))        # placeholder: vazio, sem aviso
        self.assertEqual(base.converter_data("10/01/2026 e 12/02/2026")[1], "ambigua")
        self.assertEqual(base.converter_data(2026)[1], "invalida")        # um ano solto não é data

    def test_dinheiro(self):
        casos = {"R$ 1.234,56": "1234.56", "1.234,56": "1234.56", "1234.56": "1234.56", 1234.5: "1234.50", "R$ 0,00": "0.00",
                 "(1.234,56)": "-1234.56", "R$ 10.788.082,61": "10788082.61", "-": None, "": None, "N/A": None}
        for entrada, esperado in casos.items():
            self.assertEqual(base.converter_dinheiro(entrada)[0], esperado, entrada)
        self.assertEqual(base.converter_dinheiro("R$ 1.000,00 e R$ 500,00")[1], "ambiguo")
        self.assertEqual(base.converter_dinheiro("ilíquido")[1], "invalido")

    def test_campo_fora_do_vocabulario_nao_e_gravado(self):
        r = base.converter_campo("resultado", "Zebra azul")
        self.assertIsNone(r["valor"])
        self.assertEqual(r["avisos"][0][1], "rotulo_fora_do_vocabulario")
        self.assertEqual(base.converter_campo("resultado", "Parcial procedência")["valor"], "Parcialmente procedente")
        self.assertEqual(base.converter_campo("polo_cliente", "Réu")["valor"], "passivo")

    def test_momento_com_qualificador(self):
        r = base.converter_campo("momento_atual", "CUMPRIMENTO DE SENTENÇA (HONORÁRIOS SUSPENSOS)")
        self.assertEqual(r["valor"], "CUMPRIMENTO DE SENTENÇA")
        self.assertEqual(r["extras"]["momento_qualificador"], "HONORÁRIOS SUSPENSOS")
        self.assertEqual(base.converter_campo("momento_atual", "PROCESSO ARQUIVADO (DECISÃO FAVORÁVEL)")["valor"], "PROCESSO ARQUIVADO")
        # palavra a mais no rótulo conhecido: não perde detalhe em silêncio
        r = base.converter_campo("momento_atual", "AGUARDANDO JULGAMENTO DO AGRAVO")
        self.assertIsNone(r["valor"])
        self.assertEqual(r["avisos"][0][1], "momento_fora_do_vocabulario")
        self.assertTrue(r["avisos"][0][3])        # candidatos

    def test_percentual_vira_fracao(self):
        for entrada in (0.25, "25%", 25, "25,0 %"):
            self.assertAlmostEqual(base.converter_campo("percentual_exito", entrada)["valor"], 0.25, msg=str(entrada))

    def test_sim_nao(self):
        for entrada, esperado in {"Sim": "Sim", "S": "Sim", "não": "Não", "N": "Não", True: "Sim", False: "Não"}.items():
            self.assertEqual(base.converter_campo("houve_recurso", entrada)["valor"], esperado, entrada)

    def test_andamentos_ultimo_e_fecho(self):
        texto = ("Em 02/03/2026, foi proferido despacho. Em 20/04/2026 a ré contestou, com audiência em 22/10/2026. "
                 "No dia 10/05/2026, juntada de petição. Em 18/09/2026, sem atualizações.")
        r = base.analisar_andamentos(texto)
        self.assertEqual(r["ultimo"], "2026-05-10")          # nem o fecho nem a audiência futura contam
        self.assertEqual(r["fecho"], "2026-09-18")
        self.assertNotIn("sem atualiza", r["texto"])
        self.assertEqual([a["data"] for a in r["andamentos"]], ["2026-03-02", "2026-04-20", "2026-05-10"])
        r = base.analisar_andamentos("Em 01/02/2026 distribuído. Até 30/09/2026 sem andamentos.")
        self.assertEqual((r["ultimo"], r["fecho"]), ("2026-02-01", "2026-09-30"))
        r = base.analisar_andamentos("Sem atualizações até 30/09/2026.")
        self.assertEqual(r["fecho"], "2026-09-30")
        self.assertIsNone(r["ultimo"])
        r = base.analisar_andamentos("Audiência realizada em 10/03/2026 e sentença em 05/05/2026.")
        self.assertEqual(r["ultimo"], "2026-05-05")
        self.assertIn("andamentos_sem_marcador", [a[1] for a in r["avisos"]])

    def test_numeros(self):
        n0, n1, n2 = (ficticio.numero_ficticio(i) for i in range(3))
        principal, vinculados, av = base.interpretar_numeros(f"PROCESSO Nº {n0} / AGRAVO DE INSTRUMENTO Nº {n1} / APENSO Nº {n2}", "x")
        self.assertEqual(principal, n0)
        self.assertEqual(vinculados, [{"numero": n1, "tipo": "agravo"}, {"numero": n2, "tipo": "apenso"}])
        self.assertEqual(av, [])
        principal, vinculados, av = base.interpretar_numeros(f"{n0} ({n1})", "x")
        self.assertEqual(vinculados[0]["tipo"], "apenso")
        self.assertEqual(av[0]["codigo"], "vinculo_tipo_indefinido")
        principal, vinculados, av = base.interpretar_numeros(f"{ficticio.numero_com_dv_errado(n0)} e {n1}", "x")
        self.assertIsNone(principal)
        self.assertEqual(av[0]["codigo"], "numero_dv_invalido")
        sem_mascara = n0.replace("-", "").replace(".", "")
        self.assertEqual(base.interpretar_numeros(sem_mascara, "x")[0], n0)


class TestMapeador(unittest.TestCase):
    def test_cabecalhos_do_modelo_b_casam_exatamente(self):
        regs = grade.propor_mapeamento(CABECALHO_B)
        for r in regs:
            self.assertTrue(r["aplicado"], r["coluna"])
            self.assertEqual(r["confianca"], 1.0, r["coluna"])
        destinos = {r["campo"] for r in regs}
        self.assertEqual(destinos, {*CAMPO_DA_COLUNA_B.values(), "numero", "tribunal", "ativo", "andamentos", "valor_economizado",
                                    "taxa_resolucao_dias", "percentual_exito"})

    def test_sinonimos_e_variacoes(self):
        casos = {"Nº do processo": "numero", "N° Processo": "numero", "NUMERO CNJ": "numero", "Requerente": "autores",
                 "Reclamada(s)": "reus", "Vlr. Causa (R$)": "valor_causa", "Valor da Cauza": "valor_causa",
                 "Data de distribuição": "data_ajuizamento", "Juízo": "vara", "Comarca": "municipio", "Status": "situacao",
                 "Histórico": "andamentos", "Obs.": "observacoes", "Último andamento": "ultimo_andamento",
                 "Momento atual do processo": "momento_atual", "Dep. recursais": None}
        for cab, esperado in casos.items():
            reg = grade.propor_mapeamento([cab])[0]
            if esperado:
                self.assertEqual(reg["campo"], esperado, cab)
                self.assertGreaterEqual(reg["confianca"], 0.7, cab)
            else:
                self.assertFalse(reg["aplicado"] and reg["campo"] != "depositos_recursais", cab)

    def test_cabecalho_ambiguo_nao_e_aplicado(self):
        reg = grade.propor_mapeamento(["Valor"])[0]
        self.assertIsNone(reg["campo"])
        self.assertTrue(reg["ambigua"])
        self.assertGreaterEqual(len(reg["candidatos"]), 2)
        reg = grade.propor_mapeamento(["Data"])[0]
        self.assertFalse(reg["aplicado"])

    def test_coluna_duplicada(self):
        regs = grade.propor_mapeamento(["Processo", "Valor da causa", "Valor da causa"])
        campos = [r["campo"] for r in regs]
        self.assertEqual(campos.count("valor_causa"), 1)
        self.assertEqual(regs[2]["duplicada_de"], 1)

    def test_numero_pelo_conteudo(self):
        cnjs = [ficticio.numero_ficticio(i) for i in range(5)]
        regs = grade.propor_mapeamento(["Ref. interna", "Quem"], [cnjs, ["a", "b", "c", "d", "e"]])
        self.assertEqual(regs[0]["campo"], "numero")
        self.assertIsNone(regs[1]["campo"])

    def test_tipo_do_conteudo_derruba_confianca(self):
        regs = grade.propor_mapeamento(["Data do ajuizamento"], [["abc", "def", "ghi", "jkl"]])
        self.assertLess(regs[0]["confianca"], 0.7)
        self.assertFalse(regs[0]["aplicado"])


# ================================================================ modelo A (.docx)

@unittest.skipUnless(COM_SPIKES, "precisa de spikes/ e python-docx")
class TestDocxA(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.pasta = pasta("docx")
        cls.dados = s2_gm.processos()
        cls.arquivo = cls.pasta / "modelo.docx"
        s2_gm.gerar(cls.arquivo, semente=3)
        cls.rel = leitores.ler(cls.arquivo)

    def test_formato_cabecalho_e_quantidade(self):
        self.assertEqual(leitores.detectar(self.arquivo), "docx_a")
        self.assertEqual(self.rel["formato"], "docx_a")
        self.assertEqual(self.rel["cliente"], s2_gm.CLIENTE)
        self.assertEqual(self.rel["data_base"], "2026-09-18")
        self.assertEqual(len(self.rel["processos"]), 12)
        self.assertEqual(self.rel["arquivo"], "modelo.docx")
        for chave in ("formato", "arquivo", "cliente", "data_base", "processos", "parametros", "colunas_sem_destino", "avisos"):
            self.assertIn(chave, self.rel)

    def test_numeros_multiplos_e_tipos(self):
        esperado = self.dados[1]["numeros"]
        p = self.rel["processos"][1]
        self.assertEqual(p["numero"], esperado[0])
        self.assertEqual(p["vinculados"], [{"numero": esperado[1], "tipo": "agravo"}, {"numero": esperado[2], "tipo": "apenso"}])
        self.assertTrue(all(len(q["vinculados"]) == 0 for i, q in enumerate(self.rel["processos"]) if i != 1))

    def test_campos_de_cada_bloco(self):
        for dado, p in zip(self.dados, self.rel["processos"]):
            self.assertEqual(p["numero"], dado["numeros"][0])
            self.assertEqual(valor(p, "assunto"), dado["assunto"])
            self.assertEqual(valor(p, "autores"), dado["autores"])
            self.assertEqual(valor(p, "reus"), dado["reus"])          # inclui o caso de célula mesclada na vertical
            self.assertEqual(valor(p, "data_ajuizamento"), "2026-02-15")
            self.assertEqual(valor(p, "vara"), dado["juizo"])
            self.assertEqual(valor(p, "area"), dado["area"])
            self.assertEqual(valor(p, "materia_principal"), dado["materia"])
            self.assertEqual(valor(p, "cliente"), s2_gm.CLIENTE)
            if dado["valor_causa"]:
                self.assertEqual(valor(p, "valor_causa"), ficha.parse_dinheiro(dado["valor_causa"]))
            else:
                self.assertNotIn("valor_causa", p["campos"])            # célula vazia: campo vazio, nunca inventado
            if dado["citacao"] == "-":
                self.assertNotIn("data_citacao", p["campos"])
            self.assertTrue(all(c["origem"] == "migrado" for c in p["campos"].values()))

    def test_momento_atual_e_avisos_de_vocabulario(self):
        momentos = {i: valor(p, "momento_atual") for i, p in enumerate(self.rel["processos"])}
        self.assertEqual(momentos[0], "AGUARDANDO SENTENÇA")
        self.assertIsNone(momentos[1])                             # "AGUARDANDO JULGAMENTO DO AGRAVO" não está no vocabulário
        self.assertEqual(momentos[8], "PROCESSO ARQUIVADO")          # abreviação do rótulo conhecido
        fora = avisos_de(self.rel, "momento_fora_do_vocabulario")
        self.assertEqual(len(fora), 1)
        self.assertIn("AGRAVO", fora[0]["mensagem"])
        self.assertTrue(fora[0]["candidatos"])

    def test_ultimo_andamento_ignora_datas_no_meio_da_frase(self):
        for dado, p in zip(self.dados, self.rel["processos"]):
            esperado = max(datetime.datetime.strptime(d, "%d/%m/%Y") for d, _ in dado["andamentos"]).date().isoformat()
            self.assertEqual(p["ultimo_andamento"], esperado, dado["numeros"][0])
        # os casos com audiência/prazo futuros no texto
        self.assertEqual(self.rel["processos"][2]["ultimo_andamento"], "2026-09-01")
        self.assertEqual(self.rel["processos"][10]["ultimo_andamento"], "2026-09-04")

    def test_fecho_fora_do_texto_e_texto_preservado(self):
        for i, p in enumerate(self.rel["processos"]):
            self.assertNotIn("sem atualiza", p["andamentos_texto"].lower())
            for _, texto in self.dados[i]["andamentos"]:
                self.assertIn(norm_espacos(texto), norm_espacos(p["andamentos_texto"]))
        self.assertEqual(self.rel["processos"][0]["fecho"], "2026-09-18")
        self.assertEqual(self.rel["processos"][7]["fecho"], "2026-09-18")      # fecho todo em negrito
        self.assertEqual(self.rel["processos"][9]["fecho"], "2026-09-18")      # "sem atualização" no singular
        self.assertNotIn("fecho", self.rel["processos"][8])                     # sem fecho
        self.assertIn("Anotação do advogado", self.rel["processos"][1]["andamentos_texto"])   # texto humano mantido

    def test_datas_em_negrito(self):
        for p in self.rel["processos"]:
            self.assertTrue(p["andamentos"])
            self.assertTrue(all(a["data_em_negrito"] for a in p["andamentos"]))

    def test_origem_no_arquivo(self):
        self.assertTrue(all(re.fullmatch(r"tabela \d+", p["origem_no_arquivo"]) for p in self.rel["processos"]))

    def test_fragmentacao_dos_runs_nao_muda_o_resultado(self):
        referencia = json.dumps(self.rel["processos"], sort_keys=True, ensure_ascii=False)
        for semente in (1, 2, 5, 9):
            destino = self.pasta / f"frag{semente}.docx"
            s2_gm.gerar(destino, semente=semente)
            lido = leitores.ler(destino)
            self.assertEqual(json.dumps(lido["processos"], sort_keys=True, ensure_ascii=False), referencia, semente)

    def test_variante_com_texto_na_linha_de_baixo(self):
        destino = self.pasta / "abaixo.docx"
        s2_gm.gerar(destino, variante="andamentos_abaixo")
        lido = leitores.ler(destino)
        self.assertEqual(len(lido["processos"]), 12)
        self.assertEqual(lido["processos"][0]["ultimo_andamento"], self.rel["processos"][0]["ultimo_andamento"])
        self.assertEqual(lido["processos"][0]["andamentos_texto"], self.rel["processos"][0]["andamentos_texto"])

    def test_ler_nao_altera_o_arquivo(self):
        antes = self.arquivo.read_bytes()
        leitores.ler(self.arquivo)
        self.assertEqual(self.arquivo.read_bytes(), antes)

    def test_documento_sem_quadro_resumo(self):
        n = ficticio.numero_ficticio(0)
        destino = montar_docx(self.pasta / "sem_resumo.docx", "Cliente Exemplo 01 Ltda", "18/09/2026",
                              [bloco_basico(f"PROCESSO Nº {n}")], resumo=None)
        rel = leitores.ler(destino)
        self.assertEqual(len(rel["processos"]), 1)
        self.assertIn("quadro_resumo_ausente", codigos(rel))

    def test_dv_errado_no_titulo_recusa_o_processo_e_lista(self):
        bom, ruim = ficticio.numero_ficticio(0), ficticio.numero_com_dv_errado(ficticio.numero_ficticio(1))
        destino = montar_docx(self.pasta / "dv.docx", "Cliente Exemplo 01 Ltda", "18/09/2026",
                              [bloco_basico(f"PROCESSO Nº {bom}"), bloco_basico(f"PROCESSO Nº {ruim}")],
                              resumo=[(bom, "Cobrança", "AGUARDANDO SENTENÇA", "20/04/2026"), (ruim, "Cobrança", "AGUARDANDO SENTENÇA", "20/04/2026")])
        rel = leitores.ler(destino)
        self.assertEqual([p["numero"] for p in rel["processos"]], [bom])
        av = avisos_de(rel, "numero_dv_invalido")
        self.assertEqual(len(av), 1)                 # título e quadro-resumo não repetem o aviso
        self.assertEqual(av[0]["nivel"], "erro")
        self.assertIn(ruim, av[0]["candidatos"])

    def test_numero_repetido_em_dois_blocos_e_avisado(self):
        n = ficticio.numero_ficticio(3)
        destino = montar_docx(self.pasta / "dup.docx", "Cliente Exemplo 01 Ltda", "18/09/2026",
                              [bloco_basico(f"PROCESSO Nº {n}"), bloco_basico(f"PROCESSO Nº {n}", assunto="Outro assunto")],
                              resumo=[(n, "Cobrança", "AGUARDANDO SENTENÇA", "20/04/2026")])
        rel = leitores.ler(destino)
        self.assertEqual(len(rel["processos"]), 2)        # os dois registros ficam (nada se perde)
        self.assertEqual(len(avisos_de(rel, "numero_repetido")), 1)

    def test_qualificador_do_momento_e_divergencia_titulo_resumo(self):
        n1, n2 = ficticio.numero_ficticio(0), ficticio.numero_ficticio(1)
        destino = montar_docx(
            self.pasta / "qual.docx", "Cliente Exemplo 01 Ltda", "18/09/2026",
            [bloco_basico(f"PROCESSO Nº {n1}", momento="CUMPRIMENTO DE SENTENÇA (HONORÁRIOS SUSPENSOS)"),
             bloco_basico(f"PROCESSO Nº {n2}", momento="AGUARDANDO AUDIÊNCIA")],
            resumo=[(n1, "Cobrança", "CUMPRIMENTO DE SENTENÇA (HONORÁRIOS SUSPENSOS)", "20/04/2026"),
                    (n2, "Cobrança", "AGUARDANDO SENTENÇA", "20/04/2026")])
        rel = leitores.ler(destino)
        p1, p2 = rel["processos"]
        self.assertEqual(valor(p1, "momento_atual"), "CUMPRIMENTO DE SENTENÇA")
        self.assertEqual(p1["momento_qualificador"], "HONORÁRIOS SUSPENSOS")
        self.assertEqual(valor(p2, "momento_atual"), "AGUARDANDO AUDIÊNCIA")     # vale o título
        div = avisos_de(rel, "momento_divergente")
        self.assertEqual(len(div), 1)
        self.assertEqual(len(div[0]["candidatos"]), 2)

    def test_resumo_sem_bloco_vira_processo_com_aviso(self):
        n1, n2 = ficticio.numero_ficticio(0), ficticio.numero_ficticio(1)
        destino = montar_docx(self.pasta / "so_resumo.docx", "Cliente Exemplo 01 Ltda", "18/09/2026", [bloco_basico(f"PROCESSO Nº {n1}")],
                              resumo=[(n1, "Cobrança", "AGUARDANDO SENTENÇA", "20/04/2026"),
                                      (n2, "Despejo", "AGUARDANDO AUDIÊNCIA", "10/05/2026")])
        rel = leitores.ler(destino)
        p = por_numero(rel)[n2]
        self.assertEqual(valor(p, "momento_atual"), "AGUARDANDO AUDIÊNCIA")
        self.assertEqual(p["ultimo_andamento"], "2026-05-10")
        self.assertEqual(p["andamentos_texto"], "")
        self.assertEqual(len(avisos_de(rel, "resumo_sem_bloco")), 1)

    def test_ultimo_andamento_divergente_do_resumo(self):
        n = ficticio.numero_ficticio(0)
        destino = montar_docx(self.pasta / "div.docx", "Cliente Exemplo 01 Ltda", "18/09/2026", [bloco_basico(f"PROCESSO Nº {n}")],
                              resumo=[(n, "Cobrança", "AGUARDANDO SENTENÇA", "30/06/2026")])
        rel = leitores.ler(destino)
        self.assertEqual(rel["processos"][0]["ultimo_andamento"], "2026-04-20")     # vale o texto
        self.assertEqual(len(avisos_de(rel, "ultimo_andamento_divergente")), 1)

    def test_rotulo_desconhecido_vai_para_colunas_sem_destino(self):
        n = ficticio.numero_ficticio(0)
        destino = montar_docx(self.pasta / "extra.docx", "Cliente Exemplo 01 Ltda", "18/09/2026",
                              [bloco_basico(f"PROCESSO Nº {n}", extra_campos=[("Provisão contábil", "R$ 1,00")])],
                              resumo=[(n, "Cobrança", "AGUARDANDO SENTENÇA", "20/04/2026")])
        rel = leitores.ler(destino)
        self.assertEqual([c["coluna"] for c in rel["colunas_sem_destino"]], ["Provisão contábil"])
        self.assertEqual(rel["colunas_sem_destino"][0]["amostra"], ["R$ 1,00"])

    def test_tabela_dentro_de_tabela(self):
        n = ficticio.numero_ficticio(0)
        destino = montar_docx(self.pasta / "aninhada.docx", "Cliente Exemplo 01 Ltda", "18/09/2026",
                              [bloco_basico(f"PROCESSO Nº {n}")], resumo=[(n, "Cobrança", "AGUARDANDO SENTENÇA", "20/04/2026")],
                              embrulhar=True)
        rel = leitores.ler(destino)
        self.assertEqual([p["numero"] for p in rel["processos"]], [n])
        self.assertEqual(valor(rel["processos"][0], "assunto"), "Cobrança")

    def test_data_em_negrito_so_quando_esta_em_negrito(self):
        n = ficticio.numero_ficticio(0)
        runs = [("Em ", False), ("02/03/2026", False), (", despacho. Em ", False), ("20/04/2026", True), (", contestação.", False)]
        b = bloco_basico(f"PROCESSO Nº {n}")
        b["andamentos"] = runs
        destino = montar_docx(self.pasta / "negrito.docx", "Cliente Exemplo 01 Ltda", "18/09/2026", [b],
                              resumo=[(n, "Cobrança", "AGUARDANDO SENTENÇA", "20/04/2026")])
        rel = leitores.ler(destino)
        self.assertEqual([a["data_em_negrito"] for a in rel["processos"][0]["andamentos"]], [False, True])

    def test_revisoes_do_documento_geram_aviso_e_texto_apagado_nao_conta(self):
        n = ficticio.numero_ficticio(0)
        base_docx = montar_docx(self.pasta / "rev0.docx", "Cliente Exemplo 01 Ltda", "18/09/2026", [bloco_basico(f"PROCESSO Nº {n}")],
                                resumo=[(n, "Cobrança", "AGUARDANDO SENTENÇA", "20/04/2026")])

        def com_revisao(xml):
            w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            ins = f'<w:ins xmlns:w="{w}" w:id="9" w:author="x" w:date="2026-01-01T00:00:00Z"><w:r><w:t xml:space="preserve"> Em 30/06/2026, juntada de laudo.</w:t></w:r></w:ins>'
            dele = f'<w:del xmlns:w="{w}" w:id="10" w:author="x" w:date="2026-01-01T00:00:00Z"><w:r><w:delText xml:space="preserve"> Em 31/07/2026, texto apagado.</w:delText></w:r></w:del>'
            return xml.replace("contestação.</w:t></w:r>", "contestação.</w:t></w:r>" + ins + dele, 1)

        destino = reescrever_docx(base_docx, self.pasta / "rev.docx", com_revisao)
        rel = leitores.ler(destino)
        self.assertIn("revisoes_no_documento", codigos(rel))
        self.assertEqual(rel["processos"][0]["ultimo_andamento"], "2026-06-30")
        self.assertNotIn("texto apagado", rel["processos"][0]["andamentos_texto"])

    def test_arquivo_corrompido_ou_estranho_volta_com_aviso_erro(self):
        ruim = self.pasta / "corrompido.docx"
        ruim.write_bytes(self.arquivo.read_bytes()[:300])
        rel = leitores.ler(ruim)
        self.assertEqual(rel["processos"], [])
        self.assertIn("erro", [a["nivel"] for a in rel["avisos"]])
        texto = self.pasta / "falso.docx"
        texto.write_text("isto não é um documento Word", encoding="utf-8")
        rel = leitores.ler(texto, formato="docx_a")
        self.assertEqual(rel["processos"], [])
        self.assertIn("arquivo_ilegivel", codigos(rel, "erro"))
        sem_documento = self.pasta / "sem_documento.docx"
        with zipfile.ZipFile(sem_documento, "w") as z:
            z.writestr("qualquer.txt", "x")
        self.assertIn("arquivo_ilegivel", codigos(leitores.ler(sem_documento, formato="docx_a"), "erro"))

    def test_200_processos_sem_erro_e_ida_e_volta(self):
        fichas = fichas_200()
        eventos, fichas_estado = [], []
        for f in fichas:
            g = copy.deepcopy(f)
            texto = f["linha_de_base"]["andamentos_texto"]
            for a in s2_da._parse_andamentos(texto):
                if not a["fecho"]:
                    eventos.append({"numero": f["numero"], "status": "aprovado", "data": ficha.parse_data(a["data"]), "frase": a["texto"]})
            fichas_estado.append(g)
        estado = {"cliente": "Cliente Exemplo 01 Ltda", "data_base": "2026-10-07", "fichas": fichas_estado, "eventos": eventos}
        gerado = self.pasta / "gerado200.docx"
        s2_da.gravar(None, estado, gerado)
        t0 = time.time()
        rel = leitores.ler(gerado)
        self.assertLess(time.time() - t0, 30)
        self.assertEqual(rel["formato"], "docx_a")
        self.assertEqual(len(rel["processos"]), 200)
        self.assertEqual(codigos(rel, "erro"), [])
        lidos = por_numero(rel)
        campos = ("assunto", "autores", "reus", "data_ajuizamento", "valor_causa", "data_citacao", "vara", "area",
                  "materia_principal", "momento_atual")
        for f in fichas:
            p = lidos[f["numero"]]
            self.assertEqual(lido_do_processo(p, campos), esperado_do_ficha(f, campos), f["numero"])
            self.assertEqual(p["vinculados"], f["vinculados"], f["numero"])
            # o protótipo escreve "Em DD/MM/AAAA, ..." (com vírgula); o resto do texto tem de ser o mesmo
            self.assertEqual(norm_andamento(p["andamentos_texto"]), norm_andamento(f["linha_de_base"]["andamentos_texto"]))
            self.assertEqual(p["ultimo_andamento"], f["linha_de_base"]["ultimo_andamento"])
        # ida e volta: ficha lida -> regravada pelo protótipo -> lida de novo
        fichas2, eventos2 = [], []
        for p in rel["processos"]:
            g = ficha.nova_ficha(p["numero"])
            for campo, c in p["campos"].items():
                ficha.definir(g, campo, c["valor"], c["origem"])
            for v in p["vinculados"]:
                ficha.vincular(g, v["numero"], v["tipo"])
            ficha.definir(g, "ultimo_andamento", p["ultimo_andamento"], "migrado")
            fichas2.append(g)
            for a in p["andamentos"]:
                eventos2.append({"numero": p["numero"], "status": "aprovado", "data": a["data"], "frase": a["texto"].rstrip(".")})
        destino2 = self.pasta / "regravado200.docx"
        s2_da.gravar(None, {"cliente": rel["cliente"], "data_base": rel["data_base"], "fichas": fichas2, "eventos": eventos2}, destino2)
        rel2 = leitores.ler(destino2)
        self.assertEqual(len(rel2["processos"]), 200)
        a, b = por_numero(rel), por_numero(rel2)
        for numero, p in a.items():
            q = b[numero]
            self.assertEqual({c: v["valor"] for c, v in q["campos"].items()}, {c: v["valor"] for c, v in p["campos"].items()}, numero)
            self.assertEqual(q["vinculados"], p["vinculados"])
            self.assertEqual(norm_espacos(q["andamentos_texto"]), norm_espacos(p["andamentos_texto"]), numero)
            self.assertEqual(q["ultimo_andamento"], p["ultimo_andamento"])


# ================================================================ modelo B (.xlsx)

@unittest.skipUnless(COM_SPIKES, "precisa de spikes/ e python-docx")
class TestXlsxB(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.pasta = pasta("xlsx")

    def test_modelo_do_spike_com_cache_de_formulas(self):
        destino = self.pasta / "spike.xlsx"
        s1_gm.gerar(destino, n_linhas=10)
        self.assertEqual(leitores.detectar(destino), "xlsx_b")
        rel = leitores.ler(destino)
        self.assertEqual(rel["formato"], "xlsx_b")
        self.assertEqual(len(rel["processos"]), 10)
        self.assertEqual(rel["data_base"], "2026-09-30")
        self.assertEqual(rel["parametros"]["headcount"], 350)
        self.assertEqual(rel["parametros"]["empresas_do_grupo"], ["Empresa Modelo 1 Ltda", "Empresa Modelo 2 Ltda"])
        p0 = rel["processos"][0]
        self.assertEqual(p0["numero"], s1_gm.cnj_ficticio(1))
        self.assertEqual(valor(p0, "valor_causa"), "10000.00")
        self.assertEqual(valor(p0, "data_ajuizamento"), "2019-01-15")
        self.assertEqual(p0["campos"]["probabilidade"]["origem"], "humano")       # coluna de julgamento digitada
        self.assertEqual(p0["campos"]["valor_estimado"]["origem"], "humano")
        self.assertEqual(p0["campos"]["valor_causa"]["origem"], "migrado")
        self.assertEqual(p0["campos"]["valor_economizado"]["origem"], "migrado")  # fórmula nunca é humano
        self.assertEqual(valor(p0, "valor_economizado"), "6000.00")                # valor em cache
        self.assertTrue(p0["ativo"])
        self.assertEqual(p0["ultimo_andamento"], "2019-01-15")
        self.assertNotIn("Até 30/09/2026", p0["andamentos_texto"])      # o fecho sai do histórico...
        self.assertEqual(p0["fecho"], "2026-09-30")                     # ...e a data dele fica em `fecho`
        self.assertEqual(p0["origem_no_arquivo"], "aba 'Processos', linha 2")
        self.assertEqual(rel["colunas_sem_destino"], [])
        self.assertEqual(len(avisos_de(rel, "aba_ignorada")), 3)        # Indicadores, Dashboard, Dinâmica
        self.assertEqual(avisos_de(rel, "formula_sem_valor"), [])       # as fórmulas têm valor em cache

    def test_formulas_sem_cache_avisam(self):
        origem = self.pasta / "spike_base.xlsx"
        s1_gm.gerar(origem, n_linhas=3)
        fichas = fichas_200()[:6]
        estado = {"cliente": "x", "data_base": "2026-10-07", "fichas": copy.deepcopy(fichas), "eventos": []}
        destino = self.pasta / "spike_gravado.xlsx"
        s1_xc.gravar(origem, estado, destino)       # o protótipo invalida o cache das fórmulas (padrão)
        rel = leitores.ler(destino)
        self.assertEqual(len(rel["processos"]), 9)
        sem_valor = avisos_de(rel, "formula_sem_valor")
        self.assertTrue(sem_valor)
        self.assertTrue(all(a["nivel"] == "info" for a in sem_valor))     # colunas calculadas: Ativo, Valor Economizado, Taxa
        self.assertEqual(valor(rel["processos"][-1], "valor_economizado"), None)

    def test_ida_e_volta_com_o_escritor_do_spike(self):
        origem = self.pasta / "ida_volta_base.xlsx"
        s1_gm.gerar(origem, n_linhas=1)
        fichas = fichas_200()
        estado = {"cliente": "x", "data_base": "2026-10-07", "fichas": copy.deepcopy(fichas), "eventos": []}
        destino = self.pasta / "ida_volta.xlsx"
        res = s1_xc.gravar(origem, estado, destino)
        self.assertEqual(len(res["processos_novos"]), 200)
        rel = leitores.ler(destino)
        self.assertEqual(len(rel["processos"]), 201)
        self.assertEqual(codigos(rel, "erro"), [])
        lidos = por_numero(rel)
        for f in fichas:
            p = lidos[f["numero"]]
            esperado = esperado_do_ficha(f, CAMPOS_COMPARAVEIS_B)
            self.assertEqual(lido_do_processo(p, CAMPOS_COMPARAVEIS_B), esperado, f["numero"])
        # regravar o que foi lido e ler outra vez preserva os campos
        fichas2 = []
        for p in rel["processos"]:
            g = ficha.nova_ficha(p["numero"])
            for campo, c in p["campos"].items():
                ficha.definir(g, campo, c["valor"], c["origem"])
            fichas2.append(g)
        destino2 = self.pasta / "ida_volta2.xlsx"
        s1_xc.gravar(origem, {"cliente": "x", "data_base": "2026-10-07", "fichas": fichas2, "eventos": []}, destino2)
        rel2 = leitores.ler(destino2)
        a, b = por_numero(rel), por_numero(rel2)
        self.assertEqual(set(a), set(b))
        for numero, p in a.items():
            q = b[numero]
            comuns = [c for c in CAMPOS_COMPARAVEIS_B if c in p["campos"]]
            self.assertEqual(lido_do_processo(q, comuns), lido_do_processo(p, comuns), numero)

    def test_200_processos_em_duas_abas_com_formatos_misturados(self):
        fichas = fichas_200()
        destino = gerar_xlsx_b(self.pasta / "b200.xlsx", fichas)
        t0 = time.time()
        rel = leitores.ler(destino)
        self.assertLess(time.time() - t0, 30)
        self.assertEqual(rel["formato"], "xlsx_b")
        self.assertEqual(len(rel["processos"]), 200)
        self.assertEqual(codigos(rel, "erro"), [])
        self.assertEqual(rel["data_base"], "2026-09-30")
        self.assertEqual(rel["parametros"]["headcount"], 1200)
        self.assertAlmostEqual(rel["parametros"]["fator_correcao"], 1.0523)
        self.assertEqual(rel["parametros"]["empresas_do_grupo"], ["Empresa Fictícia Alfa Ltda", "Empresa Fictícia Beta S.A."])
        self.assertEqual(rel["parametros"]["outros"], {"Observação livre": "texto qualquer"})
        lidos = por_numero(rel)
        abas = {p["origem_no_arquivo"].split(",")[0] for p in rel["processos"]}
        self.assertEqual(abas, {"aba 'Ativos'", "aba 'Arquivados'"})
        for f in fichas:
            p = lidos[f["numero"]]
            self.assertEqual(lido_do_processo(p, CAMPOS_COMPARAVEIS_B), esperado_do_ficha(f, CAMPOS_COMPARAVEIS_B), f["numero"])
            self.assertEqual(p["vinculados"], f["vinculados"], f["numero"])
            self.assertEqual(p["ativo"], f.get("ativo", True))
            texto = f["linha_de_base"]["andamentos_texto"]
            self.assertEqual(p["andamentos_texto"], texto)
            self.assertEqual(p["ultimo_andamento"], f["linha_de_base"]["ultimo_andamento"])
            for campo, c in p["campos"].items():
                esperada = "humano" if campo in ficha.CAMPOS_DE_JULGAMENTO else "migrado"
                self.assertEqual(c["origem"], esperada, (f["numero"], campo))
        self.assertEqual(len(avisos_de(rel, "linhas_ignoradas")), 2)        # marcador e total, uma vez por aba
        self.assertIn("fecho", rel["processos"][0])        # alguns textos têm "Até 30/09/2026 sem atualizações."

    def test_detecta_aba_pelo_cabecalho_e_nao_pelo_nome(self):
        from openpyxl import load_workbook
        origem = gerar_xlsx_b(self.pasta / "nomes.xlsx", fichas_200()[:30])
        wb = load_workbook(origem)
        wb["Ativos"].title = "Plan1"
        wb["Arquivados"].title = "Zebra"
        destino = self.pasta / "nomes2.xlsx"
        wb.save(destino)
        rel = leitores.ler(destino)
        self.assertEqual(len(rel["processos"]), 30)
        self.assertEqual({p["origem_no_arquivo"].split(",")[0] for p in rel["processos"]} - {"aba 'Plan1'", "aba 'Zebra'"}, set())

    def test_linha_com_dados_e_sem_numero_valido_e_recusada(self):
        from openpyxl import load_workbook
        fichas = fichas_200()[:12]
        origem = gerar_xlsx_b(self.pasta / "recusa.xlsx", fichas, abas=("Lista", "Outra"), titulo_antes=False)
        wb = load_workbook(origem)
        ws = wb["Lista"]
        ws["A2"] = ficticio.numero_com_dv_errado(ws["A2"].value.split("\n")[0])
        ws["A3"] = "sem número, só texto"
        ws["A4"] = None                                  # número vazio com a linha cheia
        destino = self.pasta / "recusa2.xlsx"
        wb.save(destino)
        rel = leitores.ler(destino)
        self.assertEqual(len(avisos_de(rel, "numero_dv_invalido")), 1)
        self.assertEqual(len(avisos_de(rel, "numero_invalido")), 1)
        self.assertEqual(len(avisos_de(rel, "linha_sem_numero")), 1)
        lidos = {p["numero"] for p in rel["processos"]}
        self.assertEqual(len(lidos), len(rel["processos"]))
        self.assertEqual(len(rel["processos"]), 12 - 3)
        self.assertTrue(avisos_de(rel, "numero_dv_invalido")[0]["onde"].endswith("linha 2"))

    def test_mesmo_numero_em_duas_abas_fica_nas_duas_com_aviso(self):
        from openpyxl import load_workbook
        fichas = fichas_200()[:20]
        origem = gerar_xlsx_b(self.pasta / "dupla.xlsx", fichas, titulo_antes=False)
        wb = load_workbook(origem)
        wb["Arquivados"]["A2"] = wb["Ativos"]["A2"].value
        destino = self.pasta / "dupla2.xlsx"
        wb.save(destino)
        rel = leitores.ler(destino)
        av = avisos_de(rel, "numero_repetido")
        self.assertEqual(len(av), 1)
        self.assertEqual(len(av[0]["candidatos"]), 1)
        self.assertEqual(sum(1 for p in rel["processos"] if p["numero"] == av[0]["candidatos"][0]), 2)
        self.assertEqual(len(rel["processos"]), 20)

    def test_numero_principal_que_tambem_e_vinculado_de_outro(self):
        from openpyxl import Workbook
        n0, n1, n2 = (ficticio.numero_ficticio(i) for i in range(3))
        wb = Workbook()
        ws = wb.active
        ws.append(["Número do Processo", "Autor(es)", "Réu(s)", "Andamentos"])
        ws.append([f"{n0}\nAgravo: {n1}", "A", "B", "Em 01/02/2026 distribuído."])
        ws.append([n1, "C", "D", "Em 03/02/2026 distribuído."])
        ws.append([n2, "E", "F", "Em 05/02/2026 distribuído."])
        destino = self.pasta / "dois_lugares.xlsx"
        wb.save(destino)
        rel = leitores.ler(destino)
        av = avisos_de(rel, "numero_em_dois_lugares")
        self.assertEqual([a["candidatos"] for a in av], [[n1]])
        self.assertEqual(len(rel["processos"]), 3)

    def test_planilha_no_estilo_da_fase_1(self):
        """Coluna 'Andamentos' como a planilha.py da Fase 1 escreve: 'Em DD/MM/AAAA texto.' (sem vírgula) + fecho 'Até ...'."""
        from openpyxl import Workbook
        n = [ficticio.numero_ficticio(i) for i in range(2)]
        wb = Workbook()
        ws = wb.active
        ws.title = "Processos"
        ws.append(["Processo", "Autor", "Andamentos"])
        ws.append([n[0], "Pessoa Fictícia 0001", "Em 01/09/2026 o juiz proferiu decisão. Até 03/10/2026 sem andamentos."])
        ws.append([n[1], "Pessoa Fictícia 0002", "Em 01/09/2026 foi proferida sentença. Em 01/10/2026 o autor apelou. Audiência: 20/11/2026."])
        destino = self.pasta / "fase1.xlsx"
        wb.save(destino)
        self.assertEqual(leitores.detectar(destino), "tabela_livre")
        rel = leitores.ler(destino)
        a, b = rel["processos"]
        self.assertEqual((a["ultimo_andamento"], a["fecho"]), ("2026-09-01", "2026-10-03"))
        self.assertEqual(a["andamentos_texto"], "Em 01/09/2026 o juiz proferiu decisão.")
        self.assertEqual((b["ultimo_andamento"], b.get("fecho")), ("2026-10-01", None))      # a audiência futura não conta
        self.assertEqual(valor(a, "autores"), "Pessoa Fictícia 0001")

    def test_modelo_b_com_poucas_colunas_e_aba_de_parametros_continua_xlsx_b(self):
        from openpyxl import Workbook
        n = [ficticio.numero_ficticio(i) for i in range(3)]
        wb = Workbook()
        ws = wb.active
        ws.title = "Processos"
        ws.append(["Número do Processo", "Autor(es)", "Réu(s)", "Vara", "Município", "Valor da Causa", "Andamentos", "Situação"])
        for i, x in enumerate(n):
            ws.append([x, "A", "B", "1ª Vara", "Cidade", 1000.0 * (i + 1), f"Em 0{i + 1}/02/2026 distribuído.", "Ativo"])
        pa = wb.create_sheet("Parâmetros")
        pa.append(["Número de funcionários", 80])
        destino = self.pasta / "poucas.xlsx"
        wb.save(destino)
        self.assertEqual(leitores.detectar(destino), "xlsx_b")
        rel = leitores.ler(destino)
        self.assertEqual(rel["parametros"], {"headcount": 80})
        self.assertEqual(len(rel["processos"]), 3)

    def test_varios_numeros_na_mesma_celula(self):
        fichas = [f for f in fichas_200() if f["vinculados"]][:5]
        self.assertTrue(fichas)
        destino = gerar_xlsx_b(self.pasta / "vinc.xlsx", fichas)
        rel = leitores.ler(destino)
        lidos = por_numero(rel)
        for f in fichas:
            self.assertEqual(lidos[f["numero"]]["vinculados"], f["vinculados"])

    def test_cabecalhos_trocados_e_colunas_extras(self):
        fichas = fichas_200()[:40]
        trocados = {"Autor(es)": "Requerente", "Réu(s)": "Reclamada", "Valor da Causa": "Vlr. Causa (R$)", "Vara": "Juízo",
                    "Data do Ajuizamento": "Data de distribuição", "Matéria Principal": "Matéria", "Andamentos": "Histórico"}
        destino = gerar_xlsx_b(self.pasta / "trocado.xlsx", fichas, renomear=trocados, extras=True)
        rel = leitores.ler(destino)
        self.assertEqual(len(rel["processos"]), 40)
        destinos = {m["coluna"]: m["campo"] for m in rel["mapeamento"] if m["aba"] == "Ativos"}
        self.assertEqual(destinos["Requerente"], "autores")
        self.assertEqual(destinos["Reclamada"], "reus")
        self.assertEqual(destinos["Vlr. Causa (R$)"], "valor_causa")
        self.assertEqual(destinos["Histórico"], "andamentos")
        sem = {(c["aba"], c["coluna"]) for c in rel["colunas_sem_destino"]}
        self.assertIn(("Ativos", "Observação interna"), sem)
        self.assertIn(("Ativos", "Código interno"), sem)
        lidos = por_numero(rel)
        for f in fichas:
            p = lidos[f["numero"]]
            self.assertEqual(valor(p, "autores"), ficha.obter(f, "autores"))
            self.assertEqual(valor(p, "valor_causa"), ficha.obter(f, "valor_causa"))
        amostra = next(c for c in rel["colunas_sem_destino"] if c["coluna"] == "Código interno")["amostra"]
        self.assertTrue(amostra[0].startswith("C-"))

    def test_celula_com_erro_do_excel_vira_vazio_com_aviso(self):
        from openpyxl import load_workbook
        fichas = fichas_200()[:8]
        origem = gerar_xlsx_b(self.pasta / "erro.xlsx", fichas, abas=("A", "B"), titulo_antes=False)
        wb = load_workbook(origem)
        wb["A"]["K2"] = "#N/D"
        destino = self.pasta / "erro2.xlsx"
        wb.save(destino)
        rel = leitores.ler(destino)
        self.assertEqual(len(avisos_de(rel, "celula_com_erro")), 1)

    def test_data_valor_e_rotulo_ilegiveis_viram_aviso_e_campo_vazio(self):
        from openpyxl import load_workbook
        fichas = fichas_200()[:8]
        origem = gerar_xlsx_b(self.pasta / "ilegivel.xlsx", fichas, abas=("A", "B"), titulo_antes=False)
        wb = load_workbook(origem)
        ws = wb["A"]
        ws["G2"], ws["K2"], ws["V2"], ws["Z2"] = "31/02/2026", "valor a apurar", "Resultado inventado", "talvez"
        destino = self.pasta / "ilegivel2.xlsx"
        wb.save(destino)
        rel = leitores.ler(destino)
        p = rel["processos"][0]
        for campo in ("data_ajuizamento", "valor_causa", "resultado", "houve_recurso"):
            self.assertNotIn(campo, p["campos"])
        for codigo in ("data_invalida", "valor_invalido", "rotulo_fora_do_vocabulario"):
            self.assertTrue(any(a["onde"].endswith("2") for a in avisos_de(rel, codigo)), codigo)
        self.assertTrue(any("Houve recurso" in a["mensagem"] for a in avisos_de(rel, "valor_invalido")))

    def test_sem_aba_de_processos(self):
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.append(["Nome", "Idade"])
        wb.active.append(["Fulano", 3])
        destino = self.pasta / "sem_processos.xlsx"
        wb.save(destino)
        self.assertEqual(leitores.detectar(destino), "desconhecido")
        rel = leitores.ler(destino, formato="xlsx_b")
        self.assertEqual(rel["processos"], [])
        self.assertIn("sem_aba_de_processos", codigos(rel, "erro"))

    def test_xlsx_corrompido_volta_com_aviso(self):
        origem = gerar_xlsx_b(self.pasta / "bom.xlsx", fichas_200()[:5])
        dados = origem.read_bytes()
        ruim = self.pasta / "ruim.xlsx"
        ruim.write_bytes(dados[:len(dados) // 2])
        rel = leitores.ler(ruim)
        self.assertEqual(rel["processos"], [])
        self.assertIn("erro", [a["nivel"] for a in rel["avisos"]])
        self.assertEqual(leitores.detectar(ruim), "desconhecido")
        rel = leitores.ler(ruim, formato="xlsx_b")
        self.assertEqual(codigos(rel, "erro"), ["arquivo_ilegivel"])

    @unittest.skipUnless(SOFFICE, "sem LibreOffice")
    def test_planilha_resalva_pelo_libreoffice_continua_legivel(self):
        import subprocess
        origem = self.pasta / "lo_origem.xlsx"
        s1_gm.gerar(origem, n_linhas=10)
        saida = self.pasta / "lo_saida"
        saida.mkdir(exist_ok=True)
        r = subprocess.run([SOFFICE, "--headless", "--convert-to", "xlsx", "--outdir", str(saida), str(origem)],
                           capture_output=True, timeout=120)
        arquivo = saida / "lo_origem.xlsx"
        if r.returncode != 0 or not arquivo.exists():
            self.skipTest("o LibreOffice não converteu")
        rel = leitores.ler(arquivo)
        self.assertEqual(len(rel["processos"]), 10)
        self.assertEqual(rel["parametros"]["headcount"], 350)
        self.assertEqual(valor(rel["processos"][0], "valor_economizado"), "6000.00")


# ================================================================ lista

class TestLista(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.pasta = pasta("lista")
        cls.fichas = fichas_200()
        cls.arquivos = ficticio.gerar_lista_bruta(cls.pasta, cls.fichas)

    def _conferir(self, chave, formato="lista"):
        esperado = self.arquivos[chave]
        self.assertEqual(leitores.detectar(esperado["arquivo"]), formato, chave)
        rel = leitores.ler(esperado["arquivo"])
        self.assertEqual(rel["formato"], "lista")
        self.assertCountEqual([p["numero"] for p in rel["processos"]], esperado["validos"], chave)
        recusados = [n for a in avisos_de(rel, "numero_dv_invalido") for n in a["candidatos"]]
        self.assertCountEqual(recusados, esperado["invalidos"], chave)
        # mesmos números que a leitura da Fase 1
        antigos, invalidos = cart.ler_lista(str(esperado["arquivo"]))
        self.assertCountEqual([r["numero"] for r in antigos], [p["numero"] for p in rel["processos"]], chave)
        return rel

    def test_texto_de_email(self):
        rel = self._conferir("txt")
        self.assertEqual(len(avisos_de(rel, "numero_repetido")), 1)       # o primeiro número é citado de novo
        self.assertEqual(avisos_de(rel, "numero_repetido")[0]["nivel"], "info")
        self.assertTrue(all(re.fullmatch(r"linha \d+", p["origem_no_arquivo"]) for p in rel["processos"]))

    def test_csv_ponto_e_virgula(self):
        rel = self._conferir("csv")
        esperado = {f["numero"]: f for f in self.fichas}
        for p in rel["processos"]:
            f = esperado[p["numero"]]
            self.assertEqual(valor(p, "cliente"), ficha.obter(f, "cliente"))
            self.assertEqual(valor(p, "polo_cliente"), ficha.obter(f, "polo_cliente"))
            self.assertEqual(valor(p, "responsavel"), ficha.obter(f, "responsavel"))
            self.assertEqual(valor(p, "parte_contraria"), ficha.obter(f, "parte_contraria"))

    def test_xlsx_simples(self):
        rel = self._conferir("xlsx")
        self.assertTrue(all("cliente" in p["campos"] for p in rel["processos"]))

    def test_xlsx_bagunçado_com_cabecalho_na_linha_3(self):
        rel = self._conferir("xlsx_bagunca")
        self.assertTrue(all("cliente" in p["campos"] for p in rel["processos"]))
        self.assertIn("Seq", [c["coluna"] for c in rel["colunas_sem_destino"]])

    def test_csv_com_virgula_e_tabulacao_e_windows_1252(self):
        n = [ficticio.numero_ficticio(i) for i in range(3)]
        casos = {"virgula.csv": ("Processo,Cliente\n{0},Cliente Exemplo 01 Ltda\n{1},Cliente Exemplo 02 Ltda\n".format(*n), "utf-8"),
                 "tab.csv": ("Processo\tCliente\n{0}\tCliente Exemplo 01 Ltda\n{1}\tCliente Exemplo 02 Ltda\n".format(*n), "utf-8"),
                 "cp1252.csv": ("Nº;Cliente\n{0};Comércio Exemplo Ação Ltda\n{1};Serviços Exemplo Ltda\n".format(*n), "cp1252")}
        for nome, (texto, cod) in casos.items():
            arq = self.pasta / nome
            arq.write_bytes(texto.encode(cod))
            rel = leitores.ler(arq)
            self.assertEqual([p["numero"] for p in rel["processos"]], n[:2], nome)
            self.assertTrue(valor(rel["processos"][0], "cliente").startswith(("Cliente", "Comércio")), nome)
        self.assertEqual(valor(leitores.ler(self.pasta / "cp1252.csv")["processos"][0], "cliente"), "Comércio Exemplo Ação Ltda")

    def test_texto_sem_numero_nao_e_lista(self):
        arq = self.pasta / "nada.txt"
        arq.write_text("Olá, segue o relatório em anexo.", encoding="utf-8")
        self.assertEqual(leitores.detectar(arq), "desconhecido")
        rel = leitores.ler(arq)
        self.assertEqual(rel["formato"], "desconhecido")
        self.assertIn("formato_nao_reconhecido", codigos(rel, "erro"))

    def test_docx_de_lista_simples(self):
        if python_docx is None:
            self.skipTest("sem python-docx")
        n = [ficticio.numero_ficticio(i) for i in range(3)]
        d = python_docx.Document()
        d.add_paragraph("Processos do mês")
        for x in n:
            d.add_paragraph(f"- {x}")
        arq = self.pasta / "lista.docx"
        d.save(str(arq))
        self.assertEqual(leitores.detectar(arq), "lista")
        self.assertEqual([p["numero"] for p in leitores.ler(arq)["processos"]], n)

    def test_coluna_de_numero_vazia_e_sem_cabecalho(self):
        from openpyxl import Workbook
        n = [ficticio.numero_ficticio(i) for i in range(4)]
        wb = Workbook()
        ws = wb.active
        ws.append(["anotações soltas"])
        for x in n:
            ws.append([f"ver {x} depois"])
        ws.append([None])
        arq = self.pasta / "sem_cab.xlsx"
        wb.save(arq)
        self.assertEqual(leitores.detectar(arq), "lista")
        self.assertEqual([p["numero"] for p in leitores.ler(arq)["processos"]], n)


# ================================================================ tabela fora do padrão

class TestTabelaLivre(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.pasta = pasta("livre")

    def _planilha(self, nome="livre.xlsx", n=30):
        from openpyxl import Workbook
        fichas = fichas_200()[:n]
        wb = Workbook()
        ws = wb.active
        ws.title = "Carteira"
        ws.append(["Ref.", "Requerente", "Requerido", "Distribuído em", "Valor atribuído à causa", "Código interno", "Nota do sócio",
                   "Data", "Status"])
        for i, f in enumerate(fichas):
            ws.append([f["numero"], ficha.obter(f, "autores"), ficha.obter(f, "reus"), ficha.obter(f, "data_ajuizamento") and
                       datetime.date.fromisoformat(ficha.obter(f, "data_ajuizamento")), float(ficha.obter(f, "valor_causa")),
                       f"INT-{i:03d}", "texto livre", "10/01/2026", ficha.obter(f, "situacao")])
        arq = self.pasta / nome
        wb.save(arq)
        return arq, fichas

    def test_propoe_mapeamento_e_lista_o_que_nao_tem_destino(self):
        arq, fichas = self._planilha()
        self.assertEqual(leitores.detectar(arq), "tabela_livre")
        rel = leitores.ler(arq)
        self.assertEqual(rel["formato"], "tabela_livre")
        mapa = {m["coluna"]: m for m in rel["mapeamento"]}
        self.assertEqual(mapa["Ref."]["campo"], "numero")                    # reconhecida pelo conteúdo (números CNJ)
        self.assertEqual(mapa["Requerente"]["campo"], "autores")
        self.assertEqual(mapa["Requerido"]["campo"], "reus")
        self.assertEqual(mapa["Distribuído em"]["campo"], "data_ajuizamento")
        self.assertEqual(mapa["Valor atribuído à causa"]["campo"], "valor_causa")
        self.assertEqual(mapa["Status"]["campo"], "situacao")
        for m in mapa.values():
            self.assertIn("confianca", m)
        sem = {c["coluna"]: c for c in rel["colunas_sem_destino"]}
        self.assertEqual(set(sem), {"Código interno", "Nota do sócio", "Data"})
        self.assertEqual(sem["Código interno"]["amostra"][0], "INT-000")
        self.assertEqual(sem["Data"]["motivo"], "ambigua")
        self.assertIn("coluna_ambigua", codigos(rel))
        self.assertEqual(len(rel["processos"]), 30)
        p = por_numero(rel)[fichas[0]["numero"]]
        self.assertEqual(valor(p, "autores"), ficha.obter(fichas[0], "autores"))
        self.assertEqual(valor(p, "valor_causa"), ficha.obter(fichas[0], "valor_causa"))

    def test_mapeamento_corrigido_pelo_usuario(self):
        arq, fichas = self._planilha("livre2.xlsx", 10)
        rel = leitores.ler(arq, formato="tabela_livre", mapeamento={"Data": "data_citacao", "Nota do sócio": "observacoes",
                                                                  "Código interno": None, "Requerente": "reus", "Requerido": "autores"})
        p = por_numero(rel)[fichas[0]["numero"]]
        self.assertEqual(valor(p, "data_citacao"), "2026-01-10")
        self.assertEqual(valor(p, "observacoes"), "texto livre")
        self.assertEqual(valor(p, "reus"), ficha.obter(fichas[0], "autores"))        # trocado de propósito
        self.assertEqual(valor(p, "autores"), ficha.obter(fichas[0], "reus"))
        sem = {c["coluna"] for c in rel["colunas_sem_destino"]}
        self.assertEqual(sem, {"Código interno"})            # a escolha "não migrar" continua listada
        mapa = {m["coluna"]: m for m in rel["mapeamento"]}
        self.assertEqual(mapa["Data"]["confianca"], 1.0)
        # a lista devolvida também serve de entrada
        rel2 = leitores.ler(arq, formato="tabela_livre", mapeamento=rel["mapeamento"])
        self.assertEqual(json.dumps(rel2["processos"], sort_keys=True), json.dumps(rel["processos"], sort_keys=True))

    def test_destino_inexistente_e_destino_repetido(self):
        arq, fichas = self._planilha("livre3.xlsx", 5)
        rel = leitores.ler(arq, formato="tabela_livre", mapeamento={"Nota do sócio": "campo_que_nao_existe", "Data": "autores"})
        self.assertIn("mapeamento_invalido", codigos(rel))
        self.assertIn("coluna_duplicada", codigos(rel))
        p = por_numero(rel)[fichas[0]["numero"]]
        self.assertEqual(valor(p, "autores"), ficha.obter(fichas[0], "autores"))     # vale a primeira coluna

    def test_csv_fora_do_padrao(self):
        n = [ficticio.numero_ficticio(i) for i in range(3)]
        arq = self.pasta / "livre.csv"
        arq.write_text("Autos;Requerente;Réu;Valor da causa;Quem cuida\n" + "".join(
            f"{x};Pessoa Fictícia {i:04d};Cliente Exemplo 01 Ltda;R$ {1000 * (i + 1)},00;Fulana\n" for i, x in enumerate(n)), encoding="utf-8")
        self.assertEqual(leitores.detectar(arq), "tabela_livre")
        rel = leitores.ler(arq)
        self.assertEqual(valor(rel["processos"][1], "valor_causa"), "2000.00")
        self.assertEqual(valor(rel["processos"][0], "reus"), "Cliente Exemplo 01 Ltda")
        self.assertEqual([c["coluna"] for c in rel["colunas_sem_destino"]], ["Quem cuida"])

    def test_destinos_possiveis_para_a_tela(self):
        destinos = {d["campo"] for d in leitores.destinos_possiveis()}
        self.assertTrue(set(ficha.CAMPOS) <= destinos)
        self.assertTrue({"numero", "andamentos", "tribunal", "ativo"} <= destinos)


# ================================================================ detecção e robustez

class TestDetectar(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.pasta = pasta("detectar")

    def test_formatos_conhecidos(self):
        if COM_SPIKES:
            s2_gm.gerar(self.pasta / "a.docx")
            self.assertEqual(leitores.detectar(self.pasta / "a.docx"), "docx_a")
        gerar_xlsx_b(self.pasta / "b.xlsx", fichas_200()[:10])
        self.assertEqual(leitores.detectar(self.pasta / "b.xlsx"), "xlsx_b")
        self.assertEqual(leitores.detectar(str(self.pasta / "b.xlsx")), "xlsx_b")       # aceita texto também

    def test_desconhecidos_com_aviso_claro(self):
        casos = {"doc.pdf": b"%PDF-1.7\n1 0 obj", "antigo.doc": b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100,
                 "foto.png": b"\x89PNG\r\n\x1a\n\x00\x00", "vazio.txt": b"", "binario.bin": bytes(range(256)) * 4}
        for nome, dados in casos.items():
            arq = self.pasta / nome
            arq.write_bytes(dados)
            formato, aviso = leitores.diagnosticar(arq)
            self.assertEqual(formato, "desconhecido", nome)
            self.assertEqual(aviso["nivel"], "erro")
            self.assertTrue(aviso["mensagem"])
            rel = leitores.ler(arq)
            self.assertEqual(rel["formato"], "desconhecido", nome)
            self.assertEqual(rel["processos"], [])
            self.assertTrue(rel["avisos"])

    def test_arquivo_inexistente_nao_levanta(self):
        rel = leitores.ler(self.pasta / "nao_existe.xlsx")
        self.assertEqual(rel["processos"], [])
        self.assertEqual(codigos(rel, "erro"), ["arquivo_inexistente"])
        rel = leitores.ler(self.pasta / "nao_existe.xlsx", formato="xlsx_b")
        self.assertEqual(codigos(rel, "erro"), ["arquivo_inexistente"])
        rel = leitores.ler(self.pasta)       # pasta, não arquivo
        self.assertTrue(codigos(rel, "erro"))

    def test_formato_forcado_invalido(self):
        arq = self.pasta / "x.txt"
        arq.write_text(ficticio.numero_ficticio(0), encoding="utf-8")
        rel = leitores.ler(arq, formato="zebra")
        self.assertEqual(codigos(rel, "erro"), ["formato_nao_reconhecido"])

    def test_arquivos_truncados_nunca_levantam_excecao(self):
        validos = [gerar_xlsx_b(self.pasta / "t.xlsx", fichas_200()[:6])]
        if COM_SPIKES:
            s2_gm.gerar(self.pasta / "t.docx", n=4)
            validos.append(self.pasta / "t.docx")
        rng = random.Random(7)
        for arq in validos:
            dados = arq.read_bytes()
            for k in range(12):
                corte = rng.randrange(1, len(dados))
                ruim = self.pasta / f"corte{k}{arq.suffix}"
                ruim.write_bytes(dados[:corte])
                rel = leitores.ler(ruim)
                self.assertIn("avisos", rel)
                self.assertEqual(rel["processos"], [] if rel["avisos"] and any(a["nivel"] == "erro" for a in rel["avisos"]) else rel["processos"])
                for formato in ("docx_a", "xlsx_b", "lista", "tabela_livre"):
                    leitores.ler(ruim, formato=formato)       # nenhum formato forçado pode levantar exceção
            # bytes aleatórios com a mesma extensão
            lixo = self.pasta / f"lixo{arq.suffix}"
            lixo.write_bytes(bytes(rng.randrange(256) for _ in range(500)))
            for formato in (None, "docx_a", "xlsx_b", "lista", "tabela_livre"):
                rel = leitores.ler(lixo, formato=formato)
                self.assertEqual(rel["processos"], [])


class TestContrato(unittest.TestCase):
    """O que CONTRATOS §4 promete de todo RelatorioLido."""

    def test_estrutura_do_relatorio_e_do_processo(self):
        destino = gerar_xlsx_b(pasta("contrato") / "b.xlsx", fichas_200()[:20])
        rel = leitores.ler(destino)
        self.assertEqual(set(rel) - {"mapeamento"}, {"formato", "arquivo", "cliente", "data_base", "processos", "parametros",
                                                     "colunas_sem_destino", "avisos"})
        for p in rel["processos"]:
            for chave in ("numero", "vinculados", "campos", "andamentos_texto", "ultimo_andamento", "origem_no_arquivo"):
                self.assertIn(chave, p)
            for campo, c in p["campos"].items():
                self.assertIn(campo, ficha.CAMPOS)
                self.assertIn(c["origem"], ("migrado", "humano"))
            self.assertTrue(cart.numeros_no_texto(p["numero"])[0])
        for a in rel["avisos"]:
            self.assertEqual(set(a), {"nivel", "codigo", "onde", "mensagem", "candidatos"})
            self.assertIn(a["nivel"], ("info", "atencao", "erro"))
            self.assertTrue(a["codigo"])
        for c in rel["colunas_sem_destino"]:
            self.assertIn("coluna", c)
            self.assertIn("amostra", c)

    def test_valores_prontos_para_a_ficha(self):
        """O que o leitor devolve entra na ficha sem perda (ficha.definir aceita e valida)."""
        destino = gerar_xlsx_b(pasta("contrato") / "b2.xlsx", fichas_200()[:60])
        rel = leitores.ler(destino)
        for p in rel["processos"]:
            g = ficha.nova_ficha(p["numero"])
            for campo, c in p["campos"].items():
                self.assertTrue(ficha.definir(g, campo, c["valor"], c["origem"]), (campo, c))
                self.assertEqual(ficha.obter(g, campo), c["valor"], campo)
            self.assertEqual(ficha.validar(g), [] if not g["campos"].get("momento_atual") else ficha.validar(g))

    def test_nenhum_numero_real_nos_arquivos_do_pacote(self):
        """Mesmo padrão do empacotar.sh: só números sintéticos permitidos nos arquivos desta entrega."""
        padrao = re.compile(r"\b[0-9]{7}-[0-9]{2}\.[0-9]{4}\.[0-9]\.[0-9]{2}\.[0-9]{4}\b")
        permitido = re.compile(r"0000000-00\.0000\.0\.00\.0000|9999999-99\.9999\.9\.99\.9999|123456[7-9]-|1234570-")
        arquivos = [*(RAIZ / "src" / "leitores").glob("*.py"), Path(__file__)]
        for arq in arquivos:
            for n, linha in enumerate(arq.read_text(encoding="utf-8").splitlines(), 1):
                for achado in padrao.findall(linha):
                    self.assertTrue(permitido.search(achado), f"{arq.name}:{n} {achado}")


if __name__ == "__main__":
    unittest.main()
