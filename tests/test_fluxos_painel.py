"""Os fluxos pelo painel (WS-14): do primeiro uso, sem relatório nenhum, até a página de Entregas, com os módulos REAIS
(leitores, consolidar, fila, escritores, qualidade) e o coletor simulado no lugar do ColetorReal. Tudo com
`app.test_client()`, dados fictícios, sem rede.

    python3 -m unittest tests/test_fluxos_painel.py -v
"""
import io
import re
import shutil
import sys
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent))
from isolamento import TMP  # noqa: E402  (antes de tudo)
import ficticio  # noqa: E402
import simulado  # noqa: E402
import comum  # noqa: E402
import ficha  # noqa: E402
from flask import Flask  # noqa: E402
from test_fluxos import ColetorAte, escrever_docx_a  # noqa: E402

TOKEN = "token-de-teste"
DATA_1, DATA_2 = "2026-07-31", "2026-08-31"


class PainelDosFluxos(unittest.TestCase):
    N, CLIENTES = 40, 2

    def setUp(self):
        from painel import assistente, base, entregas, migracao, perfil, revisao_eventos
        import acesso
        import fila
        self.ass, self.ent, self.base, self.fila_mod = assistente, entregas, base, fila
        self.carteira = ficticio.gerar_carteira(self.N, clientes=self.CLIENTES, semente=5, com_linha_de_base=True)
        nomes = sorted({ficha.obter(f, "cliente") for f in self.carteira})
        self.c1 = [f for f in self.carteira if ficha.obter(f, "cliente") == nomes[0]]
        self.c2 = [f for f in self.carteira if ficha.obter(f, "cliente") == nomes[1]]
        # nenhum relatório existe: é o primeiro uso (o painel de verdade começaria em /novo)
        shutil.rmtree(TMP / "projetos-de-teste", ignore_errors=True)
        ficticio.criar_projeto_de_teste([], nome="Provisório")
        shutil.rmtree(comum.PROJETOS_DIR / "provisorio", ignore_errors=True)
        comum.PROJETO = None
        self.addCleanup(ficticio.restaurar_comum)
        self.pasta = TMP / f"fluxos-painel-{self._testMethodName}"
        shutil.rmtree(self.pasta, ignore_errors=True)
        self.pasta.mkdir(parents=True)
        cofre = {"cert_senha": "x", "totp_secret": "y"}
        patches = [mock.patch.object(acesso, "obter", lambda chave: cofre.get(chave)),
                   mock.patch.object(assistente, "INTERVALO_DA_JANELA_S", 0.05),
                   mock.patch.object(entregas, "abrir_fila", self._abrir_fila)]
        for p in patches:
            p.start()
            self.addCleanup(p.stop)
        simulacao = simulado.ColetorSimulado(self.carteira, semente=1, taxa_falha=0.0, pasta=self.pasta / "docs-sim")
        self.coletor = ColetorAte(simulacao, DATA_1)
        assistente.FABRICA_DE_COLETOR = lambda: self.coletor
        self.addCleanup(setattr, assistente, "FABRICA_DE_COLETOR", None)
        self._zerar_execucao()
        self.addCleanup(self._parar_execucao)
        base.TAREFA.clear()
        self.app = Flask(__name__)
        self.app.config["TESTING"] = True
        token_ok = base.criar_token_ok(TOKEN)
        for tela in (base, revisao_eventos, assistente, migracao, entregas, perfil):
            tela.registrar(self.app, TOKEN, base.cabecalho, token_ok)
        self.c = self.app.test_client()

    # --- apoio

    def _abrir_fila(self):
        """A fila do painel sem pausa entre processos (a configuração real espera segundos de cada vez)."""
        return self.fila_mod.Fila(comum.PROJETO, pausa_s=(0, 0))

    def _zerar_execucao(self):
        self.ass.EXEC.update(thread=None, slug=None, fila=None, estado="parado", pedido=None, resumo=None, erro=None,
                             inicio=None, modo="continuo")
        self.ass.EXEC["log"].clear()

    def _parar_execucao(self):
        t = self.ass.EXEC["thread"]
        if t is not None and t.is_alive():
            self.ass.EXEC["pedido"] = "parada"
            self.ass._ESPERA.set()
            t.join(10)
        self._zerar_execucao()

    def get(self, caminho, seguir=True):
        r = self.c.get(caminho, follow_redirects=seguir)
        return r, r.get_data(as_text=True)

    def post(self, caminho, dados=None, seguir=True, **kw):
        r = self.c.post(caminho, data={**(dados or {}), "token": TOKEN}, follow_redirects=seguir, **kw)
        return r, r.get_data(as_text=True)

    def enviar(self, caminho, campo, arquivos, extra=None):
        dados = {"token": TOKEN, **(extra or {}), campo: [(io.BytesIO(Path(a).read_bytes()), Path(a).name) for a in arquivos]}
        r = self.c.post(caminho, data=dados, content_type="multipart/form-data", follow_redirects=True)
        return r, r.get_data(as_text=True)

    def esperar_coleta(self, segundos=60):
        t = self.ass.EXEC["thread"]
        self.assertIsNotNone(t, "a coleta não foi iniciada")
        t.join(segundos)
        self.assertFalse(t.is_alive(), "a coleta não terminou a tempo")

    def fichas(self):
        return ficha.carregar(todas=True)

    def lote_da_pagina(self, texto):
        return re.search(r"name='lote' value='([^']+)'", texto).group(1)

    def revisar(self):
        """A pessoa que revisa: a triagem em lote aprova os tranquilos; o resto é aprovado evento a evento na tela."""
        import triagem
        lista = comum.eventos()
        triagem.aplicar_lote(lista, self.fichas(), "Revisor Fictício", pct=0)
        comum.salvar_eventos(lista)
        pendentes = [e for e in comum.eventos() if e["status"] == "rascunho"]
        for e in pendentes:
            r, _ = self.post("/evento", {"id": e["id"], "acao": "aprovar", "frase": e["frase"], "conteudo": e.get("conteudo") or "",
                                         "prazo": e.get("prazo") or "", "audiencia": e.get("audiencia") or ""}, seguir=False)
            self.assertEqual(r.status_code, 302)
        self.assertEqual([e for e in comum.eventos() if e["status"] == "rascunho"], [])
        return len(pendentes)

    # --- o caminho completo

    def test_primeiro_uso_importar_elaborar_inicial_e_entregas(self):
        # 1. sem relatório nenhum o assistente abre (não manda para /novo)
        r, t = self.get("/fluxo", seguir=False)
        self.assertEqual(r.status_code, 200)
        self.assertIn("Importar relatórios existentes", t)
        # 2. importar: o .docx do cliente 1 (com histórico) e uma lista com os processos do cliente 2
        docx = escrever_docx_a(self.c1, self.pasta / "relatorio-cliente-1.docx", ficha.obter(self.c1[0], "cliente"), "2026-06-30")
        lista = self.pasta / "lista.txt"
        lista.write_text("\n".join(f["numero"] for f in self.c2), encoding="utf-8")
        r, t = self.enviar("/fluxo/importar/enviar", "arquivos", [docx, lista])
        self.assertEqual(r.request.path, "/fluxo/importar/conferir")
        self.assertIn(f"{self.N}", t)
        lote = self.lote_da_pagina(t)
        r, t = self.post("/fluxo/importar/confirmar", {"lote": lote, "nome": "Relatório do Painel"})
        self.assertIn("Relatório criado: Relatório do Painel", t)
        self.assertEqual(comum.PROJETO, "relatorio-do-painel")
        fichas = self.fichas()
        self.assertEqual(len(fichas), self.N)
        sem_relatorio = [f for f in fichas if not f.get("linha_de_base")]
        self.assertEqual({f["numero"] for f in sem_relatorio}, {f["numero"] for f in self.c2})
        # 3. elaborar o relatório inicial só dos que não têm relatório (modo imediato: pede confirmação)
        r, t = self.get("/fluxo/inicial")
        self.assertIn("sem relatório anterior", t)
        r, t = self.post("/fluxo/inicial/preparar", {"com_entregas": "1", "entregas": ["docx_a", "xlsx_b", "dashboard"],
                                                     "profundidade": "padrao", "modo_coleta": "imediato", "so_novos": "1"})
        self.assertEqual(r.request.path, "/fluxo/confirmar")
        self.assertIn("Confirmar e começar agora", t)
        r, t = self.post("/fluxo/comecar")
        self.esperar_coleta()
        r, t = self.get("/fluxo/progresso.json")
        dados = r.get_json()
        self.assertEqual(dados["estado"], "concluida", dados)
        self.assertEqual(dados["resumo"]["coletado"], len(self.c2))
        # 4. a coleta gravou a capa, os eventos (como rascunho) e o momento atual
        with_capa = [f for f in self.fichas() if f["numero"] in {c["numero"] for c in self.c2}]
        for f in with_capa:
            self.assertEqual(ficha.origem(f, "valor_causa"), "coletado")
        eventos = comum.eventos()
        self.assertTrue(eventos)
        self.assertEqual({e["status"] for e in eventos} - {"rascunho", "descartado"}, set())
        self.assertTrue(all(e["numero"] in {c["numero"] for c in self.c2} for e in eventos))
        self.assertTrue(any(ficha.obter(f, "momento_atual") for f in with_capa))
        r, t = self.get("/")
        self.assertIn("para revisar", t)
        # 5. revisão (triagem em lote + tela evento a evento) e a página de Entregas
        self.assertGreater(self.revisar(), -1)
        r, t = self.post("/entregas/gerar", {"entregas": ["docx_a", "xlsx_b", "dashboard"]})
        self.assertIn("arquivo(s) gerado(s) em saida/", t)
        rodada = next(p for p in self.ent.pasta_saida().iterdir() if p.is_dir())
        nomes = sorted(a.name for a in rodada.iterdir())
        self.assertEqual(sum(n.endswith(".docx") for n in nomes), 2)
        self.assertTrue(any(n.startswith("Planilha") for n in nomes) and any(n.startswith("Painel") for n in nomes))
        r, t = self.get("/entregas")
        self.assertIn("Arquivos entregues", t)
        self.assertIn("Qualidade da base", t)
        self.assertIn("/entregas/baixar?p=", t)
        for f in self.fichas():
            self.assertTrue(f.get("ultimo_texto_gravado"), f["numero"])
        self.assertEqual(len([e for e in comum.eventos() if e["status"] == "aprovado"]), 0)    # tudo virou relatado

    def test_atualizar_com_o_docx_enviado_gera_versao_nova_e_regenera_a_planilha(self):
        # prepara um relatório já entregue (o mês anterior)
        docx = escrever_docx_a(self.c1, self.pasta / "relatorio-cliente-1.docx", ficha.obter(self.c1[0], "cliente"), DATA_1)
        r, t = self.enviar("/fluxo/importar/enviar", "arquivos", [docx])
        r, t = self.post("/fluxo/importar/confirmar", {"lote": self.lote_da_pagina(t), "nome": "Relatório do Painel"})
        self.assertEqual(comum.PROJETO, "relatorio-do-painel")
        self.assertEqual(len(self.fichas()), len(self.c1))
        # atualizar: o painel confere o arquivo com a carteira, coleta só o posterior e para na revisão
        self.coletor.ate = DATA_2
        r, t = self.get("/fluxo/atualizar")
        self.assertIn("Conferir com a carteira", t)
        r, t = self.enviar("/fluxo/atualizar/enviar", "arquivos", [docx])
        self.assertEqual(r.request.path, "/fluxo/atualizar/conferir")
        r, t = self.post("/fluxo/atualizar/preparar", {"lote": self.lote_da_pagina(t), "com_entregas": "1",
                                                       "entregas": ["docx_a", "xlsx_b"], "profundidade": "padrao",
                                                       "modo_coleta": "imediato", "incluir_novos": "1", "incluir_sumidos": "1"})
        self.assertEqual(r.request.path, "/fluxo/confirmar")
        self.post("/fluxo/comecar")
        self.esperar_coleta()
        eventos = comum.eventos()
        self.assertTrue(eventos)
        # só o que veio depois da data-base do relatório enviado (desde é estrito)
        for e in eventos:
            self.assertGreater(ficha.parse_data(e["data"]), DATA_1, e["titulo"])
        self.revisar()
        r, t = self.post("/entregas/gerar", {"entregas": ["docx_a", "xlsx_b"]})
        self.assertIn("arquivo(s) gerado(s)", t)
        rodada = next(p for p in self.ent.pasta_saida().iterdir() if p.is_dir())
        novo_docx = next(rodada.glob("*.docx"))
        self.assertNotEqual(novo_docx.read_bytes(), docx.read_bytes())
        self.assertTrue(docx.exists())                                # o original continua intacto
        self.assertTrue(list(rodada.glob("*.xlsx")))                  # o outro formato foi regenerado a partir da ficha
        from escritores import docx_a
        est = docx_a.ler_estrutura(novo_docx)
        self.assertEqual({b["numeros"][0] for b in est["processos"]}, {f["numero"] for f in self.c1})

    def test_migrar_de_modelo_pelo_painel(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["Nº do Processo", "Cliente", "Juízo", "Valor da causa (R$)", "Pasta física nº"])
        for f in self.c1[:8]:
            ws.append([f["numero"], ficha.obter(f, "cliente"), ficha.obter(f, "vara"), float(ficha.dinheiro(ficha.obter(f, "valor_causa"))),
                       "nota"])
        tabela = self.pasta / "controle.xlsx"
        wb.save(tabela)
        r, t = self.get("/migracao", seguir=False)
        self.assertEqual(r.status_code, 200)                           # sem relatório, a migração também abre
        r, t = self.enviar("/migracao/enviar", "arquivo", [tabela])
        self.assertEqual(r.request.path, "/migracao/mapear")
        lote = self.lote_da_pagina(t)
        indices = {m.group(2): m.group(1) for m in re.finditer(r"name='map_(\d+)' aria-label='Campo para ([^']+)'", t)}
        form = {"lote": lote, "nome": "Convertido pelo Painel", "estilo_texto": "a", "modelos": ["xlsx_b", "docx_a"], "acao": "converter"}
        for coluna, i in indices.items():
            m = re.search(rf"name='map_{i}'.*?</select>", t, re.S)
            selecionado = re.search(r"<option value='([^']*)' selected", m.group(0))
            form[f"map_{i}"] = selecionado.group(1) if selecionado else ""
        r, t = self.post("/migracao/converter", form)
        self.assertEqual(r.request.path, "/entregas")
        self.assertIn("Relatório convertido: Convertido pelo Painel. 8 processo(s).", t)
        rodada = next(p for p in self.ent.pasta_saida().iterdir() if p.is_dir())
        self.assertTrue(list(rodada.glob("*.xlsx")) and list(rodada.glob("*.docx")))
        import openpyxl
        wb2 = openpyxl.load_workbook(next(rodada.glob("*.xlsx")), data_only=True)
        achatado = " ".join(str(c) for ws in wb2 for linha in ws.iter_rows(values_only=True) for c in linha if c)
        self.assertIn("Pasta física nº", achatado)               # a coluna sem destino não se perdeu

    def test_pagina_de_entregas_mostra_o_que_conferir(self):
        docx = escrever_docx_a(self.c1, self.pasta / "r.docx", ficha.obter(self.c1[0], "cliente"), DATA_1)
        r, t = self.enviar("/fluxo/importar/enviar", "arquivos", [docx])
        self.post("/fluxo/importar/confirmar", {"lote": self.lote_da_pagina(t), "nome": "Relatório do Painel"})
        alvo = next(f["numero"] for f in self.fichas() if f.get("ativo", True))
        self.coletor.base.falhar_em = {alvo: "captcha"}
        self.coletor.ate = DATA_2
        self.post("/fluxo/atualizar/preparar", {"lote": "000000000000"}, seguir=False)    # lote inexistente: 404, sem efeito
        r, t = self.enviar("/fluxo/atualizar/enviar", "arquivos", [docx])
        self.post("/fluxo/atualizar/preparar", {"lote": self.lote_da_pagina(t), "com_entregas": "1", "entregas": ["docx_a"],
                                                "profundidade": "padrao", "modo_coleta": "imediato"})
        self.post("/fluxo/comecar")
        self.esperar_coleta()
        self.revisar()
        self.post("/entregas/gerar", {"entregas": ["docx_a"]})
        r, t = self.get("/entregas")
        self.assertIn("Conferir manualmente", t)
        self.assertIn(alvo, t)
        self.assertIn("captcha", t)


if __name__ == "__main__":
    unittest.main()
