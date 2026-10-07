"""Desempenho (teste transversal do WS-13): uma carteira fictícia de 200 processos percorre
leitura -> consolidação -> qualidade -> escritores (e a fila de coleta) em tempo razoável.

Cada ETAPA mede o próprio tempo e tem um limite GENEROSO (o objetivo é pegar uma regressão grosseira, como um
algoritmo quadrático que passa a levar minutos, e não medir o computador). Etapas cujo módulo ainda não existe
são PULADAS com mensagem clara; o coordenador as reativa na integração (nada a editar aqui).

Ajustes por variável de ambiente (computador lento, carteira maior):

    RELATORIO_DESEMPENHO_N=500        número de processos (padrão 200)
    RELATORIO_DESEMPENHO_FATOR=3      multiplica todos os limites (padrão 1)
    RELATORIO_DESEMPENHO_SAIDA=r.json grava a tabela de tempos em JSON (para comparar entre execuções)

    python3 -m unittest tests/test_desempenho.py -v

Etapas (na ordem do pipeline): carteira_e_ficha, fase1_planilha_e_relatorio, lista_bruta_fase1, leitura_docx_a,
leitura_xlsx_b, leitura_lista, consolidacao, qualidade, escritor_docx_a, escritor_xlsx_b, dashboard, fila.
`test_pipeline_completo` roda todas em sequência, confere o orçamento total e, se alguma foi pulada, termina como
"pulado" listando o que falta (as etapas que rodaram continuam valendo). A tabela de tempos vai para o stderr.
"""
import copy
import json
import os
import signal
import sys
import tempfile
import time
import traceback
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent))
import transversal as T  # noqa: E402  (isolamento antes de tudo)
import comum  # noqa: E402
import ficha as fch  # noqa: E402
import ficticio  # noqa: E402
import simulado  # noqa: E402

N = int(os.environ.get("RELATORIO_DESEMPENHO_N", "200"))
FATOR = float(os.environ.get("RELATORIO_DESEMPENHO_FATOR", "1"))
# Limites em segundos para 200 processos (valores generosos; escalam com N/200 e com o FATOR).
LIMITES_BASE = {"carteira_e_ficha": 30, "fase1_planilha_e_relatorio": 60, "lista_bruta_fase1": 30,
                "leitura_docx_a": 60, "leitura_xlsx_b": 60, "leitura_lista": 30, "consolidacao": 60, "qualidade": 60,
                "escritor_docx_a": 120, "escritor_xlsx_b": 120, "dashboard": 120, "fila": 120}
TOTAL_BASE = 600


def limite(nome):
    return LIMITES_BASE[nome] * FATOR * max(1.0, N / 200)


def limite_total():
    return TOTAL_BASE * FATOR * max(1.0, N / 200)


class Pulada(unittest.SkipTest):
    """Etapa que depende de módulo que ainda não existe."""


class Estouro(Exception):
    """A etapa passou do dobro do limite e foi interrompida."""


class Prazo:
    """Interrompe a etapa (SIGALRM, onde existir) se passar de `segundos`: evita que uma fila que espera o relógio
    deixe a suíte pendurada. Sem SIGALRM (Windows), só o limite medido ao final vale."""
    def __init__(self, segundos):
        self.segundos = segundos

    def __enter__(self):
        if hasattr(signal, "setitimer") and hasattr(signal, "SIGALRM"):
            def estourou(_sig, _frame):
                raise Estouro(f"interrompida após {self.segundos:.0f}s (o dobro do limite)")
            self.anterior = signal.signal(signal.SIGALRM, estourou)
            signal.setitimer(signal.ITIMER_REAL, self.segundos)
        return self

    def __exit__(self, *a):
        if hasattr(signal, "setitimer") and hasattr(signal, "SIGALRM"):
            signal.setitimer(signal.ITIMER_REAL, 0)
            signal.signal(signal.SIGALRM, self.anterior)


def numeros_no_arquivo(caminho):
    """Números CNJ (com máscara) presentes no arquivo, inclusive por dentro de .docx/.xlsx."""
    import confidencialidade_regras as R
    return set(R.CNJ_MASCARA.findall(T.texto_do_arquivo(caminho)))


# ---------------------------------------------------------------- as etapas
# Cada etapa recebe `c` (contexto compartilhado: fichas, pasta temporária) e devolve um texto curto com o que fez.

def etapa_carteira_e_ficha(c):
    c["fichas"] = ficticio.gerar_carteira(N, clientes=5, semente=1, com_linha_de_base=True)
    assert len(c["fichas"]) == N
    erros = [f["numero"] for f in c["fichas"] if fch.validar(f)]
    assert not erros, f"{len(erros)} fichas inválidas na carteira limpa"
    with ficticio.projeto_de_teste(c["fichas"]):
        lidas = fch.carregar(todas=True)
        assert len(lidas) == N, f"ficha.carregar devolveu {len(lidas)} de {N}"
    return f"{N} fichas geradas, validadas, gravadas e relidas"


def etapa_fase1_planilha_e_relatorio(c):
    """Fluxo da Fase 1 (não pode regredir com a refatoração do WS-7): relatório HTML e planilha.gerar."""
    import datetime
    import openpyxl
    import planilha
    import relatorio
    fichas = c["fichas"]
    eventos = T.eventos_aprovados(fichas, por_ficha=3)
    modelo = c["pasta"] / "fase1-modelo.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = planilha.ABA
    ws["A1"], ws["P1"] = "Número do Processo", "Andamentos"
    for i, f in enumerate(fichas, start=2):
        ws[f"A{i}"] = f["numero"]
        ws[f"P{i}"] = "Em 01/09/2026 foi proferido despacho."
    wb.save(modelo)
    with ficticio.projeto_de_teste(fichas):
        comum.salvar_eventos(eventos)
        por_cliente = {}
        for e in eventos:
            por_cliente.setdefault(e["cliente"], []).append(e)
        paginas = [relatorio.gerar_html(cliente, evs, datetime.datetime(2026, 10, 7)) for cliente, evs in por_cliente.items()]
        assert paginas and all(len(p) > 1000 for p in paginas)
        atualizados, fora = planilha.gerar(modelo, c["pasta"] / "fase1-saida.xlsx", hoje=datetime.date(2026, 10, 7))
    assert len(atualizados) == N and not fora, f"planilha.gerar atualizou {len(atualizados)} de {N}; sem linha: {len(fora)}"
    return f"{len(paginas)} relatórios HTML e planilha da Fase 1 com {N} linhas e {len(eventos)} eventos"


def _listas_grandes(c):
    """Lista bruta com TODOS os processos (.csv com ';' e .xlsx simples) mais um número com dígito errado, que deve
    ser recusado e listado. (ficticio.gerar_lista_bruta só faz blocos pequenos; aqui o que importa é o volume.)"""
    import openpyxl
    if "listas" in c:
        return c["listas"]
    pasta = c["pasta"] / "listas"
    pasta.mkdir(exist_ok=True)
    numeros = [f["numero"] for f in c["fichas"]]
    errado = ficticio.numero_com_dv_errado(numeros[-1])
    linhas = ["Nº;Cliente;Polo;Parte contrária"] + [
        f"{f['numero']};{fch.obter(f, 'cliente') or ''};Réu;{fch.obter(f, 'parte_contraria') or ''}" for f in c["fichas"]]
    linhas.append(f"{errado};Cliente Exemplo 01 Ltda;Réu;Pessoa Fictícia 0000")
    csv = pasta / "lista-grande.csv"
    csv.write_text("\n".join(linhas) + "\n", encoding="utf-8")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Processo", "Cliente", "Polo do cliente"])
    for f in c["fichas"]:
        ws.append([f["numero"], fch.obter(f, "cliente"), fch.obter(f, "polo_cliente")])
    ws.append([errado, "Cliente Exemplo 01 Ltda", "passivo"])
    xlsx = pasta / "lista-grande.xlsx"
    wb.save(xlsx)
    c["listas"] = {"csv": csv, "xlsx": xlsx, "validos": len(set(numeros)), "invalido": errado}
    return c["listas"]


def etapa_lista_bruta_fase1(c):
    import carteira
    listas = _listas_grandes(c)
    for formato in ("csv", "xlsx"):
        registros, invalidos = carteira.ler_lista(str(listas[formato]))
        assert len(registros) == listas["validos"], f"{formato}: {len(registros)} lidos de {listas['validos']}"
        assert [i for i in invalidos if listas["invalido"] in str(i)], f"{formato}: o número com DV errado não foi listado"
    return f"lista bruta (carteira.ler_lista): {listas['validos']} números em .csv e em .xlsx, 1 recusado"


def etapa_leitura_docx_a(c):
    leitores = T.exigir("leitores", "leitor .docx")
    eventos = T.eventos_aprovados(c["fichas"], por_ficha=2)
    arquivo = T.docx_de_entrada(c["pasta"] / "entrada-a.docx", c["fichas"], eventos)
    lido = leitores.ler(arquivo)
    assert lido["formato"] == "docx_a", lido["formato"]
    assert len(lido["processos"]) == N, f"{len(lido['processos'])} processos lidos de {N}"
    erros = [a for a in lido["avisos"] if a.get("nivel") == "erro"]
    assert not erros, f"{len(erros)} avisos de erro na leitura: {erros[:2]}"
    return f"{len(lido['processos'])} processos lidos do .docx ({len(lido['avisos'])} avisos)"


def etapa_leitura_xlsx_b(c):
    leitores = T.exigir("leitores", "leitor .xlsx")
    arquivo = T.xlsx_de_entrada(c["pasta"] / "entrada-b.xlsx", N, c["fichas"])
    lido = leitores.ler(arquivo)
    assert lido["formato"] == "xlsx_b", lido["formato"]
    assert len(lido["processos"]) == N, f"{len(lido['processos'])} processos lidos de {N}"
    erros = [a for a in lido["avisos"] if a.get("nivel") == "erro"]
    assert not erros, f"{len(erros)} avisos de erro na leitura: {erros[:2]}"
    return f"{len(lido['processos'])} processos lidos do .xlsx ({len(lido['avisos'])} avisos)"


def etapa_leitura_lista(c):
    leitores = T.exigir("leitores", "leitor de lista")
    listas = _listas_grandes(c)
    for formato in ("csv", "xlsx"):
        lido = leitores.ler(listas[formato])
        assert lido["formato"] in ("lista", "tabela_livre"), lido["formato"]
        assert len(lido["processos"]) == listas["validos"], f"{formato}: {len(lido['processos'])} lidos de {listas['validos']}"
        assert any(listas["invalido"] in str(a) for a in lido["avisos"]), f"{formato}: DV errado sem aviso"
    return f"{listas['validos']} números lidos da lista bruta (.csv e .xlsx), DV errado avisado"


def etapa_consolidacao(c):
    consolidar = T.exigir("consolidar", "consolidação")
    sujas = ficticio.gerar_carteira(N, clientes=5, semente=2, com_defeitos=True)
    c["sujas"] = sujas
    fichas, avisos = consolidar.consolidar(copy.deepcopy(list(sujas)))
    assert isinstance(avisos, list) and len(fichas) <= N and len(fichas) > 0.8 * N
    return f"{N} fichas com defeitos consolidadas em {len(fichas)} ({len(avisos)} avisos)"


def etapa_qualidade(c):
    qualidade = T.exigir("qualidade", "verificador de qualidade")
    sujas = c.get("sujas") or ficticio.gerar_carteira(N, clientes=5, semente=2, com_defeitos=True)
    achados = qualidade.verificar(list(sujas), T.perfil_padrao())
    assert isinstance(achados, list) and achados, "carteira com defeitos intencionais não gerou nenhum achado"
    assert all({"codigo", "gravidade", "mensagem"} <= set(a) for a in achados[:20])
    limpa = qualidade.verificar(list(c["fichas"]), T.perfil_padrao())
    assert isinstance(limpa, list)
    return f"{len(achados)} achados na carteira suja; {len(limpa)} na limpa"


def etapa_escritor_docx_a(c):
    escritor = T.exigir("escritores.docx_a", "escritor .docx")
    fichas = c["fichas"]
    # ciclo 1: do zero a partir do modelo padrão
    estado1 = T.estado_de_teste(fichas, "2026-09-18", eventos=T.eventos_aprovados(fichas, 2))
    saida1 = c["pasta"] / "relatorio-a-ciclo1.docx"
    r1 = escritor.gravar(None, estado1, saida1)
    assert Path(r1["destino"]).exists()
    presentes = numeros_no_arquivo(saida1)
    faltam = {f["numero"] for f in fichas} - presentes
    assert not faltam, f"{len(faltam)} processos não aparecem no .docx do ciclo 1"
    # ciclo 2: atualiza o arquivo do ciclo 1 com eventos novos (o original não pode ser tocado)
    antes = saida1.read_bytes()
    estado2 = T.estado_de_teste(fichas, "2026-10-30", eventos=T.eventos_aprovados(fichas, 1, a_partir="2026-10-20"))
    saida2 = c["pasta"] / "relatorio-a-ciclo2.docx"
    r2 = escritor.gravar(saida1, estado2, saida2)
    assert Path(r2["destino"]).exists() and saida1.read_bytes() == antes, "o original foi alterado"
    assert not ({f["numero"] for f in fichas} - numeros_no_arquivo(saida2)), "processo sumiu no ciclo 2"
    return f"{N} processos em 2 ciclos (criar do zero + atualizar)"


def etapa_escritor_xlsx_b(c):
    escritor = T.exigir("escritores.xlsx_b", "escritor .xlsx")
    fichas = c["fichas"]
    estado1 = T.estado_de_teste(fichas, "2026-09-18", eventos=T.eventos_aprovados(fichas, 2))
    saida1 = c["pasta"] / "relatorio-b-ciclo1.xlsx"
    r1 = escritor.gravar(None, estado1, saida1)
    assert Path(r1["destino"]).exists()
    faltam = {f["numero"] for f in fichas} - numeros_no_arquivo(saida1)
    assert not faltam, f"{len(faltam)} processos não aparecem no .xlsx do ciclo 1"
    antes = saida1.read_bytes()
    estado2 = T.estado_de_teste(fichas, "2026-10-30", eventos=T.eventos_aprovados(fichas, 1, a_partir="2026-10-20"))
    saida2 = c["pasta"] / "relatorio-b-ciclo2.xlsx"
    r2 = escritor.gravar(saida1, estado2, saida2)
    assert Path(r2["destino"]).exists() and saida1.read_bytes() == antes, "o original foi alterado"
    assert not ({f["numero"] for f in fichas} - numeros_no_arquivo(saida2)), "processo sumiu no ciclo 2"
    c["xlsx_gerado"] = saida2
    return f"{N} processos em 2 ciclos (criar do zero + atualizar)"


def etapa_dashboard(c):
    dash = T.exigir("escritores.dashboard", "gerador de dashboards")
    import confidencialidade_regras as R
    entrada = T.xlsx_de_entrada(c["pasta"] / "entrada-dash.xlsx", N, c["fichas"])
    for modo in ("modelo", "embutido"):
        destino = c["pasta"] / f"painel-{modo}.html"
        dash.gravar(entrada, destino, T.perfil_padrao(), modo=modo)
        html = destino.read_text(encoding="utf-8")
        assert len(html) > 5000, f"dashboard {modo} pequeno demais ({len(html)} bytes)"
        assert R.referencias_externas(html) == [], f"dashboard {modo} referencia domínio externo"
    return "dashboard nos dois modos (modelo e embutido)"


def etapa_fila(c):
    fila = T.exigir("fila", "fila de coleta")
    fichas = c["fichas"]
    sem_pausa = {"coleta": {"max_documentos_por_rodada": 30, "pausa_entre_processos_s": [0, 0],
                            "pausa_entre_documentos_s": [0, 0]}}
    with ficticio.projeto_de_teste(fichas) as proj, mock.patch("time.sleep") as dormir, \
            mock.patch.object(comum, "config", lambda: sem_pausa):
        for nome in ("fila", "coletor", "trt", "rodar"):         # módulos que fizeram `from comum import config`
            m = sys.modules.get(nome)
            if m is not None and hasattr(m, "config") and m.config is not comum.config:
                stack = mock.patch.object(m, "config", lambda: sem_pausa)
                stack.start()
                c.setdefault("_patches", []).append(stack)
        try:
            f = fila.Fila(proj["slug"])
            f.enfileirar([x["numero"] for x in fichas], modo="imediato", profundidade="rapido")
            coletor = simulado.ColetorSimulado(fichas, semente=1, taxa_falha=0.1)
            fila.rodar_fila(f, coletor)
            resumo = f.resumo()
            cobertura = fila.cobertura(proj["slug"])
        finally:
            for p in c.pop("_patches", []):
                p.stop()
    assert resumo["total"] == N, resumo
    assert resumo["pendente"] == 0 and resumo["coletando"] == 0, f"a fila terminou com itens abertos: {resumo}"
    assert resumo["coletado"] + resumo["erro"] + resumo["manual"] == N, resumo
    assert resumo["coletado"] >= 0.7 * N, f"só {resumo['coletado']} de {N} coletados com 10% de falha simulada: {resumo}"
    assert isinstance(cobertura, dict) and cobertura, "cobertura vazia"
    return f"fila: {resumo['coletado']} coletados, {resumo['manual']} manuais, {resumo['erro']} com erro; {coletor.chamadas} chamadas"


ETAPAS = {nome[len("etapa_"):]: fn for nome, fn in list(globals().items()) if nome.startswith("etapa_")}
ORDEM = [n for n in LIMITES_BASE if n in ETAPAS]
assert set(ORDEM) == set(ETAPAS), "toda etapa precisa de limite em LIMITES_BASE"


# ---------------------------------------------------------------- execução e relatório

class Desempenho(unittest.TestCase):
    resultados = {}     # etapa -> {"estado": ok|pulada|falhou, "segundos", "detalhe", "erro"}
    contexto = {}

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory(prefix="desempenho-")
        cls.contexto = {"pasta": Path(cls.tmp.name)}
        cls.resultados = {}

    @classmethod
    def tearDownClass(cls):
        if cls.resultados:
            linhas = ["", f"Desempenho com {N} processos (limites x{FATOR:g}):"]
            for nome in ORDEM:
                r = cls.resultados.get(nome)
                if r:
                    linhas.append(f"  {nome:28s} {r['estado']:7s} {r['segundos']:7.2f}s / {limite(nome):.0f}s  {r['detalhe']}")
            sys.stderr.write("\n".join(linhas) + "\n")
            saida = os.environ.get("RELATORIO_DESEMPENHO_SAIDA")
            if saida:
                Path(saida).write_text(json.dumps({"n": N, "fator": FATOR, "etapas": cls.resultados}, ensure_ascii=False,
                                                  indent=2), encoding="utf-8")
        cls.tmp.cleanup()

    @classmethod
    def executar(cls, nome):
        """Roda a etapa uma única vez (as demais a reaproveitam) e guarda o resultado."""
        if nome in cls.resultados:
            return cls.resultados[nome]
        if "fichas" not in cls.contexto and nome != "carteira_e_ficha":
            cls.executar("carteira_e_ficha")
        inicio = time.perf_counter()
        try:
            with Prazo(limite(nome) * 2):
                detalhe = ETAPAS[nome](cls.contexto)
            r = {"estado": "ok", "detalhe": detalhe}
        except unittest.SkipTest as e:
            r = {"estado": "pulada", "detalhe": str(e)}
        except BaseException as e:  # noqa: BLE001  (inclui Estouro e AssertionError: viram falha da etapa)
            r = {"estado": "falhou", "detalhe": f"{type(e).__name__}: {e}", "erro": traceback.format_exc()}
        r["segundos"] = round(time.perf_counter() - inicio, 3)
        cls.resultados[nome] = r
        return r

    def conferir(self, nome):
        r = self.executar(nome)
        if r["estado"] == "pulada":
            self.skipTest(r["detalhe"])
        if r["estado"] == "falhou":
            self.fail(f"etapa {nome} falhou: {r['detalhe']}\n{r.get('erro', '')}")
        self.assertLessEqual(r["segundos"], limite(nome),
                             f"etapa {nome} levou {r['segundos']:.1f}s (limite {limite(nome):.0f}s para {N} processos; "
                             f"use RELATORIO_DESEMPENHO_FATOR em computador lento)")

    # um teste por etapa (nome com número só para a ordem de execução na saída)
    def test_01_carteira_e_ficha(self):
        self.conferir("carteira_e_ficha")

    def test_02_fase1_planilha_e_relatorio(self):
        self.conferir("fase1_planilha_e_relatorio")

    def test_03_lista_bruta_fase1(self):
        self.conferir("lista_bruta_fase1")

    def test_04_leitura_docx_a(self):
        self.conferir("leitura_docx_a")

    def test_05_leitura_xlsx_b(self):
        self.conferir("leitura_xlsx_b")

    def test_06_leitura_lista(self):
        self.conferir("leitura_lista")

    def test_07_consolidacao(self):
        self.conferir("consolidacao")

    def test_08_qualidade(self):
        self.conferir("qualidade")

    def test_09_escritor_docx_a(self):
        self.conferir("escritor_docx_a")

    def test_10_escritor_xlsx_b(self):
        self.conferir("escritor_xlsx_b")

    def test_11_dashboard(self):
        self.conferir("dashboard")

    def test_12_fila(self):
        self.conferir("fila")

    def test_99_pipeline_completo(self):
        """Todas as etapas em sequência dentro do orçamento total; etapa pulada = módulo ainda ausente."""
        for nome in ORDEM:
            self.executar(nome)
        falhas = {n: r for n, r in self.resultados.items() if r["estado"] == "falhou"}
        self.assertEqual(falhas, {}, "etapas que falharam: " + "; ".join(f"{n}: {r['detalhe']}" for n, r in falhas.items()))
        total = sum(r["segundos"] for r in self.resultados.values())
        self.assertLessEqual(total, limite_total(),
                             f"pipeline levou {total:.0f}s (orçamento {limite_total():.0f}s para {N} processos)")
        puladas = [n for n in ORDEM if self.resultados[n]["estado"] == "pulada"]
        if puladas:
            self.skipTest(f"pipeline PARCIAL: {len(ORDEM) - len(puladas)} de {len(ORDEM)} etapas rodaram; "
                          f"faltam (módulo ainda ausente): {', '.join(puladas)}; reativar na integração")


if __name__ == "__main__":
    unittest.main()
