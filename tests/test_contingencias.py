"""Migração de planilhas "desformatadas" de CONTINGÊNCIAS JURÍDICAS (passivo contingente): leitor, campos novos do modelo B,
relatório de lacunas ("o que falta para a transição"), escritor da planilha, painel HTML e as telas de migração.

A planilha de entrada é FICTÍCIA (`contingencias_ficticias.py`, gerada em tempo de execução): títulos nas linhas 1 a 3,
cabeçalho na linha 5, valores em formatos misturados, grau de perda colado à justificativa, fórmulas sem valor guardado,
rodapé com total e quadro de critérios. Sem rede, sem IA, sem dados de cliente.

    python3 -m unittest tests/test_contingencias.py -v
"""
import html
import io
import json
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from isolamento import TMP  # noqa: E402  (antes de tudo)
import contingencias_ficticias as cf  # noqa: E402
import ficticio  # noqa: E402
import ficha  # noqa: E402
import lacunas  # noqa: E402
import leitores  # noqa: E402
from leitores import base, grade  # noqa: E402

JSC = next((c for c in (shutil.which("jsc"), "/System/Library/Frameworks/JavaScriptCore.framework/Versions/A/Helpers/jsc")
            if c and Path(c).exists()), None)
PASTA_DASHBOARD = Path(__file__).resolve().parent.parent / "src" / "modelos" / "dashboard"
HARNESS = Path(__file__).resolve().parent / "fixtures" / "dashboard_jsc.js"
TOKEN = "token-de-teste"


def pasta_temporaria():
    pasta = Path(tempfile.mkdtemp(dir=TMP))
    return pasta


# ================================================================ valores

class TestValores(unittest.TestCase):
    def test_dinheiro_em_todos_os_formatos(self):
        casos = {"R$ 260.123,33": "260123.33", "R$41.000,00": "41000.00", 124229.25: "124229.25", "124229.25": "124229.25",
                 "  R$ 175.913,36 ": "175913.36", "-R$ 1.234,56": "-1234.56", "R$ -1.234,56": "-1234.56", "(1.234,56)": "-1234.56",
                 "R$ 1 234,56": "1234.56", "R$ 0,00": "0.00", 0: "0.00", 53806.655: "53806.66", "1.234,56-": "-1234.56"}
        for bruto, esperado in casos.items():
            self.assertEqual(base.converter_dinheiro(bruto), (esperado, None), repr(bruto))
        for vazio in ("-", "R$ -", "", None, "  ", "N/A", "–"):
            self.assertEqual(base.converter_dinheiro(vazio), (None, None), repr(vazio))
        self.assertEqual(base.converter_dinheiro("abc"), (None, "invalido"))
        self.assertEqual(base.converter_dinheiro("R$ 10,00 e R$ 20,00"), (None, "ambiguo"))

    def test_percentual_como_fracao(self):
        for bruto, esperado in ((0.5, 0.5), ("50%", 0.5), ("50", 0.5), ("0,5", 0.5), (1.0, 1.0), ("100%", 1.0), (0.25, 0.25), (25, 0.25)):
            self.assertEqual(base.converter_numero(bruto, percentual=True), (esperado, None), repr(bruto))
        conv = base.converter_campo("percentual_provisao", "50%")
        self.assertEqual(conv["valor"], 0.5)
        self.assertEqual(base.converter_campo("percentual_provisao", 0.5, formato_celula="0%")["valor"], 0.5)

    def test_sim_nao(self):
        for bruto, esperado in (("SIM", "Sim"), ("NÃO", "Não"), ("Sim", "Sim"), ("nao", "Não"), ("não", "Não"), ("S", "Sim")):
            self.assertEqual(base.converter_campo("deposito_judicial", bruto)["valor"], esperado, bruto)
        self.assertIsNone(base.converter_campo("deposito_judicial", "-")["valor"])
        self.assertIsNone(base.converter_campo("deposito_judicial", "")["valor"])
        conv = base.converter_campo("deposito_judicial", "talvez")
        self.assertIsNone(conv["valor"])
        self.assertEqual(conv["avisos"][0][1], "valor_invalido")

    def test_grau_e_justificativa_separados(self):
        casos = {"REMOTA - Processo extinto sem resolução do mérito": ("Remota", "Processo extinto sem resolução do mérito"),
                 "PROVÁVEL - Celebração de acordo": ("Provável", "Celebração de acordo"),
                 "POSSÍVEL ": ("Possível", ""), "possivel": ("Possível", ""), "Remoto: acordo": ("Remota", "acordo"),
                 "PROVAVEL (acordo fechado)": ("Provável", "acordo fechado"),
                 "REMOTA - não é possível recorrer": ("Remota", "não é possível recorrer")}
        for texto, esperado in casos.items():
            self.assertEqual(base.separar_probabilidade(texto), esperado, texto)
        self.assertEqual(base.separar_probabilidade("Alta"), (None, ""))
        conv = base.converter_campo("probabilidade", "REMOTA - Realizado acordo")
        self.assertEqual(conv["valor"], "Remota")
        self.assertEqual(conv["extras"]["justificativa_probabilidade"], "Realizado acordo")
        self.assertEqual(base.converter_campo("probabilidade", "Provável")["extras"], {})

    def test_campos_novos_da_ficha(self):
        novos = ("passivo_potencial", "passivo_atualizado", "ativo_potencial", "percentual_provisao", "provisao",
                 "deposito_judicial", "cnpj_processado", "pagamento_realizado", "justificativa_probabilidade")
        self.assertEqual(tuple(ficha.CAMPOS_DE_CONTINGENCIA), novos)
        tipos = {"passivo_potencial": "dinheiro", "passivo_atualizado": "dinheiro", "ativo_potencial": "dinheiro",
                 "percentual_provisao": "numero", "provisao": "dinheiro", "deposito_judicial": "sim_nao",
                 "cnpj_processado": "texto", "pagamento_realizado": "dinheiro", "justificativa_probabilidade": "texto"}
        for campo, tipo in tipos.items():
            self.assertEqual(ficha.CAMPOS[campo][1:3], ("contingencia", tipo), campo)
        f = ficha.nova_ficha(ficticio.numero_ficticio(1))
        self.assertTrue(ficha.definir(f, "percentual_provisao", "50%", "humano"))
        self.assertEqual(ficha.obter(f, "percentual_provisao"), 0.5)
        self.assertTrue(ficha.definir(f, "passivo_potencial", "R$ 1.000,50", "humano"))
        self.assertEqual(ficha.obter(f, "passivo_potencial"), "1000.50")
        self.assertTrue(ficha.definir(f, "deposito_judicial", "SIM", "humano"))
        self.assertEqual(ficha.obter(f, "deposito_judicial"), "Sim")
        self.assertEqual(ficha.validar(f), [])

    def test_derivar_deposito_e_percentual(self):
        campos = {"depositos_recursais": {"valor": "9000.00", "origem": "humano"}, "provisao": {"valor": "2500.00", "origem": "humano"},
                  "passivo_potencial": {"valor": "10000.00", "origem": "humano"}}
        self.assertEqual(ficha.derivar_contingencia(campos), ["deposito_judicial", "percentual_provisao"])
        self.assertEqual((campos["deposito_judicial"]["valor"], campos["deposito_judicial"]["origem"]), ("Sim", "derivado"))
        self.assertEqual((campos["percentual_provisao"]["valor"], campos["percentual_provisao"]["origem"]), (0.25, "derivado"))
        # o que a pessoa informou nunca é trocado; sem passivo ou com valor zerado, nada é inventado
        campos = {"depositos_recursais": {"valor": "9000.00", "origem": "humano"}, "deposito_judicial": {"valor": "Não", "origem": "humano"},
                  "provisao": {"valor": "10.00", "origem": "humano"}, "percentual_provisao": {"valor": 0.9, "origem": "humano"}}
        self.assertEqual(ficha.derivar_contingencia(campos), [])
        self.assertEqual(campos["deposito_judicial"]["valor"], "Não")
        self.assertEqual(ficha.derivar_contingencia({"depositos_recursais": {"valor": "0.00", "origem": "humano"}}), [])
        self.assertEqual(ficha.derivar_contingencia({"provisao": {"valor": "10.00", "origem": "humano"}}), [])


    def test_derivar_situacao_pelo_ativo_da_ficha(self):
        f = ficha.nova_ficha(ficticio.numero_ficticio(1))
        self.assertTrue(ficha.derivar_situacao(f))
        self.assertEqual((ficha.obter(f, "situacao"), ficha.origem(f, "situacao")), ("Ativo", "derivado"))
        g = ficha.nova_ficha(ficticio.numero_ficticio(2))
        g["ativo"] = False
        ficha.derivar_situacao(g)
        self.assertEqual(ficha.obter(g, "situacao"), "Encerrado")
        h = ficha.nova_ficha(ficticio.numero_ficticio(3))
        ficha.definir(h, "situacao", "Suspenso", "humano")        # o que a pessoa informou nunca é trocado
        self.assertFalse(ficha.derivar_situacao(h))
        self.assertEqual(ficha.obter(h, "situacao"), "Suspenso")


# ================================================================ fórmulas

class TestFormulas(unittest.TestCase):
    def grade(self, linhas, formulas):
        return grade.Grade("x", linhas, formulas=set(formulas), textos_formulas=formulas)

    def test_conta_simples_sem_valor_guardado(self):
        g = self.grade([["R$ 1.000,00", 400.5, None], ["x", "y", None]], {(0, 2): "=A1-B1"})
        self.assertEqual(grade.calcular_formula(g, 0, 2), Decimal("599.50"))

    def test_celula_em_branco_vale_zero_e_soma(self):
        g = self.grade([[10, None, None], [20, None, None], [None, None, None]],
                       {(0, 1): "=A1-B2", (2, 0): "=SOMA(A1:A2)", (2, 1): "=SUM(A1, A2, 5)", (2, 2): "=(A1+A2)*2/4"})
        self.assertEqual(grade.calcular_formula(g, 0, 1), Decimal("10.00"))
        self.assertEqual(grade.calcular_formula(g, 2, 0), Decimal("30.00"))
        self.assertEqual(grade.calcular_formula(g, 2, 1), Decimal("35.00"))
        self.assertEqual(grade.calcular_formula(g, 2, 2), Decimal("15.00"))

    def test_formula_que_depende_de_outra_fica_encadeada(self):
        g = self.grade([[100, None, None]], {(0, 1): "=A1/4", (0, 2): "=B1+A1"})
        self.assertEqual(grade.calcular_formula(g, 0, 2), Decimal("125.00"))

    def test_o_que_nao_da_para_calcular_volta_vazio(self):
        g = self.grade([["texto", 1, None]], {(0, 2): "=A1-B1", (0, 0): "=SE(B1>0;1;2)"})
        self.assertIsNone(grade.calcular_formula(g, 0, 2))                  # texto não numérico
        self.assertIsNone(grade.calcular_formula(g, 0, 0))                  # função desconhecida
        g = self.grade([[1, 0, None]], {(0, 2): "=A1/B1"})
        self.assertIsNone(grade.calcular_formula(g, 0, 2))                  # divisão por zero
        g = self.grade([[1, None]], {(0, 1): "=Outra!A1+A1"})
        self.assertIsNone(grade.calcular_formula(g, 0, 1))                  # outra aba
        g = self.grade([[None]], {(0, 0): "=A1+1"})
        self.assertIsNone(grade.calcular_formula(g, 0, 0))                  # referência circular: nunca trava
        self.assertIsNone(grade.calcular_formula(g, 0, 5))                  # não é fórmula


# ================================================================ leitura da planilha

class TestLeitura(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.pasta = pasta_temporaria()
        cls.arquivo = cf.gravar(cls.pasta / "contingencias.xlsx")
        cls.rel = leitores.ler(cls.arquivo)
        cls.por_numero = {p["numero"]: p for p in cls.rel["processos"]}

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.pasta, ignore_errors=True)

    def valor(self, numero, campo):
        c = self.por_numero[numero]["campos"].get(campo)
        return c["valor"] if c else None

    def avisos(self, codigo):
        return [a for a in self.rel["avisos"] if a["codigo"] == codigo]

    def test_formato_e_quantidade(self):
        self.assertEqual(leitores.detectar(self.arquivo), "tabela_livre")
        self.assertEqual(self.rel["formato"], "tabela_livre")
        self.assertEqual(len(self.rel["processos"]), cf.GABARITO["processos"])
        self.assertEqual(sorted(self.por_numero), sorted(cf.NUMEROS_ATIVOS + cf.NUMEROS_ARQUIVADOS))
        self.assertEqual(sum(1 for p in self.rel["processos"] if p.get("ativo") is False), cf.GABARITO["arquivados"])

    def test_mapeamento_das_colunas_de_contingencia(self):
        mapa = {(m["aba"][:9], m["coluna"]): m for m in self.rel["mapeamento"]}
        esperado = {"CNPJ EMPRESA EXEMPLO PROCESSADA": "cnpj_processado", "ATIVO POTENCIAL": "ativo_potencial",
                    "PASSIVO POTENCIAL": "passivo_potencial", "DEPÓSITO JUDICIAL REALIZADO?": "deposito_judicial",
                    "VALOR DO DEPÓSITO JUDICIAL": "depositos_recursais", "POSSIBILIDADE DE PERDA": "probabilidade",
                    "%": "percentual_provisao", "PROVISÃO CONSTITUÍDA": "provisao", "BREVE RESUMO DO CASO": "objeto",
                    "OBSERVAÇÃO": "andamentos", "AUTOR/RECLAMANTE": "autores", "RÉU/RECLAMADO": "reus", "PROCESSO Nº.": "numero"}
        for coluna, campo in esperado.items():
            m = mapa[("Relatório", coluna)]
            self.assertEqual((m["campo"], m["aplicado"]), (campo, True), coluna)
        arq = {m["coluna"]: m for m in self.rel["mapeamento"] if m["aba"].startswith("Arquivados")}
        self.assertEqual(arq["PASSIVO POTENCIAL ATUALIZADO (CORREÇÃO INPC E JUROS 1% AO MÊS)"]["campo"], "passivo_atualizado")
        self.assertEqual(arq["PAGAMENTO REALIZADO"]["campo"], "pagamento_realizado")
        self.assertEqual(arq["DATA DE AJUIZAMENTO"]["campo"], "data_ajuizamento")
        self.assertEqual(arq["VALOR ECONOMIZADO"]["campo"], "valor_economizado")
        self.assertEqual(self.rel["colunas_sem_destino"], [])

    def test_ativo_potencial_nao_e_a_coluna_ativo(self):
        cands = grade.pontuar_cabecalho("ATIVO POTENCIAL")
        self.assertEqual(cands[0], (1.0, "ativo_potencial"))
        self.assertNotEqual(grade.propor_mapeamento(["Processo", "ATIVO POTENCIAL"])[1]["campo"], "ativo")

    def test_natureza_da_acao_e_assunto_nao_classe(self):
        m = next(m for m in self.rel["mapeamento"] if m["coluna"] == "NATUREZA DA AÇÃO")
        self.assertEqual(m["campo"], "assunto")
        self.assertEqual(m["confianca"], 0.85)
        self.assertIn("materia_principal", m["candidatos"])
        self.assertNotIn("classe", m["candidatos"])
        self.assertIn("Natureza da ação", m["comentario"])
        self.assertIn("classe processual", m["comentario"])
        self.assertEqual(self.valor(cf.NUMEROS_ATIVOS[0], "assunto"), "INDENIZAÇÃO POR DANOS MORAIS E MATERIAIS")
        self.assertIsNone(self.valor(cf.NUMEROS_ATIVOS[0], "classe"))
        # classe de verdade continua sendo classe
        self.assertEqual(grade.propor_mapeamento(["Processo", "Classe processual"])[1]["campo"], "classe")

    def test_porcentagem_so_ao_lado_da_provisao(self):
        com_vizinha = grade.propor_mapeamento(["Processo", "POSSIBILIDADE DE PERDA", "%", "PROVISÃO CONSTITUÍDA"])
        self.assertEqual((com_vizinha[2]["campo"], com_vizinha[2]["aplicado"]), ("percentual_provisao", True))
        sem_vizinha = grade.propor_mapeamento(["Processo", "Cidade", "%", "Observações"])
        self.assertIsNone(sem_vizinha[2]["campo"])
        self.assertFalse(sem_vizinha[2]["aplicado"])
        escrito = grade.propor_mapeamento(["Processo", "% de provisão"])
        self.assertEqual(escrito[1]["campo"], "percentual_provisao")

    def test_cabecalhos_com_palavras_a_mais(self):
        casos = {"CNPJ DA EMPRESA FULANA PROCESSADA": "cnpj_processado", "Passivo potencial atualizado (R$)": "passivo_atualizado",
                 "VALOR DO DEPÓSITO JUDICIAL (R$)": "depositos_recursais", "Provisão constituída (R$)": "provisao",
                 "Depósito judicial realizado? (sim/não)": "deposito_judicial"}
        for cab, campo in casos.items():
            self.assertEqual(grade.pontuar_cabecalho(cab)[0][1], campo, cab)

    def test_contagens_de_campos_preenchidos(self):
        cont = {}
        for p in self.rel["processos"]:
            for campo in p["campos"]:
                cont[campo] = cont.get(campo, 0) + 1
        esperado = {c: cf.GABARITO[c] for c in ("passivo_potencial", "passivo_atualizado", "ativo_potencial", "provisao", "depositos_recursais",
                                                "probabilidade", "justificativa_probabilidade", "pagamento_realizado", "valor_economizado",
                                                "cnpj_processado")}
        esperado["percentual_provisao"] = cf.GABARITO["percentual_provisao"] - 1        # o do processo 9 só nasce na ficha (derivado)
        esperado["deposito_judicial"] = cf.GABARITO["deposito_judicial"] - 1            # idem, derivado do valor depositado
        for campo, n in esperado.items():
            self.assertEqual(cont.get(campo), n, campo)

    def test_valores_dos_ativos(self):
        n = cf.NUMEROS_ATIVOS
        v = self.valor
        self.assertEqual((v(n[0], "passivo_potencial"), v(n[0], "provisao"), v(n[0], "percentual_provisao")), ("260123.33", "130061.66", 0.5))
        self.assertEqual((v(n[0], "deposito_judicial"), v(n[0], "probabilidade")), ("Não", "Possível"))
        self.assertEqual(v(n[0], "cnpj_processado"), cf.cnpj_ficticio(1))
        self.assertEqual((v(n[1], "passivo_potencial"), v(n[1], "provisao"), v(n[1], "percentual_provisao")), ("41000.00", "41000.00", 1.0))
        self.assertEqual((v(n[2], "passivo_potencial"), v(n[2], "provisao")), ("124229.25", "62114.63"))     # número puro
        self.assertEqual((v(n[2], "deposito_judicial"), v(n[2], "depositos_recursais")), ("Sim", "70115.85"))
        self.assertEqual(v(n[3], "depositos_recursais"), "144458.62")
        self.assertEqual((v(n[4], "passivo_potencial"), v(n[4], "ativo_potencial")), ("175913.36", "53839.36"))   # espaço no fim
        self.assertEqual(v(n[4], "probabilidade"), "Remota")
        self.assertIsNone(v(n[4], "justificativa_probabilidade"))
        self.assertIsNone(v(n[5], "passivo_potencial"))                      # "-" é vazio
        self.assertIsNone(v(n[5], "provisao"))
        self.assertEqual(v(n[5], "probabilidade"), "Possível")               # minúsculas
        self.assertEqual(v(n[6], "ativo_potencial"), "-500.00")              # negativo
        self.assertEqual((v(n[6], "probabilidade"), v(n[6], "justificativa_probabilidade")), ("Possível", "Sem decisão"))   # sem acento
        self.assertEqual(v(n[7], "depositos_recursais"), "9000.00")
        self.assertIsNone(v(n[7], "deposito_judicial"))                      # o leitor não deduz: a ficha deriva do valor
        self.assertIsNone(v(n[8], "percentual_provisao"))

    def test_origem_dos_campos_de_contingencia(self):
        p = self.por_numero[cf.NUMEROS_ATIVOS[0]]["campos"]
        for campo in ("passivo_potencial", "provisao", "percentual_provisao", "deposito_judicial", "cnpj_processado"):
            self.assertEqual(p[campo]["origem"], "humano", campo)
        self.assertEqual(p["assunto"]["origem"], "migrado")
        self.assertEqual(self.por_numero[cf.NUMEROS_ARQUIVADOS[0]]["campos"]["valor_economizado"]["origem"], "migrado")   # fórmula nunca é humano

    def test_arquivados(self):
        a = cf.NUMEROS_ARQUIVADOS
        v = self.valor
        for numero in a:
            self.assertIs(self.por_numero[numero]["ativo"], False)
        self.assertEqual(v(a[0], "data_ajuizamento"), "2023-03-14")
        self.assertEqual((v(a[0], "passivo_potencial"), v(a[0], "passivo_atualizado"), v(a[0], "pagamento_realizado")),
                         ("22921.00", "41149.37", "0.00"))
        self.assertEqual(v(a[0], "probabilidade"), "Remota")
        self.assertEqual(v(a[0], "justificativa_probabilidade"), "Processo extinto sem resolução do mérito, com trânsito em julgado")
        self.assertEqual(v(a[1], "probabilidade"), "Provável")
        self.assertEqual(v(a[1], "justificativa_probabilidade"), "Celebração de acordo")
        self.assertEqual(v(a[1], "assunto"), "TRABALHISTA")                  # espaço no fim sai
        self.assertEqual((v(a[2], "deposito_judicial"), v(a[2], "depositos_recursais")), ("Sim", "7352.00"))
        self.assertIsNone(v(a[2], "pagamento_realizado"))

    def test_formulas_sem_valor_guardado_sao_calculadas(self):
        a = cf.NUMEROS_ARQUIVADOS
        self.assertEqual(self.valor(a[0], "valor_economizado"), "41149.37")      # =G4-K4 com K4 = 0
        self.assertEqual(self.valor(a[1], "valor_economizado"), "42566.01")      # R$ 48.665,84 - 6.099,83
        self.assertEqual(self.valor(a[2], "valor_economizado"), "7352.00")       # K em branco vale zero
        calculadas = self.avisos("formula_calculada")
        self.assertEqual(len(calculadas), 1)
        self.assertEqual(calculadas[0]["candidatos"], ["L4", "L5", "L6"])
        self.assertEqual(self.avisos("formula_sem_valor"), [])

    def test_titulos_totais_e_rodape_nao_viram_processo(self):
        ignoradas = {a["onde"]: a for a in self.avisos("linhas_ignoradas")}
        ativos = next(a for onde, a in ignoradas.items() if "Relatório" in onde)
        self.assertIn("5 linha(s)", ativos["mensagem"])           # total + cabeçalho do quadro + 3 critérios
        arquivados = next(a for onde, a in ignoradas.items() if "Arquivados" in onde)
        self.assertIn("1 linha(s)", arquivados["mensagem"])       # VALOR TOTAL ECONOMIZADO
        self.assertTrue(all(re.fullmatch(r"\d{7}-\d{2}\.\d{4}\.\d\.\d{2}\.\d{4}", p["numero"]) for p in self.rel["processos"]))

    def test_numero_invalido_no_meio_dos_dados_continua_sendo_erro(self):
        invalidos = self.avisos("numero_invalido")
        self.assertEqual(len(invalidos), 1)
        self.assertEqual(invalidos[0]["nivel"], "erro")
        self.assertIn("linha 8", invalidos[0]["onde"])
        self.assertIn("SEM NÚMERO AINDA", invalidos[0]["mensagem"])

    def test_soma_do_passivo_dos_ativos(self):
        total = sum(Decimal(self.valor(n, "passivo_potencial")) for n in cf.NUMEROS_ATIVOS if self.valor(n, "passivo_potencial"))
        self.assertEqual(f"{total:.2f}", cf.GABARITO["total_passivo_ativos"])

    def test_leitura_de_xlsx_sem_fechar_arquivo_e_erro_de_arquivo(self):
        rel = leitores.ler(self.pasta / "nao-existe.xlsx", formato="tabela_livre")
        self.assertEqual(rel["processos"], [])
        self.assertEqual(rel["avisos"][0]["codigo"], "arquivo_inexistente")


# ================================================================ lacunas

class TestLacunas(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.pasta = pasta_temporaria()
        cls.rel = leitores.ler(cf.gravar(cls.pasta / "contingencias.xlsx"))
        cls.lac = lacunas.lacunas(cls.rel)
        cls.campos = {c["campo"]: c for g in cls.lac["grupos"] for c in g["campos"]}

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.pasta, ignore_errors=True)

    def conferir(self, campo, preenchidos, total, origem, completar):
        c = self.campos[campo]
        self.assertEqual((c["preenchidos"], c["total"]), (preenchidos, total), campo)
        self.assertEqual(c["origem"], origem, campo)
        self.assertEqual(c["completar_por"], completar, campo)

    def test_totais_do_arquivo(self):
        self.assertEqual((self.lac["total_processos"], self.lac["ativos"], self.lac["encerrados"]), (12, 9, 3))
        self.assertEqual([g["grupo"] for g in self.lac["grupos"]],
                         ["gestao", "partes", "capa", "situacao", "julgamento", "contingencia"])
        self.assertEqual(set(self.campos), set(ficha.CAMPOS))               # todo campo do modelo B aparece, agrupado como em ficha.CAMPOS
        for g in self.lac["grupos"]:
            self.assertTrue(all(ficha.CAMPOS[c["campo"]][1] == g["grupo"] for c in g["campos"]))

    def test_o_que_o_arquivo_traz(self):
        col = lambda nome: f"do arquivo, coluna «{nome}»"      # noqa: E731
        self.conferir("autores", 12, 12, col("AUTOR/RECLAMANTE"), None)
        self.conferir("assunto", 12, 12, col("NATUREZA DA AÇÃO"), None)
        self.conferir("passivo_potencial", 11, 12, col("PASSIVO POTENCIAL"), lacunas.VOCE)
        self.conferir("passivo_atualizado", 3, 3, col("PASSIVO POTENCIAL ATUALIZADO (CORREÇÃO INPC E JUROS 1% AO MÊS)"), None)
        self.conferir("pagamento_realizado", 2, 3, col("PAGAMENTO REALIZADO"), lacunas.VOCE)
        self.conferir("data_ajuizamento", 3, 12, col("DATA DE AJUIZAMENTO"), lacunas.BUSCAVEL)
        self.conferir("cnpj_processado", 9, 12, col("CNPJ EMPRESA EXEMPLO PROCESSADA"), lacunas.VOCE)
        self.conferir("justificativa_probabilidade", 4, 12, col("POSSIBILIDADE DE PERDA"), lacunas.DEDUZIVEL)

    def test_o_que_se_aplica_so_a_parte_dos_processos(self):
        self.assertEqual(self.campos["provisao"]["aplica_a"], "risco")
        self.conferir("provisao", 7, 8, "do arquivo, coluna «PROVISÃO CONSTITUÍDA»", lacunas.VOCE)       # 8 ativos com perda provável/possível
        self.conferir("percentual_provisao", 7, 8, "do arquivo, coluna «%»", lacunas.DEDUZIVEL)
        self.conferir("depositos_recursais", 4, 4, "do arquivo, coluna «VALOR DO DEPÓSITO JUDICIAL»", None)  # só onde há depósito
        self.conferir("data_transito", 0, 3, lacunas.DEDUZIVEL, lacunas.DEDUZIVEL)                       # só os 3 encerrados

    def test_deduzidos_pelo_programa_ficam_a_confirmar(self):
        self.assertEqual(self.campos["deposito_judicial"]["a_confirmar"], 1)     # processo 8: Sim a partir do valor depositado
        self.assertEqual(self.campos["percentual_provisao"]["a_confirmar"], 1)   # processo 9: 2.500 ÷ 10.000
        self.assertEqual(self.campos["deposito_judicial"]["preenchidos"], 12)
        self.assertEqual(self.lac["resumo"]["a_confirmar"], sum(c["a_confirmar"] for c in self.campos.values()))
        self.assertEqual(self.campos["momento_atual"]["a_confirmar"], 9)         # deduzido do histórico dos ativos

    def test_o_que_o_programa_busca_ou_deduz_e_o_que_precisa_de_voce(self):
        for campo in ("vara", "municipio", "uf", "classe", "valor_causa"):
            self.conferir(campo, 0, 12, lacunas.BUSCAVEL, lacunas.BUSCAVEL)
        for campo in ("area", "materia_principal", "resultado"):
            self.conferir(campo, 0, 12, lacunas.DEDUZIVEL, lacunas.DEDUZIVEL)
        # a situação (Ativo/Encerrado) o programa já deduz do ativo/arquivado: entra "a confirmar"
        self.assertEqual((self.campos["situacao"]["preenchidos"], self.campos["situacao"]["a_confirmar"]), (12, 12))
        for campo in ("cliente", "responsavel"):
            self.conferir(campo, 0, 12, lacunas.VOCE, lacunas.VOCE)

    def test_frases_do_que_falta(self):
        falta = self.lac["o_que_falta"]
        esperadas = [
            "Vara, Município, UF e Classe: 0 de 12 no arquivo; o programa busca no PJe/jus.br ao coletar.",
            "Valor da causa: 0 de 12 no arquivo; a capa do tribunal costuma trazer; se a coleta não achar, informe.",
            "Data do ajuizamento: 3 de 12 no arquivo; faltam 9: o programa busca no PJe/jus.br ao coletar.",
            "Cliente e Responsável: 0 de 12 no arquivo; precisa de você: informe na planilha ou na ficha do processo.",
            "Passivo potencial: 11 de 12 no arquivo; faltam 1: informe o valor em risco de cada processo "
            "(é estimativa do escritório; nenhum tribunal informa).",
            "Provisão constituída: 7 de 8 no arquivo (só dos 8 processos ativos com perda provável ou possível); faltam 1: informe a "
            "provisão constituída (é decisão contábil do cliente; nenhum tribunal informa).",
            "Pagamento realizado: 2 de 3 no arquivo (só dos 3 processos encerrados); faltam 1: informe quanto foi efetivamente pago em cada "
            "processo encerrado.",
        ]
        for frase in esperadas:
            self.assertIn(frase, falta)
        grupo = next(g for g in self.lac["grupos"] if g["grupo"] == "contingencia")
        self.assertIn(esperadas[4], grupo["o_que_falta"])
        self.assertNotIn(esperadas[0], grupo["o_que_falta"])                   # cada frase fica no seu grupo
        # campo opcional nunca vira "o que falta"
        self.assertFalse(any("Apelido" in f or "Contato" in f for f in falta))

    def test_percentuais_e_resumo(self):
        contingencia = next(g for g in self.lac["grupos"] if g["grupo"] == "contingencia")
        self.assertEqual((contingencia["preenchidos"], contingencia["total"]), (sum(
            c["preenchidos"] for c in contingencia["campos"] if not c["opcional"] and c["total"]), sum(
            c["total"] for c in contingencia["campos"] if not c["opcional"] and c["total"])))
        self.assertAlmostEqual(self.campos["passivo_potencial"]["percentual"], 11 / 12)
        self.assertEqual(self.campos["passivo_atualizado"]["percentual"], 1.0)
        self.assertEqual(self.campos["vara"]["percentual"], 0.0)
        resumo = self.lac["resumo"]
        self.assertGreater(resumo["campos_completos"], 0)
        self.assertGreater(resumo["campos_vazios"], resumo["campos_completos"])
        self.assertEqual(lacunas.percentual_texto(0.456), "46%")
        self.assertEqual(lacunas.percentual_texto(None), "—")

    def test_vale_para_fichas_e_para_o_relatorio_lido(self):
        from painel import assistente
        fichas, _ = assistente.fichas_do_relatorio(self.rel)
        de_fichas = lacunas.lacunas(fichas, mapeamento=self.rel["mapeamento"])
        for g1, g2 in zip(self.lac["grupos"], de_fichas["grupos"]):
            self.assertEqual([(c["campo"], c["preenchidos"], c["total"], c["origem"]) for c in g1["campos"]],
                             [(c["campo"], c["preenchidos"], c["total"], c["origem"]) for c in g2["campos"]])
        self.assertEqual(self.lac["o_que_falta"], de_fichas["o_que_falta"])

    def test_sem_colunas_mapeadas_a_origem_diz_so_do_arquivo(self):
        from painel import assistente
        fichas, _ = assistente.fichas_do_relatorio(self.rel)
        sem_mapa = {c["campo"]: c for g in lacunas.lacunas(fichas)["grupos"] for c in g["campos"]}
        self.assertEqual(sem_mapa["passivo_potencial"]["origem"], "do arquivo")

    def test_colunas_sem_destino_passam_para_o_quadro(self):
        lac = lacunas.lacunas(self.rel, colunas_sem_destino=[{"coluna": "Pasta física nº", "aba": "A", "amostra": ["1", "2", "3", "4"]}])
        self.assertEqual(lac["colunas_sem_destino"], [{"coluna": "Pasta física nº", "aba": "A", "amostra": ["1", "2", "3"]}])

    def test_sem_processos_nao_quebra(self):
        lac = lacunas.lacunas([])
        self.assertEqual((lac["total_processos"], lac["o_que_falta"], lac["resumo"]["percentual"]), (0, [], None))
        self.assertEqual(lacunas.lacunas({"processos": [], "formato": "tabela_livre"})["total_processos"], 0)

    def test_linhas_para_a_planilha(self):
        linhas = lacunas.linhas_para_planilha(self.lac)
        self.assertEqual(set(linhas[0]), {"Grupo", "Campo", "Preenchidos", "Total", "Percentual", "Origem", "O que fazer"})
        prov = next(l for l in linhas if l["Campo"] == "Provisão constituída")
        self.assertEqual((prov["Preenchidos"], prov["Total"], prov["Grupo"]), (7, 8, "Contingência (risco, provisão e depósitos)"))
        self.assertIn("decisão contábil", prov["O que fazer"])
        resumos = [l for l in linhas if l["Campo"] == "(resumo do grupo)"]
        self.assertEqual(len(resumos), sum(1 for g in self.lac["grupos"] if g["o_que_falta"]))
        self.assertFalse(any(l["Total"] == 0 and l["Preenchidos"] == 0 for l in linhas))


# ================================================================ escritor da planilha, leitor do modelo B e painel

class TestPlanilhaEPainel(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from painel import assistente
        from escritores import xlsx_b
        cls.x = xlsx_b
        cls.pasta = pasta_temporaria()
        cls.rel = leitores.ler(cf.gravar(cls.pasta / "contingencias.xlsx"))
        cls.fichas, _ = assistente.fichas_do_relatorio(cls.rel)
        cls.lac = lacunas.lacunas(cls.fichas, mapeamento=cls.rel["mapeamento"])
        cls.estado = {"cliente": "Empresa Exemplo Comercial Ltda", "data_base": "2026-09-30", "fichas": cls.fichas, "eventos": [],
                      "perfil": {}, "parametros": {"faltas_da_migracao": cls.lac}}
        cls.destino = cls.pasta / "saida.xlsx"
        cls.res = xlsx_b.gravar(None, cls.estado, cls.destino)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.pasta, ignore_errors=True)

    def test_gravou_sem_erro_e_passa_na_verificacao(self):
        self.assertEqual([a for a in self.res["avisos"] if a["nivel"] == "erro"], [])
        self.assertEqual(self.x.validar(self.destino), [])
        self.assertEqual(len(self.res["processos_novos"]), 12)

    def test_colunas_de_contingencia_no_fim_e_visiveis_quando_ha_dado(self):
        import openpyxl
        wb = openpyxl.load_workbook(self.destino)
        ws = wb["Processos"]
        cab = [c.value for c in ws[1]]
        self.assertEqual(cab[37:], [self.x.CAMPOS_B[c] for c in self.x.CONTINGENCIA])
        ocultas = {c for c, d in ws.column_dimensions.items() if d.hidden}
        for campo in self.x.CONTINGENCIA:
            self.assertNotIn(self.x.col_letra(cab.index(self.x.CAMPOS_B[campo]) + 1), ocultas, campo)
        linha = {cab[j]: ws.cell(2, j + 1).value for j in range(len(cab))}
        self.assertEqual(linha["Passivo Potencial"], 260123.33)
        self.assertEqual(linha["Provisão Constituída"], 130061.66)
        self.assertEqual(linha["Percentual de Provisão"], 0.5)
        self.assertEqual(linha["Depósito Judicial Realizado?"], "Não")
        self.assertEqual(linha["CNPJ da Empresa Processada"], cf.cnpj_ficticio(1))
        self.assertEqual(ws.tables["tblProcessos"].ref, "A1:AT13")

    def test_sem_dado_de_contingencia_as_colunas_continuam_ocultas(self):
        import openpyxl
        fichas = ficticio.gerar_carteira(3)
        destino = self.pasta / "sem-contingencia.xlsx"
        res = self.x.gravar(None, {"cliente": "Cliente Exemplo 01 Ltda", "data_base": "2026-09-30", "fichas": fichas, "eventos": [],
                                   "perfil": {}, "parametros": {}}, destino)
        self.assertEqual([a for a in res["avisos"] if a["nivel"] == "erro"], [])
        ws = openpyxl.load_workbook(destino)["Processos"]
        cab = [c.value for c in ws[1]]
        ocultas = {c for c, d in ws.column_dimensions.items() if d.hidden}
        for campo in self.x.CONTINGENCIA:
            self.assertIn(self.x.col_letra(cab.index(self.x.CAMPOS_B[campo]) + 1), ocultas, campo)

    def test_indicadores_de_contingencia(self):
        import openpyxl
        ws = openpyxl.load_workbook(self.destino)["Indicadores"]
        formulas = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value}
        esperados = {"Passivo potencial total": "SUM(tblProcessos[Passivo Potencial])",
                     "Passivo potencial dos processos ativos": 'SUMIFS(tblProcessos[Passivo Potencial],tblProcessos[Ativo],"Sim")',
                     "Provisão constituída total": "SUM(tblProcessos[Provisão Constituída])",
                     "Passivo potencial (ativos): perda remota": 'SUMIFS(tblProcessos[Passivo Potencial],tblProcessos[Probabilidade],"Remota",tblProcessos[Ativo],"Sim")',
                     "Passivo potencial (ativos): perda possível": 'SUMIFS(tblProcessos[Passivo Potencial],tblProcessos[Probabilidade],"Possível",tblProcessos[Ativo],"Sim")',
                     "Passivo potencial (ativos): perda provável": 'SUMIFS(tblProcessos[Passivo Potencial],tblProcessos[Probabilidade],"Provável",tblProcessos[Ativo],"Sim")',
                     "Processos com depósito judicial": 'COUNTIFS(tblProcessos[[Depósito Judicial Realizado?]],"Sim")',
                     "Total de depósitos judiciais": "SUM(tblProcessos[Depósitos Recursais])",
                     "Pagamentos realizados": "SUM(tblProcessos[Pagamento Realizado])",
                     "Valor economizado": "SUM(tblProcessos[Valor Economizado])",
                     "Passivo potencial atualizado (encerrados)": "SUM(tblProcessos[Passivo Potencial Atualizado])",
                     "Ativo potencial total": "SUM(tblProcessos[Ativo Potencial])"}
        for rotulo, formula in esperados.items():
            self.assertIn(rotulo, formulas)
            if formula:
                self.assertEqual(formulas[rotulo], "=" + formula, rotulo)
        # quadro que alimenta o gráfico do painel de planilha
        self.assertIn("Passivo potencial dos ativos por probabilidade", formulas)
        d = openpyxl.load_workbook(self.destino)["Dashboard"]
        rotulos = [d.cell(3, c).value for c in range(1, 14)]
        self.assertIn("Passivo potencial", rotulos)
        self.assertIn("Provisão constituída", rotulos)

    def test_aba_faltas_da_migracao(self):
        import openpyxl
        ws = openpyxl.load_workbook(self.destino)["Faltas da migração"]
        self.assertEqual([c.value for c in ws[1]], ["Grupo", "Campo", "Preenchidos", "Total", "Percentual", "Origem", "O que fazer"])
        linhas = {(ws.cell(r, 1).value, ws.cell(r, 2).value): [ws.cell(r, j).value for j in range(3, 8)] for r in range(2, ws.max_row + 1)}
        grupo = "Contingência (risco, provisão e depósitos)"
        self.assertEqual(linhas[(grupo, "Provisão constituída")][:2], [7, 8])
        self.assertAlmostEqual(linhas[(grupo, "Provisão constituída")][2], 7 / 8)
        self.assertIn("PROVISÃO CONSTITUÍDA", linhas[(grupo, "Provisão constituída")][3])
        self.assertEqual(linhas[("Capa do processo", "Vara / Juízo")][:2], [0, 12])
        resumo = linhas[(grupo, "(resumo do grupo)")]
        self.assertIn("Passivo potencial: 11 de 12 no arquivo", resumo[4])
        self.assertEqual(ws.tables["tblFaltas"].ref.split(":")[0], "A1")

    def test_gravar_de_novo_sobre_o_resultado_nao_duplica_as_faltas(self):
        import openpyxl
        segunda = self.pasta / "segunda.xlsx"
        res = self.x.gravar(self.destino, self.estado, segunda)
        self.assertEqual([a for a in res["avisos"] if a["nivel"] == "erro"], [])
        a = openpyxl.load_workbook(self.destino)["Faltas da migração"]
        b = openpyxl.load_workbook(segunda)["Faltas da migração"]
        self.assertEqual(a.max_row, b.max_row)
        self.assertEqual(self.x.validar(segunda), [])

    def test_ida_e_volta_leitor_ficha_escritor_leitor(self):
        rel2 = leitores.ler(self.destino)
        self.assertEqual(rel2["formato"], "xlsx_b")
        self.assertEqual([a for a in rel2["avisos"] if a["nivel"] == "erro"], [])
        por_numero = {p["numero"]: p for p in rel2["processos"]}
        self.assertEqual(len(por_numero), 12)
        for f in self.fichas:
            lido = por_numero[f["numero"]]["campos"]
            for campo in ficha.CAMPOS_DE_CONTINGENCIA:
                esperado = ficha.obter(f, campo)
                if esperado in (None, ""):
                    self.assertNotIn(campo, lido, (f["numero"], campo))
                    continue
                self.assertIn(campo, lido, (f["numero"], campo))
                if isinstance(esperado, float):
                    self.assertAlmostEqual(lido[campo]["valor"], esperado, places=6, msg=campo)
                else:
                    self.assertEqual(lido[campo]["valor"], esperado, (f["numero"], campo))
            self.assertEqual(lido.get("probabilidade", {}).get("valor"), ficha.obter(f, "probabilidade"))
            self.assertEqual(lido.get("depositos_recursais", {}).get("valor"), ficha.obter(f, "depositos_recursais"))

    def test_leitor_do_modelo_b_ignora_a_aba_de_faltas(self):
        rel2 = leitores.ler(self.destino)
        self.assertFalse(any("Faltas" in str(a.get("onde")) and a["codigo"] == "numero_invalido" for a in rel2["avisos"]))
        self.assertEqual(rel2["colunas_sem_destino"], [])

    def test_modelo_padrao_atual_tem_as_abas_e_colunas_novas(self):
        import openpyxl
        wb = openpyxl.load_workbook(self.x.MODELO_PADRAO)
        self.assertIn("Faltas da migração", wb.sheetnames)
        cab = [c.value for c in wb["Processos"][1]]
        self.assertEqual(len(cab), 46)
        self.assertEqual(cab[:37], [self.x.CAMPOS_B[c] for c in self.x.COLUNAS_PADRAO + self.x.COLUNAS_EXTRAS][:37])


# ================================================================ painel HTML (modelo C)

@unittest.skipUnless(JSC, "jsc (JavaScriptCore) ausente: o painel é conferido no navegador pelos testes com Playwright")
class TestPainelDeContingencia(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from painel import assistente
        from escritores import dashboard, xlsx_b
        cls.dashboard, cls.pasta = dashboard, pasta_temporaria()
        rel = leitores.ler(cf.gravar(cls.pasta / "contingencias.xlsx"))
        fichas, _ = assistente.fichas_do_relatorio(rel)
        cls.com = cls.pasta / "com.xlsx"
        xlsx_b.gravar(None, {"cliente": "Empresa Exemplo Comercial Ltda", "data_base": "2026-09-30", "fichas": fichas, "eventos": [],
                             "perfil": {}, "parametros": {}}, cls.com)
        cls.sem = cls.pasta / "sem.xlsx"
        xlsx_b.gravar(None, {"cliente": "Cliente Exemplo 01 Ltda", "data_base": "2026-09-30", "fichas": ficticio.gerar_carteira(6),
                             "eventos": [], "perfil": {}, "parametros": {}}, cls.sem)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.pasta, ignore_errors=True)

    def rodar(self, xlsx, template):
        html_ = self.pasta / f"{xlsx.stem}-{template}.html"
        r = self.dashboard.gravar(xlsx, html_, {}, modo="embutido", template=template, data_base="2026-09-30")
        self.assertEqual([a for a in r["avisos"] if a["nivel"] == "erro"], [])
        texto = html_.read_text(encoding="utf-8")
        cfg = re.search(r'<script type="application/json" id="cfg">(.*?)</script>', texto, re.S).group(1)
        arquivo = self.pasta / f"cfg-{xlsx.stem}-{template}.json"
        arquivo.write_text(cfg, encoding="utf-8")
        saida = subprocess.run([JSC, str(HARNESS), "--", str(arquivo), template, str(PASTA_DASHBOARD) + "/"], capture_output=True,
                               text=True, timeout=120)
        self.assertEqual(saida.returncode, 0, saida.stderr + saida.stdout)
        return texto, json.loads(saida.stdout.strip().splitlines()[-1])

    def test_cartoes_e_distribuicao_por_probabilidade_com_dados(self):
        for template in ("contencioso", "carteira"):
            texto, r = self.rodar(self.com, template)
            c = r["contingencia"]
            self.assertTrue(c["tem"])
            self.assertFalse(r["bloco_oculto"])
            self.assertEqual((r["total"], r["ativos"]), (12, 9))
            self.assertAlmostEqual(c["passivoAtivos"], float(cf.GABARITO["total_passivo_ativos"]), places=2)
            self.assertAlmostEqual(r["exposicao"], c["passivoAtivos"], places=2)        # passivo lançado é a exposição
            self.assertAlmostEqual(sum(c["passivoProb"].values()), c["passivoAtivos"], places=2)
            self.assertAlmostEqual(c["passivoProb"]["Provável"], 25361.40 + 41000.00 + 9000.00 + 10000.00, places=2)
            self.assertAlmostEqual(c["passivoProb"]["Remota"], 175913.36, places=2)
            self.assertAlmostEqual(c["provisao"], 130061.66 + 41000 + 62114.63 + 25361.4 + 25000 + 9000 + 2500, places=2)
            self.assertAlmostEqual(c["pagamentos"], 6099.83, places=2)
            self.assertEqual(c["depositosQtd"], 4)
            self.assertEqual(r["kpis"], ["passivo_potencial", "provisao", "passivo_provavel", "ativo_potencial", "depositos_judiciais", "pagamentos"])
            self.assertIn("ativo_sem_passivo", r["qualidade"])

    def test_sem_dados_de_contingencia_o_bloco_some_e_nada_quebra(self):
        for template in ("contencioso", "carteira"):
            _, r = self.rodar(self.sem, template)
            self.assertFalse(r["contingencia"]["tem"])
            self.assertTrue(r["bloco_oculto"])
            self.assertEqual(r["kpis"], [])
            self.assertEqual(r["total"], 6)
            self.assertNotIn("ativo_sem_passivo", r["qualidade"])

    def test_a_pagina_traz_o_bloco_e_continua_sem_recurso_externo(self):
        texto, _ = self.rodar(self.com, "contencioso")
        self.assertIn("Contingência e provisão", texto)
        self.assertIn("passivo_prob", texto)
        self.assertNotRegex(texto, r"(src|href)=['\"]https?://")


# ================================================================ telas de migração e de conferência

class TestTelas(unittest.TestCase):
    def setUp(self):
        from flask import Flask
        from painel import assistente, base as pbase, entregas, migracao, perfil
        self.pasta = pasta_temporaria()
        self.arquivo = cf.gravar(self.pasta / "contingencias.xlsx")
        ficticio.criar_projeto_de_teste([], nome="Relatório do Teste")
        self.app = Flask(__name__)
        self.app.config["TESTING"] = True
        ok = pbase.criar_token_ok(TOKEN)
        for tela in (pbase, assistente, migracao, entregas, perfil):
            tela.registrar(self.app, TOKEN, pbase.cabecalho, ok)
        self.c = self.app.test_client()
        self.mig = migracao

    def tearDown(self):
        ficticio.restaurar_comum()
        shutil.rmtree(self.pasta, ignore_errors=True)

    def enviar(self, rota="/migracao/enviar", campo="arquivo"):
        return self.c.post(rota, data={"token": TOKEN, campo: (io.BytesIO(self.arquivo.read_bytes()), "contingencias.xlsx")},
                           content_type="multipart/form-data", follow_redirects=True)

    def texto(self, resposta):
        return html.unescape(re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", resposta.get_data(as_text=True))))

    def test_mapeamento_mostra_o_quadro_do_que_falta(self):
        r = self.enviar()
        self.assertEqual(r.request.path, "/migracao/mapear")
        html_ = r.get_data(as_text=True)
        t = self.texto(r)
        self.assertIn("O que falta para a transição", t)
        self.assertIn("12 processo(s) (9 ativo(s), 3 encerrado(s))", t)
        self.assertIn("Vara, Município, UF e Classe: 0 de 12 no arquivo; o programa busca no PJe/jus.br ao coletar.", t)
        self.assertIn("Valor da causa: 0 de 12 no arquivo; a capa do tribunal costuma trazer; se a coleta não achar, informe.", t)
        self.assertIn("Provisão constituída: 7 de 8 no arquivo (só dos 8 processos ativos com perda provável ou possível)", t)
        self.assertIn("Colunas do seu arquivo que ainda não têm destino", t)
        self.assertIn("Sem destino (0)", t)
        self.assertIn("do arquivo, coluna «PASSIVO POTENCIAL»", t)
        self.assertIn("buscável nos tribunais", t)
        self.assertIn("precisa de você", t)
        self.assertIn("class='ajuda'", html_)                       # botão (i) do quadro
        self.assertIn("<progress", html_)
        # a coluna "Natureza da ação" é lida como Assunto, com a explicação
        self.assertIn("Colunas chamadas \"Natureza da ação\" costumam trazer o assunto", t)
        self.assertNotRegex(html_, r"(src|href)=['\"]https?://")

    def test_previa_refaz_o_quadro_com_o_mapeamento_escolhido(self):
        r = self.enviar()
        lote = re.search(r"name='lote' value='([^']+)'", r.get_data(as_text=True)).group(1)
        linhas = self.mig.colunas_da_leitura(leitores.ler(self.arquivo, formato="tabela_livre"))
        form = {"token": TOKEN, "lote": lote, "nome": "Teste", "acao": "prever"}
        for i, l in enumerate(linhas):
            form[f"map_{i}"] = "" if l["coluna"] == "PASSIVO POTENCIAL" else l["campo"]      # a pessoa tira o passivo potencial
        r2 = self.c.post("/migracao/converter", data=form, follow_redirects=True)
        t = self.texto(r2)
        self.assertIn("Prévia: 12 processo(s)", t)
        self.assertIn("Passivo potencial: 0 de 12 no arquivo", t)
        self.assertIn("Sem destino (1)", t)
        self.assertIn("PASSIVO POTENCIAL", t)

    def test_converter_gera_planilha_com_as_faltas_e_painel_com_contingencia(self):
        import comum
        import openpyxl
        r = self.enviar()
        lote = re.search(r"name='lote' value='([^']+)'", r.get_data(as_text=True)).group(1)
        r2 = self.c.post("/migracao/converter", data={"token": TOKEN, "lote": lote, "nome": "Contingências convertidas",
                                                        "acao": "converter", "modelos": ["xlsx_b", "dashboard"]}, follow_redirects=True)
        self.assertEqual(r2.request.path, "/entregas")
        planilhas = sorted(Path(comum.PROJETO_DIR).rglob("saida/*/Planilha*.xlsx"))
        paineis = sorted(Path(comum.PROJETO_DIR).rglob("saida/*/Painel*.html"))
        self.assertEqual((len(planilhas), len(paineis)), (1, 1))
        wb = openpyxl.load_workbook(planilhas[0])
        self.assertIn("Faltas da migração", wb.sheetnames)
        ws = wb["Faltas da migração"]
        campos = {ws.cell(r_, 2).value: (ws.cell(r_, 3).value, ws.cell(r_, 4).value) for r_ in range(2, ws.max_row + 1)}
        self.assertEqual(campos["Provisão constituída"], (7, 8))
        self.assertEqual(campos["Vara / Juízo"], (0, 12))
        self.assertEqual(campos["Passivo potencial"], (11, 12))
        proc = wb["Processos"]
        cab = [c.value for c in proc[1]]
        from escritores import xlsx_b
        self.assertEqual(cab[37:], [xlsx_b.CAMPOS_B[c] for c in xlsx_b.CONTINGENCIA])      # contingência no fim, colunas 38 em diante
        self.assertEqual(len(cab), 46)
        self.assertEqual(proc.max_row, 13)
        html_ = paineis[0].read_text(encoding="utf-8")
        self.assertIn("Contingência e provisão", html_)
        self.assertIn("Passivo Potencial", html_)

    def test_conferencia_do_importar_mostra_o_mesmo_quadro(self):
        r = self.enviar("/fluxo/importar/enviar", "arquivos")
        self.assertEqual(r.request.path, "/fluxo/importar/conferir")
        t = self.texto(r)
        html_ = r.get_data(as_text=True)
        self.assertIn("O que falta para a transição", t)
        self.assertIn("Li 12 processo(s) de 1 arquivo(s)", t)
        self.assertIn("Data do ajuizamento: 3 de 12 no arquivo; faltam 9: o programa busca no PJe/jus.br ao coletar.", t)
        self.assertIn("Contingência (risco, provisão e depósitos)", t)
        self.assertIn("<details open>", html_)                      # planilha fora do modelo: o quadro já vem aberto

    def test_o_quadro_nunca_derruba_a_tela(self):
        from painel import lacunas_tela
        saida = lacunas_tela.quadro(object())
        self.assertIn("O que falta para a transição", saida)
        self.assertIn("Não consegui montar este quadro", saida)
        vazio = lacunas_tela.quadro([])
        self.assertIn("Nenhum processo lido", vazio)


if __name__ == "__main__":
    unittest.main()
