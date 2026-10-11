"""Campo `grau_tramitacao` ("Grau de tramitação"): vocabulário, leitura da coluna, escrita na planilha, sinal que chega à
coleta do PJe e atualização automática pelo resultado da coleta. Tudo fictício, sem rede e sem login.

    python3 -m unittest tests/test_grau_tramitacao.py -v
"""
import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent))
import isolamento  # noqa: E402,F401  (antes de tudo)
import capa  # noqa: E402
import coletor  # noqa: E402
import ficha  # noqa: E402
import ficticio  # noqa: E402
import fila  # noqa: E402
import fluxos  # noqa: E402
import leitores  # noqa: E402
import pje_trt  # noqa: E402
import taxonomia  # noqa: E402
import trt  # noqa: E402
from leitores import grade  # noqa: E402
from escritores import xlsx_b  # noqa: E402
from test_pje_trt import Base as BasePje, TIMELINE, TL2, sessao  # noqa: E402

N_TRT = ficticio.numero_ficticio(1, j=5, tr=7)


def lido(grau):
    return taxonomia.normalizar("grau_tramitacao", grau)


class Vocabulario(unittest.TestCase):
    def test_valores_canonicos_na_ordem(self):
        self.assertEqual(taxonomia.valores("grau_tramitacao"), ["1º grau", "2º grau", "TST"])
        self.assertIn("grau_tramitacao", taxonomia.VOCABULARIOS)

    def test_sinonimos(self):
        casos = {"1º grau": ["Primeiro grau", "1o grau", "primeira instância", "Vara", "1º Grau"],
                 "2º grau": ["segundo grau", "2o", "Segunda instância", "TRT", "TJ", "tribunal regional"],
                 "TST": ["Tribunal Superior do Trabalho", "terceiro grau", "3o grau", "tst", "tribunal superior"]}
        for canonico, formas in casos.items():
            for forma in formas:
                self.assertEqual(lido(forma), canonico, forma)
        self.assertIsNone(lido("grau de risco"))

    def test_campo_da_ficha(self):
        self.assertEqual(ficha.CAMPOS["grau_tramitacao"], ("Grau de tramitação", "situacao", "texto", "grau_tramitacao"))
        f = {"numero": N_TRT, "campos": {}}
        self.assertTrue(ficha.definir(f, "grau_tramitacao", "segunda instância", "humano"))
        self.assertEqual(ficha.obter(f, "grau_tramitacao"), "2º grau")
        self.assertFalse(ficha.definir(f, "grau_tramitacao", "um lugar qualquer", "humano"))      # fora do vocabulário
        self.assertEqual(ficha.obter(f, "grau_tramitacao"), "2º grau")


class LeituraDaColuna(unittest.TestCase):
    def setUp(self):
        self.pasta = Path(tempfile.mkdtemp(prefix="grau-", dir=isolamento.TMP))
        self.addCleanup(shutil.rmtree, self.pasta, ignore_errors=True)

    def ler(self, cabecalhos, linha):
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active.append(cabecalhos)
        wb.active.append(linha)
        caminho = self.pasta / "r.xlsx"
        wb.save(caminho)
        return leitores.ler(caminho)

    def test_cabecalhos_de_grau_viram_o_campo(self):
        for cab in ("Grau de Tramitação", "Instância", "Tramitação", "Grau de jurisdição"):
            confianca, destino = grade.pontuar_cabecalho(cab)[0]
            self.assertEqual((destino, confianca >= grade.LIMIAR_AUTO), ("grau_tramitacao", True), cab)

    def test_nao_colide_com_grau_de_risco_ou_probabilidade(self):
        for cab in ("Grau de risco", "Grau de Probabilidade", "Grau de probabilidade de perda", "Probabilidade"):
            self.assertEqual(grade.pontuar_cabecalho(cab)[0][1], "probabilidade", cab)
        m = grade.propor_mapeamento(["Grau de risco", "Grau de Tramitação"], [["Provável"], ["2º grau"]])
        self.assertEqual([x["campo"] for x in m], ["probabilidade", "grau_tramitacao"])
        self.assertFalse(any(x["ambigua"] for x in m))

    def test_grau_sozinho_nao_decide_sozinho(self):
        m = grade.propor_mapeamento(["Grau"], [["2º grau"]])
        self.assertIsNone(m[0]["campo"])               # ambíguo: pede confirmação em vez de chutar

    def test_planilha_com_as_duas_colunas(self):
        rel = self.ler(["Número do Processo", "Grau de Tramitação", "Grau de Risco", "Andamentos"],
                       [N_TRT, "Segunda instância", "Provável", "texto"])
        campos = rel["processos"][0]["campos"]
        self.assertEqual(campos["grau_tramitacao"]["valor"], "2º grau")
        self.assertEqual(campos["grau_tramitacao"]["origem"], "migrado")          # a coleta pode atualizá-lo
        self.assertEqual(campos["probabilidade"]["valor"], "Provável")


class EscritaNaPlanilha(unittest.TestCase):
    def setUp(self):
        self.pasta = Path(tempfile.mkdtemp(prefix="grau-", dir=isolamento.TMP))
        self.addCleanup(shutil.rmtree, self.pasta, ignore_errors=True)

    def gravar(self, fichas):
        import openpyxl
        destino = self.pasta / "saida.xlsx"
        res = xlsx_b.gravar(None, {"cliente": "Cliente Exemplo 01 Ltda", "data_base": "2026-09-30", "fichas": fichas,
                                   "eventos": [], "perfil": {}, "parametros": {}}, destino)
        self.assertEqual([a for a in res["avisos"] if a["nivel"] == "erro"], [])
        ws = openpyxl.load_workbook(destino)["Processos"]
        return ws, [c.value for c in ws[1]]

    def test_coluna_e_mecanica_extra_e_ultima(self):
        self.assertEqual(xlsx_b.CAMPOS_B["grau_tramitacao"], "Grau de Tramitação")
        self.assertEqual(xlsx_b.PAPEIS["grau_tramitacao"], "mecanico")
        self.assertIn("grau_tramitacao", xlsx_b.COLUNAS_EXTRAS)
        self.assertEqual(xlsx_b.COLUNAS_B_ORDEM[-1], "grau_tramitacao")

    def test_com_valor_a_coluna_aparece_e_e_gravada(self):
        fichas = ficticio.gerar_carteira(3)
        ficha.definir(fichas[0], "grau_tramitacao", "2º grau", "coletado")
        ws, cab = self.gravar(fichas)
        col = cab.index("Grau de Tramitação") + 1
        self.assertEqual(len(cab), 47)
        self.assertNotIn(xlsx_b.col_letra(col), {c for c, d in ws.column_dimensions.items() if d.hidden})
        valores = {ws.cell(r, 1).value: ws.cell(r, col).value for r in range(2, ws.max_row + 1)}
        self.assertEqual(valores[fichas[0]["numero"]], "2º grau")

    def test_sem_valor_a_coluna_fica_oculta(self):
        ws, cab = self.gravar(ficticio.gerar_carteira(3))
        col = cab.index("Grau de Tramitação") + 1
        self.assertIn(xlsx_b.col_letra(col), {c for c, d in ws.column_dimensions.items() if d.hidden})


class SinalNaColeta(BasePje):
    def coletar(self, proc):
        s1, _ = sessao(timeline=TIMELINE, outra=False)             # sem indício de recurso no 1º grau
        s2, _ = sessao(timeline=TL2, instancia=2, outra=True)
        s2.grau = 2
        relato, pediu = {}, []

        def segunda():
            pediu.append(2)
            return s2
        pje_trt.coletar_processo(s1, proc, {}, [], 5, 10, None, relato, segunda=segunda)
        return relato, pediu

    def test_sem_sinal_e_sem_indicio_fica_no_1_grau(self):
        relato, pediu = self.coletar(dict(self.proc))
        self.assertEqual((pediu, relato["graus_lidos"]), ([], ["1"]))

    def test_sinal_2_grau_manda_ler_o_2_grau(self):
        relato, pediu = self.coletar({**self.proc, "sinal_grau": "2º grau"})
        self.assertEqual((pediu, relato["graus_lidos"]), ([2], ["1", "2"]))

    def test_sinal_tst_le_o_2_grau_antes(self):
        relato, pediu = self.coletar({**self.proc, "sinal_grau": "TST"})
        self.assertEqual((pediu, relato["graus_lidos"]), ([2], ["1", "2"]))      # sem subida no 2º grau, o TST não é pedido

    def test_sinal_de_1_grau_nao_conta_como_recurso(self):
        _, pediu = self.coletar({**self.proc, "sinal_grau": "1º grau"})
        self.assertEqual(pediu, [])

    def test_motivo_do_sinal(self):
        self.assertEqual(trt.indicio_de_recurso([], {"sinal_grau": "2º grau"}),
                         "o relatório indica que o processo está no 2º grau")
        self.assertIn("TST", trt.indicio_de_recurso([], {"sinal_grau": "TST"}))
        self.assertIsNone(trt.indicio_de_recurso([], {"sinal_grau": "1º grau"}))

    def test_sem_login_do_2_grau_avisa_com_o_motivo(self):
        s1, _ = sessao(outra=False)
        relato = {}
        pje_trt.coletar_processo(s1, {**self.proc, "sinal_grau": "2º grau"}, {}, [], 5, 10, None, relato, segunda=lambda: None)
        self.assertIn("o relatório indica que o processo está no 2º grau", relato["avisos"][0]["mensagem"])


class DaFichaAteOColetorReal(unittest.TestCase):
    """O grau da ficha chega a `proc["sinal_grau"]` do coletor."""

    def real(self, processo):
        real = fila.ColetorReal()
        real._contexto = object()
        vistos = {}

        def falso(contexto, proc, estado, lista, historico, cota, desde=None, relato=None):
            vistos.update(proc)
            relato.update(graus_lidos=["1"])
            return 0
        with mock.patch.object(coletor, "coletar_processo", side_effect=falso), \
                mock.patch.object(capa, "_config_datajud", return_value=(False, None)), \
                mock.patch.object(fila.fch, "carregar", return_value=[]):
            real.coletar(processo, "rapido", None)
        return vistos

    def test_item_da_fila_com_grau_vira_sinal(self):
        self.assertEqual(self.real({"numero": N_TRT, "cliente": "x", "grau_tramitacao": "2º grau"})["sinal_grau"], "2º grau")
        self.assertEqual(self.real({"numero": N_TRT, "cliente": "x", "grau_tramitacao": "tribunal superior do trabalho"})["sinal_grau"], "TST")

    def test_sem_sinal_ou_1_grau_nao_ha_chave(self):
        self.assertNotIn("sinal_grau", self.real({"numero": N_TRT, "cliente": "x", "grau_tramitacao": None}))
        self.assertNotIn("sinal_grau", self.real({"numero": N_TRT, "cliente": "x", "grau_tramitacao": "1º grau"}))


class NaFila(unittest.TestCase):
    def setUp(self):
        self.fichas = ficticio.gerar_carteira(3)
        self._ctx = ficticio.projeto_de_teste(self.fichas)
        self.proj = self._ctx.__enter__()
        self.addCleanup(self._ctx.__exit__, None, None, None)

    def test_fila_guarda_o_grau_e_entrega_ao_coletor(self):
        from test_fila import nova_fila
        n = self.fichas[0]["numero"]
        f = nova_fila(self.proj)
        f.enfileirar([{"numero": n, "cliente": "x", "grau_tramitacao": "2º grau"}])
        vistos = []

        class Coletor:
            def coletar(self, processo, profundidade, desde):
                vistos.append(dict(processo))
                return {"capa": {}, "movimentos": [], "documentos": [], "erro": None}
        fila.rodar_fila(f, Coletor(), numeros=[n])
        self.assertEqual(vistos[0]["grau_tramitacao"], "2º grau")

    def test_fluxo_enfileira_o_grau_so_do_principal(self):
        from test_fila import nova_fila
        f0 = self.fichas[0]
        ficha.definir(f0, "grau_tramitacao", "2º grau", "humano")
        f0["vinculados"] = [{"numero": self.fichas[1]["numero"], "tipo": "agravo"}]
        f = nova_fila(self.proj)

        class Coletor:
            def coletar(self, processo, profundidade, desde):
                return {"capa": {}, "movimentos": [], "documentos": [], "erro": None}
        with mock.patch.object(fluxos, "processar_resultado"):
            fluxos._coletar(self.proj["slug"], [f0], profundidade="rapido", modo="imediato", coletor=Coletor(), fila=f,
                            fila_opcoes=None, data_base_do_arquivo=None, ciclo_novo=False, ao_progresso=None)
        itens = {i["numero"]: i for i in f.itens()}
        self.assertEqual(itens[f0["numero"]]["grau_tramitacao"], "2º grau")
        self.assertIsNone(itens[self.fichas[1]["numero"]]["grau_tramitacao"])


class AtualizacaoPeloResultado(unittest.TestCase):
    def setUp(self):
        self.fichas = ficticio.gerar_carteira(2)
        self.numero = self.fichas[0]["numero"]
        self._ctx = ficticio.projeto_de_teste(self.fichas)
        self.proj = self._ctx.__enter__()
        self.addCleanup(self._ctx.__exit__, None, None, None)
        alvo = mock.patch.object(capa, "_config_datajud", return_value=(False, None))
        alvo.start()
        self.addCleanup(alvo.stop)

    def processar(self, lidos=None, **extra):
        resultado = {"capa": {}, "movimentos": [], "documentos": [], "erro": None, **extra}
        if lidos is not None:
            resultado["graus"] = {"lidos": lidos, "falhos": []}
        fluxos.processar_resultado(self.proj["slug"], self.numero, resultado)
        f = next(x for x in ficha.carregar(todas=True) if x["numero"] == self.numero)
        return ficha.obter(f, "grau_tramitacao"), ficha.origem(f, "grau_tramitacao")

    def preparar(self, valor, origem):
        fs = ficha.carregar(todas=True)
        ficha.definir(next(x for x in fs if x["numero"] == self.numero), "grau_tramitacao", valor, origem)
        ficha.salvar(fs)

    def test_maior_grau_lido(self):
        self.assertEqual(self.processar(["1", "2"]), ("2º grau", "coletado"))
        self.assertEqual(self.processar(["1", "2", "3"]), ("TST", "coletado"))

    def test_so_o_1_grau_lido_grava_1_grau(self):
        self.assertEqual(self.processar(["1"]), ("1º grau", "coletado"))

    def test_processo_no_tst_pelo_datajud(self):
        self.assertEqual(self.processar(["1"], no_tst=True), ("TST", "coletado"))

    def test_coletor_que_nao_informa_graus_nao_chuta(self):
        self.assertEqual(self.processar(None), (None, None))

    def test_humano_nao_e_sobrescrito(self):
        self.preparar("1º grau", "humano")
        self.assertEqual(self.processar(["1", "2"]), ("1º grau", "humano"))
        self.preparar("TST", "humano")
        self.assertEqual(self.processar(["1"]), ("TST", "humano"))

    def test_valor_migrado_sobe_com_a_coleta(self):
        self.preparar("1º grau", "migrado")
        self.assertEqual(self.processar(["1", "2"]), ("2º grau", "coletado"))

    def test_nunca_rebaixa(self):
        self.preparar("2º grau", "migrado")          # o relatório diz 2º grau e o 2º grau não pôde ser lido agora
        self.assertEqual(self.processar(["1"]), ("2º grau", "migrado"))

    def test_idempotente(self):
        self.processar(["1", "2"])
        self.assertEqual(self.processar(["1", "2"]), ("2º grau", "coletado"))


if __name__ == "__main__":
    unittest.main()
