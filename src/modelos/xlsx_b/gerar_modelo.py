"""Gera o MODELO PADRÃO do relatório em planilha (modelo B): `modelo_padrao.xlsx`, ao lado deste script.

O modelo é SANITIZADO: só a estrutura (cabeçalhos, fórmulas, validações, gráficos); nenhum nome, número de
processo ou valor de cliente. Em produção ele é só LIDO por `escritores/xlsx_b.py` (biblioteca padrão);
o `openpyxl` é usado aqui, uma vez, para criá-lo (como o spike S1 fazia com o modelo fictício).

Abas:
  Processos           tabela do Excel `tblProcessos`: as 29 colunas do modelo B + 8 colunas extras OCULTAS
                      (momento atual, último andamento, citação, fase, classe, assunto, acordo, observações);
                      validação de dados em Situação, Probabilidade, Resultado etc.; colunas calculadas
                      Valor Economizado e Taxa de resolução (fórmulas com referência estruturada).
                      Vem com UMA linha em branco (a tabela do Excel precisa de ao menos uma); o escritor
                      a aproveita para o primeiro processo.
  Parâmetros          headcount, data de referência, fator de correção, empresas do grupo.
  Indicadores         acervo, tempo e recursos, resultados e financeiro, só com COUNTIFS/SUMIFS/AVERAGE/
                      MEDIAN (sem FILTER nem funções dinâmicas), mais quadros por situação/resultado/área.
  Dashboard           gráficos básicos ligados aos Indicadores e ao Histórico.
  Histórico           retrato mensal (CONTRATOS 9): uma linha por data-base.
  Campos não migrados colunas do arquivo antigo sem destino no modelo (nada se perde em silêncio).

O arquivo é determinístico (mesmas entradas, mesmos bytes, para a mesma versão do openpyxl).

Uso:  python3 src/modelos/xlsx_b/gerar_modelo.py [DESTINO.xlsx]
"""
import datetime
import io
import re
import sys
import zipfile
from pathlib import Path

AQUI = Path(__file__).resolve().parent
sys.path.insert(0, str(AQUI.parent.parent))     # src/

from openpyxl import Workbook                                              # noqa: E402
from openpyxl.chart import BarChart, LineChart, Reference                  # noqa: E402
from openpyxl.formatting.rule import FormulaRule                           # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side     # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation               # noqa: E402
from openpyxl.worksheet.table import Table, TableFormula, TableStyleInfo   # noqa: E402

import ficha                                                               # noqa: E402
import taxonomia                                                           # noqa: E402
from escritores import xlsx_b                                              # noqa: E402

TABELA, TABELA_HIST, TABELA_NM = "tblProcessos", "tblHistorico", "tblNaoMigrados"
DATA_FIXA = datetime.datetime(2026, 10, 7, 12, 0, 0)
ZIP_DATA = (2026, 10, 7, 12, 0, 0)
LARGURAS = {"autores": 30, "reus": 30, "vara": 32, "municipio": 18, "objeto": 40, "andamentos": 70,
            "numero": 26, "outras_partes": 26, "observacoes": 30, "materia_principal": 26, "assunto": 24,
            "momento_atual": 34, "valor_arbitrado": 18, "valor_estimado": 16}
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
BORDA = Side(style="thin", color="BFBFBF")
PROB = {"Provável": "FFC7CE", "Possível": "FFEB9C", "Remota": "C6EFCE"}


def _ref_estruturada(nome):
    """Coluna inteira: Tabela[Nome]; com caracteres especiais no nome, Tabela[[Nome]]."""
    return f"{TABELA}[[{nome}]]" if re.search(r"[^A-Za-zÀ-ÿ0-9 ]", nome) else f"{TABELA}[{nome}]"


def _esta_linha(nome):
    return f"{TABELA}[[#This Row],[{nome}]]"


def _formato(campo):
    if campo == "percentual_exito":
        return "0.0%"
    tipo = ficha.CAMPOS.get(campo, ("", "", "texto", None))[2]
    return {"data": "dd/mm/yyyy", "dinheiro": "#,##0.00", "numero": "0"}.get(tipo)


def _letra(indice):
    return xlsx_b.col_letra(indice)


def _processos(wb):
    ws = wb.active
    ws.title = "Processos"
    colunas = list(xlsx_b._COLUNAS)
    n = len(colunas)
    cabecalhos = [c[1] for c in colunas]
    indice = {c[0]: i + 1 for i, c in enumerate(colunas)}
    for i, (campo, cab, _papel, _extra, _sin) in enumerate(colunas, start=1):
        cel = ws.cell(1, i, cab)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = HEADER_FILL
        cel.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[_letra(i)].width = LARGURAS.get(campo, 16)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"
    # uma linha em branco, já formatada, com as colunas calculadas
    for campo, i in indice.items():
        cel = ws.cell(2, i)
        cel.border = Border(left=BORDA, right=BORDA, top=BORDA, bottom=BORDA)
        fmt = _formato(campo)
        if fmt:
            cel.number_format = fmt
        if campo in ("andamentos", "objeto"):
            cel.alignment = Alignment(wrap_text=True, vertical="top")
    cab = xlsx_b.CAMPOS_B
    calculadas = {
        "valor_economizado": (
            f'IF(OR({_esta_linha(cab["ativo"])}<>"Não",{_esta_linha(cab["valor_causa"])}="",'
            f'{_esta_linha(cab["valor_estimado"])}=""),"",{_esta_linha(cab["valor_causa"])}-{_esta_linha(cab["valor_estimado"])})'),
        "taxa_resolucao_dias": (
            f'IF(OR({_esta_linha(cab["data_transito"])}="",{_esta_linha(cab["data_ajuizamento"])}=""),"",'
            f'{_esta_linha(cab["data_transito"])}-{_esta_linha(cab["data_ajuizamento"])})')}
    for campo, formula in calculadas.items():
        ws.cell(2, indice[campo]).value = "=" + formula
    tabela = Table(displayName=TABELA, ref=f"A1:{_letra(n)}2")
    tabela.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    # colunas calculadas declaradas na Tabela: o Excel as propaga para as linhas novas digitadas à mão
    tabela._initialise_columns()
    for celula, coluna in zip(ws[tabela.ref][0], tabela.tableColumns):
        coluna.name = str(celula.value)
        for campo, formula in calculadas.items():
            if coluna.name == cab[campo]:
                coluna.calculatedColumnFormula = TableFormula(attr_text=formula)
    ws.add_table(tabela)
    # colunas extras ocultas até o perfil ativá-las
    for campo in xlsx_b.COLUNAS_EXTRAS:
        ws.column_dimensions[_letra(indice[campo])].hidden = True
    # formatação condicional (probabilidade e encerrados)
    p = _letra(indice["probabilidade"])
    for texto, cor in PROB.items():
        ws.conditional_formatting.add(
            f"{p}2:{p}2", FormulaRule(formula=[f'${p}2="{texto}"'], fill=PatternFill("solid", bgColor=cor)))
    a = _letra(indice["ativo"])
    ws.conditional_formatting.add(
        f"A2:{_letra(n)}2", FormulaRule(formula=[f'${a}2="Não"'], font=Font(color="7F7F7F")))
    # validação de dados (aviso, não bloqueio: a pessoa pode digitar outro valor)
    listas = {
        "situacao": list(taxonomia.SITUACAO), "probabilidade": list(taxonomia.PROBABILIDADE),
        "resultado": list(taxonomia.RESULTADO), "area": list(taxonomia.AREA),
        "ativo": ["Sim", "Não"], "houve_recurso": ["Sim", "Não"], "terceirizado": ["Sim", "Não"],
    }
    for campo, itens in listas.items():
        formula = '"' + ",".join(itens) + '"'
        assert len(formula) < 255, campo
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True,
                            errorStyle="warning", errorTitle=cab[campo], error="Valor fora da lista padrão. Manter assim?")
        col = _letra(indice[campo])
        dv.add(f"{col}2")
        ws.add_data_validation(dv)
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True,
                        showErrorMessage=True, errorStyle="warning", errorTitle=cab["percentual_exito"],
                        error="Use uma fração entre 0 e 1 (0,85 = 85%).")
    dv.add(f"{_letra(indice['percentual_exito'])}2")
    ws.add_data_validation(dv)
    return indice


def _parametros(wb):
    ws = wb.create_sheet("Parâmetros")
    ws["A1"], ws["B1"] = "Parâmetro", "Valor"
    ws["A2"], ws["A3"], ws["A4"], ws["A5"] = (
        "Número de funcionários (headcount)", "Data de referência", "Fator de correção", "Empresas do grupo")
    ws["B3"].number_format = "dd/mm/yyyy"
    ws["C2"] = "Usado em 'Litígios por 100 funcionários'."
    ws["C3"] = "Data-base do relatório; usada nos 'novos ajuizamentos em 12 meses'."
    ws["C4"] = "Multiplica o valor da causa total (vazio = sem correção)."
    ws["C5"] = "Nomes separados por ponto e vírgula."
    for c in ("A1", "B1"):
        ws[c].font = Font(bold=True)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 70


def _indicadores(wb):
    ws = wb.create_sheet("Indicadores")
    cab = xlsx_b.CAMPOS_B
    R = {k: _ref_estruturada(cab[k]) for k in cab}
    ws["A1"], ws["B1"], ws["C1"] = "Indicador", "Valor", "Como é calculado"
    linha = [2]
    pos = {}

    def secao(titulo):
        linha[0] += 1 if linha[0] > 2 else 0
        ws.cell(linha[0], 1, titulo).font = Font(bold=True, color="1F4E78")
        linha[0] += 1

    def ind(chave, rotulo, formula, nota, formato=None):
        r = linha[0]
        pos[chave] = f"B{r}"
        ws.cell(r, 1, rotulo)
        ws.cell(r, 2, formula)
        ws.cell(r, 3, nota)
        if formato:
            ws.cell(r, 2).number_format = formato
        linha[0] += 1

    P = "'Parâmetros'!"
    secao("Acervo")
    ind("total", "Total de processos", f"=COUNTA({R['numero']})", "Linhas com número de processo.", "0")
    ind("ativos", "Processos ativos", f'=COUNTIFS({R["ativo"]},"Sim")', "Coluna Ativo = Sim.", "0")
    ind("encerrados", "Processos encerrados", f'=COUNTIFS({R["ativo"]},"Não")', "Coluna Ativo = Não.", "0")
    ind("por100", "Litígios por 100 funcionários",
        f'=IF(N({P}$B$2)>0,{pos["total"]}/{P}$B$2*100,"")',
        "Total de processos ÷ número de funcionários × 100 (precisa do headcount em Parâmetros).", "0.0")
    ind("novos12", "Novos ajuizamentos em 12 meses",
        f'=IF(N({P}$B$3)>0,COUNTIFS({R["data_ajuizamento"]},">"&EDATE({P}$B$3,-12),'
        f'{R["data_ajuizamento"]},"<="&{P}$B$3),"")',
        "Ajuizados nos 12 meses até a data de referência (Parâmetros).", "0")
    secao("Tempo e recursos")
    ind("tmedio", "Tempo médio de resolução (dias)", f'=IFERROR(AVERAGE({R["taxa_resolucao_dias"]}),"")',
        "Média da Taxa de resolução (trânsito em julgado − ajuizamento).", "0.0")
    ind("tmediana", "Mediana do tempo de resolução (dias)", f'=IFERROR(MEDIAN({R["taxa_resolucao_dias"]}),"")',
        "Mediana da Taxa de resolução.", "0.0")
    ind("recursos", "Processos com recurso da empresa", f'=COUNTIFS({R["houve_recurso"]},"Sim")',
        "Coluna Houve recurso da empresa? = Sim.", "0")
    # os desfechos vêm abaixo; "decididos" soma os cinco desfechos de mérito/saída
    linha_decididos = linha[0]
    linha[0] += 1
    linha_taxa_recurso = linha[0]
    linha[0] += 1
    secao("Resultados")
    for chave, rotulo, valor in (("improc", "Improcedentes", "Improcedente"),
                                 ("arq", "Arquivadas / desistência", "Arquivado / desistência"),
                                 ("acordo", "Acordos", "Acordo"),
                                 ("parcial", "Parcialmente procedentes", "Parcialmente procedente"),
                                 ("proc", "Procedentes", "Procedente")):
        ind(chave, rotulo, f'=COUNTIFS({R["resultado"]},"{valor}")', f"Coluna Resultado = {valor}.", "0")
    ws.cell(linha_decididos, 1, "Processos decididos")
    ws.cell(linha_decididos, 2,
            "=" + "+".join(pos[k] for k in ("improc", "arq", "acordo", "parcial", "proc")))
    ws.cell(linha_decididos, 3, "Soma dos cinco desfechos abaixo (não conta extinção sem mérito nem incompetência).")
    ws.cell(linha_decididos, 2).number_format = "0"
    pos["decididos"] = f"B{linha_decididos}"
    ws.cell(linha_taxa_recurso, 1, "Taxa de recurso")
    ws.cell(linha_taxa_recurso, 2, f'=IF({pos["decididos"]}>0,{pos["recursos"]}/{pos["decididos"]},"")')
    ws.cell(linha_taxa_recurso, 3, "Processos com recurso da empresa ÷ decididos.")
    ws.cell(linha_taxa_recurso, 2).number_format = "0.0%"
    ind("exito", "Taxa de êxito",
        f'=IF({pos["decididos"]}>0,({pos["improc"]}+{pos["arq"]})/{pos["decididos"]},"")',
        "(Improcedentes + arquivadas) ÷ decididos.", "0.0%")
    ind("pacordo", "% por acordo", f'=IF({pos["decididos"]}>0,{pos["acordo"]}/{pos["decididos"]},"")',
        "Acordos ÷ decididos.", "0.0%")
    secao("Financeiro")
    causa, ativo, arb, exe = R["valor_causa"], R["ativo"], R["valor_arbitrado"], R["valor_execucao"]
    ind("causa", "Valor da causa total", f"=SUM({causa})", "Soma da coluna Valor da Causa.", "#,##0.00")
    ind("causa_corr", "Valor da causa total corrigido",
        f"=IF(N({P}$B$4)>0,{pos['causa']}*{P}$B$4,{pos['causa']})", "Valor da causa × fator de correção (Parâmetros).",
        "#,##0.00")
    ind("causa_enc", "Valor da causa dos encerrados", f'=SUMIFS({causa},{ativo},"Não")',
        "Processos com Ativo = Não.", "#,##0.00")
    ind("enc_lancado", "Encerrados com valor realizado lançado",
        f'=COUNTIFS({ativo},"Não",{arb},"<>")+COUNTIFS({ativo},"Não",{arb},"",{exe},"<>")',
        "Têm valor arbitrado ou, se este estiver vazio, valor da execução.", "0")
    ind("enc_sem", "Encerrados sem valor lançado (fora da economia)", f'={pos["encerrados"]}-{pos["enc_lancado"]}',
        "Ficam fora da economia para não inflá-la (acordo sem valor lançado etc.).", "0")
    ind("causa_lancado", "Valor da causa dos encerrados com valor lançado",
        f'=SUMIFS({causa},{ativo},"Não",{arb},"<>")+SUMIFS({causa},{ativo},"Não",{arb},"",{exe},"<>")',
        "Base da economia efetiva.", "#,##0.00")
    ind("realizado", "Valor realizado dos encerrados",
        f'=SUMIFS({arb},{ativo},"Não")+SUMIFS({exe},{ativo},"Não",{arb},"")',
        "Valor arbitrado em juízo, ou o valor da execução quando o arbitrado está vazio.", "#,##0.00")
    ind("economia", "Economia efetiva", f'={pos["causa_lancado"]}-{pos["realizado"]}',
        "Valor da causa − valor realizado, só dos encerrados com valor lançado.", "#,##0.00")
    # quadros que alimentam os gráficos
    quadros = {}
    for chave, titulo, campo, valores in (
            ("situacao", "Processos por situação", "situacao", list(taxonomia.SITUACAO)),
            ("resultado", "Processos por resultado", "resultado", list(taxonomia.RESULTADO)),
            ("area", "Processos por área", "area", list(taxonomia.AREA)),
            ("prob", "Processos por probabilidade", "probabilidade", list(taxonomia.PROBABILIDADE))):
        linha[0] += 1
        ws.cell(linha[0], 1, titulo).font = Font(bold=True, color="1F4E78")
        ws.cell(linha[0], 2, "Processos").font = Font(bold=True)
        linha[0] += 1
        ini = linha[0]
        for v in valores:
            ws.cell(linha[0], 1, v)
            ws.cell(linha[0], 2, f"=COUNTIFS({R[campo]},A{linha[0]})").number_format = "0"
            linha[0] += 1
        quadros[chave] = (titulo, ini, linha[0] - 1)
    ws["A1"].font = ws["B1"].font = ws["C1"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 90
    ws.freeze_panes = "A2"
    return ws, pos, quadros


def _historico(wb):
    ws = wb.create_sheet("Histórico")
    cabecalhos = ["Data-base", "Processos", "Ativos", "Encerrados", "Valor da causa", "Valor estimado",
                  "Valor economizado"]
    for i, c in enumerate(cabecalhos, start=1):
        cel = ws.cell(1, i, c)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = HEADER_FILL
        ws.column_dimensions[_letra(i)].width = 18
    for i, fmt in enumerate(["dd/mm/yyyy", "0", "0", "0", "#,##0.00", "#,##0.00", "#,##0.00"], start=1):
        cel = ws.cell(2, i)
        cel.number_format = fmt
        cel.border = Border(left=BORDA, right=BORDA, top=BORDA, bottom=BORDA)
    tabela = Table(displayName=TABELA_HIST, ref=f"A1:{_letra(len(cabecalhos))}2")
    tabela.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tabela)
    ws.freeze_panes = "A2"
    return ws


def _nao_migrados(wb):
    ws = wb.create_sheet("Campos não migrados")
    ws["A1"], ws["B1"], ws["C1"] = "Coluna de origem", "Processo", "Valor"
    for c in ("A1", "B1", "C1"):
        ws[c].font = Font(bold=True, color="FFFFFF")
        ws[c].fill = HEADER_FILL
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 70
    tabela = Table(displayName=TABELA_NM, ref="A1:C2")
    tabela.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(tabela)
    return ws


def _dashboard(wb, ind, pos, quadros, hist):
    ws = wb.create_sheet("Dashboard")
    ws["A1"] = "Resumo da carteira"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
    kpis = [("Processos", "total", "0"), ("Ativos", "ativos", "0"), ("Taxa de êxito", "exito", "0.0%"),
            ("Economia efetiva", "economia", "#,##0.00")]
    for i, (rotulo, chave, fmt) in enumerate(kpis):
        col = 1 + 2 * i
        ws.cell(3, col, rotulo).font = Font(bold=True, color="595959")
        v = ws.cell(4, col, f"=Indicadores!{pos[chave]}")
        v.number_format = fmt
        v.font = Font(bold=True, size=16, color="1F4E78")
        ws.column_dimensions[_letra(col)].width = 22
        ws.column_dimensions[_letra(col + 1)].width = 4

    def barras(titulo, chave, ancora):
        _t, ini, fim = quadros[chave]
        g = BarChart()
        g.title = titulo
        g.type = "bar" if chave in ("area", "resultado") else "col"
        g.add_data(Reference(ind, min_col=2, min_row=ini, max_row=fim), titles_from_data=False)
        g.set_categories(Reference(ind, min_col=1, min_row=ini, max_row=fim))
        g.legend = None
        g.y_axis.title = "Processos"
        g.x_axis.delete = False
        g.y_axis.delete = False
        g.height, g.width = 7.5, 14
        ws.add_chart(g, ancora)
    barras("Processos por situação", "situacao", "A6")
    barras("Processos por resultado", "resultado", "H6")
    barras("Processos por área", "area", "A22")
    barras("Processos por probabilidade", "prob", "H22")
    linhas = LineChart()
    linhas.title = "Evolução do acervo (retrato mensal)"
    linhas.add_data(Reference(hist, min_col=2, max_col=3, min_row=1, max_row=2), titles_from_data=True)
    linhas.set_categories(Reference(hist, min_col=1, min_row=2, max_row=2))
    linhas.x_axis.number_format = "mm/yyyy"
    linhas.x_axis.delete = False
    linhas.y_axis.delete = False
    linhas.height, linhas.width = 7.5, 14
    ws.add_chart(linhas, "A38")


def construir():
    wb = Workbook()
    wb.properties.creator = "Relatório de andamentos (modelo padrão)"
    wb.properties.created = wb.properties.modified = DATA_FIXA
    wb.properties.lastModifiedBy = wb.properties.creator
    _processos(wb)
    _parametros(wb)
    ind, pos, quadros = _indicadores(wb)
    hist = _historico(wb)
    _nao_migrados(wb)
    _dashboard(wb, ind, pos, quadros, hist)
    # ordem das abas: Processos, Parâmetros, Indicadores, Dashboard, Histórico, Campos não migrados
    ordem = ["Processos", "Parâmetros", "Indicadores", "Dashboard", "Histórico", "Campos não migrados"]
    wb._sheets = [wb[n] for n in ordem]
    wb.active = 0
    return wb


NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _compartilhar_textos(partes):
    """Troca as strings inline que o openpyxl desta versão grava por sharedStrings, como o Excel grava
    (cabeçalhos de tabela, em especial, ficam no formato que o Excel espera)."""
    tabela, ordem, total = {}, [], 0
    padrao = re.compile(r'<c r="([A-Z]+\d+)"((?: s="\d+")?) t="inlineStr"><is><t(?: [^>]*)?>(.*?)</t></is></c>', re.S)
    for nome in sorted(partes):
        if not (nome.startswith("xl/worksheets/sheet") and nome.endswith(".xml")):
            continue

        def sub(m):
            nonlocal total
            txt = m.group(3)
            if txt not in tabela:
                tabela[txt] = len(ordem)
                ordem.append(txt)
            total += 1
            return f'<c r="{m.group(1)}"{m.group(2)} t="s"><v>{tabela[txt]}</v></c>'
        partes[nome] = padrao.sub(sub, partes[nome].decode("utf-8")).encode("utf-8")
    if "xl/sharedStrings.xml" in partes or not ordem:
        return
    itens = "".join(f'<si><t xml:space="preserve">{t}</t></si>' if t != t.strip() else f"<si><t>{t}</t></si>"
                    for t in ordem)
    partes["xl/sharedStrings.xml"] = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<sst xmlns="{NS_MAIN}" count="{total}" '
        f'uniqueCount="{len(ordem)}">{itens}</sst>').encode("utf-8")
    partes["xl/_rels/workbook.xml.rels"] = partes["xl/_rels/workbook.xml.rels"].replace(
        b"</Relationships>",
        b'<Relationship Id="rIdSst1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
        b'sharedStrings" Target="sharedStrings.xml"/></Relationships>')
    partes["[Content_Types].xml"] = partes["[Content_Types].xml"].replace(
        b"</Types>",
        b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.'
        b'spreadsheetml.sharedStrings+xml"/></Types>')


def gerar(destino=None):
    """Grava o modelo padrão (determinístico) em `destino` e devolve o caminho."""
    destino = Path(destino or AQUI / "modelo_padrao.xlsx")
    buf = io.BytesIO()
    construir().save(buf)
    with zipfile.ZipFile(buf) as z:
        partes = {i.filename: z.read(i.filename) for i in z.infolist()}
    _compartilhar_textos(partes)
    carimbo = DATA_FIXA.strftime("%Y-%m-%dT%H:%M:%SZ")
    core = partes["docProps/core.xml"].decode("utf-8")
    partes["docProps/core.xml"] = re.sub(r"(<dcterms:(?:created|modified)\b[^>]*>)[^<]*",
                                         lambda m: m.group(1) + carimbo, core).encode("utf-8")
    destino.parent.mkdir(parents=True, exist_ok=True)
    ordem = ["[Content_Types].xml"] + sorted(k for k in partes if k != "[Content_Types].xml")
    with zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as out:
        for nome in ordem:
            zi = zipfile.ZipInfo(nome, ZIP_DATA)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            out.writestr(zi, partes[nome])
    return destino


if __name__ == "__main__":
    print(gerar(sys.argv[1] if len(sys.argv) > 1 else None))
