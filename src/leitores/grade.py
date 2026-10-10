"""Leitura de planilhas e tabelas (.xlsx, .csv) para os leitores `xlsx_b`, `tabela_livre` e `lista` (WS-2).

Três peças:

1. `Grade`: uma aba ou um CSV como lista de linhas (valores já "crus"), com a marca de quais células são
   fórmula e de qual é o formato de cada célula. `carregar_xlsx` e `carregar_csv` produzem grades; erro de
   arquivo (corrompido, protegido, ilegível) volta como aviso `arquivo_ilegivel`, nunca como exceção.

2. Mapeador de cabeçalhos: `propor_mapeamento` casa cada coluna com um campo da ficha (`ficha.CAMPOS`) por
   semelhança do cabeçalho (sem acento, caixa ou pontuação; abreviações; sinônimos; palavras a mais) e devolve
   a CONFIANÇA. Casamento ambíguo (dois campos quase empatados) ou repetido (duas colunas para o mesmo
   campo) não é aplicado: vira aviso e a coluna vai para `colunas_sem_destino`. Também reconhece, pelo
   conteúdo, a coluna de números CNJ quando o cabeçalho não ajuda. Destinos especiais além dos campos da
   ficha: `numero`, `andamentos` (histórico em texto), `tribunal` (deduzido do número, ignorado) e `ativo`.

3. `extrair_processos`: percorre as linhas de uma grade e devolve `ProcessoLido`. Ignora linhas-marcador e
   totais (e as lista, em um aviso `linhas_ignoradas`); recusa e lista número com dígito errado; usa o valor
   em cache das células de fórmula e avisa (`formula_sem_valor`) quando não há; nada se perde em silêncio.

Contingência (relatórios de passivo contingente): reconhece passivo potencial (e atualizado), ativo potencial,
depósito judicial (Sim/Não e valor), "%" ao lado da provisão, provisão constituída, CNPJ processado e pagamento
realizado; "NATUREZA DA AÇÃO" é lida como ASSUNTO (não como classe processual); "REMOTA - justificativa" vira o grau
(`probabilidade`) mais o texto (`justificativa_probabilidade`); fórmulas simples sem valor guardado (`=G4-K4`) são
calculadas pelo programa (aviso `formula_calculada`); tabelas de critérios, legendas e totais no rodapé não viram processo.

Limites conhecidos: só lê cabeçalho de uma linha; células mescladas contam só na célula de cima/esquerda;
`.xls` antigo e `.ods` não são lidos (o detector avisa para salvar como `.xlsx`).
"""
import csv
import functools
import io
import math
import re
import warnings
import zipfile
from decimal import Decimal, InvalidOperation
from difflib import SequenceMatcher
from pathlib import Path

import ficha

from . import base

LIMIAR_AUTO = 0.7        # a partir daqui o mapeamento é aplicado sem confirmação
LIMIAR_MIN = 0.5         # abaixo disso nem é proposto
MAX_LINHAS_CABECALHO = 40
LIMITE_BYTES = 300 * 1024 * 1024

ESPECIAIS = {"numero": "Número do processo", "andamentos": "Andamentos (histórico em texto)",
             "tribunal": "Tribunal (deduzido do número; ignorado)", "ativo": "Ativo (calculado; ignorado)"}
DERIVADOS = ("ativo", "valor_economizado", "taxa_resolucao_dias")   # colunas que costumam ser fórmula

# Colunas do modelo B (CONTRATOS/WORKSTREAMS "fatos estruturais"); usadas para reconhecer o próprio modelo B.
CABECALHOS_B = ("Número do Processo", "Autor(es)", "Réu(s)", "Vara", "Município", "Tribunal", "Data do Ajuizamento",
                "Área do Direito", "Matéria Principal", "Objeto", "Valor da Causa", "Andamentos", "Situação", "Ativo",
                "Valor Arbitrado em Juízo", "Probabilidade", "Valor Estimado", "Valor da Execução", "Custas Processuais",
                "Depósitos Recursais", "Garantias Processuais", "Resultado", "Valor Economizado",
                "Data do trânsito em julgado", "Taxa de resolução (em dias)", "Houve recurso da empresa?",
                "Percentual de êxito", "Reclamante terceirizado?", "Outra(s) Parte(s)")

# destino -> formas conhecidas do cabeçalho (o rótulo de ficha.CAMPOS entra automaticamente)
SINONIMOS = {
    "numero": ("numero do processo", "processo", "numero", "n do processo", "numero processo", "autos", "cnj",
               "numero cnj", "processo cnj", "numero unico", "numero unico do processo", "processos", "n cnj",
               "numero do processo cnj", "processo numero", "n processo"),
    "andamentos": ("andamentos", "andamento", "historico", "historico do processo", "movimentacoes", "movimentacao",
                   "movimentos", "andamentos processuais", "ultimos andamentos", "resumo dos andamentos",
                   "texto de andamentos", "atualizacoes", "relatorio de andamentos"),
    "tribunal": ("tribunal", "orgao", "justica", "tribunal de origem", "tribunal de justica"),
    "ativo": ("ativo", "processo ativo", "esta ativo", "ativo sim nao"),
    "cliente": ("cliente", "nosso cliente", "razao social", "cliente razao social", "nome do cliente", "parte cliente"),
    "apelido": ("apelido", "apelido do processo", "nome do caso", "caso", "referencia interna"),
    "responsavel": ("responsavel", "advogado", "advogado responsavel", "responsavel pelo processo", "adv responsavel",
                    "socio responsavel"),
    "contato": ("contato", "e mail", "email", "whatsapp", "telefone", "contato do cliente"),
    "observacoes": ("observacoes", "observacao", "anotacoes", "comentarios", "notas", "obs"),
    "polo_cliente": ("polo", "polo do cliente", "posicao", "posicao do cliente", "polo processual", "polo cliente"),
    "autores": ("autor es", "autor", "autores", "autora", "requerente", "requerentes", "reclamante", "reclamantes",
                "exequente", "exequentes", "demandante", "polo ativo", "parte autora", "impetrante", "parte ativa"),
    "reus": ("reu s", "reu", "reus", "re", "requerido", "requeridos", "reclamado", "reclamada", "reclamadas",
             "reclamado s", "executado", "executados", "demandado", "polo passivo", "parte re", "parte passiva",
             "impetrado"),
    "parte_contraria": ("parte contraria", "contraria", "adverso", "parte adversa", "contraparte"),
    "outras_partes": ("outra s parte s", "outras partes", "outras partes envolvidas", "litisconsortes", "demais partes"),
    "terceirizado": ("reclamante terceirizado", "terceirizado", "trabalhador terceirizado", "terceirizado s"),
    "vara": ("vara", "juizo", "vara juizo", "orgao julgador", "vara de origem", "unidade", "vara comarca"),
    "municipio": ("municipio", "cidade", "comarca", "foro", "localidade", "municipio da vara"),
    "uf": ("uf", "estado", "unidade federativa"),
    "data_ajuizamento": ("data do ajuizamento", "ajuizamento", "data de ajuizamento", "data de distribuicao",
                         "distribuicao", "data da distribuicao", "distribuido em", "ajuizado em", "data da propositura",
                         "data distribuicao", "data de entrada"),
    "data_citacao": ("data de citacao", "citacao", "data da citacao", "citado em"),
    "classe": ("classe", "classe processual", "tipo de acao", "tipo de processo", "tipo da acao"),
    "assunto": ("assunto", "assuntos", "assunto principal", "assunto do processo", "natureza da acao", "natureza",
                "natureza do processo", "natureza da demanda"),
    "area": ("area do direito", "area", "area juridica", "ramo do direito", "ramo", "area do processo"),
    "materia_principal": ("materia principal", "materia", "tese", "tese principal", "materia do processo",
                          "causa de pedir"),
    "objeto": ("objeto", "objeto da acao", "pedido", "pedidos", "pedido principal", "resumo do objeto",
               "breve resumo do caso", "resumo do caso", "sintese do caso", "resumo da acao"),
    "valor_causa": ("valor da causa", "valor causa", "valor atribuido a causa", "valor atribuido", "valor da acao",
                    "valor do processo", "valor inicial", "valor da causa atualizado"),
    "momento_atual": ("momento atual", "momento atual do processo", "momento processual", "fase atual",
                      "andamento atual", "posicao atual"),
    "situacao": ("situacao", "status", "situacao do processo", "status do processo", "situacao processual"),
    "fase": ("fase", "fase processual"),
    "ultimo_andamento": ("ultimo andamento", "data do ultimo andamento", "data ultimo andamento", "ultima movimentacao",
                         "data da ultima movimentacao", "ultima atualizacao", "data da ultima atualizacao"),
    "houve_recurso": ("houve recurso da empresa", "houve recurso", "recurso da empresa", "recorreu", "recurso interposto",
                      "recurso"),
    "resultado": ("resultado", "desfecho", "resultado do processo", "resultado final", "decisao final"),
    "probabilidade": ("probabilidade", "probabilidade do resultado", "probabilidade de perda", "risco",
                      "classificacao de risco", "classificacao do risco", "possibilidade de perda", "chance de perda",
                      "risco de perda", "prognostico", "grau de risco", "classificacao da perda"),
    "valor_arbitrado": ("valor arbitrado em juizo", "valor arbitrado", "condenacao arbitrada", "valor da condenacao",
                        "condenacao", "valor condenado"),
    "valor_estimado": ("valor estimado", "estimativa", "valor estimado da condenacao"),
    "valor_execucao": ("valor da execucao", "valor execucao", "valor executado", "valor em execucao", "execucao"),
    "valor_acordo": ("valor do acordo", "valor acordo", "valor acordado", "acordo"),
    "valor_economizado": ("valor economizado", "economia", "economia efetiva"),
    "custas": ("custas processuais", "custas", "custas judiciais"),
    "depositos_recursais": ("depositos recursais", "deposito recursal", "depositos", "deposito",
                            "valor do deposito judicial", "valor do deposito", "valor depositado", "deposito judicial valor",
                            "valor dos depositos judiciais", "depositos judiciais", "valor do deposito recursal"),
    "garantias": ("garantias processuais", "garantia", "garantias", "garantia do juizo"),
    "data_transito": ("data do transito em julgado", "transito em julgado", "data de transito", "data do transito",
                      "transito", "data do transito em julgado"),
    "taxa_resolucao_dias": ("taxa de resolucao em dias", "taxa de resolucao dias", "taxa de resolucao",
                            "tempo de resolucao", "dias para resolucao", "prazo de resolucao", "duracao em dias"),
    "percentual_exito": ("percentual de exito", "exito", "taxa de exito", "percentual exito"),
    # contingência
    "passivo_potencial": ("passivo potencial", "passivo", "valor contingenciado", "contingencia", "valor da contingencia",
                          "valor em risco", "passivo contingente", "passivo estimado", "risco financeiro"),
    "passivo_atualizado": ("passivo potencial atualizado", "passivo atualizado", "passivo potencial corrigido",
                           "contingencia atualizada", "passivo potencial atualizado correcao"),
    "ativo_potencial": ("ativo potencial", "ativo contingente", "credito potencial", "valor a receber"),
    "percentual_provisao": ("percentual de provisao", "percentual provisao", "percentual provisionado",
                            "percentual da provisao", "indice de provisao", "perc provisao", "pct provisao"),
    "provisao": ("provisao constituida", "provisao", "valor provisionado", "provisao contabil", "valor da provisao",
                 "provisionamento"),
    "deposito_judicial": ("deposito judicial realizado", "deposito judicial", "houve deposito judicial",
                          "deposito realizado", "ha deposito judicial", "deposito judicial sim nao"),
    "cnpj_processado": ("cnpj", "cnpj processado", "cnpj da empresa processada", "cnpj do reu", "cnpj da reclamada",
                        "cnpj empresa", "cnpj da empresa"),
    "pagamento_realizado": ("pagamento realizado", "pagamento", "pagamentos realizados", "valor pago",
                            "valor do pagamento", "pagamento efetuado"),
    "justificativa_probabilidade": ("justificativa", "justificativa da probabilidade",
                                    "justificativa da possibilidade de perda", "justificativa do risco",
                                    "fundamento da classificacao", "motivo da classificacao"),
}

# Casamentos EXATOS que não valem 100%: o nome é parecido, mas o campo pode não ser o mesmo. {(destino, forma): confiança}
CONFIANCA_DA_FORMA = {("assunto", "natureza da acao"): 0.85, ("assunto", "natureza"): 0.8,
                      ("assunto", "natureza do processo"): 0.8, ("assunto", "natureza da demanda"): 0.85,
                      ("materia_principal", "natureza da acao"): 0.75, ("materia_principal", "natureza"): 0.7,
                      ("passivo_potencial", "passivo"): 0.9, ("passivo_potencial", "contingencia"): 0.85,
                      ("pagamento_realizado", "pagamento"): 0.9, ("cnpj_processado", "cnpj"): 0.8,
                      ("justificativa_probabilidade", "justificativa"): 0.85}
# Formas que servem como segunda opção de OUTRO campo (aparece como candidato, nunca vence o primeiro)
FORMAS_EXTRAS = {"materia_principal": ("natureza da acao", "natureza")}
# Cabeçalhos com palavras a mais, reconhecidos por padrão: (expressão sobre a chave sem acento, destino, confiança)
PADROES = (
    (r"^cnpj\b.*\b(processad[oa]s?|reu|reus|reclamad[oa]s?|empresa|devedor[a]?|demandad[oa]s?)\b", "cnpj_processado", 0.92),
    (r"^passivo (potencial )?(atualizado|corrigido)\b", "passivo_atualizado", 0.95),
    (r"^deposito judicial realizado\b", "deposito_judicial", 0.95),
    (r"^valor (do |dos )?depositos? (judicia(l|is)|recursais?)\b", "depositos_recursais", 0.95),
    (r"^provisao (constituida|contabil)\b", "provisao", 0.95),
    (r"^pagamentos? (realizados?|efetuados?)\b", "pagamento_realizado", 0.95),
    (r"^justificativa\b.*\b(probabilidade|perda|risco|classificacao)\b", "justificativa_probabilidade", 0.92),
)
# Explicação mostrada na tela de mapeamento junto da confiança de um destino
COMENTARIOS_DO_DESTINO = {
    "assunto": ("Colunas chamadas \"Natureza da ação\" costumam trazer o assunto ou a matéria (por exemplo, indenização por danos "
                "morais), não a classe processual do tribunal (por exemplo, Procedimento Comum Cível); por isso a coluna foi "
                "ligada a Assunto, com confiança menor. Confira e troque se preferir."),
    "percentual_provisao": "Coluna \"%\" ao lado da provisão ou da possibilidade de perda: lida como percentual de provisão (0,5 = 50%).",
    "probabilidade": ("Se a célula trouxer o grau e a justificativa juntos (\"REMOTA - processo extinto...\"), o grau vai para "
                      "Probabilidade e o texto para Justificativa da probabilidade."),
}
_PALAVRAS_FRACAS = {"de", "da", "do", "das", "dos", "e", "a", "o", "as", "os", "em", "para", "ao", "por", "com", "es", "s",
                    "r"}
_PRIMEIRO = {"n": "numero", "no": "numero", "nr": "numero", "num": "numero", "nro": "numero"}
_ABREV = {"vlr": "valor", "dt": "data", "proc": "processo", "obs": "observacoes", "ult": "ultimo",
          "mun": "municipio", "qtd": "quantidade"}


def _expandir(k):
    """Chave normalizada com abreviações expandidas ('Nº proc.' -> 'numero processo'; 'Vlr causa (R$)' -> 'valor causa')."""
    saida = []
    for i, t in enumerate(k.split()):
        if i == 0 and t in _PRIMEIRO:
            t = _PRIMEIRO[t]
        t = _ABREV.get(t, t)
        if t == "r":
            continue
        saida.append(t)
    return " ".join(saida)


def _tokens(k):
    return {t for t in k.split() if t not in _PALAVRAS_FRACAS}


def _chave_com_percentual(texto):
    """Chave do cabeçalho em que o sinal % vale a palavra "percentual" ("% provisão" -> "percentual provisao")."""
    return base.chave(str(texto).replace("%", " percentual "))


def _construir_indice():
    indice = {}
    for destino in (*ficha.CAMPOS, *ESPECIAIS):
        formas = {_expandir(_chave_com_percentual(s)) for s in (*SINONIMOS.get(destino, ()), *FORMAS_EXTRAS.get(destino, ()))}
        if destino in ficha.CAMPOS:
            formas.add(_expandir(_chave_com_percentual(ficha.CAMPOS[destino][0])))
        formas.discard("")
        indice[destino] = [(f, _tokens(f)) for f in sorted(formas)]
    return indice


_INDICE = _construir_indice()
_EXATOS_B = {_expandir(base.chave(c)) for c in CABECALHOS_B}


def _construir_exatos():
    exatos = {}
    for destino, formas in _INDICE.items():
        for forma, _ in formas:
            exatos.setdefault(forma, destino)
            if exatos[forma] != destino:
                exatos[forma] = None          # forma que serve a dois destinos: não é exata
    return {f: d for f, d in exatos.items() if d}


_EXATOS = _construir_exatos()


def destino_exato(rotulo):
    """Campo da ficha (ou destino especial) cujo nome conhecido é IGUAL ao rótulo (sem acento, caixa ou pontuação); None se não houver."""
    return _EXATOS.get(_expandir(_chave_com_percentual(rotulo)))


def rotulo_do_destino(destino):
    if destino in ESPECIAIS:
        return ESPECIAIS[destino]
    return ficha.CAMPOS[destino][0] if destino in ficha.CAMPOS else destino


def destinos_possiveis():
    """Lista para a tela de mapeamento: [{"campo", "rotulo", "grupo"}] (campos da ficha + destinos especiais)."""
    saida = [{"campo": c, "rotulo": r, "grupo": "especial"} for c, r in ESPECIAIS.items()]
    saida += [{"campo": c, "rotulo": v[0], "grupo": v[1]} for c, v in ficha.CAMPOS.items()]
    return saida


def pontuar_cabecalho(cabecalho):
    """[(confianca, destino)] em ordem decrescente para um texto de cabeçalho (só >= LIMIAR_MIN)."""
    return list(_pontuar(str(cabecalho)))


def _sem_parenteses(texto):
    return re.sub(r"[(\[][^)\]]*[)\]]", " ", str(texto))


@functools.lru_cache(maxsize=8192)
def _pontuar(cabecalho):
    h = _expandir(_chave_com_percentual(cabecalho))
    if not h:
        return ()
    h_limpo = _expandir(_chave_com_percentual(_sem_parenteses(cabecalho)))      # "Valor da causa (R$)" ~ "Valor da causa"
    ht = _tokens(h)
    achados = []
    for destino, formas in _INDICE.items():
        melhor = 0.0
        for forma, ft in formas:
            if h == forma:
                melhor = max(melhor, CONFIANCA_DA_FORMA.get((destino, forma), 1.0))
                if melhor >= 1.0:
                    break
                continue
            if h_limpo and h_limpo != h and h_limpo == forma:
                melhor = max(melhor, min(0.97, CONFIANCA_DA_FORMA.get((destino, forma), 0.97)))
                continue
            if ht and ft and (ft <= ht or ht <= ft):
                melhor = max(melhor, 0.9 * math.sqrt(min(len(ft), len(ht)) / max(len(ft), len(ht))))
            else:
                sm = SequenceMatcher(None, h, forma)
                if sm.real_quick_ratio() >= 0.8 and sm.quick_ratio() >= 0.8:
                    r = sm.ratio()
                    if r >= 0.82:
                        melhor = max(melhor, r * 0.88)
        if melhor >= LIMIAR_MIN:
            achados.append((round(melhor, 3), destino))
    for padrao, destino, conf in PADROES:         # cabeçalhos com palavras a mais ("CNPJ DA EMPRESA X PROCESSADA")
        if re.search(padrao, h) or (h_limpo != h and re.search(padrao, h_limpo)):
            atual = max((c for c, d in achados if d == destino), default=0.0)
            achados = [(c, d) for c, d in achados if d != destino] + [(max(atual, conf), destino)]
    # "AUTOR/RECLAMANTE", "RÉU/RECLAMADO": cada parte é um nome conhecido do MESMO campo -> o cabeçalho inteiro também
    partes = [x for x in re.split(r"\s*[/|]\s*|\s+ou\s+", cabecalho.strip(), flags=re.I) if x.strip()]
    if len(partes) > 1:
        donos = {_EXATOS.get(_expandir(base.chave(x))) for x in partes}
        if len(donos) == 1 and None not in donos:
            dono = donos.pop()
            achados = [(c, d) for c, d in achados if d != dono] + [(0.95, dono)]
    return tuple(sorted(achados, key=lambda x: (-x[0], x[1])))


def _razao(amostras, teste):
    valores = [v for v in amostras if not base.vazio(v)]
    if not valores:
        return None
    return sum(1 for v in valores if teste(v)) / len(valores)


def _tem_cnj(valor):
    return bool(base.achar_numeros(valor)) if not isinstance(valor, (int, float)) else False


def _data_ok(valor):
    v, problema = base.converter_data(valor)
    return v is not None or problema is None


def _dinheiro_ok(valor):
    v, problema = base.converter_dinheiro(valor)
    return v is not None or problema is None


def _sim_nao_ok(valor):
    v, problema = base.converter_sim_nao(valor)
    return v is not None or problema is None


def _numero_ok(valor):
    v, problema = base.converter_numero(valor)
    return v is not None or problema is None


_VIZINHAS_DO_PERCENTUAL = {"provisao", "probabilidade", "passivo_potencial", "passivo_atualizado", "percentual_provisao"}


def _so_percentual(cabecalho):
    """O cabeçalho é só o sinal de porcentagem ("%", "% ", "perc.")?"""
    return _chave_com_percentual(cabecalho) in ("percentual", "perc", "pct", "porcentagem")


def _destino_da_vizinha(cabecalhos, i):
    """Destino (melhor palpite) da coluna preenchida mais próxima à esquerda ou, falhando, à direita de `i`."""
    for passo in (-1, 1):
        j = i + passo
        while 0 <= j < len(cabecalhos) and base.vazio(cabecalhos[j]):
            j += passo
        if 0 <= j < len(cabecalhos) and not _so_percentual(cabecalhos[j]):
            achados = pontuar_cabecalho(cabecalhos[j])
            if achados and achados[0][1] in _VIZINHAS_DO_PERCENTUAL and achados[0][0] >= LIMIAR_AUTO:
                return achados[0][1]
    return None


def propor_mapeamento(cabecalhos, amostras=None):
    """Casa colunas e campos. `cabecalhos`: textos (None/'' = coluna sem cabeçalho, ignorada); `amostras`: lista
    (uma por coluna) de valores de exemplo (opcional; melhora a decisão).
    -> [{"indice", "coluna", "campo" | None, "confianca", "candidatos": [campo...], "ambigua": bool,
         "duplicada_de": indice | None, "aplicado": bool}]"""
    amostras = amostras or [[] for _ in cabecalhos]
    registros = []
    for i, cab in enumerate(cabecalhos):
        if base.vazio(cab):
            continue
        cands = pontuar_cabecalho(cab)
        amo = amostras[i] if i < len(amostras) else []
        if _so_percentual(cab):          # coluna "%": só vale ao lado da provisão ou da possibilidade de perda
            vizinha = _destino_da_vizinha(cabecalhos, i)
            cands = [(0.85, "percentual_provisao")] if vizinha in _VIZINHAS_DO_PERCENTUAL else []
        cnj = _razao(amo, _tem_cnj)
        if cnj is not None and cnj >= 0.5 and not any(d == "numero" and s >= 0.9 for s, d in cands):
            cands = sorted([(s, d) for s, d in cands if d != "numero"] + [(0.9, "numero")], key=lambda x: (-x[0], x[1]))
        ajustados = []
        for s, d in cands:
            if d in ficha.CAMPOS and ficha.CAMPOS[d][2] in ("data", "dinheiro") and len([v for v in amo if not base.vazio(v)]) >= 3:
                r = _razao(amo, _data_ok if ficha.CAMPOS[d][2] == "data" else _dinheiro_ok)
                if r is not None and r < 0.3:
                    s = round(s * 0.5, 3)
            elif d in ficha.CAMPOS and ficha.CAMPOS[d][2] in ("sim_nao", "numero") and len([v for v in amo if not base.vazio(v)]) >= 3:
                teste = _sim_nao_ok if ficha.CAMPOS[d][2] == "sim_nao" else _numero_ok
                r = _razao(amo, teste)
                if r is not None and r < 0.3:
                    s = round(s * 0.5, 3)
            if s >= LIMIAR_MIN:
                ajustados.append((s, d))
        ajustados.sort(key=lambda x: (-x[0], x[1]))
        reg = {"indice": i, "coluna": str(cab).strip(), "campo": None, "confianca": 0.0,
               "candidatos": [d for _, d in ajustados[:3]], "ambigua": False, "duplicada_de": None, "aplicado": False}
        if ajustados:
            melhor_s, melhor_d = ajustados[0]
            reg["campo"], reg["confianca"] = melhor_d, melhor_s
            if len(ajustados) > 1 and melhor_s < 0.97 and melhor_s - ajustados[1][0] < 0.05:
                reg["ambigua"], reg["campo"] = True, None
        registros.append(reg)
    # um campo por coluna: a de maior confiança ganha; a outra vira "duplicada"
    donos = {}
    for reg in sorted((r for r in registros if r["campo"]), key=lambda r: (-r["confianca"], r["indice"])):
        if reg["campo"] in donos:
            reg["duplicada_de"], reg["campo"] = donos[reg["campo"]], None
        else:
            donos[reg["campo"]] = reg["indice"]
    for reg in registros:
        reg["aplicado"] = bool(reg["campo"]) and reg["confianca"] >= LIMIAR_AUTO
    _reconhecer_historico_pelo_conteudo(registros, amostras)
    return registros


_DATA_NO_TEXTO = re.compile(r"\b\d{2}/\d{2}/\d{4}\b")


def _parece_historico(amostra):
    """A coluna parece o HISTÓRICO de andamentos pelo conteúdo? Pelo menos 3 células preenchidas e metade delas com
    2 datas ou mais e 25 palavras ou mais (texto corrido: 'Em 10/05/2026, foi ...'). Serve para colunas chamadas
    'Observação', 'Comentários' etc. que o escritório usa para o histórico."""
    cheias = [str(v) for v in amostra if not base.vazio(v)]
    if len(cheias) < 3:
        return False
    bons = sum(1 for t in cheias if len(_DATA_NO_TEXTO.findall(t)) >= 2 and len(t.split()) >= 25)
    return bons / len(cheias) >= 0.5


def _reconhecer_historico_pelo_conteudo(registros, amostras):
    """Sem uma coluna de andamentos reconhecida pelo nome, a coluna cujo CONTEÚDO é um histórico datado vira
    `andamentos` (confiança 0,9; a tela de mapeamento deixa corrigir). Só vale para colunas que estavam sem destino ou
    com destino de texto livre (observações, apelido, objeto)."""
    if any(r["campo"] == "andamentos" and r["confianca"] >= 0.9 for r in registros):
        return
    candidatas = [r for r in registros if r["indice"] < len(amostras) and _parece_historico(amostras[r["indice"]])
                  and r["campo"] in (None, "observacoes", "apelido", "objeto", "situacao")]
    if not candidatas:
        return
    melhor = max(candidatas, key=lambda r: sum(len(str(v).split()) for v in amostras[r["indice"]] if not base.vazio(v)))
    for r in registros:
        if r["campo"] == "andamentos" and r is not melhor:
            r["campo"], r["aplicado"] = None, False
    melhor.update(campo="andamentos", confianca=0.9, aplicado=True, ambigua=False, duplicada_de=None)


def aplicar_mapeamento_do_usuario(registros, mapeamento, avisos, onde):
    """Sobrepõe as escolhas do usuário. `mapeamento`: {cabeçalho: campo | None} ou lista de {"coluna","campo"}."""
    if not mapeamento:
        return registros
    if not isinstance(mapeamento, dict):
        mapeamento = {m.get("coluna"): m.get("campo") for m in mapeamento if isinstance(m, dict)}
    escolhas = {base.chave(k): v for k, v in mapeamento.items()}
    validos = set(ficha.CAMPOS) | set(ESPECIAIS)
    for reg in registros:
        k = base.chave(reg["coluna"])
        if k not in escolhas:
            continue
        campo = escolhas[k] or None
        if campo is not None and campo not in validos:
            avisos.append(base.aviso("atencao", "mapeamento_invalido", onde,
                                     f"O destino {campo!r} escolhido para a coluna {reg['coluna']!r} não existe; escolha mantida como automática.", [campo]))
            continue
        reg.update(campo=campo, confianca=1.0 if campo else 0.0, ambigua=False, duplicada_de=None, aplicado=bool(campo),
                   decisao="usuario")
    donos = {}
    for reg in registros:                 # dois destinos iguais escolhidos à mão: vale o primeiro
        if reg["campo"] and reg["aplicado"]:
            if reg["campo"] in donos:
                avisos.append(base.aviso("atencao", "coluna_duplicada", onde,
                                         f"As colunas {donos[reg['campo']]['coluna']!r} e {reg['coluna']!r} foram ligadas ao mesmo campo "
                                         f"({rotulo_do_destino(reg['campo'])}); só a primeira será lida.", [reg["coluna"]]))
                reg.update(campo=None, aplicado=False, duplicada_de=donos[reg["campo"]]["indice"])
            else:
                donos[reg["campo"]] = reg
    return registros


# ---------------------------------------------------------------- grades

class Grade:
    """Uma aba (ou um CSV): `linhas[i][j]` é a célula da linha i+1, coluna j+1."""

    def __init__(self, nome, linhas, formulas=None, formatos=None, epoch=None, oculta=False, textos_formulas=None):
        self.nome = nome
        self.linhas = linhas
        self.formulas = formulas or set()
        self.textos_formulas = textos_formulas or {}        # {(i, j): "=G4-K4"}, para recalcular fórmula sem valor guardado
        self.formatos = formatos or {}
        self.epoch = epoch
        self.oculta = oculta

    def celula(self, i, j):
        if i < 0 or i >= len(self.linhas):
            return None
        linha = self.linhas[i]
        return linha[j] if 0 <= j < len(linha) else None

    def largura(self):
        return max((len(l) for l in self.linhas), default=0)


def _ler_bytes(caminho):
    tam = Path(caminho).stat().st_size
    if tam > LIMITE_BYTES:
        return None, base.aviso("erro", "arquivo_muito_grande", Path(caminho).name,
                                f"O arquivo tem {tam // (1024 * 1024)} MB; acima do limite de leitura ({LIMITE_BYTES // (1024 * 1024)} MB).")
    return Path(caminho).read_bytes(), None


def carregar_xlsx(caminho, max_linhas=None):
    """Grades de todas as abas de uma planilha. -> (grades, avisos). `max_linhas`: leitura leve (detecção)."""
    with warnings.catch_warnings():        # o openpyxl avisa de extensões que ignora (gráficos, formatação condicional x14...)
        warnings.simplefilter("ignore")
        return _carregar_xlsx(caminho, max_linhas)


def _carregar_xlsx(caminho, max_linhas):
    from openpyxl import load_workbook
    nome = Path(caminho).name
    avisos = []
    wb_v = wb_f = None
    try:
        if Path(caminho).stat().st_size > LIMITE_BYTES:
            raise ValueError("arquivo grande demais")
        wb_v = load_workbook(str(caminho), read_only=True, data_only=True)
        wb_f = load_workbook(str(caminho), read_only=True, data_only=False) if max_linhas is None else None
    except Exception as e:      # noqa: BLE001 -- qualquer falha de abertura é "arquivo ilegível", não erro de programa
        if wb_v is not None:
            wb_v.close()
        return [], [base.aviso("erro", "arquivo_ilegivel", nome,
                               "Não foi possível abrir a planilha (arquivo corrompido, protegido por senha ou "
                               f"não é um .xlsx): {type(e).__name__}.")]
    grades = []
    try:
        for ws in wb_v.worksheets:
            try:
                ws.reset_dimensions()
                linhas, formatos = [], {}
                for i, row in enumerate(ws.iter_rows(max_row=max_linhas)):
                    linhas.append([c.value for c in row])
                    for j, c in enumerate(row):
                        fmt = getattr(c, "number_format", None)
                        if fmt and fmt != "General":
                            formatos[(i, j)] = fmt
                formulas, textos = set(), {}
                if wb_f is not None:
                    wf = wb_f[ws.title]
                    wf.reset_dimensions()
                    for i, row in enumerate(wf.iter_rows()):
                        for j, c in enumerate(row):
                            if getattr(c, "data_type", None) == "f":
                                formulas.add((i, j))
                                if isinstance(c.value, str):
                                    textos[(i, j)] = c.value
                grades.append(Grade(ws.title, linhas, formulas, formatos, getattr(wb_v, "epoch", None),
                                    getattr(ws, "sheet_state", "visible") != "visible", textos))
            except Exception as e:      # noqa: BLE001
                avisos.append(base.aviso("erro", "arquivo_ilegivel", f"aba {ws.title!r}",
                                         f"A aba não pôde ser lida: {type(e).__name__}."))
    finally:
        wb_v.close()
        if wb_f is not None:
            wb_f.close()
    return grades, avisos


def decodificar_texto(dados):
    """bytes -> texto (UTF-8 com ou sem BOM; senão Windows-1252, comum em CSV do Excel brasileiro)."""
    for cod in ("utf-8-sig", "cp1252"):
        try:
            return dados.decode(cod)
        except UnicodeDecodeError:
            continue
    return dados.decode("latin-1", errors="replace")


def carregar_csv(caminho):
    """Grade de um .csv (separador `,`, `;`, tabulação ou `|`; UTF-8 ou Windows-1252). -> (grades, avisos)."""
    nome = Path(caminho).name
    dados, erro = _ler_bytes(caminho)
    if erro:
        return [], [erro]
    if b"\x00" in dados[:4096]:
        return [], [base.aviso("erro", "arquivo_ilegivel", nome, "O arquivo parece binário, não um CSV de texto.")]
    texto = decodificar_texto(dados)
    amostra = texto[:8192]
    try:
        dialeto = csv.Sniffer().sniff(amostra, delimiters=",;\t|")
    except csv.Error:
        primeira = amostra.splitlines()[0] if amostra.splitlines() else ""
        sep = max(",;\t|", key=primeira.count)

        class dialeto(csv.excel):        # noqa: N801
            delimiter = sep
    try:
        linhas = [[c.strip() for c in l] for l in csv.reader(io.StringIO(texto), dialeto)]
    except csv.Error as e:
        return [], [base.aviso("erro", "arquivo_ilegivel", nome, f"O CSV não pôde ser lido ({e}).")]
    return [Grade(Path(caminho).stem or "csv", linhas)], []


def _linha_tem_cnj(linha):
    return any(isinstance(v, str) and base.achar_numeros(v) for v in linha)


def achar_cabecalho(grade, max_linhas=MAX_LINHAS_CABECALHO):
    """(indice_da_linha, registros_de_mapeamento) do cabeçalho mais provável, ou None. É cabeçalho a linha cujo
    mapeamento inclui a coluna `numero` (pelo nome ou, se nenhum cabeçalho serve, pelo conteúdo da coluna logo
    abaixo) e que não tem número CNJ nas próprias células."""
    melhor = None
    for i, linha in enumerate(grade.linhas[:max_linhas]):
        textos = [v for v in linha if isinstance(v, str) and v.strip()]
        if not textos or _linha_tem_cnj(linha):
            continue
        amostras = [[grade.celula(r, j) for r in range(i + 1, min(i + 31, len(grade.linhas)))] for j in range(len(linha))]
        regs = propor_mapeamento(linha, amostras)
        num = next((r for r in regs if r["campo"] == "numero" and r["aplicado"]), None)
        if num is None:
            continue
        pont = sum(r["confianca"] for r in regs if r["campo"] and r["confianca"] >= LIMIAR_MIN)
        if melhor is None or pont > melhor[0]:
            melhor = (pont, i, regs)
    return (melhor[1], melhor[2]) if melhor else None


def cabecalho_exato_b(grade, indice, regs):
    """Quantas colunas do cabeçalho são, literalmente, colunas do modelo B."""
    linha = grade.linhas[indice]
    return sum(1 for v in linha if isinstance(v, str) and _expandir(base.chave(v)) in _EXATOS_B)


def grades_de_processos(grades):
    """[(grade, indice_cabecalho, registros)] das grades que têm cabeçalho com `numero` e ao menos um número CNJ abaixo."""
    saida = []
    for g in grades:
        achado = achar_cabecalho(g)
        if achado is None:
            continue
        i, regs = achado
        num = next(r for r in regs if r["campo"] == "numero" and r["aplicado"])
        if any(base.achar_numeros(g.celula(r, num["indice"])) for r in range(i + 1, len(g.linhas))
               if isinstance(g.celula(r, num["indice"]), str)):
            saida.append((g, i, regs))
    return saida


LISTA_CAMPOS = {"numero", "cliente", "polo_cliente", "parte_contraria", "responsavel", "contato", "observacoes", "apelido",
                "situacao", "tribunal"}


def classificar_grades(grades):
    """'xlsx_b' | 'tabela_livre' | 'lista' | None para as grades de uma planilha ou CSV.
    xlsx_b: cabeçalho com pelo menos 10 colunas do modelo B (ou 6, se a pasta de trabalho tem aba de parâmetros ou de
    indicadores, como o próprio modelo B com poucas colunas ativas); tabela_livre: cabeçalho com `numero` e dois ou mais
    campos que não são de lista; lista: o resto (inclusive sem cabeçalho, desde que haja números CNJ)."""
    com_proc = grades_de_processos(grades)
    do_modelo = any(t in base.chave(g.nome) for g in grades for t in ("param", "indicador"))
    for g, i, regs in com_proc:
        exatas = cabecalho_exato_b(g, i, regs)
        if exatas >= 10 or (exatas >= 6 and do_modelo):
            return "xlsx_b"
    for g, i, regs in com_proc:
        outros = {r["campo"] for r in regs if r["campo"] and r["aplicado"]} - LISTA_CAMPOS
        if len(outros) >= 2 or "andamentos" in outros:
            return "tabela_livre"
    if com_proc:
        return "lista"
    for g in grades:
        if any(_linha_tem_cnj(l) for l in g.linhas):
            return "lista"
    return None


# ---------------------------------------------------------------- fórmulas simples

class _FormulaNaoCalculavel(Exception):
    """A fórmula usa algo que o cálculo simples não conhece (outra aba, função, texto): o valor fica vazio."""


_TOKEN_FORMULA = re.compile(
    r"\s*(?:(?P<num>\d+(?:[.,]\d+)?)|(?P<ref>\$?[A-Za-z]{1,3}\$?\d+(?::\$?[A-Za-z]{1,3}\$?\d+)?)"
    r"|(?P<func>[A-Za-zÀ-ÿ]+)\s*\(|(?P<op>[-+*/(),;]))")


class _AvaliadorDeFormula:
    """Soma, subtração, multiplicação, divisão, parênteses, SUM/SOMA e referências a células da MESMA aba. Nada de eval."""

    def __init__(self, grade, texto, profundidade):
        from openpyxl.utils import column_index_from_string
        self._coluna = column_index_from_string
        self.grade, self.profundidade = grade, profundidade
        corpo = texto.strip()
        if not corpo.startswith("=") or "!" in corpo or "[" in corpo or '"' in corpo:
            raise _FormulaNaoCalculavel(texto)
        self.tokens, pos, corpo = [], 0, corpo[1:]
        while pos < len(corpo):
            m = _TOKEN_FORMULA.match(corpo, pos)
            if not m or m.end() == pos:
                if corpo[pos:].strip() == "":
                    break
                raise _FormulaNaoCalculavel(texto)
            pos = m.end()
            tipo = m.lastgroup
            self.tokens.append((tipo, m.group(tipo)))
        self.p = 0

    def _ver(self):
        return self.tokens[self.p] if self.p < len(self.tokens) else (None, None)

    def _tomar(self):
        t = self._ver()
        self.p += 1
        return t

    def avaliar(self):
        valor = self._soma()
        if self.p != len(self.tokens):
            raise _FormulaNaoCalculavel("sobrou")
        return valor

    def _soma(self):
        v = self._produto()
        while self._ver() in (("op", "+"), ("op", "-")):
            op = self._tomar()[1]
            w = self._produto()
            v = v + w if op == "+" else v - w
        return v

    def _produto(self):
        v = self._fator()
        while self._ver() in (("op", "*"), ("op", "/")):
            op = self._tomar()[1]
            w = self._fator()
            if op == "*":
                v = v * w
            else:
                if w == 0:
                    raise _FormulaNaoCalculavel("divisão por zero")
                v = v / w
        return v

    def _fator(self):
        tipo, valor = self._tomar()
        if tipo == "op" and valor in "+-":
            f = self._fator()
            return -f if valor == "-" else f
        if tipo == "op" and valor == "(":
            v = self._soma()
            if self._tomar() != ("op", ")"):
                raise _FormulaNaoCalculavel("parêntese")
            return v
        if tipo == "num":
            return Decimal(valor.replace(",", "."))
        if tipo == "ref":
            valores = self._celulas(valor)
            if len(valores) != 1 or ":" in valor:
                raise _FormulaNaoCalculavel("intervalo fora de SOMA")
            return valores[0]
        if tipo == "func":
            if base.chave(valor) not in ("sum", "soma"):
                raise _FormulaNaoCalculavel(valor)
            total, primeiro = Decimal(0), True
            while True:
                if self._ver() == ("op", ")"):
                    self._tomar()
                    break
                if not primeiro:
                    if self._tomar()[1] not in (",", ";"):
                        raise _FormulaNaoCalculavel("argumentos")
                primeiro = False
                seguinte = self.tokens[self.p + 1] if self.p + 1 < len(self.tokens) else (None, None)
                if self._ver()[0] == "ref" and seguinte in (("op", ","), ("op", ";"), ("op", ")")):
                    total += sum(self._celulas(self._tomar()[1]), Decimal(0))
                else:
                    total += self._soma()
            return total
        raise _FormulaNaoCalculavel("expressão")

    def _celulas(self, ref):
        partes = ref.replace("$", "").upper().split(":")
        pontos = []
        for parte in partes:
            m = re.fullmatch(r"([A-Z]{1,3})(\d+)", parte)
            if not m:
                raise _FormulaNaoCalculavel(ref)
            pontos.append((int(m.group(2)) - 1, self._coluna(m.group(1)) - 1))
        (i1, j1), (i2, j2) = pontos[0], pontos[-1]
        i1, i2, j1, j2 = min(i1, i2), max(i1, i2), min(j1, j2), max(j1, j2)
        if (i2 - i1 + 1) * (j2 - j1 + 1) > 5000:
            raise _FormulaNaoCalculavel("intervalo grande")
        return [self._numero(i, j) for i in range(i1, i2 + 1) for j in range(j1, j2 + 1)]

    def _numero(self, i, j):
        v = self.grade.celula(i, j)
        if base.vazio(v) and (i, j) in self.grade.textos_formulas:
            if self.profundidade >= 6:
                raise _FormulaNaoCalculavel("fórmulas encadeadas demais")
            v = calcular_formula(self.grade, i, j, self.profundidade + 1)
            if v is None:
                raise _FormulaNaoCalculavel("célula de fórmula sem valor")
        if base.vazio(v):
            return Decimal(0)                  # célula em branco vale zero em uma conta
        if isinstance(v, bool):
            raise _FormulaNaoCalculavel("lógico")
        if isinstance(v, (int, float, Decimal)):
            return Decimal(str(v))
        valor, problema = base.converter_dinheiro(v)
        if valor is None:
            raise _FormulaNaoCalculavel("texto")
        return Decimal(valor)


def calcular_formula(grade, i, j, profundidade=0):
    """Valor (Decimal) de uma fórmula simples (`=G4-K4`, `=SOMA(A2:A9)`), ou None se não der para calcular sem inventar.
    Só olha células da mesma aba; função desconhecida, texto, outra aba e divisão por zero voltam None."""
    texto = grade.textos_formulas.get((i, j))
    if not texto:
        return None
    try:
        return _AvaliadorDeFormula(grade, texto, profundidade).avaliar().quantize(Decimal("0.01"))
    except (_FormulaNaoCalculavel, InvalidOperation, ArithmeticError, ValueError):
        return None


# ---------------------------------------------------------------- extração

def col_letra(j):
    from openpyxl.utils import get_column_letter
    return get_column_letter(j + 1)


_ERROS_EXCEL = {"#N/A", "#N/D", "#VALUE!", "#VALOR!", "#REF!", "#DIV/0!", "#NAME?", "#NOME?", "#NULL!", "#NUM!", "#NÚM!"}
_TOTAL = re.compile(r"^(sub ?)?total|^totais|^soma\b|^media\b|^valor total|^criterios?\b|^legenda\b")


def _repete_cabecalho(cabecalho, linha):
    """True se a linha repete o cabeçalho (duas tabelas empilhadas na mesma aba, cada uma com o seu cabeçalho)."""
    textos = [c for c in cabecalho if isinstance(c, str) and c.strip()]
    iguais = sum(1 for a, b in zip(cabecalho, linha)
                 if isinstance(a, str) and isinstance(b, str) and a.strip() and base.chave(a) == base.chave(b))
    return iguais >= max(2, int(0.6 * len(textos)))


def _amostra(grade, i0, j, n=3):
    saida = []
    for r in range(i0, len(grade.linhas)):
        v = grade.celula(r, j)
        if not base.vazio(v):
            t = base.limpar_texto(v.isoformat() if hasattr(v, "isoformat") else v)
            saida.append(t[:80])
            if len(saida) == n:
                break
    return saida


def _traduzir(avisos_conv, onde):
    return [base.aviso(n, c, onde, m, cand) for n, c, m, cand in avisos_conv]


def extrair_processos(grade, indice_cab, registros, cliente_padrao=None, avisos=None):
    """Lê as linhas abaixo do cabeçalho. -> (processos, colunas_sem_destino)
    `registros`: mapeamento (de `propor_mapeamento`, já ajustado). Os avisos vão para a lista `avisos`."""
    avisos = avisos if avisos is not None else []
    nome = grade.nome
    onde_aba = f"aba {nome!r}"
    aplicados = [r for r in registros if r["campo"] and r["aplicado"]]
    por_coluna = {r["indice"]: r["campo"] for r in aplicados}
    col_num = next((j for j, c in por_coluna.items() if c == "numero"), None)
    processos, ignoradas = [], []
    lidas = []             # linhas (índices) que viraram processo
    com_erro = []          # células com valor de erro do Excel
    calculadas = []        # células de fórmula sem valor guardado que o programa calculou ("L12")
    resolvidas = set()     # (linha, coluna) das fórmulas calculadas
    rodape = False         # depois de uma linha de total ou de um quadro de critérios, texto sem número é rodapé, não erro
    i0 = indice_cab + 1
    for i in range(i0, len(grade.linhas)):
        linha = list(grade.linhas[i])
        numero_linha = i + 1
        for j in por_coluna:               # fórmula sem valor guardado (=G4-K4): calcula quando é conta simples
            if (i, j) in grade.formulas and j < len(linha) and base.vazio(linha[j]) and ficha.CAMPOS.get(
                    por_coluna[j], ("", "", ""))[2] in ("dinheiro", "numero"):
                calculado = calcular_formula(grade, i, j)
                if calculado is not None:
                    linha[j] = calculado
                    resolvidas.add((i, j))
        nao_vazias = [j for j, v in enumerate(linha) if not base.vazio(v)]
        if not nao_vazias:
            continue
        onde_linha = f"{onde_aba}, linha {numero_linha}"
        bruto_num = grade.celula(i, col_num)
        texto_num = "" if base.vazio(bruto_num) else base.limpar_texto(bruto_num)
        primeira = next((base.limpar_texto(linha[j]) for j in nao_vazias), "")
        if not base.achar_numeros(texto_num) and _repete_cabecalho(grade.linhas[indice_cab], linha):
            ignoradas.append(f"linha {numero_linha}: cabeçalho repetido")
            continue
        if not texto_num or not base.achar_numeros(texto_num):
            eh_total = any(_TOTAL.match(base.chave(t)) for t in (texto_num, primeira) if t)
            if eh_total:
                rodape = True
            if eh_total or len(nao_vazias) <= 2:
                ignoradas.append(f"linha {numero_linha}: {(texto_num or primeira)[:50]}")
            elif rodape:                    # quadro de critérios / legenda abaixo dos totais: não é processo nem erro
                ignoradas.append(f"linha {numero_linha}: rodapé ({(texto_num or primeira)[:40]})")
            elif not texto_num:
                avisos.append(base.aviso("atencao", "linha_sem_numero", onde_linha,
                                         "A linha tem dados mas a coluna do número do processo está vazia; linha não lida.", []))
            else:
                avisos.append(base.aviso("erro", "numero_invalido", onde_linha,
                                         f"{texto_num[:60]!r} não tem o formato de número de processo (NNNNNNN-DD.AAAA.J.TR.OOOO); linha não lida.", [texto_num[:60]]))
            continue
        principal, vinculados, av_num = base.interpretar_numeros(texto_num, onde_linha)
        avisos.extend(av_num)
        if principal is None:
            continue
        rodape = False                      # voltou a ter processo: o que veio antes era total de bloco, não o fim da tabela
        campos, extras = {}, {}
        aba_de_encerrados = bool(re.search(r"arquivad|encerrad|baixad", base.chave(nome or "")))     # aba "Arquivados": processo inativo
        andamentos_texto, ultimo_texto, fecho, lista_andamentos = "", None, None, []
        for j in nao_vazias:
            destino = por_coluna.get(j)
            if destino is None or destino == "numero":
                continue
            valor = linha[j]
            e_formula = (i, j) in grade.formulas
            if isinstance(valor, str) and valor.strip().upper() in _ERROS_EXCEL:
                com_erro.append(f"{col_letra(j)}{numero_linha}")
                continue
            onde_cel = f"{onde_aba}, célula {col_letra(j)}{numero_linha}"
            if destino == "tribunal":
                continue
            if destino == "ativo":
                v, problema = base.converter_sim_nao(valor)
                if v is not None:
                    extras["ativo"] = v == "Sim"
                continue
            if destino == "andamentos":
                analise = base.analisar_andamentos(valor if isinstance(valor, str) else str(valor))
                andamentos_texto, ultimo_texto, fecho = analise["texto"], analise["ultimo"], analise["fecho"]
                lista_andamentos = [{"data": a["data"], "texto": a["texto"]} for a in analise["andamentos"]]
                avisos.extend(base.aviso(n, c, onde_cel, m) for n, c, m in analise["avisos"])
                continue
            conv = base.converter_campo(destino, valor, grade.formatos.get((i, j)), grade.epoch)
            if destino == "situacao" and conv["valor"] is None and not base.vazio(valor):
                alt = base.converter_campo("momento_atual", valor)
                if alt["valor"] is not None and "momento_atual" not in campos:
                    campos["momento_atual"] = base.montar_campo("momento_atual", alt["valor"], "migrado")
                    avisos.append(base.aviso("info", "situacao_era_momento", onde_cel,
                                             f"A coluna de situação traz {base.limpar_texto(valor)!r}, que é um momento atual; "
                                             "lido como momento atual do processo.", [alt["valor"]]))
                    continue
            avisos.extend(_traduzir(conv["avisos"], onde_cel))
            extras.update(conv["extras"])
            if (i, j) in resolvidas:
                calculadas.append(f"{col_letra(j)}{numero_linha}")
            if conv["valor"] is not None:
                campos[destino] = base.montar_campo(destino, conv["valor"], base.origem_do_campo(destino, e_formula))
        justificativa = extras.pop("justificativa_probabilidade", None)      # "REMOTA - texto": o texto vai para o seu campo
        if justificativa and "justificativa_probabilidade" not in campos:
            campos["justificativa_probabilidade"] = base.montar_campo("justificativa_probabilidade", justificativa, "migrado")
        lidas.append(i)
        ultimo = ultimo_texto
        if ultimo is None and "ultimo_andamento" in campos:
            ultimo = campos["ultimo_andamento"]["valor"]
        if cliente_padrao and "cliente" not in campos:
            campos["cliente"] = base.montar_campo("cliente", cliente_padrao, "migrado")
        processos.append(base.processo_lido(
            principal, campos, vinculados, andamentos_texto, ultimo,
            f"aba {nome!r}, linha {numero_linha}" if nome else f"linha {numero_linha}",
            ativo=extras.get("ativo") if "ativo" in extras else (False if aba_de_encerrados else None),
            fecho=fecho, andamentos=lista_andamentos,
            momento_qualificador=extras.get("momento_qualificador")))
    if ignoradas:
        avisos.append(base.aviso("info", "linhas_ignoradas", onde_aba,
                                 f"{len(ignoradas)} linha(s) ignorada(s) por serem marcadores ou totais.", ignoradas[:20]))
    # fórmulas sem valor em cache: só avisa quando NENHUMA célula de fórmula da coluna tem valor guardado (fórmula que
    # devolve "" de vez em quando é normal; arquivo salvo sem cache deixa a coluna toda vazia)
    faltando = {}
    for j in por_coluna:
        formulas_col = [i for i in lidas if (i, j) in grade.formulas]
        vazias = [i for i in formulas_col if base.vazio(grade.celula(i, j)) and (i, j) not in resolvidas]
        if formulas_col and len(vazias) == len(formulas_col):
            faltando[j] = [f"{col_letra(j)}{i + 1}" for i in vazias]
    for j, celulas in faltando.items():
        destino = por_coluna[j]
        derivada = destino in DERIVADOS
        avisos.append(base.aviso("info" if derivada else "atencao", "formula_sem_valor", f"{onde_aba}, coluna {col_letra(j)}",
                                 f"{len(celulas)} célula(s) de fórmula da coluna {grade.celula(indice_cab, j)!r} não têm valor "
                                 "calculado guardado no arquivo (salvo sem cache); foram lidas como vazias."
                                 + (" É uma coluna calculada: o valor se refaz ao abrir no Excel." if derivada else ""),
                                 celulas[:10]))
    if calculadas:
        avisos.append(base.aviso("info", "formula_calculada", onde_aba,
                                 f"{len(calculadas)} célula(s) de fórmula não tinham valor guardado no arquivo; o programa fez a conta "
                                 "(soma, subtração, multiplicação ou divisão simples). Confira os valores.", calculadas[:10]))
    if com_erro:
        avisos.append(base.aviso("atencao", "celula_com_erro", onde_aba,
                                 f"{len(com_erro)} célula(s) com erro do Excel (#N/D, #REF!...) foram lidas como vazias.", com_erro[:10]))
    # colunas que ficaram sem destino
    sem_destino = []
    for r in registros:
        if r["campo"] and r["aplicado"]:
            continue
        if r["campo"] in ESPECIAIS and r["aplicado"]:
            continue
        amostra = _amostra(grade, i0, r["indice"])
        if not amostra:
            continue
        motivo = ("ambigua" if r["ambigua"] else "duplicada" if r["duplicada_de"] is not None else
                  "baixa_confianca" if r["campo"] else "sem_campo_equivalente")
        sem_destino.append({"coluna": r["coluna"], "amostra": amostra, "aba": nome, "motivo": motivo,
                            "candidatos": r["candidatos"]})
        onde = f"{onde_aba}, coluna {col_letra(r['indice'])}"
        if r["ambigua"]:
            avisos.append(base.aviso("atencao", "coluna_ambigua", onde,
                                     f"A coluna {r['coluna']!r} poderia ser {' ou '.join(rotulo_do_destino(c) for c in r['candidatos'][:2])}; "
                                     "não foi lida. Indique o destino na tela de mapeamento.", r["candidatos"]))
        elif r["duplicada_de"] is not None:
            avisos.append(base.aviso("atencao", "coluna_duplicada", onde,
                                     f"A coluna {r['coluna']!r} repete um campo já lido pela coluna {col_letra(r['duplicada_de'])}; só a primeira foi lida.", r["candidatos"]))
        elif r["campo"]:
            avisos.append(base.aviso("atencao", "coluna_baixa_confianca", onde,
                                     f"A coluna {r['coluna']!r} lembra {rotulo_do_destino(r['campo'])} (confiança {r['confianca']:.0%}), "
                                     "mas não foi lida sem confirmação.", [r["campo"]]))
    return processos, sem_destino


def registros_publicos(grade, registros, indice_cab=0):
    """Mapeamento no formato exposto em RelatorioLido["mapeamento"]. Chave aditiva `comentario`: por que o destino foi
    escolhido, quando isso merece explicação (por exemplo, "Natureza da ação" lida como Assunto, não como Classe)."""
    return [{"aba": grade.nome, "coluna": r["coluna"], "campo": r["campo"], "rotulo_campo": rotulo_do_destino(r["campo"]) if r["campo"] else None,
             "confianca": r["confianca"], "aplicado": r["aplicado"], "candidatos": r["candidatos"],
             "ambigua": r["ambigua"], "amostra": _amostra(grade, indice_cab + 1, r["indice"]), "indice": r["indice"],
             "comentario": _comentario_do_destino(r)} for r in registros]


def _comentario_do_destino(reg):
    campo = reg.get("campo")
    if not campo or campo not in COMENTARIOS_DO_DESTINO:
        return ""
    if campo == "assunto" and not base.chave(reg["coluna"]).startswith("natureza"):
        return ""
    if campo == "percentual_provisao" and not _so_percentual(reg["coluna"]):
        return ""
    return COMENTARIOS_DO_DESTINO[campo]


def carregar(caminho):
    """Grades de um .xlsx ou de um CSV, pelo conteúdo (zip = planilha). -> (grades, avisos)"""
    with open(caminho, "rb") as f:
        cabeca = f.read(4)
    if cabeca[:2] == b"PK":
        return carregar_xlsx(caminho)
    return carregar_csv(caminho)


def ler_grades(grades, mapeamento=None, cliente_padrao=None, avisos=None, candidatas=None, lidas_de_outro_modo=()):
    """Processa todas as grades que têm cabeçalho e números. -> (processos, sem_destino, mapeamento_publico, avisos)
    `candidatas`: resultado de `grades_de_processos` se já calculado; `lidas_de_outro_modo`: abas (ids) que o chamador
    leu por outro caminho (parâmetros) e que não merecem o aviso de aba ignorada."""
    avisos = avisos if avisos is not None else []
    processos, sem_destino, publico = [], [], []
    if candidatas is None:
        candidatas = grades_de_processos(grades)
    usadas = {id(g) for g, _, _ in candidatas} | set(lidas_de_outro_modo)
    for g in grades:
        if id(g) not in usadas:
            avisos.append(base.aviso("info", "aba_ignorada", f"aba {g.nome!r}",
                                     "Sem cabeçalho com coluna de número de processo (ou sem números abaixo); aba não lida como lista de processos."))
    for g, i, regs in candidatas:
        regs = aplicar_mapeamento_do_usuario(regs, mapeamento, avisos, f"aba {g.nome!r}")
        publico.extend(registros_publicos(g, regs, i))
        procs, sd = extrair_processos(g, i, regs, cliente_padrao, avisos)
        processos.extend(procs)
        sem_destino.extend(sd)
    return processos, sem_destino, publico, avisos
