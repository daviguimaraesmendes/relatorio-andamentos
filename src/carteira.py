"""Carteira de processos do relatório: importa listas e mantém carteira.json.

A lista pode vir em .csv, .xlsx, .txt/.md ou colada no terminal:
  - planilha com cabeçalho: coluna de processo ('processo', 'numero', 'nº'...)
    e, se houver, 'cliente', 'polo', 'parte contrária', 'responsável', 'contato';
  - qualquer outro texto: todo número no padrão CNJ encontrado é importado
    (o cliente vem de --cliente).

Cada número tem o dígito verificador conferido (número com erro de digitação
é recusado) e o tribunal é deduzido do próprio número. Com --completar-djen,
o polo do cliente e a parte contrária são preenchidos a partir das
publicações do DJEN (quando o processo tiver alguma).

Uso:
    python carteira.py importar lista.xlsx
    python carteira.py importar lista.txt --cliente "RAZÃO SOCIAL LTDA" --completar-djen
    pbpaste | python carteira.py importar - --cliente "RAZÃO SOCIAL LTDA"
    python carteira.py listar
    python carteira.py completar-djen        # tenta preencher polo/parte contrária que faltam
"""
import csv
import datetime
import io
import re
import sys
from pathlib import Path

import comum
from comum import load_json, normalizar, save_json
CNJ = re.compile(r"(?<!\d)(\d{7})-?(\d{2})\.?(\d{4})\.?(\d)\.?(\d{2})\.?(\d{4})(?!\d)")

# Código do tribunal (TR) na Justiça Estadual, na ordem da Resolução CNJ que
# criou a numeração única (ordem alfabética do nome do estado).
UFS = ["AC", "AL", "AP", "AM", "BA", "CE", "DF", "ES", "GO", "MA", "MT", "MS", "MG", "PA", "PB",
       "PR", "PE", "PI", "RJ", "RN", "RS", "RO", "RR", "SC", "SE", "SP", "TO"]
POLOS = {"a": "ativo", "ativo": "ativo", "autor": "ativo", "autora": "ativo", "requerente": "ativo",
         "exequente": "ativo", "impetrante": "ativo", "reclamante": "ativo", "apelante": "ativo",
         "p": "passivo", "passivo": "passivo", "reu": "passivo", "re": "passivo", "requerido": "passivo",
         "requerida": "passivo", "executado": "passivo", "executada": "passivo", "impetrado": "passivo",
         "reclamado": "passivo", "reclamada": "passivo", "apelado": "passivo", "apelada": "passivo"}
SINONIMOS = {
    "numero": ("processo", "numero", "n", "no", "numero do processo", "n do processo", "autos"),
    "cliente": ("cliente", "parte", "nosso cliente", "razao social"),
    "polo_cliente": ("polo", "polo do cliente", "posicao"),
    "parte_contraria": ("parte contraria", "contraria", "adverso", "parte adversa", "contraparte"),
    "responsavel": ("responsavel", "advogado", "advogado responsavel"),
    "contato": ("contato", "email", "e-mail", "whatsapp", "telefone"),
}


def gravar_plano(p, campo, valor, origem="humano"):
    """Grava um campo "plano" (cliente, polo_cliente, parte_contraria, responsavel...) num item da carteira mantendo em
    sincronia a chave plana (usada pelo código da Fase 1) e o registro v2 (`campos`, que a ficha lê primeiro).
    Sem isso, uma edição na tela de cadastro ficaria escondida atrás do valor antigo de `campos`."""
    p[campo] = valor
    campos = p.get("campos")
    if isinstance(campos, dict):
        if valor in (None, ""):
            campos.pop(campo, None)
        else:
            campos[campo] = {"valor": valor, "origem": origem, "em": datetime.datetime.now().isoformat(timespec="seconds")}


def mascara(m):
    return f"{m[0]}-{m[1]}.{m[2]}.{m[3]}.{m[4]}.{m[5]}"


def dv_correto(m):
    return int(m[1]) == 98 - int(m[0] + m[2] + m[3] + m[4] + m[5] + "00") % 97


def tribunal(m):
    j, tr = m[3], int(m[4])
    if j == "8" and 1 <= tr <= 27:
        return f"TJ{UFS[tr - 1]}"
    return {"4": f"TRF{tr}", "5": f"TRT{tr}", "1": "STF", "3": "STJ", "6": f"TRE{tr}", "9": f"TJM{tr}"}.get(j, f"J{j}.{m[4]}")


def numeros_no_texto(texto):
    """[(numero_mascarado, tribunal)] + lista de números com dígito errado."""
    validos, invalidos, vistos = [], [], set()
    for m in CNJ.findall(texto):
        num = mascara(m)
        if num in vistos:
            continue
        vistos.add(num)
        (validos if dv_correto(m) else invalidos).append((num, tribunal(m)))
    return validos, [n for n, _ in invalidos]


def _linhas_tabela(caminho):
    if caminho.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        ws = load_workbook(caminho, read_only=True, data_only=True).active
        return [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]
    texto = caminho.read_text(encoding="utf-8-sig")
    dialeto = csv.Sniffer().sniff(texto[:2000], delimiters=",;\t")
    return list(csv.reader(io.StringIO(texto), dialeto))


def _mapear_cabecalho(cabecalho):
    mapa = {}
    for i, nome in enumerate(cabecalho):
        n = re.sub(r"[^a-z ]", "", normalizar(nome)).strip()
        for campo, sinonimos in SINONIMOS.items():
            if n in sinonimos and campo not in mapa:
                mapa[campo] = i
    return mapa if "numero" in mapa else None


def ler_lista(origem, cliente_padrao=""):
    """Volta (registros, invalidos). Cada registro: numero, tribunal e o que
    mais a lista trouxer."""
    if origem == "-":
        texto, caminho = sys.stdin.read(), None
    else:
        caminho = Path(origem)
        texto = None if caminho.suffix.lower() in (".xlsx", ".csv") else caminho.read_text(encoding="utf-8")
    registros, invalidos = [], []
    if caminho is not None and texto is None:
        linhas = _linhas_tabela(caminho)
        mapa = _mapear_cabecalho(linhas[0]) if linhas else None
        if mapa:
            for linha in linhas[1:]:
                val = lambda campo: (linha[mapa[campo]].strip() if campo in mapa and mapa[campo] < len(linha) else "")
                validos, ruins = numeros_no_texto(val("numero"))
                invalidos += ruins
                for num, trib in validos:
                    polo = POLOS.get(normalizar(val("polo_cliente")), "")
                    registros.append({"numero": num, "tribunal": trib, "cliente": val("cliente") or cliente_padrao,
                                      "polo_cliente": polo, "parte_contraria": val("parte_contraria"),
                                      "responsavel": val("responsavel"), "contato": val("contato")})
            return registros, invalidos
        texto = "\n".join(" ".join(l) for l in linhas)
    validos, invalidos = numeros_no_texto(texto)
    return [{"numero": n, "tribunal": t, "cliente": cliente_padrao} for n, t in validos], invalidos


def mesclar(registros, origem):
    """Inclui os novos; nos já existentes, só preenche campo vazio (nunca
    sobrescreve o que foi ajustado à mão)."""
    carteira = load_json(comum.CARTEIRA_FILE, [])
    por_numero = {p["numero"]: p for p in carteira}
    novos = 0
    hoje = datetime.date.today().isoformat()
    for r in registros:
        atual = por_numero.get(r["numero"])
        if atual is None:
            base = {"numero": r["numero"], "tribunal": r["tribunal"], "cliente": "", "polo_cliente": "",
                    "parte_contraria": "", "responsavel": "", "contato": "", "ativo": True,
                    "origem": origem, "incluido_em": hoje}
            base.update({k: v for k, v in r.items() if v})
            carteira.append(base)
            por_numero[r["numero"]] = base
            novos += 1
        else:
            for k, v in r.items():
                if v and not atual.get(k):
                    atual[k] = v
    save_json(comum.CARTEIRA_FILE, carteira)
    return novos


def clientes():
    return load_json(comum.CLIENTES_FILE, {}).get("clientes", [])


def _sem_sufixo(nome):
    n = re.sub(r"[^a-z0-9]+", " ", normalizar(nome))
    return re.sub(r"\b(ltda|s a|sa|eireli|me|epp)\b", " ", n).split()


def nome_bate(nome_cliente, nome_publicado):
    """Mesmo nome, tolerando sufixo societário, pontuação e o corte de nome
    longo que o DJEN faz ('SINDICATO TRABALHADORES TRANSPORTES RODOVI')."""
    a, b = " ".join(_sem_sufixo(nome_cliente)), " ".join(_sem_sufixo(nome_publicado))
    if not a or not b:
        return False
    return a == b or (len(b) >= 15 and a.startswith(b)) or (len(a) >= 15 and b.startswith(a))


def variacoes_do_cliente(nome):
    for c in clientes():
        if normalizar(c["nome"]) == normalizar(nome):
            return [c["nome"], *c.get("variacoes", [])]
    return [nome] if nome else []


def completar_com_djen(apenas=None):
    """Polo do cliente e parte contrária a partir dos destinatários das
    publicações do DJEN. Só preenche o que está vazio."""
    import djen
    carteira = load_json(comum.CARTEIRA_FILE, [])
    feitos = sem_cliente = 0
    for p in carteira:
        if (apenas and p["numero"] not in apenas) or (p.get("polo_cliente") and p.get("parte_contraria")):
            continue
        if not p.get("cliente"):
            sem_cliente += 1  # sem saber quem é o cliente, não dá para achar o polo dele
            continue
        polos = djen.partes_do_processo(p["numero"])
        nomes = variacoes_do_cliente(p["cliente"])
        do_cliente = {polo for nome, polo in polos if any(nome_bate(v, nome) for v in nomes)}
        if len(do_cliente) != 1:
            continue  # cliente não aparece, ou aparece nos dois polos: deixa para preencher à mão
        polo = do_cliente.pop()
        contrarios = sorted({nome for nome, pl in polos if pl and pl != polo})
        if not p.get("polo_cliente"):
            gravar_plano(p, "polo_cliente", {"A": "ativo", "P": "passivo"}.get(polo, ""), "coletado")
        if not p.get("parte_contraria") and contrarios:
            gravar_plano(p, "parte_contraria", "; ".join(contrarios), "coletado")
        feitos += 1
    save_json(comum.CARTEIRA_FILE, carteira)
    print(f"{feitos} processo(s) completado(s) com dados do DJEN.")
    return feitos, sem_cliente


def listar():
    carteira = load_json(comum.CARTEIRA_FILE, [])
    for p in carteira:
        falta = [c for c in ("cliente", "polo_cliente", "parte_contraria") if not p.get(c)]
        print(f"{p['numero']}  {p['tribunal']:6}  {p.get('cliente') or '?':30.30}  {p.get('polo_cliente') or '?':8}"
              + (f"  falta: {', '.join(falta)}" if falta else ""))
    print(f"{len(carteira)} processo(s).")


if __name__ == "__main__":
    args = sys.argv[1:]
    if args[:1] == ["importar"] and len(args) > 1:
        cliente = args[args.index("--cliente") + 1] if "--cliente" in args else ""
        regs, ruins = ler_lista(args[1], cliente)
        n = mesclar(regs, "lista")
        print(f"{len(regs)} número(s) lido(s), {n} novo(s) na carteira.")
        if ruins:
            print("Recusados (dígito verificador não confere, provável erro de digitação):", *ruins, sep="\n  ")
        sem_cliente = [r["numero"] for r in regs if not r.get("cliente")]
        if sem_cliente:
            print(f"Atenção: {len(sem_cliente)} processo(s) sem cliente. Use --cliente ou a coluna 'cliente'.")
        if "--completar-djen" in args:
            completar_com_djen({r["numero"] for r in regs})
    elif args[:1] == ["completar-djen"]:
        completar_com_djen()
    elif args[:1] == ["listar"]:
        listar()
    else:
        print(__doc__)
