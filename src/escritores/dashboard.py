"""Gerador de dashboards HTML (modelo C): um arquivo .html autônomo e offline.

    import escritores.dashboard as dashboard
    r = dashboard.gravar(Path("relatorio.xlsx"), Path("painel.html"), perfil, modo="embutido",
                         template="contencioso", data_base="2026-10-07")
    r["destino"], r["avisos"]

Dois modos (mesmo código de leitura e de cálculo no navegador; mesmo .xlsx, mesmos números):

    modelo     (padrão) a página abre vazia e o usuário ARRASTA o .xlsx para ela. O .xlsx passado a
               `gravar` não é usado (pode ser None): quem lê a planilha é o navegador, com o SheetJS embutido.
    embutido   a página abre já preenchida: este módulo lê o .xlsx (openpyxl, valores em cache) e grava a
               grade das abas dentro do HTML. Sem SheetJS na página (só o Chart.js); sem botão de arquivo.
               Por padrão o texto da coluna "Andamentos" NÃO vai no HTML (só a data do último andamento
               que ele traz), para a página enviada ao cliente não carregar o histórico inteiro:
               `incluir_andamentos=True` para embutir o texto.

Dois templates (src/modelos/dashboard/): "contencioso" (KPIs, roscas e barras, desfechos, economia com as
regras de integridade, causa-raiz, série histórica, tabela) e "carteira" (momento atual, área, matéria,
tribunal e prazos de último andamento, no estilo do relatório em texto). Sem `template`: "contencioso" se o
perfil lista `empresas_do_grupo`, senão "carteira".

Opções de `gravar` (todas opcionais):
    modo, template, titulo, cliente, data_base (ISO ou DD/MM/AAAA; data de referência dos prazos e dos 12
    meses; sem ela vale a aba Parâmetros da planilha e, por fim, a data de hoje), empresas_do_grupo (substitui
    a do perfil), historico (lista de retratos mensais, CONTRATOS §9) ou pasta_historico (pasta com os
    AAAA-MM-DD.json), incluir_andamentos.

O HTML final não tem CDN, fonte externa, imagem remota nem nenhuma requisição de rede: SheetJS e Chart.js
vão dentro da página (ver modelos/dashboard/BIBLIOTECAS.md, com versões e hashes). Erro esperado (planilha
ilegível, bibliotecas ausentes, destino igual à origem) volta como aviso de nível "erro" e `destino: None`,
sem arquivo; exceção só para erro de programação.

Colunas reconhecidas (por similaridade de cabeçalho, sem acento nem caixa): as 29 colunas do modelo B
("Número do Processo", "Valor da Causa", ...) e os rótulos de `ficha.CAMPOS` ("Matéria principal",
"Momento atual do processo", "Valor do acordo", ...). Abas: as que têm cabeçalho com o número do processo
(pelo cabeçalho, não pelo nome), "Parâmetros" (empresas do grupo, data de referência, headcount) e
"Histórico" (uma linha por data-base, com colunas de totais).

Contingência: quando a planilha tem as colunas Passivo Potencial, Provisão Constituída, Ativo Potencial, Pagamento Realizado ou
Depósito Judicial Realizado?, os dois modelos de painel ganham o bloco "Contingência e provisão" (cartões de passivo potencial dos
processos ativos, provisão e cobertura, perda provável, ativo potencial, depósitos e pagamentos, e o passivo e a provisão por
probabilidade Provável/Possível/Remota). O passivo potencial lançado passa a ser a exposição do processo. Sem esses dados o bloco
fica oculto e nada muda nos números de sempre.
"""
import datetime
import hashlib
import json
import re
import sys
import unicodedata
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))  # src/ (taxonomia, comum)
import taxonomia  # noqa: E402

PASTA_MODELOS = Path(__file__).resolve().parent.parent / "modelos" / "dashboard"
VERSAO_MODELO = 1

# arquivo -> (nome, versão, SHA-256). Atualize junto com modelos/dashboard/BIBLIOTECAS.md.
BIBLIOTECAS = {
    "xlsx": ("xlsx.full.min.js", "SheetJS 0.18.5", "c9506197caf809a075b6dee1da0d36fb19da7158ffe8a88e7b0c96c5d8623c99"),
    "chart": ("chart.umd.min.js", "Chart.js 4.4.1", "81ffafe13c37e1b25793b020d446f4d9739b949dadb7f9f79d709a0cad781c2f"),
}
TEMPLATES = {
    "contencioso": {"titulo": "Painel do contencioso", "sobretitulo": "Contencioso · painel executivo", "js": "contencioso.js"},
    "carteira": {"titulo": "Painel da carteira de processos", "sobretitulo": "Carteira · acompanhamento processual", "js": "carteira.js"},
}
MODOS = ("modelo", "embutido")


def _aviso(nivel, codigo, onde, mensagem, candidatos=None):
    return {"nivel": nivel, "codigo": codigo, "onde": onde, "mensagem": mensagem, "candidatos": candidatos or []}


def _resultado(destino=None, avisos=None, **extra):
    r = {"destino": destino, "processos_atualizados": [], "processos_novos": [], "ignorados": [], "mudancas": [],
         "textos_gravados": {}, "avisos": avisos or []}
    r.update(extra)
    return r


def _norm(texto):
    s = unicodedata.normalize("NFD", str(texto or ""))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", re.sub(r"[ºª°]", "", s)).strip().lower()


# ---------------------------------------------------------------- bibliotecas e modelos

def verificar_bibliotecas(pasta=PASTA_MODELOS, so=None):
    """Avisos sobre as bibliotecas embutidas: arquivo ausente (erro) ou com hash diferente do registrado (atenção)."""
    avisos = []
    for chave, (nome, versao, sha) in BIBLIOTECAS.items():
        if so and chave not in so:
            continue
        arq = Path(pasta) / nome
        if not arq.is_file():
            avisos.append(_aviso("erro", "bibliotecas_ausentes", nome,
                                 f"A biblioteca {versao} ({nome}) não está em {pasta}. O painel não usa CDN: coloque o "
                                 "arquivo nessa pasta (versão e hash em BIBLIOTECAS.md) e gere de novo."))
        elif hashlib.sha256(arq.read_bytes()).hexdigest() != sha:
            avisos.append(_aviso("atencao", "biblioteca_diferente", nome,
                                 f"{nome} não confere com o hash registrado em BIBLIOTECAS.md ({versao}). "
                                 "Se a troca foi de propósito, atualize o registro."))
    return avisos


def _ler_modelo(nome):
    return (PASTA_MODELOS / nome).read_text(encoding="utf-8")


# ---------------------------------------------------------------- leitura do .xlsx (modo embutido)

_DATA_NO_TEXTO = re.compile(r"\d{1,2}/\d{1,2}/\d{4}")


def _ultima_data_do_texto(texto):
    """Maior DD/MM/AAAA válida do texto (mesma regra do navegador) ou None."""
    melhor = None
    for m in _DATA_NO_TEXTO.findall(str(texto or "")):
        d, mes, a = (int(x) for x in m.split("/"))
        try:
            dt = datetime.date(a, mes, d)
        except ValueError:
            continue
        if 1900 < a < 2200 and (melhor is None or dt > melhor):
            melhor = dt
    return melhor


def _celula(v):
    from openpyxl.utils.datetime import to_excel
    if v is None or isinstance(v, bool):
        return v
    if isinstance(v, (datetime.datetime, datetime.date)):
        return float(to_excel(v))
    if isinstance(v, Decimal):
        v = float(v)
    if isinstance(v, (int, float)):
        return None if v != v or v in (float("inf"), float("-inf")) else v
    if isinstance(v, str):
        return v
    return str(v)


def grade_do_xlsx(caminho, incluir_andamentos=False):
    """Grade das abas, no formato que a página lê: [{"nome", "linhas": [[célula, ...], ...]}].

    Valores em cache (data_only); fórmula sem cache vira vazio, como no SheetJS. Datas viram número de série
    do Excel. Sem `incluir_andamentos`, a coluna "Andamentos" é trocada por "Em DD/MM/AAAA." (só a data do
    último andamento que o texto traz)."""
    from openpyxl import load_workbook
    wb = load_workbook(str(caminho), data_only=True, read_only=True)
    grade = []
    try:
        for ws in wb.worksheets:
            linhas = [[_celula(c) for c in lin] for lin in ws.iter_rows(values_only=True)]
            if not incluir_andamentos:
                _minimizar_andamentos(linhas)
            for lin in linhas:
                while lin and lin[-1] in (None, ""):
                    lin.pop()
            while linhas and not linhas[-1]:
                linhas.pop()
            grade.append({"nome": ws.title, "linhas": linhas})
    finally:
        wb.close()
    return grade


def _minimizar_andamentos(linhas):
    for i, lin in enumerate(linhas[:20]):
        cols = [j for j, c in enumerate(lin) if isinstance(c, str) and _norm(c) == "andamentos"]
        if not cols:
            continue
        for outra in linhas[i + 1:]:
            for j in cols:
                if j < len(outra) and outra[j] not in (None, ""):
                    d = _ultima_data_do_texto(outra[j])
                    outra[j] = f"Em {d:%d/%m/%Y}." if d else None
        return


_ROTULOS_NUMERO = {"numero do processo", "n do processo", "no do processo", "numero", "processo", "autos"}
_ROTULOS_CONHECIDOS = {"autor(es)", "reu(s)", "vara", "municipio", "tribunal", "valor da causa", "area do direito", "materia principal",
                       "momento atual do processo", "situacao", "resultado", "valor estimado", "probabilidade", "ativo", "objeto",
                       "data do ajuizamento", "cliente", "polo do cliente"}


def _linha_marcador(numero, linha):
    """Total, subtotal, "não alterar" ou contador de rodapé (mesma regra do navegador): não é processo."""
    if any(isinstance(c, str) and (_norm(c).startswith(("total", "subtotal")) or "nao alterar" in _norm(c)) for c in linha):
        return True
    s = str(numero if numero is not None else "").strip()
    return bool(re.fullmatch(r"\d+([.,]\d+)?", s)) and len(re.sub(r"\D", "", s)) < 7


def _contar_processos(grade):
    """Quantos processos distintos a grade tem (aproximação do que a página reconhece), para o aviso de planilha vazia."""
    vistos = set()
    for aba in grade:
        for i, lin in enumerate(aba["linhas"][:20]):
            rot = [_norm(c) for c in lin if isinstance(c, str)]
            col = next((j for j, c in enumerate(lin) if isinstance(c, str) and _norm(c) in _ROTULOS_NUMERO), None)
            if col is None or sum(1 for r in rot if r in _ROTULOS_CONHECIDOS) < 2:
                continue
            for outra in aba["linhas"][i + 1:]:
                v = outra[col] if col < len(outra) else None
                if _linha_marcador(v, outra):
                    continue
                if v not in (None, "") and re.search(r"\d", str(v)):
                    vistos.add(re.sub(r"\D", "", str(v)) if re.search(r"\d{7}", str(v)) else _norm(v))
            break
    return len(vistos)


# ---------------------------------------------------------------- histórico (retratos mensais)

def _ler_retratos(historico, pasta_historico, avisos):
    retratos = list(historico or [])
    if pasta_historico:
        pasta = Path(pasta_historico)
        for arq in sorted(pasta.glob("*.json")) if pasta.is_dir() else []:
            try:
                retratos.append(json.loads(arq.read_text(encoding="utf-8")))
            except (OSError, ValueError):
                avisos.append(_aviso("atencao", "retrato_ilegivel", arq.name, f"Não consegui ler o retrato {arq.name}; ele ficou fora da série histórica."))
    validos = []
    for r in retratos:
        if isinstance(r, dict) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(r.get("data_base", ""))) and isinstance(r.get("totais"), dict):
            validos.append({"data_base": r["data_base"], "totais": {k: r["totais"].get(k) for k in
                            ("processos", "ativos", "encerrados", "valor_causa", "valor_estimado", "valor_economizado")}})
        else:
            avisos.append(_aviso("atencao", "retrato_invalido", str(r.get("data_base", "?")) if isinstance(r, dict) else "?",
                                 "Retrato sem data_base ou sem totais; ficou fora da série histórica."))
    return sorted({r["data_base"]: r for r in validos}.values(), key=lambda r: r["data_base"])


# ---------------------------------------------------------------- montagem do HTML

def _json_para_script(obj):
    """JSON seguro dentro de <script>: sem '</script', sem '<!--' e sem separadores de linha Unicode."""
    return (json.dumps(obj, ensure_ascii=False, separators=(",", ":")).replace("<", "\\u003c")
            .replace("\u2028", "\\u2028").replace("\u2029", "\\u2029"))


def _escapar(texto):
    return str(texto).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


def _data_iso(valor):
    if not valor:
        return None
    s = str(valor).strip()
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", s)
    try:
        if m:
            return datetime.date(int(m[3]), int(m[2]), int(m[1])).isoformat()
        return datetime.date.fromisoformat(s[:10]).isoformat()
    except ValueError:
        return None


def montar_html(modo, template, config, titulo):
    """HTML completo (sem gravar). `config` é o que a página lê em <script id="cfg">."""
    info = TEMPLATES[template]
    bibs = []
    bibs_cfg = ["chart"] if modo == "embutido" else ["xlsx", "chart"]
    for chave in bibs_cfg:
        nome, versao, _ = BIBLIOTECAS[chave]
        codigo = (PASTA_MODELOS / nome).read_text(encoding="utf-8")
        bibs.append(f'<script data-biblioteca="{_escapar(versao)}">\n{codigo}\n</script>')
    valores = {
        "MODO": modo, "TEMPLATE": template, "TITULO": _escapar(titulo), "SOBRETITULO": _escapar(info["sobretitulo"]),
        "ETIQUETA_INICIAL": "Solte a planilha de acompanhamento abaixo para gerar os indicadores." if modo == "modelo" else "",
        "CARREGADOR_OCULTO": "hidden" if modo == "embutido" else "", "TROCAR_OCULTO": "hidden" if modo == "embutido" else "",
        "CSS": _ler_modelo("estilo.css"), "CORPO": '<div id="conteudo"></div>', "CONFIG": _json_para_script(config),
        "BIBLIOTECAS": "\n".join(bibs), "NUCLEO": _ler_modelo("nucleo.js"), "TEMPLATE_JS": _ler_modelo(info["js"]),
    }
    return re.sub(r"\{\{([A-Z_]+)\}\}", lambda m: valores[m.group(1)], _ler_modelo("base.html"))


def gravar(xlsx, destino, perfil, **opcoes):
    """Gera o dashboard em `destino` (.html). Ver o cabeçalho do módulo para as opções. Devolve o Resultado do
    CONTRATOS §5, com os campos extras `modo`, `template`, `processos_no_painel` (só no modo embutido) e `bytes`."""
    perfil = perfil or {}
    avisos = []
    modo = opcoes.get("modo", "modelo")
    parametros = dict(perfil.get("parametros") or {})
    grupo = list(opcoes.get("empresas_do_grupo") if opcoes.get("empresas_do_grupo") is not None else parametros.get("empresas_do_grupo") or [])
    template = opcoes.get("template") or ("contencioso" if grupo else "carteira")
    base = _resultado(modo=modo, template=template, processos_no_painel=None, bytes=0)
    if modo not in MODOS:
        avisos.append(_aviso("erro", "modo_invalido", "modo", f"Modo {modo!r} desconhecido; use um de {', '.join(MODOS)}.", list(MODOS)))
    if template not in TEMPLATES:
        avisos.append(_aviso("erro", "template_invalido", "template", f"Modelo de painel {template!r} desconhecido.", list(TEMPLATES)))
    if avisos:
        return {**base, "avisos": avisos}
    destino = Path(destino)
    if xlsx is not None and Path(xlsx).resolve() == destino.resolve():
        return {**base, "avisos": [_aviso("erro", "destino_igual_origem", str(destino), "O destino não pode ser a própria planilha.")]}
    avisos += verificar_bibliotecas(so=("chart",) if modo == "embutido" else None)
    if any(a["nivel"] == "erro" for a in avisos):
        return {**base, "avisos": avisos}

    dados = None
    if modo == "embutido":
        if xlsx is None or not Path(xlsx).is_file():
            return {**base, "avisos": avisos + [_aviso("erro", "planilha_inexistente", str(xlsx), "O modo embutido precisa da planilha (.xlsx) para ler os dados.")]}
        try:
            grade = grade_do_xlsx(xlsx, bool(opcoes.get("incluir_andamentos")))
        except Exception as e:  # arquivo corrompido, não é .xlsx, protegido por senha...
            return {**base, "avisos": avisos + [_aviso("erro", "planilha_ilegivel", Path(xlsx).name, f"Não consegui abrir a planilha: {e}")]}
        n = _contar_processos(grade)
        base["processos_no_painel"] = n
        if n == 0:
            avisos.append(_aviso("atencao", "planilha_sem_processos", Path(xlsx).name,
                                 "Não achei, em nenhuma aba, o cabeçalho com \"Número do Processo\" e colunas conhecidas; o painel abrirá vazio."))
        if not opcoes.get("incluir_andamentos"):
            avisos.append(_aviso("info", "andamentos_minimizados", "Andamentos",
                                 "O texto da coluna \"Andamentos\" não foi embutido (só a data do último andamento). Use incluir_andamentos=True para embutir."))
        dados = {"sheets": grade}
    elif xlsx is not None and not Path(xlsx).is_file():
        avisos.append(_aviso("atencao", "planilha_inexistente", str(xlsx), "A planilha informada não existe; no modo modelo ela não é usada (o usuário a arrasta para a página)."))

    data_base = _data_iso(opcoes.get("data_base"))
    if opcoes.get("data_base") and not data_base:
        avisos.append(_aviso("atencao", "data_base_invalida", "data_base", f"Data de referência inválida ({opcoes.get('data_base')!r}); a página usará a da planilha ou a de hoje."))
    cliente = opcoes.get("cliente") or ""
    titulo = opcoes.get("titulo") or (f"{TEMPLATES[template]['titulo']} · {cliente}" if cliente else TEMPLATES[template]["titulo"])
    config = {
        "versao": VERSAO_MODELO, "modo": modo, "template": template, "cliente": cliente, "titulo": titulo, "data_base": data_base,
        "empresas_do_grupo": grupo, "parametros": {"headcount": parametros.get("headcount")},
        "materias_fora_do_ranking": [k for k, v in taxonomia.MATERIA.items() if not v[2]],
        "momentos": {k: [v[0], bool(v[1])] for k, v in taxonomia.MOMENTO_ATUAL.items()},
        "historico": _ler_retratos(opcoes.get("historico"), opcoes.get("pasta_historico"), avisos),
        "dados": dados, "fonte": Path(xlsx).name if (dados and xlsx) else "",
    }
    html = montar_html(modo, template, config, titulo)
    destino.parent.mkdir(parents=True, exist_ok=True)
    tmp = destino.with_name(destino.name + ".tmp")
    tmp.write_text(html, encoding="utf-8")
    tmp.replace(destino)
    return {**base, "destino": destino, "avisos": avisos, "bytes": len(html.encode("utf-8"))}
