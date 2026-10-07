"""Testes do SPIKE S1 (gravação cirúrgica em .xlsx). Só dados fictícios, sem rede.

    python3 -m unittest spikes/s1_xlsx/test_s1.py -v

Os testes que abrem o LibreOffice (/usr/bin/soffice) são pulados se ele não existir.
Eles provam que o LibreOffice abre e recalcula; NÃO provam nada sobre Excel ou Google Sheets.
"""
import csv
import datetime
import hashlib
import io
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
import warnings
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import gerar_modelo as g  # noqa: E402
import xlsx_cirurgico as x  # noqa: E402

import openpyxl  # noqa: E402

warnings.filterwarnings("ignore", module="openpyxl")
SOFFICE = shutil.which("soffice") or "/usr/bin/soffice"
TEM_LO = Path(SOFFICE).exists()
M = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def partes(caminho):
    with zipfile.ZipFile(caminho) as z:
        return {i.filename: z.read(i.filename) for i in z.infolist()}


def sha(b):
    return hashlib.sha256(b).hexdigest()


class Base(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp(prefix="s1-"))
        cls.completo = cls.tmp / "completo.xlsx"
        cls.sujo = cls.tmp / "sujo.xlsx"
        cls.pre = cls.tmp / "pre.xlsx"
        g.gerar(cls.completo, "completo")
        g.gerar(cls.sujo, "sujo")
        g.gerar(cls.pre, "preformatado")
        cls.perfil_lo = cls.tmp / "perfil-lo"

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def saida(self, nome):
        return self.tmp / nome

    def lo(self, arquivo, formato, filtro=None):
        """Converte com o LibreOffice; devolve a pasta de saída."""
        pasta = self.tmp / ("lo-" + Path(arquivo).stem + "-" + formato.split(":")[0])
        shutil.rmtree(pasta, ignore_errors=True)
        pasta.mkdir()
        p = subprocess.run(
            [SOFFICE, f"-env:UserInstallation=file://{self.perfil_lo}", "--headless", "--convert-to",
             filtro or formato, "--outdir", str(pasta), str(arquivo)],
            capture_output=True, text=True, timeout=240)
        self.assertEqual(p.returncode, 0, p.stderr)
        return pasta

    def lo_valores(self, arquivo):
        """Abre no LibreOffice, deixa ele salvar de novo como xlsx e lê os valores calculados."""
        pasta = self.lo(arquivo, "xlsx")
        return openpyxl.load_workbook(pasta / Path(arquivo).name, data_only=True)

    @staticmethod
    def linhas_novas(a, b):
        return [g.linha_ficticia(i) for i in range(a, b)]


class TestInserir200(Base):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.novas = cls.linhas_novas(10, 210)
        cls.dest = cls.tmp / "200.xlsx"
        cls.res = x.inserir_linhas(cls.completo, cls.dest, "Processos", cls.novas)

    def test_a_resultado(self):
        r = self.res
        self.assertEqual((r.linhas_inseridas, r.primeira_linha, r.ultima_linha), (200, 12, 211))
        self.assertEqual(r.avisos, [])

    def test_b_partes_nao_editadas_identicas(self):
        antes, depois = partes(self.completo), partes(self.dest)
        removidas = set(antes) - set(depois)
        self.assertEqual(set(depois) - set(antes), set())          # nenhuma parte nova
        self.assertEqual(removidas, {"xl/calcChain.xml"})          # única remoção: calcChain
        alteradas = {n for n in depois if antes[n] != depois[n]}
        # cada parte que mudou, e por quê:
        justificativa = {
            "xl/worksheets/sheet1.xml": "aba Processos: linhas novas, dimension, CF/DV, cache das fórmulas",
            "xl/tables/table1.xml": "ref e autoFilter da Tabela",
            "xl/workbook.xml": "intervalos nomeados e fullCalcOnLoad",
            "xl/_rels/workbook.xml.rels": "relação do calcChain removida",
            "[Content_Types].xml": "tipo do calcChain removido",
            "xl/charts/chart2.xml": "série do gráfico que aponta para a coluna da tabela",
            "xl/worksheets/sheet3.xml": "Indicadores: fórmulas A1 que cobrem a coluna inteira + cache",
        }
        self.assertEqual(alteradas, set(justificativa), f"mudaram: {sorted(alteradas)}")
        self.assertEqual(sorted(alteradas), self.res.partes_alteradas)
        # tudo o mais é idêntico byte a byte: estilos, textos, tema, outras abas, gráfico 1,
        # desenhos, tabela dinâmica (cache e definição), propriedades
        for n in set(antes) - alteradas - removidas:
            self.assertEqual(antes[n], depois[n], n)
        for essencial in ("xl/styles.xml", "xl/sharedStrings.xml", "xl/theme/theme1.xml", "xl/charts/chart1.xml",
                          "xl/drawings/drawing1.xml", "xl/pivotTables/pivotTable1.xml",
                          "xl/pivotCache/pivotCacheDefinition1.xml", "xl/worksheets/sheet2.xml",
                          "xl/worksheets/sheet4.xml", "xl/worksheets/sheet5.xml", "docProps/core.xml"):
            self.assertIn(essencial, depois)
            self.assertEqual(antes[essencial], depois[essencial], essencial)

    def test_b2_linhas_antigas_do_xml_intactas(self):
        """A linha 1 (cabeçalho) e as células que não são fórmula das linhas antigas voltam idênticas."""
        a = partes(self.completo)["xl/worksheets/sheet1.xml"].decode()
        d = partes(self.dest)["xl/worksheets/sheet1.xml"].decode()
        self.assertEqual(re.search(r'<row r="1".*?</row>', a, re.S).group(0),
                         re.search(r'<row r="1".*?</row>', d, re.S).group(0))
        for n in range(2, 12):
            ca = re.findall(r'<c r="[A-Z]+%d"[^>]*?(?<!<f)>(?:(?!<f>).)*?</c>' % n,
                            re.search(rf'<row r="{n}".*?</row>', a, re.S).group(0))
            for celula in ca:
                if "<f>" not in celula:
                    self.assertIn(celula, d)

    def test_c_openpyxl_carrega(self):
        wb = openpyxl.load_workbook(self.dest)
        ws = wb["Processos"]
        self.assertEqual(ws.max_row, 211)
        self.assertEqual(ws.tables["tblProcessos"].ref, "A1:AC211")
        self.assertEqual(ws["A211"].value, g.cnj_ficticio(210))
        self.assertEqual(ws["G12"].number_format, "dd/mm/yyyy")           # estilo copiado da linha-modelo
        self.assertEqual(ws["K12"].number_format, "#,##0.00")
        self.assertEqual(ws["G12"].value.date(), self.novas[0]["Data do Ajuizamento"])
        self.assertEqual(ws["W12"].value[:3], "=IF")                      # fórmula propagada
        self.assertEqual(ws["Y12"].value, '=IF(X12="","",X12-G12)')       # A1 deslocada
        self.assertEqual(ws["N211"].value, '=IF(OR(M211="Encerrado",M211="Arquivado"),"Não","Sim")')
        self.assertEqual(wb.sheetnames, ["Processos", "Parâmetros", "Indicadores", "Dashboard", "Dinâmica"])

    def test_d_estruturas_acompanham(self):
        d = partes(self.dest)
        s = d["xl/worksheets/sheet1.xml"].decode()
        self.assertIn('<dimension ref="A1:AC211"/>', s)
        sq = re.findall(r'sqref="([^"]*)"', s)
        for esperado in ("A2:AC211", "M2:N211", "K2:K211", "O2:O211 Q2:Q211", "P2:P211", "M2:M211", "AA2:AA211"):
            self.assertIn(esperado, sq)
        wb = d["xl/workbook.xml"].decode()
        self.assertIn("Processos!$A$1:$AC$211", wb)
        self.assertIn("Processos!$K$2:$K$211", wb)
        self.assertIn("'Parâmetros'!$B$3", wb)                       # nome que não aponta para a tabela: igual
        self.assertIn('fullCalcOnLoad="1"', wb)
        self.assertNotIn("calcChain", d["xl/_rels/workbook.xml.rels"].decode())
        self.assertNotIn("calcChain", d["[Content_Types].xml"].decode())
        self.assertIn("$K$2:$K$211", d["xl/charts/chart2.xml"].decode())
        self.assertIn("$B$12:$B$18", d["xl/charts/chart1.xml"].decode())   # gráfico da outra aba: intacto
        ind = d["xl/worksheets/sheet3.xml"].decode()
        self.assertIn("Processos!$K$2:$K$211,Processos!$P$2:$P$211", ind)
        self.assertIn("COUNTIFS(Processos!$M$2:$M$211,A12)", ind)
        self.assertIn("tblProcessos[Valor da Causa]", ind)           # referência estruturada: não se mexe

    def test_e_xml_coerente(self):
        self.assertEqual(x.validar(self.dest), [])
        d = partes(self.dest)
        for nome in ("xl/worksheets/sheet1.xml", "xl/worksheets/sheet3.xml", "xl/tables/table1.xml",
                     "xl/workbook.xml", "xl/charts/chart2.xml", "[Content_Types].xml", "xl/_rels/workbook.xml.rels"):
            ET.fromstring(d[nome])                                    # bem-formado
        raiz = ET.fromstring(d["xl/worksheets/sheet1.xml"])
        linhas = raiz.find(f"{M}sheetData")
        numeros = [int(r.get("r")) for r in linhas]
        self.assertEqual(numeros, list(range(1, 212)))
        refs = [c.get("r") for r in linhas for c in r]
        self.assertEqual(len(refs), len(set(refs)))                  # refs de célula únicas
        tab = ET.fromstring(d["xl/tables/table1.xml"])
        self.assertEqual(tab.get("ref"), f"A1:AC{len(numeros)}")
        self.assertEqual(tab.find(f"{M}autoFilter").get("ref"), tab.get("ref"))
        ids = [c.get("id") for c in tab.find(f"{M}tableColumns")]
        self.assertEqual(len(ids), len(set(ids)))
        self.assertEqual(len(ids), 29)

    @unittest.skipUnless(TEM_LO, "LibreOffice ausente")
    def test_f_libreoffice_pdf_e_csv(self):
        pasta = self.lo(self.dest, "pdf")
        pdf = (pasta / "200.pdf").read_bytes()
        self.assertTrue(pdf.startswith(b"%PDF") and len(pdf) > 5000)
        pasta = self.lo(self.dest, "csv", "csv:Text - txt - csv (StarCalc):44,34,76,1,,0,false,true,false,false,false,-1")
        arquivos = {p.name.split("-", 1)[1]: p for p in pasta.glob("*.csv")}
        self.assertEqual(len(arquivos), 5)                            # uma por aba
        linhas = list(csv.reader(io.StringIO(next(p for n, p in arquivos.items() if n.startswith("Processos")
                                                  ).read_text("utf-8"))))
        self.assertEqual(len(linhas), 211)
        self.assertEqual(linhas[-1][0], g.cnj_ficticio(210))

    @unittest.skipUnless(TEM_LO, "LibreOffice ausente")
    def test_g_libreoffice_recalcula_formulas(self):
        wb = self.lo_valores(self.dest)
        todas = self.linhas_novas(0, 210)
        exp = g.indicadores_esperados(todas)
        ind = wb["Indicadores"]
        v = {r: ind.cell(r, 2).value for r in range(2, 19)}
        self.assertEqual(v[2], exp["total"])
        self.assertEqual(v[3], exp["ativos"])
        self.assertAlmostEqual(v[4], exp["valor_causa_ativos"], 2)
        self.assertAlmostEqual(v[5], exp["prazo_medio"], 6)
        self.assertAlmostEqual(v[6], exp["economia"], 2)
        self.assertAlmostEqual(v[7], exp["causa_provavel"], 2)
        self.assertAlmostEqual(v[8], exp["total"] / 350 * 1000, 6)
        self.assertEqual([v[12 + i] for i in range(7)], exp["por_situacao"])
        p = wb["Processos"]
        for i in (0, 9, 10, 11, 150, 209):                            # antigas e novas
            d = g.derivados(todas[i])
            r = 2 + i
            self.assertEqual(p[f"N{r}"].value, d["Ativo"])
            for col, nome in (("W", "Valor Economizado"), ("Y", "Taxa de resolução (em dias)")):
                esperado = d[nome]
                atual = p[f"{col}{r}"].value
                if esperado == "":
                    self.assertIn(atual, (None, ""))
                else:
                    self.assertAlmostEqual(float(atual), float(esperado), 4)

    @unittest.skipUnless(TEM_LO, "LibreOffice ausente")
    def test_h_libreoffice_tabela_dinamica_acompanha(self):
        pasta = self.lo(self.dest, "csv", "csv:Text - txt - csv (StarCalc):44,34,76,1,,0,false,true,false,false,false,-1")
        din = next(p for p in pasta.glob("*.csv") if "-Din" in p.name).read_text("utf-8")
        self.assertIn(",210,", din)                                   # a dinâmica reflete as 210 linhas
        pasta = self.lo(self.dest, "ods")
        with zipfile.ZipFile(pasta / "200.ods") as z:
            self.assertIn("data-pilot-table", z.read("content.xml").decode())


    @unittest.skipUnless(TEM_LO, "LibreOffice ausente")
    def test_i_fullcalconload_sozinho_nao_basta_no_libreoffice(self):
        """Achado do spike: o LibreOffice ignora fullCalcOnLoad e confia no cache das fórmulas. Por isso o
        protótipo também apaga o cache (invalidar_cache=True, padrão). Este teste documenta o comportamento:
        se algum dia o LibreOffice passar a honrar fullCalcOnLoad, ele falha e a decisão pode ser revista."""
        dest = self.tmp / "sem-invalidar.xlsx"
        ed = x.Edicao(self.completo, "Processos", invalidar_cache=False)
        ed.inserir_linhas(self.novas)
        ed.salvar(dest)
        self.assertIn('fullCalcOnLoad="1"', partes(dest)["xl/workbook.xml"].decode())
        wb = self.lo_valores(dest)
        self.assertEqual(wb["Indicadores"]["B2"].value, 10)           # velho: 10 em vez de 210
        self.assertEqual(wb["Processos"]["W211"].value, 524604.5)     # as fórmulas das linhas novas calculam


class TestSujo(Base):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.dest = cls.tmp / "sujo200.xlsx"
        cls.res = x.inserir_linhas(cls.sujo, cls.dest, "Processos", cls.linhas_novas(10, 110))

    def test_sujo_tem_as_sujeiras(self):
        p = partes(self.sujo)
        sheet = p["xl/worksheets/sheet1.xml"].decode()
        self.assertIn('t="shared"', sheet)
        self.assertNotIn("tableStyleInfo", p["xl/tables/table1.xml"].decode())
        self.assertGreater(p["xl/sharedStrings.xml"].count(b"<si>"), 6000)
        self.assertIn('<row r="30"', sheet)                           # linhas vazias formatadas no fim

    def test_insercao_e_validacao(self):
        self.assertEqual(self.res.linhas_inseridas, 100)
        self.assertEqual(x.validar(self.dest), [])
        antes, depois = partes(self.sujo), partes(self.dest)
        self.assertEqual(antes["xl/sharedStrings.xml"], depois["xl/sharedStrings.xml"])   # 6.000 textos: nem tocados
        self.assertEqual(antes["xl/worksheets/sheet6.xml"], depois["xl/worksheets/sheet6.xml"])
        self.assertNotIn("tableStyleInfo", depois["xl/tables/table1.xml"].decode())
        self.assertIn('ref="A1:AC111"', depois["xl/tables/table1.xml"].decode())

    def test_linhas_vazias_do_fim_sao_reaproveitadas(self):
        s = partes(self.dest)["xl/worksheets/sheet1.xml"].decode()
        self.assertEqual(s.count('<row r="12"'), 1)
        self.assertEqual(s.count('<row r="30"'), 1)
        numeros = [int(n) for n in re.findall(r'<row r="(\d+)"', s)]
        self.assertEqual(numeros, sorted(set(numeros)))
        self.assertEqual(max(numeros), 111)

    def test_celulas_sem_estilo_e_formula_compartilhada(self):
        wb = openpyxl.load_workbook(self.dest)
        ws = wb["Processos"]
        self.assertEqual(ws["Y12"].value, '=IF(X12="","",X12-G12)')   # veio da fórmula compartilhada (mestre Y2)
        self.assertEqual(ws["Y111"].value, '=IF(X111="","",X111-G111)')
        s = partes(self.dest)["xl/worksheets/sheet1.xml"].decode()
        # linha 12 já existia (vazia e formatada): célula mantém o estilo que tinha; linha 41 é nova e a
        # linha-modelo (última de dados) tem a coluna AC sem estilo: a célula nova também fica sem
        self.assertIn(' s="', re.search(r'<c r="AC12"[^>]*?/?>', s).group(0))
        m41 = re.search(r'<c r="AC41"[^>]*?/?>', s)
        self.assertTrue(m41 is None or ' s="' not in m41.group(0))
        self.assertIn(' s="', re.search(r'<c r="AB41"[^>]*?>', s).group(0))   # AB tem estilo na linha-modelo
        self.assertIn(' s="', re.search(r'<c r="A12"[^>]*?>', s).group(0))

    def test_cf_e_dv_do_sujo(self):
        s = partes(self.dest)["xl/worksheets/sheet1.xml"].decode()
        self.assertIn('sqref="P2:P111"', s)
        self.assertIn('sqref="O2:O111 Q2:Q111"', s)

    @unittest.skipUnless(TEM_LO, "LibreOffice ausente")
    def test_libreoffice(self):
        wb = self.lo_valores(self.dest)
        exp = g.indicadores_esperados(self.linhas_novas(0, 110))
        ind = wb["Indicadores"]
        self.assertEqual(ind["B2"].value, exp["total"])
        self.assertEqual(ind["B3"].value, exp["ativos"])
        self.assertAlmostEqual(ind["B6"].value, exp["economia"], 2)
        self.assertEqual(wb["Processos"]["Y2"].value in (None, ""), True)   # fórmula compartilhada recalcula


class TestEscreverCelulas(Base):
    def test_humanos_nunca_sobrescritos(self):
        dest = self.saida("humanos.xlsx")
        humanos = ["Probabilidade", "Valor Estimado", "Garantias Processuais"]
        r = x.escrever_celulas(self.completo, dest, "Processos", {
            (2, "Probabilidade"): "Provável",        # humano com valor: ignorado
            (2, "Garantias Processuais"): 1500.0,    # humano vazio: grava
            (2, "Objeto"): "Objeto novo",            # não humano: grava por cima
            (2, "Ativo"): "Sim",                     # fórmula: ignorada
            (3, "Valor Estimado"): 1.0,              # humano com valor: ignorado
            (3, "AA"): 0.55,                         # coluna por letra
        }, humanos=humanos)
        self.assertEqual(r.celulas_escritas, 3)
        motivos = dict(r.ignoradas)
        self.assertEqual(motivos["P2"], "campo humano já preenchido")
        self.assertEqual(motivos["Q3"], "campo humano já preenchido")
        self.assertEqual(motivos["N2"], "tem fórmula")
        self.assertEqual(x.validar(dest), [])
        ws = openpyxl.load_workbook(dest)["Processos"]
        orig = openpyxl.load_workbook(self.completo)["Processos"]
        self.assertEqual(ws["P2"].value, orig["P2"].value)
        self.assertEqual(ws["Q3"].value, orig["Q3"].value)
        self.assertEqual(ws["U2"].value, 1500.0)
        self.assertEqual(ws["U2"].number_format, "#,##0.00")          # estilo da célula preservado
        self.assertEqual(ws["J2"].value, "Objeto novo")
        self.assertEqual(ws["AA3"].value, 0.55)
        self.assertEqual(ws["AA3"].number_format, orig["AA3"].number_format)
        self.assertEqual(ws["N2"].value, orig["N2"].value)
        # partes não editadas: iguais
        a, d = partes(self.completo), partes(dest)
        for n in a:
            if n not in ("xl/worksheets/sheet1.xml", "xl/worksheets/sheet3.xml", "xl/workbook.xml",
                         "xl/_rels/workbook.xml.rels", "[Content_Types].xml", "xl/calcChain.xml"):
                self.assertEqual(a[n], d[n], n)

    def test_sem_preservar_sobrescreve(self):
        dest = self.saida("humanos2.xlsx")
        r = x.escrever_celulas(self.completo, dest, "Processos", {(2, "Probabilidade"): "Remota"},
                               humanos=["Probabilidade"], preservar_humanos=False)
        self.assertEqual(r.ignoradas, [])
        self.assertEqual(openpyxl.load_workbook(dest)["Processos"]["P2"].value, "Remota")

    def test_limpar_e_data(self):
        dest = self.saida("limpar.xlsx")
        r = x.escrever_celulas(self.completo, dest, "Processos", {
            (2, "Objeto"): x.LIMPAR, (2, "Data do Ajuizamento"): datetime.date(2020, 2, 29)})
        ws = openpyxl.load_workbook(dest)["Processos"]
        self.assertIsNone(ws["J2"].value)
        self.assertEqual(ws["G2"].value.date() if hasattr(ws["G2"].value, "date") else ws["G2"].value,
                         datetime.date(2020, 2, 29))
        self.assertEqual(r.avisos, [])

    def test_data_em_celula_sem_formato_de_data_avisa(self):
        r = x.escrever_celulas(self.completo, self.saida("d.xlsx"), "Processos",
                               {(2, "Município"): datetime.date(2020, 1, 1)})
        self.assertTrue(any("sem formato de data" in a for a in r.avisos))

    def test_texto_com_caracteres_especiais(self):
        dest = self.saida("esp.xlsx")
        texto = "linha1\nlinha2 & <b>\"aspas\"</b> \x01ilegal  espaços  "
        x.escrever_celulas(self.completo, dest, "Processos", {(4, "Andamentos"): texto})
        v = openpyxl.load_workbook(dest)["Processos"]["L4"].value
        self.assertEqual(v, texto.replace("\x01", ""))
        self.assertEqual(x.validar(dest), [])


class TestAcrescentarColuna(Base):
    def test_coluna_calculada_estruturada(self):
        dest = self.saida("col.xlsx")
        f = f'IF({g.TABELA}[[#This Row],[Valor da Causa]]>20000,"Alto","Baixo")'
        r = x.acrescentar_coluna(self.completo, dest, "Processos", "Faixa de valor", formula=f)
        self.assertEqual(x.validar(dest), [])
        d = partes(dest)
        tab = d["xl/tables/table1.xml"].decode()
        self.assertIn('ref="A1:AD11"', tab)
        self.assertIn('<tableColumns count="30">', tab)
        self.assertIn('name="Faixa de valor"><calculatedColumnFormula>', tab)
        self.assertIn("A1:AD11", re.search(r'<autoFilter[^>]*>', tab).group(0))
        self.assertIn("Processos!$A$1:$AD$11", d["xl/workbook.xml"].decode())          # nome da tabela inteira
        self.assertIn("Processos!$K$2:$K$11", d["xl/workbook.xml"].decode())           # coluna só: igual
        self.assertIn('<dimension ref="A1:AD11"/>', d["xl/worksheets/sheet1.xml"].decode())
        self.assertIn('count="30"', d["xl/pivotCache/pivotCacheDefinition1.xml"].decode())   # cache + dinâmica
        self.assertIn('<pivotFields count="30">', d["xl/pivotTables/pivotTable1.xml"].decode())
        wb = openpyxl.load_workbook(dest)
        ws = wb["Processos"]
        self.assertEqual(ws["AD1"].value, "Faixa de valor")
        self.assertEqual(ws["AD1"].font.b, ws["AC1"].font.b)           # estilo do cabeçalho copiado
        self.assertEqual(ws["AD11"].value[:3], "=IF")
        self.assertEqual(ws.tables["tblProcessos"].ref, "A1:AD11")
        self.assertEqual(r.avisos, [])

    def test_valores_e_a1(self):
        dest = self.saida("col2.xlsx")
        r = x.acrescentar_coluna(self.completo, dest, "Processos", "Observação interna",
                                 valores=[f"obs {i}" for i in range(10)])
        r2 = x.acrescentar_coluna(self.completo, self.saida("col3.xlsx"), "Processos", "Dobro",
                                  formula="K2*2")
        self.assertTrue(any("referência A1" in a for a in r2.avisos))
        self.assertEqual(x.validar(self.saida("col3.xlsx")), [])
        ws = openpyxl.load_workbook(dest)["Processos"]
        self.assertEqual([ws.cell(r, 30).value for r in (2, 11)], ["obs 0", "obs 9"])
        ws3 = openpyxl.load_workbook(self.saida("col3.xlsx"))["Processos"]
        self.assertEqual(ws3["AD7"].value, "=K7*2")                    # A1 relativa desloca por linha

    def test_idempotente(self):
        d1, d2 = self.saida("i1.xlsx"), self.saida("i2.xlsx")
        x.acrescentar_coluna(self.completo, d1, "Processos", "Observação interna", valores=["a"] * 10)
        r = x.acrescentar_coluna(d1, d2, "Processos", "Observação interna", valores=["b"] * 10)
        ws = openpyxl.load_workbook(d2)["Processos"]
        self.assertEqual(ws.max_column, 30)
        self.assertEqual(ws["AD2"].value, "a")                         # já tinha valor: não sobrescreve
        self.assertTrue(any("já existe" in a for a in r.avisos))

    @unittest.skipUnless(TEM_LO, "LibreOffice ausente")
    def test_libreoffice_coluna_calculada_e_linhas_depois(self):
        d1, d2 = self.saida("c1.xlsx"), self.saida("c2.xlsx")
        f = f'IF({g.TABELA}[[#This Row],[Valor da Causa]]>20000,"Alto","Baixo")'
        x.acrescentar_coluna(self.completo, d1, "Processos", "Faixa de valor", formula=f)
        # depois, 50 linhas novas: a coluna calculada nova deve propagar (vem da Tabela)
        x.inserir_linhas(d1, d2, "Processos", self.linhas_novas(10, 60))
        self.assertEqual(x.validar(d2), [])
        wb = self.lo_valores(d2)
        ws = wb["Processos"]
        todas = self.linhas_novas(0, 60)
        for i in (0, 5, 10, 59):
            esperado = "Alto" if todas[i]["Valor da Causa"] > 20000 else "Baixo"
            self.assertEqual(ws.cell(2 + i, 30).value, esperado)
        self.assertEqual(ws.max_row, 61)

    def test_dinamica_com_cache_salvo_nao_suportada(self):
        com_registros = self.saida("reg.xlsx")
        p = partes(self.completo)
        p["xl/pivotCache/pivotCacheDefinition1.xml"] = p["xl/pivotCache/pivotCacheDefinition1.xml"].replace(
            b"<cacheSource", b'<cacheSource', 1)
        rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.'
                'openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.'
                'openxmlformats.org/officeDocument/2006/relationships/pivotCacheRecords" '
                'Target="pivotCacheRecords1.xml"/></Relationships>').encode()
        p["xl/pivotCache/_rels/pivotCacheDefinition1.xml.rels"] = rels
        p["xl/pivotCache/pivotCacheRecords1.xml"] = b'<pivotCacheRecords xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="0"/>'
        with zipfile.ZipFile(com_registros, "w") as z:
            for n, d in p.items():
                z.writestr(n, d)
        with self.assertRaises(x.NaoSuportado):
            x.acrescentar_coluna(com_registros, self.saida("reg2.xlsx"), "Processos", "Nova")
        self.assertFalse(self.saida("reg2.xlsx").exists())


class TestSegurancaEErros(Base):
    def test_coluna_desconhecida_nao_descarta_dado(self):
        dest = self.saida("e1.xlsx")
        with self.assertRaises(x.ColunaDesconhecida):
            x.inserir_linhas(self.completo, dest, "Processos", [{"Número do Processo": "x", "Coluna Fantasma": 1}])
        self.assertFalse(dest.exists())

    def test_mapeamento_de_colunas_e_nomes_tolerantes(self):
        dest = self.saida("e2.xlsx")
        x.inserir_linhas(self.completo, dest, "Processos",
                         [{"numero": "1234567-00.2020.8.06.0001", "reus": "Ré Teste", "recurso": "Sim"}],
                         colunas={"numero": "Número do Processo", "reus": "reu(s)",
                                  "recurso": "Houve recurso da empresa"})
        ws = openpyxl.load_workbook(dest)["Processos"]
        self.assertEqual((ws["A12"].value, ws["C12"].value, ws["Z12"].value),
                         ("1234567-00.2020.8.06.0001", "Ré Teste", "Sim"))

    def test_conteudo_abaixo_da_tabela(self):
        com = self.saida("abaixo.xlsx")
        x.escrever_celulas(self.completo, com, "Processos", {(13, "A"): "nota do usuário"})
        with self.assertRaises(x.ConflitoDeConteudo):
            x.inserir_linhas(com, self.saida("abaixo2.xlsx"), "Processos", self.linhas_novas(10, 20))

    def test_linha_de_totais_nao_suportada(self):
        p = partes(self.completo)
        p["xl/tables/table1.xml"] = p["xl/tables/table1.xml"].replace(b"<table ", b'<table totalsRowCount="1" ', 1)
        arq = self.saida("tot.xlsx")
        with zipfile.ZipFile(arq, "w") as z:
            for n, d in p.items():
                z.writestr(n, d)
        with self.assertRaises(x.NaoSuportado):
            x.inserir_linhas(arq, self.saida("tot2.xlsx"), "Processos", self.linhas_novas(10, 11))

    def test_original_nunca_alterado(self):
        antes = sha(self.completo.read_bytes())
        x.inserir_linhas(self.completo, self.saida("o.xlsx"), "Processos", self.linhas_novas(10, 15))
        self.assertEqual(sha(self.completo.read_bytes()), antes)
        with self.assertRaises(x.ErroXlsx):
            x.inserir_linhas(self.completo, self.completo, "Processos", self.linhas_novas(10, 11))
        self.assertEqual(sha(self.completo.read_bytes()), antes)

    def test_texto_acima_do_limite_do_excel(self):
        with self.assertRaises(x.ErroXlsx):
            x.escrever_celulas(self.completo, self.saida("big.xlsx"), "Processos", {(2, "Andamentos"): "a" * 32768})
        x.escrever_celulas(self.completo, self.saida("big.xlsx"), "Processos", {(2, "Andamentos"): "a" * 32767})

    def test_aba_inexistente(self):
        with self.assertRaises(x.ErroXlsx):
            x.inserir_linhas(self.completo, self.saida("n.xlsx"), "Nao existe", [{"A": 1}])

    def test_arquivo_que_nao_e_xlsx(self):
        ruim = self.saida("ruim.xlsx")
        ruim.write_bytes(b"isto nao e um zip")
        with self.assertRaises(x.ErroXlsx):
            x.inserir_linhas(ruim, self.saida("r.xlsx"), "Processos", [{"A": 1}])

    def test_saida_deterministica(self):
        a, b = self.saida("det1.xlsx"), self.saida("det2.xlsx")
        x.inserir_linhas(self.completo, a, "Processos", self.linhas_novas(10, 60))
        x.inserir_linhas(self.completo, b, "Processos", self.linhas_novas(10, 60))
        self.assertEqual(a.read_bytes(), b.read_bytes())

    def test_gerador_deterministico(self):
        a, b = self.saida("g1.xlsx"), self.saida("g2.xlsx")
        g.gerar(a)
        g.gerar(b)
        self.assertEqual(a.read_bytes(), b.read_bytes())
        self.assertEqual(a.read_bytes(), self.completo.read_bytes())

    def test_composicao_100_mais_100_igual_a_200(self):
        d1, d2, d3 = self.saida("k1.xlsx"), self.saida("k2.xlsx"), self.saida("k3.xlsx")
        x.inserir_linhas(self.completo, d1, "Processos", self.linhas_novas(10, 110))
        x.inserir_linhas(d1, d2, "Processos", self.linhas_novas(110, 210))
        x.inserir_linhas(self.completo, d3, "Processos", self.linhas_novas(10, 210))
        a, b = partes(d2), partes(d3)
        self.assertEqual(set(a), set(b))
        for n in a:
            self.assertEqual(a[n], b[n], n)

    def test_dinamica_por_intervalo_e_refresh(self):
        p = partes(self.completo)
        cache = p["xl/pivotCache/pivotCacheDefinition1.xml"].decode()
        cache = cache.replace('<worksheetSource name="tblProcessos"/>',
                              '<worksheetSource ref="A1:AC11" sheet="Processos"/>').replace(' refreshOnLoad="1"', "")
        self.assertNotIn("refreshOnLoad", cache)
        p["xl/pivotCache/pivotCacheDefinition1.xml"] = cache.encode()
        arq = self.saida("pivref.xlsx")
        with zipfile.ZipFile(arq, "w") as z:
            for n, d in p.items():
                z.writestr(n, d)
        dest = self.saida("pivref2.xlsx")
        x.inserir_linhas(arq, dest, "Processos", self.linhas_novas(10, 30))
        novo = partes(dest)["xl/pivotCache/pivotCacheDefinition1.xml"].decode()
        self.assertIn('ref="A1:AC31"', novo)
        self.assertIn('refreshOnLoad="1"', novo)
        self.assertEqual(x.validar(dest), [])

    def test_validar_pega_problemas(self):
        p = partes(self.completo)
        p["xl/tables/table1.xml"] = p["xl/tables/table1.xml"].replace(b'ref="A1:AC11"', b'ref="A1:AC50"', 1)
        arq = self.saida("ruim_tab.xlsx")
        with zipfile.ZipFile(arq, "w") as z:
            for n, d in p.items():
                z.writestr(n, d)
        probs = x.validar(arq)
        self.assertTrue(any("autoFilter" in s for s in probs), probs)


class TestGravarContrato(Base):
    def test_gravar_estado(self):
        existente = g.cnj_ficticio(1)
        novo = g.cnj_ficticio(900)
        estado = {"cliente": "Cliente Fictício", "data_base": "2026-10-07", "eventos": [], "perfil": {}, "parametros": {},
                  "fichas": [
                      {"numero": existente, "campos": {"garantias": {"valor": "1500.00", "origem": "humano"},
                                                       "probabilidade": {"valor": "Remota", "origem": "sugerido"}}},
                      {"numero": novo, "campos": {"autores": {"valor": "Autor Novo"}, "valor_causa": {"valor": "9000.50"},
                                                  "data_ajuizamento": {"valor": "2026-09-01"},
                                                  "situacao": {"valor": "Em andamento"}}},
                  ]}
        dest = self.saida("contrato.xlsx")
        r = x.gravar(self.completo, estado, dest)
        self.assertEqual(set(r), {"destino", "processos_atualizados", "processos_novos", "mudancas", "avisos"})
        self.assertEqual(r["processos_novos"], [novo])
        self.assertEqual(r["processos_atualizados"], [existente])
        ws = openpyxl.load_workbook(dest)["Processos"]
        self.assertEqual(ws["U2"].value, 1500.0)                       # célula vazia: preenchida
        self.assertEqual(ws["P2"].value, "Remota" if False else openpyxl.load_workbook(self.completo)["Processos"]["P2"].value)
        self.assertEqual(ws["A12"].value, novo)
        self.assertEqual(ws["K12"].value, 9000.5)
        self.assertEqual(ws["G12"].value, datetime.datetime(2026, 9, 1))
        self.assertEqual(x.validar(dest), [])
        # idempotente: gravar o mesmo estado de novo não duplica
        dest2 = self.saida("contrato2.xlsx")
        r2 = x.gravar(dest, estado, dest2)
        self.assertEqual((r2["processos_novos"], r2["processos_atualizados"]), ([], []))
        self.assertEqual(openpyxl.load_workbook(dest2)["Processos"].max_row, 12)


class TestAlternativaPreformatada(Base):
    """Alternativa 'template com linhas pré-formatadas': a tabela já cobre 300 linhas e só se preenchem células."""

    def test_preencher_sem_mudar_estrutura(self):
        dest = self.saida("pre_cheio.xlsx")
        novas = self.linhas_novas(10, 210)
        celulas = {}
        for k, l in enumerate(novas):
            for cab, v in l.items():
                if cab not in g.FORMULAS and v is not None:
                    celulas[(12 + k, cab)] = v
        r = x.escrever_celulas(self.pre, dest, "Processos", celulas)
        self.assertEqual(x.validar(dest), [])
        a, d = partes(self.pre), partes(dest)
        alteradas = {n for n in d if a[n] != d[n]}
        # estrutura intacta: tabela, nomes, gráficos, CF/DV não mudam; só as abas com cache de fórmula
        self.assertEqual(alteradas, {"xl/worksheets/sheet1.xml", "xl/worksheets/sheet3.xml", "xl/workbook.xml",
                                     "xl/_rels/workbook.xml.rels", "[Content_Types].xml"})
        self.assertEqual(r.celulas_escritas, len(celulas))

    @unittest.skipUnless(TEM_LO, "LibreOffice ausente")
    def test_efeito_colateral_das_linhas_em_branco(self):
        """O custo da alternativa: as fórmulas das linhas vazias contam como processo 'Sim' (ativo)."""
        dest = self.saida("pre_cheio2.xlsx")
        celulas = {(12 + k, cab): v for k, l in enumerate(self.linhas_novas(10, 60))
                   for cab, v in l.items() if cab not in g.FORMULAS and v is not None}
        x.escrever_celulas(self.pre, dest, "Processos", celulas)
        wb = self.lo_valores(dest)
        real = g.indicadores_esperados(self.linhas_novas(0, 60))
        ind = wb["Indicadores"]
        self.assertEqual(ind["B2"].value, real["total"])               # COUNTA da coluna de número: certo
        self.assertEqual(ind["B3"].value, real["ativos"] + 240)       # 240 linhas em branco viram "ativo": errado
        self.assertGreater(ind["B3"].value, real["ativos"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
