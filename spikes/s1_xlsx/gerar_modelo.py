"""SPIKE S1: gera a planilha FICTÍCIA "modelo B" (e variantes "suja" e "pré-formatada")
usada para testar a gravação cirúrgica (xlsx_cirurgico.py).

Tudo aqui é fictício e determinístico (sem aleatoriedade, sem data de hoje): rodar duas
vezes produz arquivos idênticos byte a byte. Os números de processo são calculados em
tempo de execução (com dígito verificador CNJ correto) e NUNCA ficam como literais no
repositório (regra do empacotar.sh).

O openpyxl só é usado aqui, para GERAR o modelo (ele cria gráfico novo sem problema; o que
ele não faz é preservar gráfico de arquivo existente). Depois o pacote é pós-processado no
XML para ficar parecido com um arquivo salvo pelo Excel:
  - calcChain.xml (o openpyxl não escreve),
  - valores em cache nas fórmulas (o openpyxl escreve fórmulas sem valor),
  - calcPr sem fullCalcOnLoad (o openpyxl liga essa opção; o Excel não),
  - coluna calculada declarada na Tabela (calculatedColumnFormula),
  - tabela dinâmica montada à mão (cache + definição), com a tabela como origem.

Variantes:
  completo       modelo de referência (10 linhas)
  sujo           células sem estilo, sharedStrings grande, linhas vazias formatadas no fim,
                 fórmula compartilhada (t="shared"), tabela sem formatação (sem tableStyleInfo)
  preformatado   a alternativa "template com linhas pré-formatadas": tabela cobre 300 linhas,
                 as vazias já têm estilo e fórmulas

Uso:  python3 gerar_modelo.py DESTINO.xlsx [completo|sujo|preformatado]
"""
import datetime
import io
import re
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

CABECALHO = [
    "Número do Processo", "Autor(es)", "Réu(s)", "Vara", "Município", "Tribunal",
    "Data do Ajuizamento", "Área do Direito", "Matéria Principal", "Objeto", "Valor da Causa",
    "Andamentos", "Situação", "Ativo", "Valor Arbitrado em Juízo", "Probabilidade",
    "Valor Estimado", "Valor da Execução", "Custas Processuais", "Depósitos Recursais",
    "Garantias Processuais", "Resultado", "Valor Economizado", "Data do trânsito em julgado",
    "Taxa de resolução (em dias)", "Houve recurso da empresa?", "Percentual de êxito",
    "Reclamante terceirizado?", "Outra(s) Parte(s)",
]
assert len(CABECALHO) == 29
COL = {nome: i + 1 for i, nome in enumerate(CABECALHO)}   # nome -> índice (1 = A)
FORMULAS = ("Ativo", "Valor Economizado", "Taxa de resolução (em dias)")
HUMANOS = ("Probabilidade", "Valor Estimado", "Depósitos Recursais", "Garantias Processuais",
           "Valor Arbitrado em Juízo", "Percentual de êxito", "Reclamante terceirizado?")
SITUACOES = ["Em andamento", "Aguardando sentença", "Em recurso", "Cumprimento de sentença",
             "Encerrado", "Arquivado", "Suspenso"]
PROBABILIDADES = ["Remota", "Possível", "Provável"]
ENCERRADAS = ("Encerrado", "Arquivado")
TABELA = "tblProcessos"
DATA_FIXA = datetime.datetime(2026, 10, 7, 12, 0, 0)
ZIP_DATA = (2026, 10, 7, 12, 0, 0)

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


# ------------------------------------------------------------------ dados fictícios

def cnj_ficticio(n):
    """Número CNJ fictício n (n >= 1) com dígito verificador correto.
    Sequencial fixo 1234567 e varia a vara (OOOO) e o ano: nunca é um processo real."""
    seq, ano, j, tr, vara = 1234567, 2018 + n % 8, 8, 6, n
    dv = 98 - int(f"{seq:07d}{ano:04d}{j}{tr:02d}{vara:04d}00") % 97
    return f"{seq:07d}-{dv:02d}.{ano:04d}.{j}.{tr:02d}.{vara:04d}"


def linha_ficticia(i):
    """Linha de exemplo i (0, 1, ...) como {cabeçalho: valor}, SEM as colunas de fórmula."""
    n = i + 1
    situacao = SITUACOES[i % len(SITUACOES)]
    ajuiz = datetime.date(2019, 1, 15) + datetime.timedelta(days=37 * i)
    encerrado = situacao in ENCERRADAS
    d = {
        "Número do Processo": cnj_ficticio(n),
        "Autor(es)": f"Autor Fictício {n:03d}",
        "Réu(s)": f"Empresa Modelo {1 + i % 3} Ltda",
        "Vara": f"{1 + i % 9}ª Vara do Trabalho de Cidade {1 + i % 4}",
        "Município": f"Cidade {1 + i % 4}",
        "Tribunal": "TRT-0" + str(1 + i % 3),
        "Data do Ajuizamento": ajuiz,
        "Área do Direito": ["Trabalhista", "Cível", "Consumidor"][i % 3],
        "Matéria Principal": ["Horas extras", "Dano moral", "Reversão de justa causa", "Vínculo"][i % 4],
        "Objeto": f"Objeto fictício do processo {n}; texto com acentuação e & <símbolos>.",
        "Valor da Causa": 10000.0 + 2500.5 * i,
        "Andamentos": f"Em {ajuiz:%d/%m/%Y}, distribuído. Até 30/09/2026 sem andamentos.",
        "Situação": situacao,
        "Valor Arbitrado em Juízo": (5000.0 + 100 * i) if encerrado else None,
        "Probabilidade": PROBABILIDADES[i % 3],
        "Valor Estimado": 4000.0 + 1000 * (i % 5),
        "Valor da Execução": (3000.0 + i) if i % 4 == 0 else None,
        "Custas Processuais": 120.0 + i,
        "Depósitos Recursais": 2000.0 if i % 3 == 0 else None,
        "Garantias Processuais": None,
        "Resultado": ("Acordo" if i % 2 else "Improcedente") if encerrado else None,
        "Data do trânsito em julgado": (ajuiz + datetime.timedelta(days=200 + 3 * i)) if encerrado else None,
        "Houve recurso da empresa?": "Sim" if i % 3 == 1 else "Não",
        "Percentual de êxito": round(0.1 * (i % 10), 2),
        "Reclamante terceirizado?": "Sim" if i % 5 == 0 else "Não",
        "Outra(s) Parte(s)": None,
    }
    return d


def derivados(l):
    """Valores esperados das 3 colunas de fórmula para uma linha (dict com os campos base)."""
    ativo = "Não" if l["Situação"] in ENCERRADAS else "Sim"
    vc, ve = l.get("Valor da Causa"), l.get("Valor Estimado")
    econ = "" if vc in (None, "") or ve in (None, "") else vc - ve
    tr, aj = l.get("Data do trânsito em julgado"), l.get("Data do Ajuizamento")
    taxa = "" if not tr else (tr - aj).days
    return {"Ativo": ativo, "Valor Economizado": econ, "Taxa de resolução (em dias)": taxa}


def indicadores_esperados(linhas):
    """Valores esperados da aba Indicadores (independentes do LibreOffice/Excel).
    `linhas`: lista de dicts com os campos base (derivados são calculados aqui)."""
    ls = [{**l, **derivados(l)} for l in linhas]
    ativos = [l for l in ls if l["Ativo"] == "Sim"]
    enc = [l["Taxa de resolução (em dias)"] for l in ls if l["Ativo"] == "Não"
           and l["Taxa de resolução (em dias)"] != ""]
    exitos = [l["Percentual de êxito"] for l in ls if l.get("Percentual de êxito") is not None]
    return {
        "total": len(ls),
        "ativos": len(ativos),
        "valor_causa_ativos": sum(l["Valor da Causa"] for l in ativos),
        "prazo_medio": (sum(enc) / len(enc)) if enc else 0,
        "economia": sum(l["Valor Economizado"] for l in ls if l["Valor Economizado"] != ""),
        "causa_provavel": sum(l["Valor da Causa"] for l in ls if l["Probabilidade"] == "Provável"),
        "exito_medio": (sum(exitos) / len(exitos)) if exitos else 0,
        "por_situacao": [sum(1 for l in ls if l["Situação"] == s) for s in SITUACOES],
    }


# ------------------------------------------------------------------ construção (openpyxl)

def _estilos():
    borda = Side(style="thin", color="BFBFBF")
    return {
        "cab": dict(font=Font(bold=True, color="FFFFFF"), fill=PatternFill("solid", fgColor="1F4E78"),
                    alignment=Alignment(wrap_text=True, vertical="center")),
        "corpo": dict(border=Border(left=borda, right=borda, top=borda, bottom=borda)),
    }


def _formato(nome):
    if nome in ("Data do Ajuizamento", "Data do trânsito em julgado"):
        return "dd/mm/yyyy"
    if nome == "Percentual de êxito":
        return "0.0%"
    if nome == "Taxa de resolução (em dias)":
        return "0"
    if nome.startswith("Valor") or nome in ("Custas Processuais", "Depósitos Recursais", "Garantias Processuais"):
        return "#,##0.00"
    return None


def _formula(nome, r):
    """Fórmula de cada coluna calculada na linha r (misto de referência estruturada e A1)."""
    t = TABELA
    if nome == "Valor Economizado":   # referência estruturada (declarada como coluna calculada)
        vc, ve = f"{t}[[#This Row],[Valor da Causa]]", f"{t}[[#This Row],[Valor Estimado]]"
        return f'=IF(OR({vc}="",{ve}=""),"",{vc}-{ve})'
    if nome == "Taxa de resolução (em dias)":   # A1 (não declarada: caso "fórmula só na linha")
        return f'=IF(X{r}="","",X{r}-G{r})'
    if nome == "Ativo":
        return f'=IF(OR(M{r}="Encerrado",M{r}="Arquivado"),"Não","Sim")'
    raise KeyError(nome)


def _preencher_linha(ws, r, dados, est, com_formula=True, sem_estilo=False):
    for nome, c in COL.items():
        cel = ws.cell(r, c)
        if nome in FORMULAS:
            if com_formula:
                cel.value = _formula(nome, r)
        elif dados is not None and dados.get(nome) is not None:
            cel.value = dados[nome]
        if sem_estilo:
            continue
        cel.border = est["corpo"]["border"]
        f = _formato(nome)
        if f:
            cel.number_format = f


def construir(variante="completo", n_linhas=10):
    est = _estilos()
    wb = Workbook()
    wb.properties.creator = "Gerador fictício S1"
    wb.properties.created = DATA_FIXA
    wb.properties.modified = DATA_FIXA
    wb.properties.lastModifiedBy = "Gerador fictício S1"

    ws = wb.active
    ws.title = "Processos"
    pre = 290 if variante == "preformatado" else 0
    sujo = variante == "sujo"

    for nome, c in COL.items():
        cel = ws.cell(1, c, nome)
        cel.font, cel.fill, cel.alignment = est["cab"]["font"], est["cab"]["fill"], est["cab"]["alignment"]
        ws.column_dimensions[cel.column_letter].width = 16 if nome != "Andamentos" else 50
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"

    linhas = [linha_ficticia(i) for i in range(n_linhas)]
    for i, l in enumerate(linhas):
        # sujo: a coluna AC e a linha 7 ficam sem nenhum estilo
        _preencher_linha(ws, 2 + i, l, est, sem_estilo=False)
    ultima = 1 + n_linhas + pre
    for r in range(2 + n_linhas, ultima + 1):   # pré-formatado: linhas vazias DENTRO da tabela, com fórmulas
        _preencher_linha(ws, r, None, est)

    tab = Table(displayName=TABELA, ref=f"A1:AC{ultima}")
    tab.tableStyleInfo = None if sujo else TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)

    if sujo:
        # linhas vazias formatadas depois da tabela (como quem formatou "até a linha 40")
        for r in range(ultima + 1, ultima + 30):
            for c in range(1, 30):
                ws.cell(r, c).border = est["corpo"]["border"]
        for c in range(1, 30):                       # linha 7 inteira sem estilo
            ws.cell(7, c)._style = ws.cell(1000, 1)._style
        for r in range(2, ultima + 1):                # coluna AC sem estilo
            ws.cell(r, 29)._style = ws.cell(1000, 1)._style
        ws.cell(1000, 1)._style = ws.cell(1001, 1)._style

    # ---- formatação condicional (3 tipos, uma com vários intervalos)
    fim = ultima
    ws.conditional_formatting.add(
        f"A2:AC{fim}", FormulaRule(formula=['$P2="Provável"'], font=Font(color="9C0006"),
                                   fill=PatternFill("solid", bgColor="FFC7CE")))
    ws.conditional_formatting.add(
        f"M2:N{fim}", FormulaRule(formula=['OR($M2="Encerrado",$M2="Arquivado")'],
                                  fill=PatternFill("solid", bgColor="D9D9D9")))
    ws.conditional_formatting.add(f"K2:K{fim}", ColorScaleRule(start_type="min", start_color="FFFFFF",
                                                              end_type="max", end_color="63BE7B"))
    ws.conditional_formatting.add(f"O2:O{fim} Q2:Q{fim}", DataBarRule(start_type="min", end_type="max",
                                                                      color="638EC6"))
    # ---- validação de dados
    dv1 = DataValidation(type="list", formula1='"Remota,Possível,Provável"', allow_blank=True)
    dv1.error, dv1.errorTitle = "Escolha um valor da lista", "Probabilidade"
    dv1.add(f"P2:P{fim}")
    dv2 = DataValidation(type="list", formula1="'Parâmetros'!$E$2:$E$8", allow_blank=True)
    dv2.add(f"M2:M{fim}")
    dv3 = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True)
    dv3.add(f"AA2:AA{fim}")
    for dv in (dv1, dv2, dv3):
        ws.add_data_validation(dv)

    # ---- Parâmetros
    pa = wb.create_sheet("Parâmetros")
    pa["A1"], pa["B1"] = "Parâmetro", "Valor"
    pa["A2"], pa["B2"] = "Headcount (nº de funcionários)", 350
    pa["A3"], pa["B3"] = "Data de referência", datetime.date(2026, 9, 30)
    pa["B3"].number_format = "dd/mm/yyyy"
    pa["A4"], pa["B4"] = "Empresas do grupo", "Empresa Modelo 1 Ltda; Empresa Modelo 2 Ltda"
    pa["E1"] = "Situações"
    for i, s in enumerate(SITUACOES):
        pa.cell(2 + i, 5, s)
    pa.column_dimensions["A"].width = 34

    # ---- Indicadores
    ind = wb.create_sheet("Indicadores")
    t = TABELA
    metricas = [
        ("Indicador", "Valor"),
        ("Total de processos", f"=COUNTA({t}[Número do Processo])"),
        ("Processos ativos", f'=COUNTIFS({t}[Ativo],"Sim")'),
        ("Valor total da causa (ativos)", f'=SUMIFS({t}[Valor da Causa],{t}[Ativo],"Sim")'),
        ("Prazo médio de resolução (dias)",
         f'=IFERROR(AVERAGEIFS({t}[Taxa de resolução (em dias)],{t}[Ativo],"Não"),0)'),
        ("Economia total", f"=SUM({t}[Valor Economizado])"),
        ("Valor da causa com probabilidade Provável",   # referência A1 de propósito
         f'=SUMIFS(Processos!$K$2:$K${fim},Processos!$P$2:$P${fim},"Provável")'),
        ("Processos por 1.000 funcionários", "=B2/'Parâmetros'!B2*1000"),
        ("Êxito médio", f"=IFERROR(AVERAGE({t}[Percentual de êxito]),0)"),
    ]
    for i, (a, b) in enumerate(metricas):
        ind.cell(1 + i, 1, a)
        ind.cell(1 + i, 2, b)
    ind["A11"], ind["B11"] = "Situação", "Processos"
    for i, s in enumerate(SITUACOES):
        ind.cell(12 + i, 1, s)
        ind.cell(12 + i, 2, f"=COUNTIFS(Processos!$M$2:$M${fim},A{12 + i})")
    ind.column_dimensions["A"].width = 44
    ind["B4"].number_format = ind["B6"].number_format = ind["B7"].number_format = "#,##0.00"

    # ---- Dashboard (gráficos novos: aqui o openpyxl serve)
    da = wb.create_sheet("Dashboard")
    g1 = BarChart()
    g1.title = "Processos por situação"
    g1.add_data(Reference(ind, min_col=2, min_row=12, max_row=18), titles_from_data=False)
    g1.set_categories(Reference(ind, min_col=1, min_row=12, max_row=18))
    g1.legend = None
    da.add_chart(g1, "A1")
    g2 = BarChart()
    g2.title = "Valor da causa por processo"
    g2.add_data(Reference(ws, min_col=COL["Valor da Causa"], min_row=1, max_row=fim), titles_from_data=True)
    g2.set_categories(Reference(ws, min_col=1, min_row=2, max_row=fim))
    da.add_chart(g2, "A18")

    # ---- intervalos nomeados
    wb.defined_names["ProcessosDados"] = DefinedName("ProcessosDados", attr_text=f"Processos!$A$1:$AC${fim}")
    wb.defined_names["ValoresCausa"] = DefinedName("ValoresCausa", attr_text=f"Processos!$K$2:$K${fim}")
    wb.defined_names["DataReferencia"] = DefinedName("DataReferencia", attr_text="'Parâmetros'!$B$3")

    # ---- aba da tabela dinâmica (células de saída; a definição entra no pós-processamento)
    di = wb.create_sheet("Dinâmica")
    cont = {}
    for l in linhas:
        cont.setdefault(l["Situação"], [0, 0.0])
        cont[l["Situação"]][0] += 1
        cont[l["Situação"]][1] += l["Valor da Causa"]
    di["A3"], di["B3"], di["C3"] = "Rótulos de Linha", "Qtde", "Valor da causa"
    for i, s in enumerate(sorted(cont)):
        di.cell(4 + i, 1, s)
        di.cell(4 + i, 2, cont[s][0])
        di.cell(4 + i, 3, cont[s][1])
    k = 4 + len(cont)
    di.cell(k, 1, "Total Geral")
    di.cell(k, 2, sum(v[0] for v in cont.values()))
    di.cell(k, 3, sum(v[1] for v in cont.values()))
    return wb, linhas, fim


# ------------------------------------------------------------------ pós-processamento do pacote

def _num(v):
    return repr(float(v)) if isinstance(v, float) else str(v)


def _injetar_cache(xml, valores):
    """valores: {ref: valor}; troca <v></v>/<v/> das fórmulas pelo valor calculado em Python."""
    def sub(m):
        ref = m.group(1)
        if ref not in valores:
            return m.group(0)
        v = valores[ref]
        attrs = re.sub(r'\st="[^"]*"', "", m.group(2))
        if isinstance(v, str):
            return f'<c r="{ref}"{attrs} t="str"><f>{m.group(3)}</f><v>{v}</v></c>'
        return f'<c r="{ref}"{attrs}><f>{m.group(3)}</f><v>{_num(v)}</v></c>'
    return re.sub(r'<c r="([A-Z]+\d+)"([^>]*)><f>(.*?)</f><v\s*(?:/>|></v>)</c>', sub, xml, flags=re.S)


def _pivot_partes(linhas, fim):
    situ = sorted({l["Situação"] for l in linhas})
    campos = []
    for nome in CABECALHO:
        if nome == "Situação":
            itens = "".join(f'<s v="{s}"/>' for s in situ)
            campos.append(f'<cacheField name="{nome}" numFmtId="0"><sharedItems count="{len(situ)}">{itens}</sharedItems></cacheField>')
        elif nome == "Valor da Causa":
            campos.append(f'<cacheField name="{nome}" numFmtId="0"><sharedItems containsSemiMixedTypes="0" containsString="0" containsNumber="1" minValue="0" maxValue="1000000"/></cacheField>')
        else:
            campos.append(f'<cacheField name="{nome}" numFmtId="0"><sharedItems/></cacheField>')
    cache = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
             f'<pivotCacheDefinition xmlns="{NS_MAIN}" saveData="0" refreshOnLoad="1" createdVersion="6" '
             f'refreshedVersion="6" minRefreshableVersion="3" recordCount="0">'
             f'<cacheSource type="worksheet"><worksheetSource name="{TABELA}"/></cacheSource>'
             f'<cacheFields count="{len(campos)}">{"".join(campos)}</cacheFields></pivotCacheDefinition>')
    ci = COL["Situação"] - 1
    vi = COL["Valor da Causa"] - 1
    pcampos = []
    for i, _ in enumerate(CABECALHO):
        if i == ci:
            itens = "".join(f'<item x="{k}"/>' for k in range(len(situ))) + '<item t="default"/>'
            pcampos.append(f'<pivotField axis="axisRow" showAll="0"><items count="{len(situ) + 1}">{itens}</items></pivotField>')
        elif i in (0, vi):
            pcampos.append('<pivotField dataField="1" showAll="0"/>')
        else:
            pcampos.append('<pivotField showAll="0"/>')
    n = len(situ)
    linhas_i = "".join(f'<i><x v="{k}"/></i>' if k else '<i><x/></i>' for k in range(n)) + '<i t="grand"><x/></i>'
    tabela = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
              f'<pivotTableDefinition xmlns="{NS_MAIN}" name="DinamicaSituacao" cacheId="1" applyNumberFormats="0" '
              f'applyBorderFormats="0" applyFontFormats="0" applyPatternFormats="0" applyAlignmentFormats="0" '
              f'applyWidthHeightFormats="1" dataCaption="Valores" updatedVersion="6" minRefreshableVersion="3" '
              f'useAutoFormatting="1" itemPrintTitles="1" createdVersion="6" indent="0" outline="1" outlineData="1" '
              f'multipleFieldFilters="0"><location ref="A3:C{4 + n}" firstHeaderRow="0" firstDataRow="1" firstDataCol="1"/>'
              f'<pivotFields count="{len(pcampos)}">{"".join(pcampos)}</pivotFields>'
              f'<rowFields count="1"><field x="{ci}"/></rowFields><rowItems count="{n + 1}">{linhas_i}</rowItems>'
              f'<colFields count="1"><field x="-2"/></colFields><colItems count="2"><i><x/></i><i i="1"><x v="1"/></i></colItems>'
              f'<dataFields count="2"><dataField name="Qtde" fld="0" subtotal="count" baseField="0" baseItem="0"/>'
              f'<dataField name="Valor da causa" fld="{vi}" baseField="0" baseItem="0" numFmtId="4"/></dataFields>'
              f'<pivotTableStyleInfo name="PivotStyleLight16" showRowHeaders="1" showColHeaders="1" showRowStripes="0" '
              f'showColStripes="0" showLastColumn="1"/></pivotTableDefinition>')
    return cache, tabela


def _posprocessar(dados, linhas, fim, variante, pivot=True):
    """dados: {nome_da_parte: bytes} do openpyxl (em ordem). Devolve a lista [(nome, bytes)]."""
    partes = {k: v for k, v in dados.items()}
    wbx = partes["xl/workbook.xml"].decode("utf-8")
    nomes = re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="(rId\d+)"', wbx)
    rels = partes["xl/_rels/workbook.xml.rels"].decode("utf-8")
    alvo_da = {}
    for nome, rid in nomes:
        tgt = re.search(rf'<Relationship[^>]*Id="{rid}"[^>]*/>', rels).group(0)
        alvo_da[nome] = "xl/" + re.search(r'Target="/?(?:xl/)?([^"]+)"', tgt).group(1)

    # ---- valores em cache das fórmulas (como num arquivo salvo pelo Excel)
    exp = indicadores_esperados(linhas)
    proc, ind = {}, {}
    for i, l in enumerate(linhas):
        r = 2 + i
        for nome, v in derivados(l).items():
            proc[f"{_col_letra(COL[nome])}{r}"] = v
    for r in range(2 + len(linhas), fim + 1):   # linhas pré-formatadas vazias
        proc[f"N{r}"], proc[f"W{r}"], proc[f"Y{r}"] = "Sim", "", ""
    ind.update({"B2": exp["total"], "B3": exp["ativos"], "B4": exp["valor_causa_ativos"],
                "B5": exp["prazo_medio"], "B6": exp["economia"], "B7": exp["causa_provavel"],
                "B8": exp["total"] / 350 * 1000, "B9": exp["exito_medio"]})
    for i, v in enumerate(exp["por_situacao"]):
        ind[f"B{12 + i}"] = v
    for aba, vals in (("Processos", proc), ("Indicadores", ind)):
        p = alvo_da[aba]
        partes[p] = _injetar_cache(partes[p].decode("utf-8"), vals).encode("utf-8")

    # ---- tabela: coluna calculada declarada (só "Valor Economizado"); sujo: sem estilo já feito
    tp = next(k for k in partes if k.startswith("xl/tables/table"))
    tx = partes[tp].decode("utf-8")
    t = TABELA
    vc, ve = f"{t}[[#This Row],[Valor da Causa]]", f"{t}[[#This Row],[Valor Estimado]]"
    form = f'IF(OR({vc}="",{ve}=""),"",{vc}-{ve})'
    tx = re.sub(r'(<tableColumn [^>]*name="Valor Economizado"[^>]*?)\s*/>',
                lambda m: f'{m.group(1)}><calculatedColumnFormula>{form}</calculatedColumnFormula></tableColumn>', tx)
    partes[tp] = tx.encode("utf-8")

    # ---- fórmula compartilhada (t="shared") na coluna Y, só na variante suja
    if variante == "sujo":
        p = alvo_da["Processos"]
        xml = partes[p].decode("utf-8")
        ultima = 1 + len(linhas)
        for r in range(2, ultima + 1):
            alvo = re.compile(rf'(<c r="Y{r}"[^>]*>)<f>[^<]*</f>')
            if r == 2:
                xml = alvo.sub(rf'\1<f t="shared" ref="Y2:Y{ultima}" si="0">IF(X2="","",X2-G2)</f>', xml)
            else:
                xml = alvo.sub(r'\1<f t="shared" si="0"/>', xml)
        partes[p] = xml.encode("utf-8")

    # ---- workbook.xml: calcPr como o Excel; pivotCaches
    wbx = re.sub(r"<calcPr[^>]*/>", '<calcPr calcId="191029"/>', wbx)
    ct = partes["[Content_Types].xml"].decode("utf-8")
    novos_ct = []
    if pivot:
        cache, tabela = _pivot_partes(linhas, fim)
        partes["xl/pivotCache/pivotCacheDefinition1.xml"] = cache.encode("utf-8")
        partes["xl/pivotTables/pivotTable1.xml"] = tabela.encode("utf-8")
        partes["xl/pivotTables/_rels/pivotTable1.xml.rels"] = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotCacheDefinition" '
            'Target="../pivotCache/pivotCacheDefinition1.xml"/></Relationships>').encode("utf-8")
        pd = alvo_da["Dinâmica"]
        rp = pd.replace("worksheets/", "worksheets/_rels/") + ".rels"
        partes[rp] = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotTable" '
            'Target="../pivotTables/pivotTable1.xml"/></Relationships>').encode("utf-8")
        rels = rels.replace("</Relationships>",
                            '<Relationship Id="rIdPiv1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotCacheDefinition" '
                            'Target="pivotCache/pivotCacheDefinition1.xml"/></Relationships>')
        wbx = re.sub(r"(<calcPr[^>]*/>)", r'\1<pivotCaches><pivotCache xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" cacheId="1" r:id="rIdPiv1"/></pivotCaches>', wbx)
        novos_ct += [
            '<Override PartName="/xl/pivotCache/pivotCacheDefinition1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.pivotCacheDefinition+xml"/>',
            '<Override PartName="/xl/pivotTables/pivotTable1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.pivotTable+xml"/>']

    # ---- calcChain sintético (fórmulas das abas Processos e Indicadores)
    idx = {n: i + 1 for i, (n, _) in enumerate(nomes)}
    cc = []
    for aba, regex in (("Processos", r'<c r="([A-Z]+\d+)"[^>]*><f'), ("Indicadores", r'<c r="([A-Z]+\d+)"[^>]*><f')):
        for ref in re.findall(regex, partes[alvo_da[aba]].decode("utf-8")):
            cc.append(f'<c r="{ref}" i="{idx[aba]}"/>')
    partes["xl/calcChain.xml"] = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                                  f'<calcChain xmlns="{NS_MAIN}">{"".join(cc)}</calcChain>').encode("utf-8")
    rels = rels.replace("</Relationships>",
                        '<Relationship Id="rIdCalc1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain" '
                        'Target="calcChain.xml"/></Relationships>')
    novos_ct.append('<Override PartName="/xl/calcChain.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/>')
    partes["xl/workbook.xml"] = wbx.encode("utf-8")
    partes["xl/_rels/workbook.xml.rels"] = rels.encode("utf-8")
    partes["[Content_Types].xml"] = ct.replace("</Types>", "".join(novos_ct) + "</Types>").encode("utf-8")

    _compartilhar_textos(partes)
    ordem = ["[Content_Types].xml"] + sorted(k for k in partes if k != "[Content_Types].xml")
    return [(k, partes[k]) for k in ordem]


def _compartilhar_textos(partes):
    """Troca as strings inline do openpyxl por sharedStrings (como o Excel grava): o modelo de
    teste precisa ter xl/sharedStrings.xml para exercitar a leitura de cabeçalhos compartilhados."""
    tabela, ordem = {}, []
    padrao = re.compile(r'<c r="([A-Z]+\d+)"((?: s="\d+")?) t="inlineStr"><is><t(?: [^>]*)?>(.*?)</t></is></c>', re.S)
    total = 0
    for nome in sorted(partes):
        if not (nome.startswith("xl/worksheets/sheet") and nome.endswith(".xml")):
            continue

        def sub(m):
            nonlocal total
            txt = m.group(3)
            if txt not in tabela:
                tabela[txt] = len(ordem)
                ordem.append(txt)
            total += 1
            return f'<c r="{m.group(1)}"{m.group(2)} t="s"><v>{tabela[txt]}</v></c>'
        partes[nome] = padrao.sub(sub, partes[nome].decode("utf-8")).encode("utf-8")
    itens = "".join(
        f'<si><t xml:space="preserve">{t}</t></si>' if t != t.strip() else f"<si><t>{t}</t></si>" for t in ordem)
    partes["xl/sharedStrings.xml"] = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<sst xmlns="{NS_MAIN}" count="{total}" '
        f'uniqueCount="{len(ordem)}">{itens}</sst>').encode("utf-8")
    partes["xl/_rels/workbook.xml.rels"] = partes["xl/_rels/workbook.xml.rels"].replace(
        b"</Relationships>", b'<Relationship Id="rIdSst1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/></Relationships>')
    partes["[Content_Types].xml"] = partes["[Content_Types].xml"].replace(
        b"</Types>", b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')


def _col_letra(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def gerar(destino, variante="completo", n_linhas=10, pivot=True):
    """Gera o modelo em `destino`. Devolve a lista de linhas base (dicts) usada."""
    wb, linhas, fim = construir(variante, n_linhas)
    if variante == "sujo":
        aux = wb.create_sheet("Auxiliar")
        for i in range(6000):
            aux.cell(1 + i, 1, f"texto auxiliar fictício {i:05d} " + "x" * (i % 17))
    buf = io.BytesIO()
    wb.save(buf)
    with zipfile.ZipFile(buf) as z:
        dados = {i.filename: z.read(i.filename) for i in z.infolist()}
    partes = _posprocessar(dados, linhas, fim, variante, pivot)
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as out:
        for nome, conteudo in partes:
            zi = zipfile.ZipInfo(nome, ZIP_DATA)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            out.writestr(zi, conteudo)
    return linhas


if __name__ == "__main__":
    dest = sys.argv[1] if len(sys.argv) > 1 else "modelo_b.xlsx"
    var = sys.argv[2] if len(sys.argv) > 2 else "completo"
    gerar(dest, var)
    print(f"{dest} gerado ({var})")
